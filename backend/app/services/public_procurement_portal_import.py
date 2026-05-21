import io
import hashlib
import re
import unicodedata
from copy import copy, deepcopy
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from fastapi import HTTPException
from openpyxl.cell.cell import MergedCell
from openpyxl import Workbook, load_workbook


DOCUMENT_SECTION_LABELS = {
    "budget": "Presupuesto",
    "apus": "APUs",
    "resources": "Recursos",
    "vae": "VAE / desagregación",
}

ANALYSIS_CACHE_LIMIT = 12

SEMANTIC_STOPWORDS = {
    "a",
    "al",
    "con",
    "de",
    "del",
    "el",
    "en",
    "la",
    "las",
    "lo",
    "los",
    "para",
    "por",
    "sin",
    "un",
    "una",
    "uno",
    "y",
}

MANUAL_ROLE_SECTION_MAP = {
    "budget": {"budget": True, "apus": False, "resources": False, "vae": False},
    "apus_resources": {"budget": False, "apus": True, "resources": True, "vae": False},
    "integrated": {"budget": True, "apus": True, "resources": True, "vae": False},
    "vae": {"budget": False, "apus": True, "resources": True, "vae": True},
}

MANUAL_ROLE_KIND_MAP = {
    "budget": "presupuesto_only",
    "apus_resources": "tecnico_parcial",
    "integrated": "presupuesto_apus_integrados",
    "vae": "apus_vae",
}

DEFAULT_RESOURCE_UNIT_BY_TYPE = {
    "equipo": "Hora",
    "mano_obra": "Hora",
    "transporte": "Hora",
}

SOCE_BUDGET_UNIT_PATTERN = r"m3-km|m3/km|m2|m3|kg|Kg|u|glb|gl|m|lt|pto"
SOCE_NUMERIC_TOLERANCE = Decimal("0.02")


def _a_num(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = re.sub(r"[^\d.,-]", "", str(value).strip())
    if not cleaned:
        return None
    has_dot = "." in cleaned
    has_comma = "," in cleaned
    if has_dot and has_comma:
        if cleaned.rfind(".") > cleaned.rfind(","):
            cleaned = cleaned.replace(",", "")
        else:
            cleaned = cleaned.replace(".", "").replace(",", ".")
    elif has_comma:
        left, *rest = cleaned.split(",")
        if len(rest) == 1 and 1 <= len(rest[0]) <= 5 and len(left) <= 6:
            cleaned = f"{left}.{rest[0]}"
        else:
            cleaned = cleaned.replace(",", "")
    elif has_dot and cleaned.count(".") > 1:
        cleaned = cleaned.replace(".", "", cleaned.count(".") - 1)
    try:
        return float(cleaned)
    except ValueError:
        return None


def _clean_text(value):
    if not value:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _fold_latin_text(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value or "")
    return "".join(char for char in normalized if not unicodedata.combining(char))


def _normalize_signature_fragment(value: str) -> str:
    normalized = _fold_latin_text(_clean_text(value)).lower()
    normalized = re.sub(r"[^a-z0-9]+", "_", normalized)
    return normalized.strip("_")


def _normalize_semantic_tokens(value: str) -> list[str]:
    normalized = _fold_latin_text(_clean_text(value)).lower()
    normalized = re.sub(r"[^a-z0-9\s]+", " ", normalized)
    normalized = re.sub(r"\s+", " ", normalized).strip()
    tokens: list[str] = []
    for token in normalized.split():
        if len(token) <= 1 or token in SEMANTIC_STOPWORDS:
            continue
        token = re.sub(r"(\d+)[a-z]+$", r"\1", token)
        if token.endswith("es") and len(token) > 4:
            token = token[:-2]
        elif token.endswith("s") and len(token) > 4:
            token = token[:-1]
        tokens.append(token)
    deduped: list[str] = []
    seen: set[str] = set()
    for token in tokens:
        if token in seen:
            continue
        seen.add(token)
        deduped.append(token)
    return deduped


def _build_semantic_phrase_key(*values) -> str:
    tokens: list[str] = []
    seen: set[str] = set()
    for value in values:
        for token in _normalize_semantic_tokens(value or ""):
            if token in seen:
                continue
            seen.add(token)
            tokens.append(token)
    return "_".join(tokens[:8])


def _coerce_decimal(value) -> Decimal:
    if value in (None, ""):
        return Decimal("0.00")
    try:
        return Decimal(str(value))
    except Exception:
        return Decimal("0.00")


def _money_decimal(value) -> Decimal:
    numeric = _a_num(value)
    if numeric is None:
        return Decimal("0")
    return Decimal(str(numeric)).quantize(Decimal("0.0001"))


def _nearly_equal(left, right, tolerance: Decimal = SOCE_NUMERIC_TOLERANCE) -> bool:
    return abs(_money_decimal(left) - _money_decimal(right)) <= tolerance


def _extract_percentage_number(value) -> float | None:
    if value in (None, ""):
        return None
    match = re.search(r"(\d+(?:[.,]\d+)?)\s*%", str(value))
    if not match:
        return None
    try:
        return float(match.group(1).replace(",", "."))
    except ValueError:
        return None


def _coerce_percentage_ratio(value) -> float | None:
    numeric = _a_num(value)
    if numeric is None:
        return None
    if numeric > 1:
        return numeric / 100
    return numeric


def _build_unique_sheet_name(workbook, preferred_name: str) -> str:
    safe_name = (preferred_name or "Hoja")[:31].strip() or "Hoja"
    if safe_name not in workbook.sheetnames:
        return safe_name
    base_name = safe_name[:28].rstrip("_") or "Hoja"
    counter = 2
    while True:
        candidate = f"{base_name}_{counter}"[:31]
        if candidate not in workbook.sheetnames:
            return candidate
        counter += 1


def _quote_excel_sheet_name(value: str) -> str:
    safe_value = str(value or "").replace("'", "''")
    return f"'{safe_value}'"


def _clear_sheet_values(ws, *, start_row: int = 1, start_col: int = 1, end_col: int | None = None) -> None:
    max_col = end_col or ws.max_column
    for row_index in range(start_row, ws.max_row + 1):
        for column_index in range(start_col, max_col + 1):
            cell = ws.cell(row=row_index, column=column_index)
            if isinstance(cell, MergedCell):
                continue
            cell.value = None


def _copy_row_style(ws, source_row: int, target_row: int, *, end_col: int = 13) -> None:
    source_height = ws.row_dimensions[source_row].height
    if source_height is not None:
        ws.row_dimensions[target_row].height = source_height
    for column_index in range(1, end_col + 1):
        source_cell = ws.cell(row=source_row, column=column_index)
        target_cell = ws.cell(row=target_row, column=column_index)
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format
        if source_cell.font:
            target_cell.font = copy(source_cell.font)
        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)
        if source_cell.border:
            target_cell.border = copy(source_cell.border)
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)
        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)


def _build_full_apu_export_entries(analysis: dict, apus: list[dict], resources: list[dict]) -> list[dict]:
    rows = list(analysis.get("rubros") or [])
    normalized_apus = [dict(apu or {}) for apu in apus]
    normalized_resources = [dict(resource or {}) for resource in resources]
    resources_by_temp_id: dict[str, list[dict]] = {}
    resources_by_code: dict[str, list[dict]] = {}
    for resource in normalized_resources:
        temp_id = str(resource.get("apu_temp_id") or "").strip()
        if temp_id:
            resources_by_temp_id.setdefault(temp_id, []).append(resource)
        parent_code = str(resource.get("parent_apu_codigo") or "").strip()
        if parent_code:
            resources_by_code.setdefault(parent_code, []).append(resource)

    apus_by_code: dict[str, list[dict]] = {}
    for apu in normalized_apus:
        code = str(apu.get("codigo") or "").strip()
        if code:
            apus_by_code.setdefault(code, []).append(apu)

    used_temp_ids: set[str] = set()
    key_counts: dict[str, int] = {}
    export_entries: list[dict] = []

    for index, row in enumerate(rows, start=1):
        row_code = str(row.get("codigo") or "").strip()
        row_description = str(row.get("descripcion") or "").strip()
        row_unit = str(row.get("unidad") or "").strip()
        candidate = None
        for apu in apus_by_code.get(row_code, []):
            temp_id = str(apu.get("temp_id") or "").strip()
            if temp_id and temp_id in used_temp_ids:
                continue
            candidate = apu
            break
        if candidate is None:
            row_keys = _build_alignment_keys(row_code, row_description, row_unit, row.get("capitulo"))
            for apu in normalized_apus:
                temp_id = str(apu.get("temp_id") or "").strip()
                if temp_id and temp_id in used_temp_ids:
                    continue
                apu_keys = _build_alignment_keys(apu.get("codigo"), apu.get("descripcion"), apu.get("unidad"), apu.get("chapter"))
                if row_keys & apu_keys:
                    candidate = apu
                    break

        temp_id = str((candidate or {}).get("temp_id") or "").strip()
        if temp_id:
            used_temp_ids.add(temp_id)

        base_key = row_code or str((candidate or {}).get("codigo") or "").strip() or f"RUBRO_{index}"
        key_counts[base_key] = key_counts.get(base_key, 0) + 1
        export_key = base_key if key_counts[base_key] == 1 else f"{base_key}_{key_counts[base_key] - 1}"
        entry_resources = []
        if temp_id:
            entry_resources = list(resources_by_temp_id.get(temp_id) or [])
        if not entry_resources and row_code:
            entry_resources = list(resources_by_code.get(row_code) or [])
        export_entries.append({
            "item_index": index,
            "export_key": export_key,
            "sheet_code": export_key,
            "original_code": row_code or str((candidate or {}).get("codigo") or "").strip() or export_key,
            "descripcion": str((candidate or {}).get("descripcion") or row_description or "APU detectado").strip(),
            "unidad": str((candidate or {}).get("unidad") or row_unit or "-").strip() or "-",
            "cantidad": row.get("cantidad") or 0,
            "budget_unit_price": _a_num(row.get("precio_u")) or 0,
            "budget_total_price": _a_num(row.get("precio_total")) or 0,
            "chapter": row.get("capitulo") or "",
            "source_apu": candidate or {},
            "resources": entry_resources,
            "status": str((candidate or {}).get("status") or ("Pendiente" if not entry_resources else "Listo")).strip() or "Pendiente",
            "vae_percentage": (candidate or {}).get("vae_percentage"),
            "cpc_code": next((str(resource.get("cpc_code") or "").strip() for resource in entry_resources if str(resource.get("cpc_code") or "").strip()), ""),
        })

    return export_entries


def _populate_catalog_sheet(ws, items: list[dict]) -> None:
    ws.delete_rows(3, max(0, ws.max_row - 2))
    seen_codes: set[str] = set()
    row_index = 3
    item_index = 1
    for resource in items:
        metrics = _extract_detected_resource_metrics(resource)
        code = str(resource.get("codigo") or "").strip()
        if not code or code in seen_codes:
            continue
        seen_codes.add(code)
        ws.cell(row=row_index, column=1, value=item_index)
        ws.cell(row=row_index, column=2, value=code)
        ws.cell(row=row_index, column=3, value=metrics.get("label") or resource.get("descripcion") or code)
        ws.cell(row=row_index, column=4, value=metrics.get("unit") or resource.get("section_label") or DEFAULT_RESOURCE_UNIT_BY_TYPE.get(str(resource.get("resource_type") or ""), "-"))
        ws.cell(row=row_index, column=5, value=_a_num(metrics.get("price")) or 0)
        ws.cell(row=row_index, column=6, value=resource.get("cpc_code") or "")
        row_index += 1
        item_index += 1


def _build_apu_sheet(
    workbook,
    template_sheet,
    apu_entry: dict,
    *,
    default_vae_percentage: float | None = None,
) -> dict:
    preferred_name = f"APU_{_normalize_signature_fragment(apu_entry.get('sheet_code') or apu_entry.get('original_code') or 'APU')[:24]}" or "APU_GENERADO"
    ws = workbook.copy_worksheet(template_sheet)
    ws.title = _build_unique_sheet_name(workbook, preferred_name)
    _clear_sheet_values(ws, start_row=1, start_col=1, end_col=13)

    ws["A1"] = "ANÁLISIS DE PRECIOS UNITARIOS"
    ws["I1"] = "DETERMINACIÓN DEL VAE DEL RUBRO"
    ws["A2"] = "Código:"
    ws["C2"] = apu_entry.get("export_key") or apu_entry.get("original_code") or "Sin código"
    ws["A3"] = "Descripción:"
    ws["C3"] = apu_entry.get("descripcion") or "APU detectado"
    ws["A4"] = "Unidad:"
    ws["C4"] = apu_entry.get("unidad") or "-"
    ws["A5"] = "Fecha:"
    ws["C5"] = "—"

    grouped_resources = {
        "EQUIPO Y HERRAMIENTA": [resource for resource in apu_entry.get("resources") or [] if resource.get("resource_type") == "equipo"],
        "MATERIALES": [resource for resource in apu_entry.get("resources") or [] if resource.get("resource_type") == "material"],
        "TRANSPORTE": [resource for resource in apu_entry.get("resources") or [] if resource.get("resource_type") == "transporte"],
        "MANO DE OBRA": [resource for resource in apu_entry.get("resources") or [] if resource.get("resource_type") == "mano_obra"],
    }
    section_templates = {
        "EQUIPO Y HERRAMIENTA": {"label_row": 7, "header_row": 8, "resource_row": 9, "subtotal_row": 10, "catalog": "CAT_EQUIPO"},
        "MATERIALES": {"label_row": 12, "header_row": 13, "resource_row": 9, "subtotal_row": 14, "catalog": "CAT_MATERIALES"},
        "TRANSPORTE": {"label_row": 16, "header_row": 17, "resource_row": 9, "subtotal_row": 18, "catalog": "CAT_TRANSPORTE"},
        "MANO DE OBRA": {"label_row": 20, "header_row": 21, "resource_row": 22, "subtotal_row": 24, "catalog": "CAT_MANO_OBRA"},
    }

    current_row = 7
    section_subtotals: dict[str, int] = {}
    section_resource_ranges: dict[str, tuple[int, int] | None] = {}
    direct_formula_rows: list[int] = []

    for section_label in ["EQUIPO Y HERRAMIENTA", "MATERIALES", "TRANSPORTE", "MANO DE OBRA"]:
        template_info = section_templates[section_label]
        resources = list(grouped_resources.get(section_label) or [])
        _copy_row_style(ws, template_info["label_row"], current_row)
        ws.cell(row=current_row, column=1, value=section_label)
        current_row += 1

        _copy_row_style(ws, template_info["header_row"], current_row)
        header_values = [template_sheet.cell(template_info["header_row"], col).value for col in range(1, 14)]
        for col_index, value in enumerate(header_values, start=1):
            ws.cell(row=current_row, column=col_index, value=value)
        current_row += 1

        resource_start_row = current_row
        prototype_resource_row = template_info["resource_row"]
        catalog_name = template_info["catalog"]
        if resources:
            for resource in resources:
                _copy_row_style(ws, prototype_resource_row, current_row)
                metrics = _extract_detected_resource_metrics(resource)
                code = str(resource.get("codigo") or "").strip() or f"RECURSO_{current_row}"
                quantity = _a_num(metrics.get("quantity")) or 1
                price = _a_num(metrics.get("price")) or 0
                total = _a_num(metrics.get("total"))
                rendim = quantity if section_label == "MATERIALES" else (_a_num(resource.get("rendimiento")) or 1)
                if section_label == "EQUIPO Y HERRAMIENTA":
                    rendim = _a_num(resource.get("rendimiento")) or quantity or 1
                if section_label == "TRANSPORTE":
                    rendim = _a_num(resource.get("distance")) or _a_num(resource.get("distancia")) or 1
                if section_label == "MANO DE OBRA":
                    rendim = _a_num(resource.get("rendimiento")) or 1

                ws.cell(row=current_row, column=1, value=code)
                ws.cell(row=current_row, column=2, value=f'=IFERROR(VLOOKUP(A{current_row},{catalog_name}!$B:$C,2,0),"")')
                if section_label == "MANO DE OBRA":
                    ws.cell(row=current_row, column=3, value=f'=IFERROR(VLOOKUP(A{current_row},{catalog_name}!$B:$D,3,0),"")')
                else:
                    ws.cell(row=current_row, column=3, value=f'=IFERROR(VLOOKUP(A{current_row},{catalog_name}!$B:$D,3,0),"")')
                ws.cell(row=current_row, column=4, value=quantity)
                ws.cell(row=current_row, column=5, value=f'=IFERROR(VLOOKUP(A{current_row},{catalog_name}!$B:$E,4,0),0)')
                if section_label == "MATERIALES":
                    ws.cell(row=current_row, column=6, value="")
                    if total is not None and quantity:
                        ws.cell(row=current_row, column=7, value=f'=ROUND(D{current_row}*E{current_row},2)')
                    else:
                        ws.cell(row=current_row, column=7, value=f'=ROUND(D{current_row}*E{current_row},2)')
                else:
                    ws.cell(row=current_row, column=6, value=rendim)
                    ws.cell(row=current_row, column=7, value=f'=ROUND(D{current_row}*E{current_row}*F{current_row},2)')
                ws.cell(row=current_row, column=9, value="")
                ws.cell(row=current_row, column=10, value=f'=IFERROR(LEFT(VLOOKUP(A{current_row},{catalog_name}!$B:$F,5,0),9),"")')
                ws.cell(row=current_row, column=11, value=f'=IF(L{current_row}<=0,"NP",IF(L{current_row}>=1,"EP","ND"))')
                ws.cell(row=current_row, column=12, value=f'=IFERROR(VLOOKUP(IF(LEN(J{current_row})>9,LEFT(J{current_row},9),J{current_row}),UMBRAL_SERCOP!$A$3:$D$29785,4,0),0)')
                ws.cell(row=current_row, column=13, value=f'=IFERROR(I{current_row}*L{current_row},0)')
                direct_formula_rows.append(current_row)
                current_row += 1
            resource_end_row = current_row - 1
            section_resource_ranges[section_label] = (resource_start_row, resource_end_row)
        else:
            section_resource_ranges[section_label] = None

        subtotal_row = current_row
        _copy_row_style(ws, template_info["subtotal_row"], subtotal_row)
        ws.cell(row=subtotal_row, column=1, value=f"Subtotal {section_label}:")
        if section_resource_ranges[section_label]:
            start_row, end_row = section_resource_ranges[section_label]
            ws.cell(row=subtotal_row, column=7, value=f'=ROUND(SUM(G{start_row}:G{end_row}),2)')
            ws.cell(row=subtotal_row, column=13, value=f'=SUM(M{start_row}:M{end_row})')
        else:
            ws.cell(row=subtotal_row, column=7, value='=0')
            ws.cell(row=subtotal_row, column=13, value='=0')
        section_subtotals[section_label] = subtotal_row
        current_row += 2

    summary_start_row = current_row
    for offset, source_row in enumerate(range(26, 35), start=0):
        _copy_row_style(ws, source_row, summary_start_row + offset)
        for col_index in range(1, 14):
            source_value = template_sheet.cell(source_row, col_index).value
            if source_value is not None:
                ws.cell(row=summary_start_row + offset, column=col_index, value=source_value)

    equipment_subtotal_row = section_subtotals.get("EQUIPO Y HERRAMIENTA", summary_start_row)
    materials_subtotal_row = section_subtotals.get("MATERIALES", summary_start_row + 1)
    transport_subtotal_row = section_subtotals.get("TRANSPORTE", summary_start_row + 2)
    labor_subtotal_row = section_subtotals.get("MANO DE OBRA", summary_start_row + 3)

    direct_total_row = summary_start_row + 6
    indirect_total_row = summary_start_row + 7
    unit_total_row = summary_start_row + 8
    vae_total_row = summary_start_row + 6

    ws.cell(row=summary_start_row + 1, column=7, value=f'=G{equipment_subtotal_row}')
    ws.cell(row=summary_start_row + 2, column=7, value=f'=G{materials_subtotal_row}')
    ws.cell(row=summary_start_row + 3, column=7, value=f'=G{transport_subtotal_row}')
    ws.cell(row=summary_start_row + 4, column=7, value=f'=G{labor_subtotal_row}')
    fallback_direct_cost = _a_num(apu_entry.get("budget_unit_price")) or 0
    ws.cell(
        row=direct_total_row,
        column=7,
        value=f'=IF(ROUND(G{summary_start_row + 1}+G{summary_start_row + 2}+G{summary_start_row + 3}+G{summary_start_row + 4},2)=0,{fallback_direct_cost},ROUND(G{summary_start_row + 1}+G{summary_start_row + 2}+G{summary_start_row + 3}+G{summary_start_row + 4},2))',
    )
    ws.cell(row=indirect_total_row, column=6, value='=PARAMETROS!B3')
    ws.cell(row=indirect_total_row, column=7, value=f'=ROUND(G{direct_total_row}*PARAMETROS!B3,2)')
    ws.cell(row=unit_total_row, column=7, value=f'=ROUND(G{direct_total_row}+G{indirect_total_row},2)')

    vae_formula_terms = [f"M{row_ref}" for row_ref in direct_formula_rows]
    if vae_formula_terms:
        ws.cell(row=vae_total_row, column=13, value=f'=SUM({",".join(vae_formula_terms)})')
    else:
        ws.cell(row=vae_total_row, column=13, value=_coerce_percentage_ratio(apu_entry.get("vae_percentage") or default_vae_percentage) or 0)

    for section_label, subtotal_row in section_subtotals.items():
        ws.cell(row=subtotal_row, column=9, value=f'=IFERROR(G{subtotal_row}/$G${direct_total_row},0)')
        ws.cell(row=subtotal_row, column=10, value=f'Subtotal VAE {section_label}:')

    for row_ref in direct_formula_rows:
        ws.cell(row=row_ref, column=9, value=f'=IFERROR(G{row_ref}/$G${direct_total_row},0)')

    return {
        "sheet_name": ws.title,
        "direct_total_row": direct_total_row,
        "unit_total_row": unit_total_row,
        "vae_total_row": vae_total_row,
    }


def _build_budget_line_signature(row: dict) -> str:
    code = _normalize_signature_fragment(row.get("codigo") or "")
    description = _normalize_signature_fragment(row.get("descripcion") or "")
    semantic_description = _build_semantic_phrase_key(row.get("descripcion"), row.get("capitulo"))
    unit = _normalize_signature_fragment(row.get("unidad") or "")
    chapter = _normalize_signature_fragment(row.get("capitulo") or "")
    for fragment in (code, description, unit, chapter):
        if fragment:
            if semantic_description and unit and fragment in {code, description}:
                return f"{fragment}_{semantic_description}_{unit}"
            return fragment
    if semantic_description:
        return f"{semantic_description}_{unit}" if unit else semantic_description
    return f"budget_line_{_normalize_signature_fragment(row.get('nro') or '') or 'sin_codigo'}"


def _build_apu_signature(apu: dict) -> str:
    code = _normalize_signature_fragment(apu.get("codigo") or "")
    description = _normalize_signature_fragment(apu.get("descripcion") or "")
    semantic_description = _build_semantic_phrase_key(apu.get("descripcion"), apu.get("chapter"), apu.get("capitulo"))
    unit = _normalize_signature_fragment(apu.get("unidad") or "")
    chapter = _normalize_signature_fragment(apu.get("chapter") or apu.get("capitulo") or "")
    for fragment in (
        f"{code}_{unit}" if code and unit else "",
        f"{code}_{semantic_description}_{unit}" if code and semantic_description and unit else "",
        f"{description}_{unit}" if description and unit else "",
        f"{semantic_description}_{unit}" if semantic_description and unit else "",
        code,
        description,
        semantic_description,
        chapter,
    ):
        if fragment:
            return fragment
    return _normalize_signature_fragment(apu.get("temp_id") or apu.get("source_signature") or "apu_sin_firma")


def _build_resource_signature(resource: dict) -> str:
    resource_type = _normalize_signature_fragment(resource.get("type") or resource.get("resource_type") or "")
    description = _normalize_signature_fragment(resource.get("descripcion") or resource.get("name") or "")
    unit = _normalize_signature_fragment(resource.get("unidad") or "")
    parent_apu = _normalize_signature_fragment(resource.get("parent_apu_codigo") or resource.get("parent_apu") or "")
    for fragment in (
        f"{resource_type}_{description}_{unit}_{parent_apu}" if description else "",
        f"{resource_type}_{description}_{unit}" if description else "",
        description,
    ):
        if fragment:
            return fragment
    return _normalize_signature_fragment(resource.get("temp_id") or "resource_sin_firma")


def _build_vae_signature(entry: dict) -> str:
    label = _normalize_signature_fragment(entry.get("label") or entry.get("descripcion") or entry.get("name") or "")
    value = _normalize_signature_fragment(entry.get("value") or entry.get("codigo") or "")
    for fragment in (f"{label}_{value}" if label and value else "", label, value):
        if fragment:
            return fragment
    return _normalize_signature_fragment(entry.get("temp_id") or "vae_sin_firma")


def _extract_detected_resource_metrics(resource: dict) -> dict:
    raw_description = _clean_text(resource.get("descripcion") or "")
    if not raw_description:
        return {
            "label": "",
            "unit": resource.get("unidad") or None,
            "quantity": resource.get("cantidad") or None,
            "price": resource.get("precio") or None,
            "total": resource.get("total") or None,
            "weight_percentage": None,
        }

    percentage_match = re.search(r"(\d+(?:[.,]\d+)?)\s*%$", raw_description)
    weight_percentage = percentage_match.group(1).replace(",", ".") if percentage_match else None
    working = raw_description[:percentage_match.start()].strip() if percentage_match else raw_description

    number_matches = list(re.finditer(r"\d+(?:[.,]\d+)?", working))
    numbers = [match.group(0).replace(".", "").replace(",", ".") if "," in match.group(0) else match.group(0) for match in number_matches]

    unit = resource.get("unidad") or None
    quantity = resource.get("cantidad") or None
    price = resource.get("precio") or None
    total = resource.get("total") or None

    if len(numbers) >= 1 and quantity in (None, ""):
        quantity = numbers[-4] if len(numbers) >= 4 else numbers[0]
    if len(numbers) >= 2 and price in (None, ""):
        price = numbers[-3] if len(numbers) >= 4 else (numbers[1] if len(numbers) > 1 else None)
    if len(numbers) >= 1 and total in (None, ""):
        total = numbers[-1]

    unit_match = re.search(r"\b(Hora|hora|m3|m2|km|kg|lt|gl|und|ml|m|l|u)\b", working)
    if unit in (None, "") and unit_match:
        unit = unit_match.group(1)
    if unit in (None, ""):
        unit = DEFAULT_RESOURCE_UNIT_BY_TYPE.get(str(resource.get("resource_type") or ""))

    label = working
    if unit_match:
        unit_index = unit_match.start()
        label = working[:unit_index].strip() or working
    parenthetical_label_match = re.match(r"^(.+?\([^)]+\))\s+\d", working)
    if parenthetical_label_match:
        label = parenthetical_label_match.group(1).strip()
    if number_matches:
        first_number_index = number_matches[0].start()
        if (
            first_number_index > 0
            and not parenthetical_label_match
            and (not unit_match or first_number_index < unit_match.start())
        ):
            label_candidate = working[:first_number_index].strip()
            if len(label_candidate) >= 3:
                label = label_candidate
    label = re.sub(r"\s+", " ", label).strip(" -")
    if unit and label.lower().endswith(f" {str(unit).lower()}"):
        label = label[: -(len(str(unit)) + 1)].strip(" -")
    trailing_tokens = label.split()
    while (
        len(trailing_tokens) > 1
        and trailing_tokens[-1].lower() in {"a", "al", "con", "de", "del", "en", "para", "por", "sin", "y"}
    ):
        trailing_tokens.pop()
    label = " ".join(trailing_tokens).strip(" -")

    weak_label = (
        not label
        or len(label) < 4
        or len(re.findall(r"[A-Za-zÁÉÍÓÚáéíóúÑñ]", label)) < 3
        or re.fullmatch(r"(hora|m3|m2|m|km|kg|lt|l|gl|u|und|ml)(?:\s+\d+(?:[.,]\d+)*)?", label.lower() or "") is not None
    )
    if weak_label:
        fallback_prefix = {
            "equipo": "Equipo",
            "material": "Material",
            "mano_obra": "Mano de obra",
            "transporte": "Transporte",
        }.get(str(resource.get("resource_type") or ""), "Recurso")
        code = str(resource.get("codigo") or "").strip()
        label = f"{fallback_prefix} {code}".strip() if code else fallback_prefix

    return {
        "label": label or raw_description,
        "unit": unit,
        "quantity": quantity,
        "price": price,
        "total": total,
        "weight_percentage": weight_percentage,
    }


def _normalize_detected_resource(resource: dict) -> dict:
    normalized = dict(resource or {})
    metrics = _extract_detected_resource_metrics(normalized)
    normalized["descripcion"] = metrics.get("label") or normalized.get("descripcion") or ""
    normalized["unidad"] = metrics.get("unit") or normalized.get("unidad")
    normalized["cantidad"] = metrics.get("quantity") or normalized.get("cantidad")
    normalized["precio"] = metrics.get("price") or normalized.get("precio")
    normalized["total"] = metrics.get("total") or normalized.get("total")
    normalized["weight_percentage"] = metrics.get("weight_percentage") or normalized.get("weight_percentage")
    return normalized


def _merge_unique_items(items: list[dict], signature_builder) -> tuple[list[dict], int]:
    merged: list[dict] = []
    seen: set[str] = set()
    duplicates = 0
    for item in items:
        candidate = dict(item or {})
        signature = signature_builder(candidate)
        candidate["source_signature"] = candidate.get("source_signature") or signature
        if signature in seen:
            duplicates += 1
            continue
        seen.add(signature)
        merged.append(candidate)
    return merged, duplicates


def _build_alignment_keys(*values) -> set[str]:
    normalized_parts = [_normalize_signature_fragment(value or "") for value in values]
    semantic_key = _build_semantic_phrase_key(*values[:2])
    keys: set[str] = set()
    if normalized_parts[0]:
        keys.add(normalized_parts[0])
    if normalized_parts[1]:
        keys.add(normalized_parts[1])
    if normalized_parts[0] and normalized_parts[2]:
        keys.add(f"{normalized_parts[0]}_{normalized_parts[2]}")
    if normalized_parts[1] and normalized_parts[2]:
        keys.add(f"{normalized_parts[1]}_{normalized_parts[2]}")
    if semantic_key:
        keys.add(semantic_key)
        if normalized_parts[2]:
            keys.add(f"{semantic_key}_{normalized_parts[2]}")
    return {key for key in keys if key}


def _build_budget_apu_match_key(description, unit, unit_price) -> str:
    description_key = _build_semantic_phrase_key(description)
    unit_key = _normalize_signature_fragment(unit or "")
    price_key = f"{_money_decimal(unit_price):.2f}"
    return f"{description_key}|{unit_key}|{price_key}"


def _build_unit_price_match_key(unit, unit_price) -> str:
    unit_key = _normalize_signature_fragment(unit or "")
    price_key = f"{_money_decimal(unit_price):.2f}"
    return f"{unit_key}|{price_key}"


def _semantic_overlap_score(left, right) -> int:
    left_tokens = set(_normalize_semantic_tokens(left or ""))
    right_tokens = set(_normalize_semantic_tokens(right or ""))
    if not left_tokens or not right_tokens:
        return 0
    return len(left_tokens & right_tokens)


def _resource_composition_signature(resource: dict) -> tuple:
    return (
        _normalize_signature_fragment(resource.get("codigo") or resource.get("descripcion") or ""),
        _normalize_signature_fragment(resource.get("resource_type") or ""),
        _normalize_signature_fragment(resource.get("unidad") or ""),
        f"{_money_decimal(resource.get('cantidad')):.4f}",
        f"{_money_decimal(resource.get('precio')):.4f}",
        f"{_money_decimal(resource.get('rendimiento')):.4f}",
        f"{_money_decimal(resource.get('total') or resource.get('subtotal')):.4f}",
        _normalize_signature_fragment(resource.get("cod_cpc_codigo") or ""),
    )


def _is_valid_resource_code(value: str) -> bool:
    code = _clean_text(value).upper()
    if code in {"COSTOS", "COSTO", "DOS", "TRES", "CUATRO", "CINCO", "SON", "DE", "LOS", "UNIDOS"}:
        return False
    return re.search(r"\d", code) is not None and len(code) >= 3


def _apu_compositions_conflict(first: dict, second: dict) -> bool:
    if _normalize_signature_fragment(first.get("descripcion") or "") != _normalize_signature_fragment(second.get("descripcion") or ""):
        return True
    if _normalize_signature_fragment(first.get("unidad") or "") != _normalize_signature_fragment(second.get("unidad") or ""):
        return True
    if not _nearly_equal(first.get("unit_price") or first.get("precio_unitario"), second.get("unit_price") or second.get("precio_unitario")):
        return True

    first_resources = {
        _clean_text(resource.get("codigo")): resource
        for resource in list(first.get("resources") or [])
        if _is_valid_resource_code(resource.get("codigo"))
    }
    second_resources = {
        _clean_text(resource.get("codigo")): resource
        for resource in list(second.get("resources") or [])
        if _is_valid_resource_code(resource.get("codigo"))
    }
    if set(first_resources) != set(second_resources):
        return True

    for code in first_resources:
        left = first_resources[code]
        right = second_resources[code]
        if left.get("_parse_quality") == "incomplete" or right.get("_parse_quality") == "incomplete":
            continue
        for field in ("cantidad", "precio", "rendimiento", "total", "subtotal"):
            if left.get(field) in (None, "") or right.get(field) in (None, ""):
                continue
            if not _nearly_equal(left.get(field), right.get(field)):
                return True
        left_cpc = _clean_text(left.get("cod_cpc_codigo"))
        right_cpc = _clean_text(right.get("cod_cpc_codigo"))
        if left_cpc and right_cpc and left_cpc != right_cpc:
            return True
    return False


def _apu_composition_signature(apu: dict) -> tuple:
    resources = sorted(_resource_composition_signature(resource) for resource in list(apu.get("resources") or []))
    return (
        _normalize_signature_fragment(apu.get("descripcion") or ""),
        _normalize_signature_fragment(apu.get("unidad") or ""),
        f"{_money_decimal(apu.get('unit_price') or apu.get('precio_unitario')):.2f}",
        tuple(resources),
    )


def _build_apu_identity_key(apu: dict) -> str:
    code = _clean_text(apu.get("codigo"))
    if code:
        return f"codigo:{code}"
    description = _normalize_signature_fragment(apu.get("descripcion") or "")
    unit = _normalize_signature_fragment(apu.get("unidad") or "")
    price = f"{_money_decimal(apu.get('unit_price') or apu.get('precio_unitario') or apu.get('precio_u')):.2f}"
    if description and unit:
        return f"funcional:{description}|{unit}|{price}"
    return ""


def _dedupe_apus_by_identity(apus: list[dict], resources: list[dict]) -> tuple[list[dict], list[dict], list[dict]]:
    by_identity: dict[str, dict] = {}
    ordered: list[dict] = []
    incidents: list[dict] = []
    temp_id_remap: dict[str, str] = {}
    for apu in apus:
        identity = _build_apu_identity_key(apu)
        if not identity:
            ordered.append(apu)
            continue
        previous = by_identity.get(identity)
        if previous is None:
            by_identity[identity] = apu
            ordered.append(apu)
            continue
        if _apu_compositions_conflict(previous, apu):
            if identity.startswith("funcional:"):
                ordered.append(apu)
                continue
            visible_code = _clean_text(apu.get("codigo")) or _clean_text(apu.get("descripcion")) or identity
            incidents.append(
                {
                    "severity": "blocking",
                    "code": "duplicate_apu_code_divergent",
                    "message": f"El APU {visible_code} aparece repetido con composición técnica diferente.",
                    "apu_code": visible_code,
                    "first_temp_id": previous.get("temp_id"),
                    "duplicate_temp_id": apu.get("temp_id"),
                }
            )
            continue
        duplicate_temp_id = _clean_text(apu.get("temp_id"))
        canonical_temp_id = _clean_text(previous.get("temp_id"))
        if duplicate_temp_id and canonical_temp_id and duplicate_temp_id != canonical_temp_id:
            temp_id_remap[duplicate_temp_id] = canonical_temp_id

    remapped_resources: list[dict] = []
    for resource in resources:
        current = dict(resource)
        current_apu_temp_id = _clean_text(current.get("apu_temp_id"))
        if current_apu_temp_id in temp_id_remap:
            current["apu_temp_id"] = temp_id_remap[current_apu_temp_id]
        remapped_resources.append(current)

    return ordered, remapped_resources, incidents


def _align_budget_rows_with_apus(budget_rows: list[dict], apus: list[dict]) -> tuple[list[dict], list[dict], list[dict]]:
    code_index: dict[str, dict] = {}
    functional_index: dict[str, list[dict]] = {}
    unit_price_index: dict[str, list[dict]] = {}
    for apu in apus:
        code = _clean_text(apu.get("codigo"))
        if code:
            code_index.setdefault(code, apu)
        unit_price = apu.get("unit_price") or apu.get("precio_unitario")
        functional_index.setdefault(_build_budget_apu_match_key(apu.get("descripcion"), apu.get("unidad"), unit_price), []).append(apu)
        unit_price_index.setdefault(_build_unit_price_match_key(apu.get("unidad"), unit_price), []).append(apu)

    aligned_rows: list[dict] = []
    pending_apus: list[dict] = []
    incidents: list[dict] = []
    for index, row in enumerate(budget_rows, start=1):
        current = dict(row)
        strategy = current.get("match_strategy") or "codigo"
        matched_apu = None
        if strategy == "descripcion_unidad_precio":
            key = _build_budget_apu_match_key(current.get("descripcion"), current.get("unidad"), current.get("precio_u") or current.get("precio_unitario"))
            candidates = functional_index.get(key) or []
            if len(candidates) == 1:
                matched_apu = candidates[0]
                current["match_kind"] = "descripcion_unidad_precio"
            elif len(candidates) > 1:
                matched_apu = candidates[0]
                current["match_kind"] = "descripcion_unidad_precio_primer_candidato"
                current["ignored_duplicate_apu_candidates"] = max(0, len(candidates) - 1)
            else:
                price_candidates = unit_price_index.get(_build_unit_price_match_key(current.get("unidad"), current.get("precio_u") or current.get("precio_unitario"))) or []
                if len(price_candidates) == 1:
                    matched_apu = price_candidates[0]
                    current["match_kind"] = "unidad_precio_unico"
                elif len(price_candidates) > 1:
                    manual_best, manual_candidates = _resolve_manual_budget_candidate(current, price_candidates)
                    if manual_best:
                        matched_apu = manual_best
                        current["match_kind"] = "busqueda_manual_descripcion_unidad_precio"
                    elif manual_candidates:
                        matched_apu = manual_candidates[0]
                        current["match_kind"] = "busqueda_manual_primer_candidato"
                        current["ignored_duplicate_apu_candidates"] = max(0, len(manual_candidates) - 1)
                    else:
                        best = _resolve_best_semantic_candidate(current, price_candidates)
                        if best:
                            matched_apu = best
                            current["match_kind"] = "unidad_precio_semantic_tiebreak"
        else:
            matched_apu = code_index.get(_clean_text(current.get("codigo")))

        if matched_apu:
            matched_apu["codigo"] = matched_apu.get("codigo") or current.get("codigo") or f"RUB-{index:04d}"
            matched_apu["budget_row_nro"] = current.get("nro")
            matched_apu["budget_cod_cpc_apu"] = current.get("cod_cpc_apu") or current.get("codigo_cpc_apu")
            current["matched_apu_temp_id"] = matched_apu.get("temp_id")
            aligned_rows.append(current)
            continue

        pending_apu = {
            "temp_id": f"auto_empty_apu_{index}",
            "source_signature": _build_budget_line_signature(current),
            "origin": "generated_empty_from_budget_unmatched",
            "status": "Pendiente sin recursos",
            "codigo": current.get("codigo") or current.get("nro") or f"RUB-{index:04d}",
            "descripcion": current.get("descripcion") or f"APU pendiente {index}",
            "unidad": current.get("unidad"),
            "chapter": current.get("capitulo") or "General",
            "quantity": current.get("cantidad"),
            "unit_price": current.get("precio_u") or current.get("precio_unitario"),
            "precio_unitario": current.get("precio_u") or current.get("precio_unitario"),
            "budget_total": f"{_coerce_decimal(current.get('precio_total')):.2f}",
            "cod_cpc_apu": current.get("cod_cpc_apu") or current.get("codigo_cpc_apu"),
            "resources": [],
            "resource_count": 0,
            "nested_apu_links": [],
            "nested_apu_count": 0,
        }
        pending_apus.append(pending_apu)
        current["matched_apu_temp_id"] = pending_apu["temp_id"]
        aligned_rows.append(current)

    return aligned_rows, pending_apus, incidents


def _resolve_best_semantic_candidate(row: dict, candidates: list[dict]) -> dict | None:
    scored = [
        (_semantic_overlap_score(row.get("descripcion"), candidate.get("descripcion")), candidate)
        for candidate in candidates
    ]
    scored = sorted(scored, key=lambda item: item[0], reverse=True)
    if not scored or scored[0][0] <= 0:
        return None
    if len(scored) > 1 and scored[0][0] == scored[1][0]:
        return None
    return scored[0][1]


def _resolve_manual_budget_candidate(row: dict, candidates: list[dict]) -> tuple[dict | None, list[dict]]:
    row_description = _normalize_signature_fragment(row.get("descripcion") or "")
    if not row_description:
        return None, []
    exact_candidates: list[dict] = []
    for candidate in candidates:
        candidate_description = _normalize_signature_fragment(candidate.get("descripcion") or "")
        if not candidate_description:
            continue
        if row_description == candidate_description or candidate_description in row_description or row_description in candidate_description:
            exact_candidates.append(candidate)
    if len(exact_candidates) == 1:
        return exact_candidates[0], exact_candidates
    return None, exact_candidates


def _coerce_existing_analysis(existing_analysis: dict | None) -> dict | None:
    if not isinstance(existing_analysis, dict):
        return None

    source_files = list(existing_analysis.get("source_files") or [])
    rubros = list(existing_analysis.get("rubros") or [])
    generated_apus = list(existing_analysis.get("generated_apus") or [])
    sample_apus = list(existing_analysis.get("sample_apus") or [])
    sample_resources = list(existing_analysis.get("sample_resources") or [])
    sample_vae_entries = list(existing_analysis.get("sample_vae_entries") or [])
    classification = dict(existing_analysis.get("classification") or {})
    summary = dict(existing_analysis.get("summary") or {})
    analysis_bundle = dict(existing_analysis.get("analysis_bundle") or {})
    bundle_summary = dict(analysis_bundle.get("summary") or {})

    normalized = {
        "source_filename": existing_analysis.get("source_filename") or "bundle_preexistente",
        "source_format": existing_analysis.get("source_format") or "bundle_preexistente",
        "budget_title": existing_analysis.get("budget_title") or "Bundle existente",
        "items_count": len(rubros),
        "chapters_count": len(set(str(row.get("capitulo") or "").strip() for row in rubros if str(row.get("capitulo") or "").strip())),
        "total_amount": existing_analysis.get("total_amount") or "0.00",
        "currency": existing_analysis.get("currency") or "USD",
        "detected_at": existing_analysis.get("detected_at") or datetime.now(timezone.utc).isoformat(),
        "rubros": rubros,
        "chapters": list(existing_analysis.get("chapters") or []),
        "chapter_breakdown": list(existing_analysis.get("chapter_breakdown") or _build_chapter_breakdown(rubros)),
        "dominant_chapter": existing_analysis.get("dominant_chapter"),
        "sample_rubros": list(existing_analysis.get("sample_rubros") or _build_sample_rubros(rubros)),
        "classification": classification,
        "generated_apus": generated_apus,
        "sample_apus": sample_apus or list(analysis_bundle.get("sample_apus") or []),
        "sample_resources": sample_resources or list(analysis_bundle.get("sample_resources") or []),
        "sample_vae_entries": sample_vae_entries or list(analysis_bundle.get("sample_vae_entries") or []),
        "source_files": source_files,
        "summary": {
            "has_budget": bool(summary.get("has_budget")) or bool(bundle_summary.get("has_budget")) or bool(rubros),
            "has_apus": bool(summary.get("has_apus")) or bool(bundle_summary.get("has_apus")) or bool(sample_apus) or bool(generated_apus),
            "has_resources": bool(summary.get("has_resources")) or bool(bundle_summary.get("has_resources")) or bool(sample_resources),
            "has_vae": bool(summary.get("has_vae")) or bool(bundle_summary.get("has_vae")) or bool(sample_vae_entries),
            "has_location_hint": bool(summary.get("has_location_hint")),
            "source_kind": summary.get("source_kind") or "archivo_base",
            "dominant_chapter": existing_analysis.get("dominant_chapter"),
            "document_kind": summary.get("document_kind") or classification.get("document_kind") or "bundle_parcial",
            "document_kind_counts": dict(summary.get("document_kind_counts") or {}),
            "coverage_level": summary.get("coverage_level") or bundle_summary.get("coverage_level"),
            "dedupe_summary": dict(summary.get("dedupe_summary") or bundle_summary.get("dedupe_summary") or {}),
            "warnings": [item for item in list(summary.get("warnings") or bundle_summary.get("warnings") or []) if item],
        },
    }
    return normalized


def _reconcile_budget_apu_alignment(
    budget_rows: list[dict],
    *,
    extracted_apus: list[dict],
    pending_apus: list[dict],
) -> dict:
    real_apu_keys: set[str] = set()
    pending_apu_keys: set[str] = set()
    real_apu_temp_ids = {str(apu.get("temp_id") or "") for apu in extracted_apus if apu.get("temp_id")}
    pending_apu_temp_ids = {str(apu.get("temp_id") or "") for apu in pending_apus if apu.get("temp_id")}

    for apu in extracted_apus:
        real_apu_keys.update(
            _build_alignment_keys(
                apu.get("codigo"),
                apu.get("descripcion"),
                apu.get("unidad"),
            )
        )
    for apu in pending_apus:
        pending_apu_keys.update(
            _build_alignment_keys(
                apu.get("codigo"),
                apu.get("descripcion"),
                apu.get("unidad"),
            )
        )

    matched_real = 0
    matched_pending = 0
    unmatched = 0
    unmatched_rows: list[dict] = []

    for row in budget_rows:
        matched_temp_id = str(row.get("matched_apu_temp_id") or "")
        if matched_temp_id and matched_temp_id in real_apu_temp_ids:
            matched_real += 1
            continue
        if matched_temp_id and matched_temp_id in pending_apu_temp_ids:
            matched_pending += 1
            continue
        budget_keys = _build_alignment_keys(
            row.get("codigo"),
            row.get("descripcion"),
            row.get("unidad"),
        )
        if budget_keys & real_apu_keys:
            matched_real += 1
            continue
        if budget_keys & pending_apu_keys:
            matched_pending += 1
            continue
        unmatched += 1
        if len(unmatched_rows) < 5:
            unmatched_rows.append(
                {
                    "codigo": row.get("codigo") or row.get("nro"),
                    "descripcion": row.get("descripcion"),
                    "unidad": row.get("unidad"),
                    "capitulo": row.get("capitulo") or "General",
                }
            )

    total_budget_rows = len(budget_rows)
    if total_budget_rows == 0:
        coverage_ratio = Decimal("0")
    else:
        coverage_ratio = (Decimal(matched_real + matched_pending) / Decimal(total_budget_rows)) * Decimal("100")

    return {
        "budget_rows_total": total_budget_rows,
        "matched_real_apus": matched_real,
        "matched_pending_apus": matched_pending,
        "unmatched_budget_rows": unmatched,
        "coverage_ratio": f"{coverage_ratio.quantize(Decimal('0.01'))}",
        "sample_unmatched_budget_rows": unmatched_rows,
    }


def _reconcile_pending_with_real_apus(
    pending_apus: list[dict],
    extracted_apus: list[dict],
    extracted_resources: list[dict],
) -> tuple[list[dict], list[dict], list[dict], dict]:
    pending_by_signature = {
        _build_apu_signature(apu): dict(apu)
        for apu in pending_apus
        if _build_apu_signature(apu)
    }
    temp_id_remap: dict[str, str] = {}
    matched_signatures: set[str] = set()
    reconciled_real_apus: list[dict] = []

    for apu in extracted_apus:
        current = dict(apu)
        signature = _build_apu_signature(current)
        bootstrap = pending_by_signature.get(signature)
        if bootstrap:
            old_temp_id = current.get("temp_id")
            current["temp_id"] = bootstrap.get("temp_id") or current.get("temp_id")
            current["source_signature"] = bootstrap.get("source_signature") or current.get("source_signature") or signature
            current["origin"] = "reconciled_from_budget_bootstrap"
            current["status"] = "Enriquecido"
            current["bootstrap_seed_preserved"] = True
            current["bootstrap_origin_temp_id"] = bootstrap.get("temp_id")
            current["bootstrap_budget_total"] = bootstrap.get("budget_total")
            current["bootstrap_quantity"] = bootstrap.get("quantity")
            current["bootstrap_unit_price"] = bootstrap.get("unit_price")
            current["bootstrap_resource_count"] = bootstrap.get("resource_count", 0)
            if old_temp_id and current.get("temp_id") and old_temp_id != current.get("temp_id"):
                temp_id_remap[old_temp_id] = current["temp_id"]
            matched_signatures.add(signature)
        reconciled_real_apus.append(current)

    reconciled_resources: list[dict] = []
    for resource in extracted_resources:
        current = dict(resource)
        current["apu_temp_id"] = temp_id_remap.get(current.get("apu_temp_id"), current.get("apu_temp_id"))
        reconciled_resources.append(current)

    remaining_pending = [
        dict(apu)
        for apu in pending_apus
        if _build_apu_signature(apu) not in matched_signatures
    ]

    reconciliation_summary = {
        "bootstrap_promoted_to_real": len(matched_signatures),
        "pending_apus_remaining": len(remaining_pending),
        "real_apus_retained": len(reconciled_real_apus),
        "temp_id_remaps": len(temp_id_remap),
    }

    return remaining_pending, reconciled_real_apus, reconciled_resources, reconciliation_summary


def _build_pending_apus_from_budget(rows: list[dict]) -> list[dict]:
    generated_apus: list[dict] = []
    seen_signatures: set[str] = set()
    for index, row in enumerate(rows, start=1):
        signature = _build_budget_line_signature(row)
        if signature in seen_signatures:
            continue
        seen_signatures.add(signature)
        generated_apus.append(
            {
                "temp_id": f"auto_apu_{index}",
                "source_signature": signature,
                "origin": "generated_from_budget_only",
                "status": "Pendiente",
                "codigo": _clean_text(row.get("codigo")) or _clean_text(row.get("nro")) or f"AUTO-{index}",
                "descripcion": _clean_text(row.get("descripcion")) or f"APU pendiente {index}",
                "unidad": _clean_text(row.get("unidad")),
                "chapter": _clean_text(row.get("capitulo")) or "General",
                "quantity": row.get("cantidad"),
                "unit_price": row.get("precio_u"),
                "budget_total": f"{_coerce_decimal(row.get('precio_total')):.2f}",
                "resources": [],
                "resource_count": 0,
                "nested_apu_links": [],
                "nested_apu_count": 0,
            }
        )
    return generated_apus


def _build_bundle_summary(
    source_files: list[dict],
    budget_rows: list[dict],
    *,
    has_budget: bool,
    has_apus: bool,
    has_resources: bool,
    has_vae: bool,
    pending_apus: list[dict] | None = None,
    warnings: list[str] | None = None,
) -> dict:
    pending_apus = list(pending_apus or [])
    warnings = list(warnings or [])
    if has_budget and has_apus and has_resources and has_vae:
        coverage_level = "presupuesto + apus + recursos + vae"
    elif has_budget and has_apus and has_resources:
        coverage_level = "presupuesto + apus + recursos"
    elif has_budget and (has_apus or pending_apus):
        coverage_level = "presupuesto + apus_base"
    elif has_budget:
        coverage_level = "solo_presupuesto"
    else:
        coverage_level = "sin_cobertura"

    return {
        "coverage_level": coverage_level,
        "has_budget": has_budget,
        "has_apus": has_apus or bool(pending_apus),
        "has_real_apus": has_apus,
        "has_pending_apus": bool(pending_apus),
        "has_resources": has_resources,
        "has_vae": has_vae,
        "budget_items_count": len(budget_rows),
        "apus_count": len(pending_apus) if pending_apus else 0,
        "real_apus_count": 0,
        "pending_apus_count": len(pending_apus),
        "resources_count": 0,
        "nested_apus_supported": True,
        "source_count": len(source_files),
        "warnings": warnings,
    }


def _build_analysis_bundle(
    analysis: dict,
    *,
    pending_apus: list[dict] | None = None,
    extracted_apus: list[dict] | None = None,
    extracted_resources: list[dict] | None = None,
    apu_links: list[dict] | None = None,
    vae_entries: list[dict] | None = None,
    apu_reconciliation: dict | None = None,
) -> dict:
    summary = dict(analysis.get("summary") or {})
    warnings = [item for item in list(summary.get("warnings") or []) if item]
    pending_apus = list(pending_apus or [])
    extracted_apus = list(extracted_apus or [])
    extracted_resources = list(extracted_resources or [])
    apu_links = list(apu_links or [])
    vae_entries = list(vae_entries or [])
    apu_reconciliation = dict(apu_reconciliation or {})
    combined_apus = extracted_apus or pending_apus
    budget_apu_alignment = _reconcile_budget_apu_alignment(
        list(analysis.get("rubros") or []),
        extracted_apus=extracted_apus,
        pending_apus=pending_apus,
    )
    bundle_summary = _build_bundle_summary(
        list(analysis.get("source_files") or []),
        list(analysis.get("rubros") or []),
        has_budget=bool(summary.get("has_budget")),
        has_apus=bool(summary.get("has_apus")) or bool(extracted_apus),
        has_resources=bool(summary.get("has_resources")) or bool(extracted_resources),
        has_vae=bool(summary.get("has_vae")),
        pending_apus=combined_apus,
        warnings=warnings,
    )
    bundle_summary["real_apus_count"] = len(extracted_apus)
    bundle_summary["pending_apus_count"] = len(pending_apus)
    bundle_summary["apus_count"] = len(combined_apus)
    bundle_summary["resources_count"] = len(extracted_resources)
    bundle_summary["nested_apu_link_count"] = len(apu_links)
    bundle_summary["pending_nested_apu_link_count"] = sum(1 for link in apu_links if link.get("status") == "pending_reference")
    bundle_summary["ambiguous_nested_apu_link_count"] = sum(1 for link in apu_links if link.get("status") == "ambiguous_reference")
    bundle_summary["semantic_match_count"] = sum(1 for link in apu_links if link.get("match_kind") == "semantic_descripcion")
    bundle_summary["vae_entries_count"] = len(vae_entries)
    bundle_summary["vae_detail_count"] = sum(1 for entry in vae_entries if entry.get("kind") != "section_header")
    bundle_summary["vae_percentage_count"] = sum(1 for entry in vae_entries if entry.get("kind") == "percentage")
    bundle_summary["dedupe_summary"] = dict(summary.get("dedupe_summary") or {})
    bundle_summary["budget_apu_alignment"] = budget_apu_alignment
    bundle_summary["apu_reconciliation"] = apu_reconciliation
    auto_warnings = list(warnings)
    if bundle_summary["ambiguous_nested_apu_link_count"] > 0:
        auto_warnings.append(
            f"Se detectaron {bundle_summary['ambiguous_nested_apu_link_count']} referencias APU ambiguas; conviene validarlas manualmente antes de confiar en la estructura técnica."
        )
    if bundle_summary["pending_nested_apu_link_count"] > 0:
        auto_warnings.append(
            f"Quedan {bundle_summary['pending_nested_apu_link_count']} referencias APU pendientes de enlace dentro del bundle."
        )
    if budget_apu_alignment.get("unmatched_budget_rows"):
        auto_warnings.append(
            f"Persisten {budget_apu_alignment['unmatched_budget_rows']} rubros sin correspondencia técnica visible entre presupuesto y APUs."
        )
    if bundle_summary["has_budget"] and not bundle_summary["has_resources"] and bundle_summary["has_apus"]:
        auto_warnings.append("El bundle ya tiene APUs, pero todavía no consolida recursos suficientes para una lectura técnica completa.")
    unique_warnings: list[str] = []
    seen_warnings: set[str] = set()
    for warning in auto_warnings:
        normalized_warning = _clean_text(warning)
        if not normalized_warning or normalized_warning in seen_warnings:
            continue
        seen_warnings.add(normalized_warning)
        unique_warnings.append(normalized_warning)
    bundle_summary["warnings"] = unique_warnings
    bundle_summary["blocking_incidents"] = list(summary.get("blocking_incidents") or bundle_summary.get("blocking_incidents") or [])
    bundle_summary["non_blocking_incidents"] = list(summary.get("non_blocking_incidents") or bundle_summary.get("non_blocking_incidents") or [])
    bundle_summary["requires_superadmin_consent"] = bool(summary.get("requires_superadmin_consent")) or bool(bundle_summary["blocking_incidents"])
    bundle_summary["parser_profile"] = summary.get("parser_profile") or bundle_summary.get("parser_profile")
    return {
        "document_sources": list(analysis.get("source_files") or []),
        "budget_items": list(analysis.get("rubros") or []),
        "apus": combined_apus,
        "resources": extracted_resources,
        "apu_links": apu_links,
        "vae_entries": vae_entries,
        "portal_metadata": {
            "budget_title": analysis.get("budget_title"),
            "currency": analysis.get("currency") or "USD",
            "dominant_chapter": analysis.get("dominant_chapter"),
        },
        "origin_metadata": {
            "source_filename": analysis.get("source_filename"),
            "source_format": analysis.get("source_format"),
            "classification": analysis.get("classification"),
        },
        "ownership_metadata": {},
        "warnings": unique_warnings,
        "confidence_map": {
            "document_kind": (analysis.get("classification") or {}).get("confidence", "media"),
        },
        "sample_apus": combined_apus[:5],
        "sample_resources": extracted_resources[:8],
        "sample_vae_entries": vae_entries[:5],
        "sample_unmatched_budget_rows": budget_apu_alignment.get("sample_unmatched_budget_rows") or [],
        "sample_reconciled_apus": [apu for apu in combined_apus if apu.get("bootstrap_seed_preserved")][:5],
        "sample_pending_nested_apu_links": [link for link in apu_links if link.get("status") == "pending_reference"][:5],
        "sample_ambiguous_nested_apu_links": [link for link in apu_links if link.get("status") == "ambiguous_reference"][:5],
        "summary": bundle_summary,
    }


def _build_chapter_breakdown(rows: list[dict]) -> list[dict]:
    grouped: dict[str, dict] = {}
    for row in rows:
        chapter_name = _clean_text(row.get("capitulo")) or "General"
        current = grouped.setdefault(
            chapter_name,
            {
                "name": chapter_name,
                "items_count": 0,
                "total_amount": Decimal("0.00"),
            },
        )
        current["items_count"] += 1
        current["total_amount"] += Decimal(str(row.get("precio_total") or 0))

    return [
        {
            "name": item["name"],
            "items_count": item["items_count"],
            "total_amount": f"{item['total_amount']:.2f}",
        }
        for item in sorted(grouped.values(), key=lambda value: (-value["total_amount"], value["name"]))
    ]


def _build_sample_rubros(rows: list[dict], limit: int = 5) -> list[dict]:
    sample = []
    for row in rows[:limit]:
        sample.append(
            {
                "nro": row.get("nro"),
                "descripcion": row.get("descripcion"),
                "capitulo": row.get("capitulo") or "General",
                "unidad": row.get("unidad"),
                "precio_total": f"{Decimal(str(row.get('precio_total') or 0)):.2f}",
            }
        )
    return sample


def _parse_budget_row(row):
    normalized = [_clean_text(cell) if cell else "" for cell in row]
    nonempty = [cell for cell in normalized if cell]
    if len(nonempty) < 2:
        return None

    chapter_name = ""
    chapter_total = None
    is_chapter = False
    if len(normalized) >= 7 and normalized[0] and not normalized[1] and not _a_num(normalized[4]) and not _a_num(normalized[5]) and _a_num(normalized[6]):
        chapter_name = normalized[2]
        chapter_total = _a_num(normalized[6])
        is_chapter = bool(chapter_total and chapter_name and not _a_num(chapter_name))
    elif len(nonempty) == 2 and not re.match(r"^\d", nonempty[0]):
        chapter_name = nonempty[0]
        chapter_total = _a_num(nonempty[1])
        is_chapter = bool(chapter_total and chapter_name and not _a_num(chapter_name))
    elif not normalized[0] and len(normalized) > 2 and normalized[2]:
        chapter_name = normalized[2]
        chapter_total = _a_num(normalized[-1])
        is_chapter = bool(chapter_total and chapter_name and not _a_num(chapter_name))

    if is_chapter:
        return {"tipo": "capitulo", "descripcion": chapter_name.strip(), "total": chapter_total}

    item_number = normalized[0] if normalized else ""
    if not re.match(r"^\d{1,4}$", item_number.split(".")[0] if item_number else ""):
        return None

    if len(normalized) >= 7:
        code = normalized[1]
        description = normalized[2]
        unit = normalized[3]
        quantity = _a_num(normalized[4])
        unit_price = _a_num(normalized[5])
        total_price = _a_num(normalized[6])
    else:
        return None

    if total_price is None:
        return None

    return {
        "tipo": "rubro",
        "nro": item_number,
        "codigo": code,
        "descripcion": description,
        "unidad": unit,
        "cantidad": quantity,
        "precio_u": unit_price,
        "precio_total": total_price,
        "capitulo": "",
    }


def _parse_budget_pdf(pdf_bytes: bytes) -> tuple[str, list[dict]]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise HTTPException(
            status_code=400,
            detail="La lectura PDF del portal no está disponible en este entorno. Usa XLSX o habilita el parser PDF en backend.",
        ) from exc

    rows: list[dict] = []
    title = ""
    current_chapter = ""

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        first_text = pdf.pages[0].extract_text() if pdf.pages else ""
        if "RESUMEN DESAGREGACION TECNOLOGICA" in _fold_latin_text(first_text or "").upper():
            return _parse_soce_vae_budget_pdf(pdf_bytes)
        for page_index, page in enumerate(pdf.pages):
            page_text_preview = page.extract_text() or ""
            if rows and "ANALISIS DE PRECIOS UNITARIOS" in _fold_latin_text(page_text_preview).upper():
                break
            tables = page.extract_tables() or []
            if tables:
                for table in tables:
                    for row_index, row in enumerate(table):
                        if not row:
                            continue
                        normalized = [str(cell).strip() if cell else "" for cell in row]
                        if page_index == 0 and row_index < 3 and not title:
                            joined = " ".join(cell for cell in normalized if cell)
                            if joined and not any(token in joined.upper() for token in ["NRO", "CODIGO", "RUBRO", "UNIDAD", "PRECIO"]):
                                title = joined
                        parsed = _parse_budget_row(normalized)
                        if parsed is None:
                            continue
                        if parsed["tipo"] == "capitulo":
                            current_chapter = parsed["descripcion"]
                        else:
                            parsed["capitulo"] = current_chapter
                            rows.append(parsed)
                continue

            text = page_text_preview
            for line in text.splitlines():
                cleaned = _clean_text(line)
                match = re.match(
                    r"^(\d{1,4})\s+([A-Z0-9]{3,15})\s+(.+?)\s+(m2|m3|m|kg|u|gl|kg|pto|glb|m3-km|m3/km)\s+([\d,]+\.?\d*)\s+([\d,]+\.?\d+)\s+([\d,]+\.?\d+)$",
                    cleaned,
                    re.IGNORECASE,
                )
                if match:
                    rows.append({
                        "tipo": "rubro",
                        "nro": match.group(1),
                        "codigo": match.group(2),
                        "descripcion": match.group(3).strip(),
                        "unidad": match.group(4),
                        "cantidad": _a_num(match.group(5)),
                        "precio_u": _a_num(match.group(6)),
                        "precio_total": _a_num(match.group(7)),
                        "capitulo": current_chapter,
                    })
    if not rows:
        vae_title, vae_rows = _parse_soce_vae_budget_pdf(pdf_bytes)
        if vae_rows:
            return vae_title, vae_rows

    return title, rows


def _parse_soce_vae_budget_pdf(pdf_bytes: bytes) -> tuple[str, list[dict]]:
    """Parsea el resumen VAE SOCE/SERCOP donde el codigo visible es CPC del rubro."""
    try:
        import pdfplumber
    except ImportError:
        return "", []

    lines: list[str] = []
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for raw_line in text.splitlines():
                cleaned = _clean_text(raw_line)
                if cleaned:
                    lines.append(cleaned)

    return _parse_soce_vae_budget_lines(lines)


def _parse_soce_vae_budget_lines(lines: list[str]) -> tuple[str, list[dict]]:
    rows: list[dict] = []
    title = "Resumen desagregacion tecnologica"
    row_pattern = re.compile(
        rf"^(\d{{1,4}})\s+([0-9]{{6,13}})(.+?)\s+([\d.,]+)\s+({SOCE_BUDGET_UNIT_PATTERN})\s+([\d.,]+)\s+(.+)$",
        re.IGNORECASE,
    )

    in_summary = False
    for raw_line in lines:
        line = _clean_text(raw_line)
        upper_line = _fold_latin_text(line).upper()
        if "ANALISIS DE PRECIOS UNITARIOS" in upper_line:
            break
        if "RESUMEN DESAGREGACION TECNOLOGICA" in upper_line:
            in_summary = True
            continue
        if not in_summary and not rows:
            continue
        match = row_pattern.match(line)
        if not match:
            continue
        tail_numbers = re.findall(r"\d+(?:[.,]\d+)?", match.group(7))
        quantity_value = _a_num(match.group(4))
        unit_price_value = _a_num(match.group(6))
        total_price = _a_num(tail_numbers[0]) if tail_numbers else None
        expected_total = (quantity_value or 0) * (unit_price_value or 0)
        if expected_total > 0 and (total_price is None or total_price <= 0 or abs(total_price - expected_total) > max(1.0, expected_total * 0.05)):
            total_price = expected_total
        relative_weight = _a_num(tail_numbers[1]) if len(tail_numbers) > 1 else None
        vae_percentage = _a_num(tail_numbers[-2]) if len(tail_numbers) >= 2 else None
        weighted_vae = _a_num(tail_numbers[-1]) if tail_numbers else None
        description = _clean_soce_vae_budget_description(match.group(3), match.group(2))
        row_number = match.group(1)
        rows.append(
            {
                "tipo": "rubro",
                "nro": row_number,
                "codigo": f"RUB-{int(row_number):04d}",
                "cod_cpc_apu": match.group(2),
                "codigo_cpc_apu": match.group(2),
                "descripcion": description or f"Rubro SOCE/SERCOP {row_number}",
                "unidad": match.group(5),
                "cantidad": quantity_value,
                "precio_u": unit_price_value,
                "precio_unitario": unit_price_value,
                "precio_total": total_price,
                "peso_relativo": relative_weight,
                "vae_percentage": vae_percentage,
                "vae_ponderado": weighted_vae,
                "capitulo": "Resumen VAE",
                "match_strategy": "descripcion_unidad_precio",
            }
        )

    return title, rows


def _parse_budget_lines(lines: list[str]) -> tuple[str, list[dict]]:
    joined_preview = "\n".join(lines[:80])
    if "RESUMEN DESAGREGACION TECNOLOGICA" in _fold_latin_text(joined_preview).upper():
        return _parse_soce_vae_budget_lines(lines)

    rows: list[dict] = []
    current_chapter = ""
    for raw_line in lines:
        cleaned = _clean_text(raw_line)
        upper_line = _fold_latin_text(cleaned).upper()
        if rows and "ANALISIS DE PRECIOS UNITARIOS" in upper_line:
            break
        match = re.match(
            r"^(\d{1,4})\s+([A-Z0-9]{3,15})\s+(.+?)\s+(m2|m3|m|kg|u|gl|kg|pto|glb|m3-km|m3/km)\s+([\d,]+\.?\d*)\s+([\d,]+\.?\d+)\s+([\d,]+\.?\d+)$",
            cleaned,
            re.IGNORECASE,
        )
        if match:
            rows.append({
                "tipo": "rubro",
                "nro": match.group(1),
                "codigo": match.group(2),
                "descripcion": match.group(3).strip(),
                "unidad": match.group(4),
                "cantidad": _a_num(match.group(5)),
                "precio_u": _a_num(match.group(6)),
                "precio_total": _a_num(match.group(7)),
                "capitulo": current_chapter,
            })
    return "", rows


def _clean_soce_vae_budget_description(value: str, cpc_code: str) -> str:
    text = _clean_text(value)
    if not text:
        return ""
    text = re.sub(r"\b[A-ZÁÉÍÓÚÑ](?:\s+[a-záéíóúñ]){2,}\b", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    upper = _fold_latin_text(text).upper()
    # En este perfil suele aparecer primero la descripcion CPC en mayusculas y luego la descripcion del rubro.
    words = text.split()
    best_start = 0
    for idx, word in enumerate(words):
        if any(char.islower() for char in word) or any(char in word for char in "áéíóúñ"):
            best_start = max(0, idx - 1)
            break
    cleaned = " ".join(words[best_start:]).strip()
    if len(cleaned) < 8 and len(words) > 2:
        cleaned = " ".join(words[-8:])
    return cleaned or text.replace(cpc_code, "").strip()


def _parse_budget_xlsx(xlsx_bytes: bytes) -> tuple[str, list[dict]]:
    workbook = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    sheet = None
    for name in workbook.sheetnames:
        normalized = name.lower().replace(" ", "").replace("_", "")
        if normalized in {"formulario1", "presupuesto", "rubros", "tabla"}:
            sheet = workbook[name]
            break
    if sheet is None:
        sheet = workbook[workbook.sheetnames[0]]

    title = ""
    rows: list[dict] = []
    current_chapter = ""
    header_row = 6
    for row_idx in range(1, min(sheet.max_row + 1, 15)):
        values = [str(sheet.cell(row=row_idx, column=column).value or "").upper() for column in range(1, 8)]
        if any(token in value for value in values for token in ("RUBRO", "DESCRIPCION", "DESCRIPCIÓN", "CANTIDAD")):
            header_row = row_idx
            break

    for row_idx in range(header_row + 1, sheet.max_row + 1):
        item_number = sheet.cell(row=row_idx, column=1).value
        description = sheet.cell(row=row_idx, column=4).value
        quantity = sheet.cell(row=row_idx, column=5).value
        unit = sheet.cell(row=row_idx, column=6).value
        unit_price = sheet.cell(row=row_idx, column=7).value
        total_price = sheet.cell(row=row_idx, column=8).value

        if item_number is None and description is None:
            continue

        description_text = _clean_text(description) if description else ""
        if description_text and (quantity is None or quantity == "") and (unit_price is None or unit_price == ""):
            current_chapter = description_text
            continue

        if total_price is None:
            continue

        try:
            total_value = float(total_price)
        except (TypeError, ValueError):
            continue

        quantity_value = None
        if quantity is not None:
            try:
                quantity_value = float(quantity)
            except (TypeError, ValueError):
                quantity_value = _a_num(quantity)

        price_value = None
        if unit_price is not None:
            try:
                price_value = float(unit_price)
            except (TypeError, ValueError):
                price_value = _a_num(unit_price)

        rows.append({
            "tipo": "rubro",
            "nro": str(item_number).strip() if item_number is not None else "",
            "codigo": "",
            "descripcion": description_text,
            "unidad": _clean_text(unit) if unit else "",
            "cantidad": quantity_value,
            "precio_u": price_value,
            "precio_total": total_value,
            "capitulo": current_chapter,
        })

    return title, rows


def _parse_budget_xls(xls_bytes: bytes) -> tuple[str, list[dict]]:
    try:
        import xlrd
    except ImportError as exc:
        raise HTTPException(
            status_code=400,
            detail="La lectura XLS del portal no está disponible en este entorno. Instala xlrd en backend para habilitarla.",
        ) from exc

    workbook = xlrd.open_workbook(file_contents=xls_bytes)
    sheet = None
    for name in workbook.sheet_names():
        normalized = name.lower().replace(" ", "").replace("_", "")
        if normalized in {"formulario1", "presupuesto", "rubros", "tabla"}:
            sheet = workbook.sheet_by_name(name)
            break
    if sheet is None:
        sheet = workbook.sheet_by_index(0)

    title = ""
    rows: list[dict] = []
    current_chapter = ""
    header_row = 5
    for row_idx in range(0, min(sheet.nrows, 14)):
        values = [str(sheet.cell_value(row_idx, column) or "").upper() for column in range(0, min(sheet.ncols, 8))]
        if any(token in value for value in values for token in ("RUBRO", "DESCRIPCION", "DESCRIPCIÓN", "CANTIDAD")):
            header_row = row_idx
            break

    for row_idx in range(header_row + 1, sheet.nrows):
        item_number = sheet.cell_value(row_idx, 0) if sheet.ncols > 0 else None
        description = sheet.cell_value(row_idx, 3) if sheet.ncols > 3 else None
        quantity = sheet.cell_value(row_idx, 4) if sheet.ncols > 4 else None
        unit = sheet.cell_value(row_idx, 5) if sheet.ncols > 5 else None
        unit_price = sheet.cell_value(row_idx, 6) if sheet.ncols > 6 else None
        total_price = sheet.cell_value(row_idx, 7) if sheet.ncols > 7 else None

        if item_number in (None, "") and description in (None, ""):
            continue

        description_text = _clean_text(description) if description else ""
        if description_text and quantity in (None, "") and unit_price in (None, ""):
            current_chapter = description_text
            continue

        if total_price in (None, ""):
            continue

        total_value = _a_num(total_price)
        if total_value is None:
            continue

        rows.append({
            "tipo": "rubro",
            "nro": str(item_number).strip() if item_number is not None else "",
            "codigo": "",
            "descripcion": description_text,
            "unidad": _clean_text(unit) if unit else "",
            "cantidad": _a_num(quantity),
            "precio_u": _a_num(unit_price),
            "precio_total": total_value,
            "capitulo": current_chapter,
        })

    return title, rows


def _extract_pdf_signal_lines(content: bytes, max_pages: int | None = 16) -> list[str]:
    try:
        import pdfplumber
    except ImportError:
        return []

    fragments: list[str] = []
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            pages = pdf.pages if max_pages is None else pdf.pages[:max_pages]
            for page in pages:
                text = page.extract_text() or ""
                for line in text.splitlines():
                    cleaned = _clean_text(line)
                    if cleaned:
                        fragments.append(cleaned)
    except Exception:
        return []
    return fragments


def _extract_workbook_signal_lines(content: bytes, extension: str) -> list[str]:
    fragments: list[str] = []
    if extension in {".xlsx", ".xlsm"}:
        try:
            workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            for sheet_name in workbook.sheetnames[:4]:
                fragments.append(sheet_name)
                sheet = workbook[sheet_name]
                for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 12), values_only=True):
                    cells = [_clean_text(cell) for cell in row if _clean_text(cell)]
                    if cells:
                        fragments.append(" ".join(cells[:8]))
        except Exception:
            return []
    elif extension == ".xls":
        try:
            import xlrd

            workbook = xlrd.open_workbook(file_contents=content)
            for sheet_name in workbook.sheet_names()[:4]:
                fragments.append(sheet_name)
                sheet = workbook.sheet_by_name(sheet_name)
                for row_idx in range(0, min(sheet.nrows, 12)):
                    row_values = []
                    for col_idx in range(0, min(sheet.ncols, 8)):
                        value = _clean_text(sheet.cell_value(row_idx, col_idx))
                        if value:
                            row_values.append(value)
                    if row_values:
                        fragments.append(" ".join(row_values))
        except Exception:
            raw_text = content.decode("latin1", errors="ignore")
            for match in re.findall(r"[A-Za-zÁÉÍÓÚáéíóúÑñ0-9%./:_ -]{8,}", raw_text):
                cleaned = _clean_text(match)
                if cleaned:
                    fragments.append(cleaned)
    return fragments


def _extract_signal_lines(content: bytes, extension: str) -> list[str]:
    if extension == ".pdf":
        return _extract_pdf_signal_lines(content, max_pages=None)
    return _extract_workbook_signal_lines(content, extension)


def _extract_signal_text(content: bytes, extension: str) -> str:
    return " ".join(_extract_signal_lines(content, extension))


def _normalize_resource_type(section_name: str) -> str:
    section = _clean_text(section_name).lower()
    if "mano" in section:
        return "mano_obra"
    if "material" in section:
        return "material"
    if "equipo" in section:
        return "equipo"
    if "transporte" in section:
        return "transporte"
    return "recurso"


def _parse_detected_resource_line(line: str, current_section: str, current_apu: dict, resource_index: int) -> dict | None:
    upper = line.upper()
    if any(token in upper for token in ["CODIGO DESCRIPCION", "CÓDIGO DESCRIPCIÓN", "NUMERO S.R.H.", "TARIFA/U", "PESO RELATIVO", "DESCRIPCIÓN CANTIDAD", "DESCRIPCION CANTIDAD"]):
        return None
    if upper.startswith("SUBTOTAL ") or upper.startswith("TOTAL ") or upper.startswith("INDIRECTOS") or upper.startswith("OTROS INDIRECTOS"):
        return None

    resource_type = _normalize_resource_type(current_section)
    base_payload = {
        "temp_id": f"detected_resource_{resource_index}",
        "apu_temp_id": current_apu["temp_id"],
        "resource_type": resource_type,
        "section_label": current_section,
    }

    vae_tail = re.search(r"\s([0-9]{6,13})\s+(NP|EP|ND)\s+([\d.,]+%?)\s+([\d.,]+%?)$", line, re.IGNORECASE)
    if vae_tail:
        left = line[:vae_tail.start()].strip()
        number_matches = list(re.finditer(r"\d+(?:[.,]\d+)?%?", left))
        numbers = [match.group(0) for match in number_matches]
        description = left[: number_matches[0].start()].strip() if number_matches else left
        unit = DEFAULT_RESOURCE_UNIT_BY_TYPE.get(resource_type, "u")
        unit_match = re.search(rf"\s({SOCE_BUDGET_UNIT_PATTERN}|Hora)\s+\d", left, re.IGNORECASE)
        if unit_match:
            unit = unit_match.group(1)
            description = left[:unit_match.start()].strip()
        payload = {
            **base_payload,
            "codigo": f"{resource_type[:3].upper()}-{resource_index:05d}",
            "descripcion": description,
            "unidad": unit,
            "cantidad": _a_num(numbers[0]) if len(numbers) >= 1 else None,
            "precio": _a_num(numbers[1]) if len(numbers) >= 2 else None,
            "rendimiento": _a_num(numbers[-3]) if len(numbers) >= 3 else None,
            "subtotal": _a_num(str(numbers[-2]).replace("%", "")) if len(numbers) >= 2 else None,
            "weight_percentage": _a_num(str(numbers[-1]).replace("%", "")) if numbers else None,
            "cod_cpc_codigo": vae_tail.group(1),
            "np_ep_nd": vae_tail.group(2).upper(),
            "vae_percentage": _a_num(str(vae_tail.group(3)).replace("%", "")),
            "vae_ponderado": _a_num(str(vae_tail.group(4)).replace("%", "")),
        }
        return _normalize_detected_resource(payload)

    code_match = re.match(r"^([A-Z0-9]{3,20})\s+(.+)$", line)
    if not code_match:
        return None

    code = code_match.group(1)
    if not _is_valid_resource_code(code):
        return None
    rest = code_match.group(2).strip()
    numeric_tail = re.findall(r"[-]?\d+(?:[.,]\d+)?%?", rest)
    unit = DEFAULT_RESOURCE_UNIT_BY_TYPE.get(resource_type, "u")
    quantity = price = rendimiento = subtotal = None
    description = rest
    if len(numeric_tail) >= 3:
        last_numeric = numeric_tail[-5:] if len(numeric_tail) >= 5 else (numeric_tail[-4:] if len(numeric_tail) >= 4 else numeric_tail[-3:])
        quantity = _a_num(last_numeric[0])
        price = _a_num(last_numeric[1])
        if len(last_numeric) >= 5:
            rendimiento = _a_num(last_numeric[2])
            subtotal = _a_num(str(last_numeric[3]).replace("%", ""))
        elif len(last_numeric) == 4:
            rendimiento = _a_num(last_numeric[2])
            subtotal = _a_num(str(last_numeric[3]).replace("%", ""))
        else:
            subtotal = _a_num(str(last_numeric[2]).replace("%", ""))
        split_match = re.search(rf"\s({SOCE_BUDGET_UNIT_PATTERN}|Hora)\s+{re.escape(str(last_numeric[0]))}", rest, re.IGNORECASE)
        if split_match:
            unit = split_match.group(1)
            description = rest[:split_match.start()].strip()
        else:
            first_tail = str(last_numeric[0])
            tail_index = rest.rfind(first_tail)
            if tail_index > 0:
                description = rest[:tail_index].strip()

    payload = {
        **base_payload,
        "codigo": code,
        "descripcion": description,
        "unidad": unit,
        "cantidad": quantity,
        "precio": price,
        "rendimiento": rendimiento,
        "subtotal": subtotal,
        "is_nested_apu": bool(re.search(r"(?<![A-Z0-9])APU(?:S)?(?![A-Z0-9])", description.upper())),
        "_parse_quality": "complete" if len(numeric_tail) >= 3 else "incomplete",
    }
    return _normalize_detected_resource(payload)


def _extract_apu_resource_pairs(lines: list[str]) -> tuple[list[dict], list[dict]]:
    apus: list[dict] = []
    resources: list[dict] = []
    current_apu: dict | None = None
    current_section = ""

    for raw_line in lines:
        line = _clean_text(raw_line)
        if not line:
            continue

        upper = line.upper()
        if "ANALISIS DE PRECIOS UNITARIOS" in upper or "ANÁLISIS DE PRECIOS UNITARIOS" in upper:
            if current_apu:
                apus.append(current_apu)
            current_apu = {
                "temp_id": f"detected_apu_{len(apus) + 1}",
                "origin": "detected_from_source",
                "status": "Detectado",
                "codigo": "",
                "descripcion": "",
                "unidad": "",
                "chapter": "",
                "resources": [],
                "resource_count": 0,
                "nested_apu_links": [],
                "nested_apu_count": 0,
            }
            current_section = ""
            continue

        if current_apu is None:
            continue

        if upper.startswith("CODIGO:") or upper.startswith("CÓDIGO:"):
            current_apu["codigo"] = line.split(":", 1)[-1].strip()
            continue
        if upper.startswith("DESCRIP.:") or upper.startswith("DESCRIPCION:") or upper.startswith("DESCRIPCIÓN:"):
            current_apu["descripcion"] = line.split(":", 1)[-1].strip()
            continue
        if upper.startswith("RUBRO:"):
            continue
        if " UNIDAD:" in upper and not current_apu.get("descripcion"):
            desc, unit = re.split(r"\s+UNIDAD:\s*", line, maxsplit=1, flags=re.IGNORECASE)
            current_apu["descripcion"] = desc.strip()
            current_apu["unidad"] = unit.strip()
            continue
        if upper.startswith("UNIDAD:"):
            current_apu["unidad"] = line.split(":", 1)[-1].strip()
            continue
        if upper.startswith("COSTO TOTAL DEL RUBRO") or upper.startswith("PRECIO UNITARIO TOTAL") or upper.startswith("VALOR OFERTADO"):
            value = _a_num(line.rsplit(" ", 1)[-1])
            if value is not None:
                current_apu["unit_price"] = value
                current_apu["precio_unitario"] = value
            continue

        if (
            upper in {"EQUIPO Y HERRAMIENTA", "EQUIPOS", "MATERIALES", "TRANSPORTE", "MANO DE OBRA"}
            or upper.startswith("EQUIPOS ")
            or upper.startswith("MATERIALES ")
            or upper.startswith("TRANSPORTE ")
            or upper.startswith("MANO DE OBRA ")
        ):
            current_section = upper.split(" PESO ", 1)[0].title()
            continue

        if upper.startswith("SUBTOTAL DE "):
            current_section = ""
            continue

        if current_section:
            resource = _parse_detected_resource_line(line, current_section, current_apu, len(resources) + 1)
            if resource:
                resources.append(resource)
                current_apu["resources"].append(resource)
                current_apu["resource_count"] += 1
                if resource.get("is_nested_apu"):
                    current_apu["nested_apu_links"].append(resource["codigo"])
                    current_apu["nested_apu_count"] += 1

    if current_apu:
        apus.append(current_apu)

    apu_codes = {str(apu.get("codigo") or "").strip(): apu["temp_id"] for apu in apus if str(apu.get("codigo") or "").strip()}
    apu_signatures: dict[str, list[str]] = {}
    for apu in apus:
        raw_signature = _normalize_signature_fragment(str(apu.get("descripcion") or ""))
        semantic_signature = _build_semantic_phrase_key(apu.get("descripcion"), apu.get("chapter"), apu.get("capitulo"))
        for candidate in (raw_signature, semantic_signature):
            if candidate:
                apu_signatures.setdefault(candidate, []).append(apu["temp_id"])
    for resource in resources:
        code = str(resource.get("codigo") or "").strip()
        desc_signature = _normalize_signature_fragment(str(resource.get("descripcion") or ""))
        semantic_desc_signature = _build_semantic_phrase_key(resource.get("descripcion"))
        matched_apu = apu_codes.get(code)
        if matched_apu is None and desc_signature and len(apu_signatures.get(desc_signature) or []) == 1:
            matched_apu = apu_signatures[desc_signature][0]
        if matched_apu is None and semantic_desc_signature and len(apu_signatures.get(semantic_desc_signature) or []) == 1:
            matched_apu = apu_signatures[semantic_desc_signature][0]
        if matched_apu:
            resource["is_nested_apu"] = True
            resource["nested_apu_target"] = matched_apu

    resources_by_apu: dict[str, list[dict]] = {}
    for resource in resources:
        resources_by_apu.setdefault(resource["apu_temp_id"], []).append(resource)

    for apu in apus:
        apu_resources = resources_by_apu.get(apu["temp_id"], [])
        apu["nested_apu_links"] = [
            item["nested_apu_target"]
            for item in apu_resources
            if item.get("is_nested_apu") and item.get("nested_apu_target")
        ]
        apu["nested_apu_count"] = len(apu["nested_apu_links"])
        apu["resources"] = apu_resources[:8]
        apu["resource_count"] = len(apu_resources)
    return apus, resources


def _resolve_nested_apu_links(apus: list[dict], resources: list[dict]) -> tuple[list[dict], list[dict], list[dict]]:
    apu_codes = {str(apu.get("codigo") or "").strip(): apu for apu in apus if str(apu.get("codigo") or "").strip()}
    apu_signatures: dict[str, list[dict]] = {}
    for apu in apus:
        raw_signature = _normalize_signature_fragment(str(apu.get("descripcion") or ""))
        semantic_signature = _build_semantic_phrase_key(apu.get("descripcion"), apu.get("chapter"), apu.get("capitulo"))
        for candidate in (raw_signature, semantic_signature):
            if candidate:
                apu_signatures.setdefault(candidate, []).append(apu)
    links: list[dict] = []
    for resource in resources:
        target_apu = None
        match_kind = "descripcion"
        code = str(resource.get("codigo") or "").strip()
        if not resource.get("is_nested_apu"):
            continue
        desc_signature = _normalize_signature_fragment(str(resource.get("descripcion") or ""))
        semantic_desc_signature = _build_semantic_phrase_key(resource.get("descripcion"))
        if code:
            target_apu = apu_codes.get(code)
            if target_apu is not None:
                match_kind = "codigo"
        if target_apu is None and desc_signature:
            literal_candidates = apu_signatures.get(desc_signature) or []
            if len(literal_candidates) == 1:
                target_apu = literal_candidates[0]
                match_kind = "descripcion"
            elif len(literal_candidates) > 1:
                resource["nested_apu_status"] = "ambiguous_reference"
                resource["nested_apu_candidates"] = [candidate.get("temp_id") for candidate in literal_candidates[:5]]
                links.append(
                    {
                        "parent_apu_temp_id": resource.get("apu_temp_id"),
                        "child_apu_temp_id": None,
                        "child_reference_code": code or None,
                        "child_reference_label": resource.get("descripcion"),
                        "match_kind": "ambiguous_descripcion",
                        "status": "ambiguous_reference",
                        "candidate_count": len(literal_candidates),
                    }
                )
                continue
        if target_apu is None and semantic_desc_signature:
            semantic_candidates = apu_signatures.get(semantic_desc_signature) or []
            if len(semantic_candidates) == 1:
                target_apu = semantic_candidates[0]
                match_kind = "semantic_descripcion"
            elif len(semantic_candidates) > 1:
                resource["nested_apu_status"] = "ambiguous_reference"
                resource["nested_apu_candidates"] = [candidate.get("temp_id") for candidate in semantic_candidates[:5]]
                links.append(
                    {
                        "parent_apu_temp_id": resource.get("apu_temp_id"),
                        "child_apu_temp_id": None,
                        "child_reference_code": code or None,
                        "child_reference_label": resource.get("descripcion"),
                        "match_kind": "ambiguous_semantic_reference",
                        "status": "ambiguous_reference",
                        "candidate_count": len(semantic_candidates),
                    }
                )
                continue
        if target_apu is None:
            if resource.get("is_nested_apu"):
                resource["nested_apu_status"] = "pending_reference"
                links.append(
                    {
                        "parent_apu_temp_id": resource.get("apu_temp_id"),
                        "child_apu_temp_id": None,
                        "child_reference_code": code or None,
                        "child_reference_label": resource.get("descripcion"),
                        "match_kind": "pending_semantic_reference",
                        "status": "pending_reference",
                    }
                )
            continue
        resource["is_nested_apu"] = True
        resource["nested_apu_target"] = target_apu.get("temp_id")
        resource["nested_apu_status"] = "resolved"
        links.append(
            {
                "parent_apu_temp_id": resource.get("apu_temp_id"),
                "child_apu_temp_id": target_apu.get("temp_id"),
                "child_reference_code": code or None,
                "child_reference_label": resource.get("descripcion"),
                "match_kind": match_kind,
                "status": "resolved",
            }
        )

    resources_by_apu: dict[str, list[dict]] = {}
    for resource in resources:
        resources_by_apu.setdefault(resource["apu_temp_id"], []).append(resource)

    for apu in apus:
        apu_resources = resources_by_apu.get(apu.get("temp_id"), [])
        apu["nested_apu_links"] = [item["nested_apu_target"] for item in apu_resources if item.get("nested_apu_target")]
        apu["nested_apu_count"] = len(apu["nested_apu_links"])
        apu["pending_nested_apu_count"] = sum(1 for item in apu_resources if item.get("nested_apu_status") == "pending_reference")
        apu["ambiguous_nested_apu_count"] = sum(1 for item in apu_resources if item.get("nested_apu_status") == "ambiguous_reference")
        apu["resources"] = apu_resources[:8]
        apu["resource_count"] = len(apu_resources)

    return apus, resources, links


def _parse_vae_value(line: str) -> str | None:
    percentage_match = re.search(r"(\d{1,3}(?:[.,]\d{1,2})?)\s*%", line)
    if percentage_match:
        return f"{percentage_match.group(1)}%"
    ratio_match = re.search(r"\b(\d+(?:[.,]\d{1,4})?)\b", line)
    if ratio_match:
        return ratio_match.group(1)
    return None


def _extract_vae_entries(lines: list[str]) -> list[dict]:
    entries: list[dict] = []
    active = False
    seen_labels: set[str] = set()
    for line in lines:
        cleaned = _clean_text(line)
        upper = cleaned.upper()
        if not upper:
            continue
        if "VAE" in upper or "DESAGREGACIÓN TECNOLÓGICA" in upper or "DESAGREGACION TECNOLOGICA" in upper:
            active = True
            if cleaned not in seen_labels:
                seen_labels.add(cleaned)
                entries.append(
                    {
                        "temp_id": f"vae_{len(entries) + 1}",
                        "label": cleaned,
                        "kind": "section_header",
                        "value": None,
                    }
                )
            continue
        if not active:
            continue
        if len(entries) >= 12:
            break
        if any(token in upper for token in ["ANALISIS DE PRECIOS UNITARIOS", "ANÁLISIS DE PRECIOS UNITARIOS"]) and len(entries) > 1:
            break
        if any(token in upper for token in ["FORMULARIO", "RESUMEN", "DESAGREG", "TECNOLOG", "%", "VAE", "CUADRILLA", "INCIDENCIA"]):
            if cleaned in seen_labels:
                continue
            seen_labels.add(cleaned)
            kind = "detail"
            if "%" in upper:
                kind = "percentage"
            elif "FORMULARIO" in upper:
                kind = "formulario"
            elif "RESUMEN" in upper:
                kind = "summary"
            entries.append(
                {
                    "temp_id": f"vae_{len(entries) + 1}",
                    "label": cleaned,
                    "kind": kind,
                    "value": _parse_vae_value(cleaned),
                }
            )
    return entries


def _detect_document_sections(signal_text: str) -> dict:
    upper_text = signal_text.upper()
    has_budget = any(token in upper_text for token in ["PRESUPUESTO", "RUBRO", "CANTIDAD", "PRECIO UNITARIO", "PRECIO TOTAL"])
    has_apus = (
        "ANALISIS DE PRECIOS UNITARIOS" in upper_text
        or "ANÁLISIS DE PRECIOS UNITARIOS" in upper_text
        or re.search(r"(?<![A-Z0-9])APU(?:S)?(?![A-Z0-9])", upper_text) is not None
    )
    has_resources = any(token in upper_text for token in ["MATERIALES", "MANO DE OBRA", "EQUIPO", "RECURSOS"])
    has_vae = any(token in upper_text for token in ["VAE", "DESAGREGACION TECNOLOGICA", "DESAGREGACIÓN TECNOLÓGICA"])
    return {
        "budget": has_budget,
        "apus": has_apus,
        "resources": has_resources,
        "vae": has_vae,
    }


def _resolve_document_kind(sections: dict) -> str:
    has_budget = bool(sections.get("budget"))
    has_apus = bool(sections.get("apus"))
    has_resources = bool(sections.get("resources"))
    has_vae = bool(sections.get("vae"))
    if has_budget and has_apus and has_resources:
        return "presupuesto_apus_integrados"
    if has_apus and has_vae:
        return "apus_vae"
    if has_budget and not has_apus and not has_vae:
        return "presupuesto_only"
    if has_apus or has_resources or has_vae:
        return "tecnico_parcial"
    return "indeterminado"


def _build_file_classification(filename: str, content: bytes, extension: str, budget_rows: list[dict], signal_text: str | None = None) -> dict:
    signal_text = signal_text if signal_text is not None else _extract_signal_text(content, extension)

    sections = _detect_document_sections(signal_text)
    if budget_rows:
        sections["budget"] = True
    document_kind = _resolve_document_kind(sections)
    detected_sections = [key for key, value in sections.items() if value]
    warnings: list[str] = []
    if document_kind == "presupuesto_only":
        warnings.append("Solo se detectó presupuesto; el sistema deberá bootstrapear APUs pendientes sin recursos.")
    if document_kind == "indeterminado":
        warnings.append("La fuente no pudo clasificarse con alta confianza; conviene validar manualmente su rol.")
    confidence = "alta" if document_kind in {"presupuesto_apus_integrados", "apus_vae", "presupuesto_only"} else "media"
    return {
        "document_kind": document_kind,
        "detected_sections": detected_sections,
        "detected_section_labels": [DOCUMENT_SECTION_LABELS.get(section, section) for section in detected_sections],
        "confidence": confidence,
        "warnings": warnings,
        "source_filename": filename,
        "source_format": extension.lstrip("."),
    }


def _apply_manual_role_hint(classification: dict, role_hint: str | None, has_budget_rows: bool) -> dict:
    normalized_role = _clean_text(role_hint).lower().replace("-", "_")
    if normalized_role not in MANUAL_ROLE_SECTION_MAP:
        return classification

    forced_sections = dict(MANUAL_ROLE_SECTION_MAP[normalized_role])
    if has_budget_rows:
        forced_sections["budget"] = True

    detected_sections = [key for key, value in forced_sections.items() if value]
    document_kind = MANUAL_ROLE_KIND_MAP[normalized_role]
    if normalized_role == "apus_resources" and has_budget_rows:
        document_kind = "presupuesto_apus_integrados"

    updated = dict(classification)
    updated["manual_role"] = normalized_role
    updated["document_kind"] = document_kind
    updated["detected_sections"] = detected_sections
    updated["detected_section_labels"] = [DOCUMENT_SECTION_LABELS.get(section, section) for section in detected_sections]
    updated["confidence"] = "manual"
    warnings = [item for item in list(updated.get("warnings") or []) if item]
    warnings.append(f"Clasificación ajustada manualmente como `{normalized_role}`.")
    updated["warnings"] = warnings
    return updated


class PublicProcurementPortalImportService:
    def __init__(self) -> None:
        self._analysis_cache: dict[tuple[str, str, str], dict] = {}

    def _analysis_cache_key(self, filename: str, content: bytes, role_hint: str | None) -> tuple[str, str, str]:
        safe_name = Path(filename or "portal_fuente").name
        fingerprint = hashlib.sha256(content or b"").hexdigest()
        return safe_name, fingerprint, role_hint or ""

    def _get_cached_analysis(self, key: tuple[str, str, str]) -> dict | None:
        cached = self._analysis_cache.get(key)
        return deepcopy(cached) if cached is not None else None

    def _set_cached_analysis(self, key: tuple[str, str, str], analysis: dict) -> None:
        if len(self._analysis_cache) >= ANALYSIS_CACHE_LIMIT:
            self._analysis_cache.pop(next(iter(self._analysis_cache)), None)
        self._analysis_cache[key] = deepcopy(analysis)

    def analyze_upload(self, filename: str, content: bytes, role_hint: str | None = None) -> dict:
        safe_name = Path(filename or "portal_fuente").name
        cache_key = self._analysis_cache_key(safe_name, content, role_hint)
        cached = self._get_cached_analysis(cache_key)
        if cached is not None:
            return cached
        extension = Path(safe_name).suffix.lower()
        parse_warning = None
        if extension == ".pdf":
            signal_lines = _extract_signal_lines(content, extension)
            budget_title, budget_rows = _parse_budget_lines(signal_lines)
            if not budget_rows:
                budget_title, budget_rows = _parse_budget_pdf(content)
            source_format = "pdf_presupuesto"
        elif extension in {".xlsx", ".xlsm"}:
            budget_title, budget_rows = _parse_budget_xlsx(content)
            signal_lines = _extract_signal_lines(content, extension)
            source_format = "excel_presupuesto"
        elif extension == ".xls":
            try:
                budget_title, budget_rows = _parse_budget_xls(content)
            except HTTPException as exc:
                budget_title, budget_rows = Path(safe_name).stem, []
                parse_warning = exc.detail
            signal_lines = _extract_signal_lines(content, extension)
            source_format = "excel_presupuesto"
        else:
            raise HTTPException(
                status_code=400,
                detail="El importador integrado de Portal de compras públicas admite actualmente archivos PDF, XLS y XLSX.",
            )

        signal_text = " ".join(signal_lines)
        classification = _build_file_classification(safe_name, content, extension, budget_rows, signal_text=signal_text)
        classification = _apply_manual_role_hint(classification, role_hint, bool(budget_rows))
        extracted_apus, extracted_resources = _extract_apu_resource_pairs(signal_lines)
        extracted_apus, extracted_resources, duplicate_apu_incidents = _dedupe_apus_by_identity(extracted_apus, extracted_resources)
        extracted_apus, extracted_resources, apu_links = _resolve_nested_apu_links(extracted_apus, extracted_resources)
        vae_entries = _extract_vae_entries(signal_lines)
        if extracted_apus:
            if "apus" not in classification["detected_sections"]:
                classification["detected_sections"].append("apus")
                classification["detected_section_labels"].append(DOCUMENT_SECTION_LABELS["apus"])
            classification["document_kind"] = (
                "presupuesto_apus_integrados" if budget_rows else ("apus_vae" if "vae" in classification["detected_sections"] else "tecnico_parcial")
            )
        if extracted_resources and "resources" not in classification["detected_sections"]:
            classification["detected_sections"].append("resources")
            classification["detected_section_labels"].append(DOCUMENT_SECTION_LABELS["resources"])
        if vae_entries and "vae" not in classification["detected_sections"]:
            classification["detected_sections"].append("vae")
            classification["detected_section_labels"].append(DOCUMENT_SECTION_LABELS["vae"])
            if extracted_apus:
                classification["document_kind"] = "apus_vae"
        if parse_warning:
            classification["warnings"] = [*list(classification.get("warnings") or []), parse_warning]
        if not budget_rows and not any(section in classification["detected_sections"] for section in ["apus", "resources", "vae"]):
            raise HTTPException(
                status_code=400,
                detail="No se encontraron rubros válidos ni secciones técnicas reconocibles en el archivo cargado para Portal de compras públicas.",
            )
        aligned_budget_rows, unmatched_pending_apus, match_incidents = _align_budget_rows_with_apus(budget_rows, extracted_apus)
        budget_rows = aligned_budget_rows
        should_generate_pending_apus = bool(budget_rows) and "apus" not in classification["detected_sections"]
        pending_apus = _build_pending_apus_from_budget(budget_rows) if should_generate_pending_apus else unmatched_pending_apus
        if extracted_apus and unmatched_pending_apus:
            pending_apus = unmatched_pending_apus
        parser_profile = "soce_sercop_pdf_integrado_codigo_v1"
        if any((row.get("match_strategy") == "descripcion_unidad_precio") for row in budget_rows):
            parser_profile = "soce_sercop_pdf_vae_resumen_v1"
        blocking_incidents = [*duplicate_apu_incidents, *match_incidents]
        non_blocking_incidents = []
        if pending_apus:
            non_blocking_incidents.append(
                {
                    "severity": "warning",
                    "code": "empty_apus_created_from_budget",
                    "message": f"Se crearán {len(pending_apus)} APUs vacíos porque no existe correspondencia técnica en los documentos.",
                    "count": len(pending_apus),
                }
            )
        apu_reconciliation = {
            "bootstrap_promoted_to_real": 0,
            "pending_apus_remaining": len(pending_apus),
            "real_apus_retained": len(extracted_apus),
            "temp_id_remaps": 0,
            "parser_profile": parser_profile,
            "blocking_incidents_count": len(blocking_incidents),
        }

        chapters = [row.get("capitulo") for row in budget_rows if row.get("capitulo")]
        total_amount = sum(Decimal(str(row.get("precio_total") or 0)) for row in budget_rows)
        unit_count = sum(1 for row in budget_rows if row.get("precio_total") is not None)
        chapter_breakdown = _build_chapter_breakdown(budget_rows)
        dominant_chapter = chapter_breakdown[0]["name"] if chapter_breakdown else None

        analysis = {
            "source_filename": safe_name,
            "source_format": source_format,
            "budget_title": budget_title or Path(safe_name).stem,
            "items_count": unit_count,
            "chapters_count": len(set(chapters)),
            "total_amount": f"{total_amount:.2f}",
            "currency": "USD",
            "detected_at": datetime.now(timezone.utc).isoformat(),
            "rubros": budget_rows,
            "chapters": sorted(set(chapters)),
            "chapter_breakdown": chapter_breakdown,
            "dominant_chapter": dominant_chapter,
            "sample_rubros": _build_sample_rubros(budget_rows),
            "classification": classification,
            "generated_apus": pending_apus,
            "sample_apus": extracted_apus[:5],
            "sample_resources": extracted_resources[:8],
            "sample_vae_entries": vae_entries[:5],
            "source_files": [
                {
                    "filename": safe_name,
                    "format": source_format,
                    "document_kind": classification["document_kind"],
                    "detected_sections": classification["detected_sections"],
                    "detected_section_labels": classification["detected_section_labels"],
                    "confidence": classification["confidence"],
                    "warnings": classification["warnings"],
                    "manual_role": classification.get("manual_role"),
                    "sample_apus_count": len(extracted_apus),
                    "sample_resources_count": len(extracted_resources),
                    "nested_apu_link_count": len(apu_links),
                    "vae_entries_count": len(vae_entries),
                }
            ],
            "summary": {
                "has_budget": True,
                "has_apus": "apus" in classification["detected_sections"] or bool(extracted_apus),
                "has_resources": "resources" in classification["detected_sections"] or bool(extracted_resources),
                "has_vae": "vae" in classification["detected_sections"] or bool(vae_entries),
                "has_location_hint": False,
                "source_kind": "archivo_base",
                "dominant_chapter": dominant_chapter,
                "document_kind": classification["document_kind"],
                "detected_sections": classification["detected_sections"],
                "coverage_level": "presupuesto + apus_base" if pending_apus else "solo_presupuesto",
                "warnings": classification["warnings"],
                "manual_role": classification.get("manual_role"),
                "parser_profile": parser_profile,
                "blocking_incidents": blocking_incidents,
                "non_blocking_incidents": non_blocking_incidents,
                "requires_superadmin_consent": bool(blocking_incidents),
            },
        }
        analysis["analysis_bundle"] = _build_analysis_bundle(
            analysis,
            pending_apus=pending_apus,
            extracted_apus=extracted_apus,
            extracted_resources=extracted_resources,
            apu_links=apu_links,
            vae_entries=vae_entries,
            apu_reconciliation=apu_reconciliation,
        )
        analysis["summary"]["coverage_level"] = analysis["analysis_bundle"]["summary"]["coverage_level"]
        analysis["summary"]["blocking_incidents"] = blocking_incidents
        analysis["summary"]["non_blocking_incidents"] = non_blocking_incidents
        analysis["summary"]["requires_superadmin_consent"] = bool(blocking_incidents)
        analysis["summary"]["parser_profile"] = parser_profile
        analysis["analysis_bundle"]["summary"]["blocking_incidents"] = blocking_incidents
        analysis["analysis_bundle"]["summary"]["non_blocking_incidents"] = non_blocking_incidents
        analysis["analysis_bundle"]["summary"]["requires_superadmin_consent"] = bool(blocking_incidents)
        analysis["analysis_bundle"]["summary"]["parser_profile"] = parser_profile
        self._set_cached_analysis(cache_key, analysis)
        return analysis

    def analyze_uploads(
        self,
        files: list[tuple[str, bytes]],
        source_role_hints: dict[str, str] | None = None,
        existing_analysis: dict | None = None,
    ) -> dict:
        normalized_files = [(filename, content) for filename, content in files if content]
        if not normalized_files:
            raise HTTPException(
                status_code=400,
                detail="Debes subir al menos un archivo válido para analizar el portal de compras públicas.",
            )

        source_role_hints = source_role_hints or {}
        analyses = [
            self.analyze_upload(filename, content, role_hint=source_role_hints.get(Path(filename or "").name))
            for filename, content in normalized_files
        ]
        coerced_existing = _coerce_existing_analysis(existing_analysis)
        if coerced_existing:
            existing_summary = dict(coerced_existing.get("summary") or {})
            existing_warnings = [item for item in list(existing_summary.get("warnings") or []) if item]
            existing_warnings.append("Se reutilizó el bundle existente como base para una reimportación incremental.")
            existing_summary["warnings"] = existing_warnings
            coerced_existing["summary"] = existing_summary
            analyses.insert(0, coerced_existing)
        if len(analyses) == 1:
            single = dict(analyses[0])
            return single

        combined_rubros: list[dict] = []
        chapters: set[str] = set()
        budget_titles: list[str] = []
        source_files: list[dict] = []
        total_amount = Decimal("0.00")
        warnings: list[str] = []
        document_kind_counts: dict[str, int] = {}
        has_budget = False
        has_apus = False
        has_resources = False
        has_vae = False
        generated_apus: list[dict] = []
        extracted_apus: list[dict] = []
        extracted_resources: list[dict] = []
        extracted_vae_entries: list[dict] = []
        duplicate_budget_rows = 0
        blocking_incidents: list[dict] = []
        non_blocking_incidents: list[dict] = []

        for analysis in analyses:
            budget_titles.append(str(analysis.get("budget_title") or "").strip())
            source_files.extend(list(analysis.get("source_files") or []))
            summary = dict(analysis.get("summary") or {})
            has_budget = has_budget or bool(summary.get("has_budget"))
            has_apus = has_apus or bool(summary.get("has_apus"))
            has_resources = has_resources or bool(summary.get("has_resources"))
            has_vae = has_vae or bool(summary.get("has_vae"))
            warnings.extend([item for item in list(summary.get("warnings") or []) if item])
            blocking_incidents.extend(list(summary.get("blocking_incidents") or []))
            non_blocking_incidents.extend(list(summary.get("non_blocking_incidents") or []))
            document_kind = str(summary.get("document_kind") or "indeterminado")
            document_kind_counts[document_kind] = document_kind_counts.get(document_kind, 0) + 1
            bundle = dict(analysis.get("analysis_bundle") or {})
            generated_apus.extend(list(analysis.get("generated_apus") or []))
            extracted_apus.extend(list(bundle.get("apus") or analysis.get("sample_apus") or []))
            extracted_resources.extend(list(bundle.get("resources") or analysis.get("sample_resources") or []))
            extracted_vae_entries.extend(list(bundle.get("vae_entries") or analysis.get("sample_vae_entries") or []))
            for row in list(analysis.get("rubros") or []):
                enriched = dict(row)
                enriched["source_filename"] = analysis.get("source_filename")
                combined_rubros.append(enriched)
                chapter_name = str(enriched.get("capitulo") or "").strip()
                if chapter_name:
                    chapters.add(chapter_name)
            total_amount += Decimal(str(analysis.get("total_amount") or 0))

        combined_rubros, duplicate_budget_rows = _merge_unique_items(combined_rubros, _build_budget_line_signature)
        generated_apus, duplicate_generated_apus = _merge_unique_items(generated_apus, _build_apu_signature)
        extracted_apus, duplicate_extracted_apus = _merge_unique_items(extracted_apus, _build_apu_signature)
        extracted_resources, duplicate_resources = _merge_unique_items(extracted_resources, _build_resource_signature)
        extracted_vae_entries, duplicate_vae_entries = _merge_unique_items(extracted_vae_entries, _build_vae_signature)
        generated_apus, extracted_apus, extracted_resources, apu_reconciliation = _reconcile_pending_with_real_apus(
            generated_apus,
            extracted_apus,
            extracted_resources,
        )
        dedupe_total = (
            duplicate_budget_rows
            + duplicate_generated_apus
            + duplicate_extracted_apus
            + duplicate_resources
            + duplicate_vae_entries
        )
        if dedupe_total:
            warnings.append(
                "Se conciliaron elementos duplicados entre fuentes para evitar sumar rubros o APUs repetidos."
            )
        if coerced_existing:
            warnings.append(
                "La nueva lectura se fusionó con el bundle previo para preservar trabajo técnico ya analizado."
            )

        chapter_breakdown = _build_chapter_breakdown(combined_rubros)
        dominant_chapter = chapter_breakdown[0]["name"] if chapter_breakdown else None
        source_names = [item["filename"] for item in source_files if item.get("filename")]
        source_summary = ", ".join(source_names[:2])
        if len(source_names) > 2:
            source_summary = f"{source_summary}, +{len(source_names) - 2} más"

        source_document_kinds = sorted(document_kind_counts.keys())
        if has_budget and has_apus and has_resources:
            document_kind = "bundle_completo"
        elif has_budget and (has_apus or generated_apus):
            document_kind = "bundle_presupuesto_apus"
        elif has_budget:
            document_kind = "presupuesto_only"
        else:
            document_kind = source_document_kinds[0] if len(source_document_kinds) == 1 else "bundle_parcial"

        merged_analysis = {
            "source_filename": source_summary or f"{len(source_files)} archivos técnicos",
            "source_format": "multi_fuente",
            "budget_title": " / ".join([title for title in budget_titles if title][:2]) or "Importación combinada",
            "items_count": len(combined_rubros),
            "chapters_count": len(chapters),
            "total_amount": f"{total_amount:.2f}",
            "currency": "USD",
            "detected_at": datetime.now(timezone.utc).isoformat(),
            "rubros": combined_rubros,
            "chapters": sorted(chapters),
            "chapter_breakdown": chapter_breakdown,
            "dominant_chapter": dominant_chapter,
            "sample_rubros": _build_sample_rubros(combined_rubros),
            "sample_apus": extracted_apus[:5],
            "sample_resources": extracted_resources[:8],
            "sample_vae_entries": extracted_vae_entries[:5],
            "source_files": source_files,
            "generated_apus": generated_apus,
            "classification": {
                "document_kind": document_kind,
                "detected_sections": [key for key, value in {
                    "budget": has_budget,
                    "apus": has_apus,
                    "resources": has_resources,
                    "vae": has_vae,
                }.items() if value],
                "detected_section_labels": [
                    DOCUMENT_SECTION_LABELS.get(section, section)
                    for section in [key for key, value in {
                        "budget": has_budget,
                        "apus": has_apus,
                        "resources": has_resources,
                        "vae": has_vae,
                    }.items() if value]
                ],
                "confidence": "media",
                "warnings": warnings,
            },
            "summary": {
                "has_budget": has_budget,
                "has_apus": has_apus,
                "has_resources": has_resources,
                "has_vae": has_vae,
                "has_location_hint": False,
                "source_kind": "archivo_base",
                "dominant_chapter": dominant_chapter,
                "source_count": len(source_files),
                "document_kind": document_kind,
                "document_kind_counts": document_kind_counts,
                "dedupe_summary": {
                    "budget_rows_removed": duplicate_budget_rows,
                    "generated_apus_removed": duplicate_generated_apus,
                    "real_apus_removed": duplicate_extracted_apus,
                    "resources_removed": duplicate_resources,
                    "vae_entries_removed": duplicate_vae_entries,
                    "merge_mode": "incremental_signature_reconciliation" if coerced_existing else "signature_reconciliation",
                },
                "warnings": warnings,
                "blocking_incidents": blocking_incidents,
                "non_blocking_incidents": non_blocking_incidents,
                "requires_superadmin_consent": bool(blocking_incidents),
            },
        }
        resolved_apus, resolved_resources, apu_links = _resolve_nested_apu_links(extracted_apus, extracted_resources)
        merged_analysis["analysis_bundle"] = _build_analysis_bundle(
            merged_analysis,
            pending_apus=generated_apus,
            extracted_apus=resolved_apus,
            extracted_resources=resolved_resources,
            apu_links=apu_links,
            vae_entries=extracted_vae_entries,
            apu_reconciliation=apu_reconciliation,
        )
        merged_analysis["summary"]["coverage_level"] = merged_analysis["analysis_bundle"]["summary"]["coverage_level"]
        merged_analysis["summary"]["blocking_incidents"] = blocking_incidents
        merged_analysis["summary"]["non_blocking_incidents"] = non_blocking_incidents
        merged_analysis["summary"]["requires_superadmin_consent"] = bool(blocking_incidents)
        merged_analysis["analysis_bundle"]["summary"]["blocking_incidents"] = blocking_incidents
        merged_analysis["analysis_bundle"]["summary"]["non_blocking_incidents"] = non_blocking_incidents
        merged_analysis["analysis_bundle"]["summary"]["requires_superadmin_consent"] = bool(blocking_incidents)
        return merged_analysis

    def build_excel_bytes(self, analysis: dict, article_title: str, export_context: dict | None = None) -> tuple[bytes, str]:
        payload = dict(analysis or {})
        rows = list(payload.get("rubros") or [])
        generated_apus = list(payload.get("generated_apus") or [])
        analysis_bundle = dict(payload.get("analysis_bundle") or {})
        export_context = dict(export_context or {})
        if not rows:
            raise HTTPException(status_code=400, detail="No existe información importada suficiente para generar el Excel técnico.")
        repo_root = Path(__file__).resolve().parents[3]
        template_path = repo_root / "docs" / "adicionales" / "APU_Presupuesto_Completo.xlsx"
        workbook = load_workbook(template_path) if template_path.exists() else Workbook()
        template_apu_sheet_names = [name for name in workbook.sheetnames if name.startswith("APU_")]
        template_apu_sheet = workbook[template_apu_sheet_names[0]] if template_apu_sheet_names else None
        if template_apu_sheet is None:
            raise HTTPException(status_code=400, detail="La plantilla `APU_Presupuesto_Completo.xlsx` no contiene una hoja modelo `APU_*` para generar el Excel del portal.")

        full_apus = list(analysis_bundle.get("apus") or analysis_bundle.get("sample_reconciled_apus") or analysis_bundle.get("sample_apus") or generated_apus)
        full_resources = list(analysis_bundle.get("resources") or analysis_bundle.get("sample_resources") or [])
        full_vae_entries = list(analysis_bundle.get("vae_entries") or analysis_bundle.get("sample_vae_entries") or [])
        export_entries = _build_full_apu_export_entries(payload, full_apus, full_resources)
        if not export_entries:
            raise HTTPException(status_code=400, detail="No existe un conjunto exportable de APUs para construir el Excel del portal.")

        detected_vae_percentages = [
            value
            for value in (_extract_percentage_number(entry.get("value") or entry.get("label")) for entry in full_vae_entries)
            if value is not None
        ]
        default_vae_percentage = detected_vae_percentages[0] if len(detected_vae_percentages) == 1 else None

        if "PRESUPUESTO" not in workbook.sheetnames:
            workbook.active.title = "PRESUPUESTO"
        budget_ws = workbook["PRESUPUESTO"]
        budget_ws["A2"] = article_title or "PRESUPUESTO"
        budget_capacity = 224
        budget_extra_rows = max(0, len(export_entries) - budget_capacity)
        if budget_extra_rows:
            budget_ws.insert_rows(228, budget_extra_rows)
        if budget_ws.max_row >= 4:
            _clear_sheet_values(budget_ws, start_row=4, start_col=1, end_col=8)
        budget_data_row_style = 4
        budget_total_row_style = 229 + budget_extra_rows
        budget_warning_title_style = 231 + budget_extra_rows
        budget_warning_text_style = 232 + budget_extra_rows

        budget_row = 4
        for entry in export_entries:
            _copy_row_style(budget_ws, budget_data_row_style, budget_row, end_col=8)
            budget_ws.cell(row=budget_row, column=1, value=entry.get("item_index"))
            budget_ws.cell(row=budget_row, column=2, value=entry.get("original_code") or entry.get("export_key"))
            budget_ws.cell(row=budget_row, column=3, value=f'=IFERROR(VLOOKUP("{entry.get("export_key")}",INDICE_APUs!$C:$D,2,0),"")')
            budget_ws.cell(row=budget_row, column=4, value=f'=IFERROR(VLOOKUP("{entry.get("export_key")}",INDICE_APUs!$C:$E,3,0),"")')
            budget_ws.cell(row=budget_row, column=5, value=entry.get("cantidad") or 0)
            budget_ws.cell(row=budget_row, column=6, value=f'=IFERROR(VLOOKUP("{entry.get("export_key")}",INDICE_APUs!$C:$H,6,0),0)')
            budget_ws.cell(row=budget_row, column=7, value=f'=ROUND(E{budget_row}*F{budget_row},2)')
            budget_ws.cell(row=budget_row, column=8, value=entry.get("status") or "")
            budget_row += 1

        budget_total_row = budget_row + 1
        _copy_row_style(budget_ws, budget_total_row_style, budget_total_row, end_col=8)
        budget_ws.cell(row=budget_total_row, column=1, value="TOTAL DEL PRESUPUESTO (sin IVA):")
        budget_ws.cell(row=budget_total_row, column=7, value=f'=ROUND(SUMPRODUCT((ROW(G4:G{budget_row - 1})>0)*G4:G{budget_row - 1}),2)')
        _copy_row_style(budget_ws, budget_warning_title_style, budget_total_row + 2, end_col=8)
        budget_ws.cell(row=budget_total_row + 2, column=1, value="⚠️  AVISO DE RESPONSABILIDAD")
        _copy_row_style(budget_ws, budget_warning_text_style, budget_total_row + 3, end_col=8)
        budget_ws.cell(
            row=budget_total_row + 3,
            column=1,
            value="AVISO: Los valores presentados son referenciales y han sido generados automáticamente a partir de los documentos proporcionados. Pueden existir errores de cálculo, interpretación o extracción de datos. Es responsabilidad exclusiva del usuario verificar, validar y corregir toda la información antes de su uso oficial o presentación. El generador queda exento de toda responsabilidad.",
        )

        if "DESAGREGACIÓN CONSOLIDADA" in workbook.sheetnames:
            summary_ws = workbook["DESAGREGACIÓN CONSOLIDADA"]
            summary_ws["B2"] = article_title or "PRESUPUESTO"
            summary_ws["B3"] = export_context.get("location_label") or "-"
            summary_ws["B4"] = export_context.get("exported_at") or datetime.now(timezone.utc).strftime("%Y-%m-%d")
            summary_capacity = 224
            summary_extra_rows = max(0, len(export_entries) - summary_capacity)
            if summary_extra_rows:
                summary_ws.insert_rows(231, summary_extra_rows)
            _clear_sheet_values(summary_ws, start_row=7, start_col=1, end_col=11)
            summary_data_row_style = 7
            summary_total_row_style = 232 + summary_extra_rows
            summary_ppem_row_style = 233 + summary_extra_rows
            summary_notice_row_style = 235 + summary_extra_rows
            summary_notice_text_style = 236 + summary_extra_rows
            for row_index, entry in enumerate(export_entries, start=7):
                _copy_row_style(summary_ws, summary_data_row_style, row_index, end_col=11)
                summary_ws.cell(row=row_index, column=1, value=entry.get("item_index"))
                summary_ws.cell(row=row_index, column=2, value=entry.get("cpc_code") or "")
                summary_ws.cell(row=row_index, column=3, value=f'=IFERROR(VLOOKUP(IF(LEN(B{row_index})>9,LEFT(B{row_index},9),B{row_index}),UMBRAL_SERCOP!$A:$B,2,0),"")')
                summary_ws.cell(row=row_index, column=4, value=f'=IFERROR(VLOOKUP("{entry.get("export_key")}",INDICE_APUs!$C:$D,2,0),"")')
                summary_ws.cell(row=row_index, column=5, value=f'=IFERROR(VLOOKUP("{entry.get("export_key")}",INDICE_APUs!$C:$E,3,0),"")')
                summary_ws.cell(row=row_index, column=6, value=entry.get("cantidad") or 0)
                summary_ws.cell(row=row_index, column=7, value=f'=IFERROR(VLOOKUP("{entry.get("export_key")}",INDICE_APUs!$C:$H,6,0),0)')
                summary_ws.cell(row=row_index, column=8, value=f'=ROUND(F{row_index}*G{row_index},2)')
                summary_ws.cell(row=row_index, column=9, value="")
                summary_ws.cell(row=row_index, column=10, value=0)
                summary_ws.cell(row=row_index, column=11, value=f'=I{row_index}*J{row_index}')

            summary_last_data_row = 6 + len(export_entries)
            total_project_row = summary_last_data_row + 2
            ppem_row = total_project_row + 1
            _copy_row_style(summary_ws, summary_total_row_style, total_project_row, end_col=11)
            summary_ws.cell(row=total_project_row, column=1, value="MONTO TOTAL DEL PROYECTO ($):")
            summary_ws.cell(row=total_project_row, column=8, value=f'=ROUND(SUM(H7:H{summary_last_data_row}),2)')
            summary_ws.cell(row=total_project_row, column=9, value="TOTAL PESO RELATIVO:")
            summary_ws.cell(row=total_project_row, column=11, value=f'=SUM(I7:I{summary_last_data_row})')
            _copy_row_style(summary_ws, summary_ppem_row_style, ppem_row, end_col=11)
            summary_ws.cell(row=ppem_row, column=1, value="PORCENTAJE DE PARTICIPACION ECUATORIANA MINIMO DEL PROYECTO (PPEM = SAEPi):")
            summary_ws.cell(row=ppem_row, column=9, value=f'=SUM(K7:K{summary_last_data_row})')
            _copy_row_style(summary_ws, summary_notice_row_style, ppem_row + 2, end_col=11)
            summary_ws.cell(row=ppem_row + 2, column=1, value="Aviso: Los valores son referenciales y calculados automaticamente. Verifique antes de uso oficial.")
            _copy_row_style(summary_ws, summary_notice_text_style, ppem_row + 3, end_col=11)
            summary_ws.cell(
                row=ppem_row + 3,
                column=1,
                value="AVISO: Los valores presentados son referenciales y han sido generados automáticamente a partir de los documentos proporcionados. Pueden existir errores de cálculo, interpretación o extracción de datos. Es responsabilidad exclusiva del usuario verificar, validar y corregir toda la información antes de su uso oficial o presentación. El generador queda exento de toda responsabilidad.",
            )

        if "PARAMETROS" in workbook.sheetnames:
            parameters_ws = workbook["PARAMETROS"]
            parameters_ws["B3"] = export_context.get("indirect_cost_percentage", 0.2)

        if "INDICE_APUs" in workbook.sheetnames:
            index_ws = workbook["INDICE_APUs"]
            index_capacity = 264
            index_extra_rows = max(0, len(export_entries) - index_capacity)
            if index_extra_rows:
                index_ws.insert_rows(268, index_extra_rows)
            index_ws.delete_rows(3, max(0, index_ws.max_row - 2))
            index_ws.cell(row=3, column=1, value="── APUs PRINCIPALES ──")

        categorized_resources = {
            "CAT_EQUIPO": [resource for resource in full_resources if resource.get("resource_type") == "equipo"],
            "CAT_MATERIALES": [resource for resource in full_resources if resource.get("resource_type") == "material"],
            "CAT_TRANSPORTE": [resource for resource in full_resources if resource.get("resource_type") == "transporte"],
            "CAT_MANO_OBRA": [resource for resource in full_resources if resource.get("resource_type") == "mano_obra"],
        }
        for sheet_name, items in categorized_resources.items():
            if sheet_name not in workbook.sheetnames:
                continue
            _populate_catalog_sheet(workbook[sheet_name], items)

        for sheet_name in template_apu_sheet_names[1:]:
            workbook.remove(workbook[sheet_name])
        apu_sheet_details: list[dict] = []
        for entry in export_entries:
            sheet_meta = _build_apu_sheet(
                workbook,
                template_apu_sheet,
                entry,
                default_vae_percentage=default_vae_percentage,
            )
            apu_sheet_details.append({**entry, **sheet_meta})

        if "INDICE_APUs" in workbook.sheetnames:
            index_ws = workbook["INDICE_APUs"]
            for row_index, entry in enumerate(apu_sheet_details, start=4):
                sheet_ref = _quote_excel_sheet_name(entry["sheet_name"])
                index_ws.cell(row=row_index, column=1, value=entry.get("item_index"))
                index_ws.cell(row=row_index, column=2, value="PRINCIPAL")
                index_ws.cell(row=row_index, column=3, value=entry.get("export_key"))
                index_ws.cell(row=row_index, column=4, value=f"={sheet_ref}!C3")
                index_ws.cell(row=row_index, column=5, value=f"={sheet_ref}!C4")
                index_ws.cell(row=row_index, column=6, value=f"={sheet_ref}!G{entry['direct_total_row']}")
                index_ws.cell(row=row_index, column=7, value=f'=ROUND(F{row_index}*PARAMETROS!B3,2)')
                index_ws.cell(row=row_index, column=8, value=f'=ROUND(F{row_index}+G{row_index},2)')
                index_ws.cell(row=row_index, column=9, value=f"={sheet_ref}!M{entry['vae_total_row']}")

        if "DESAGREGACIÓN CONSOLIDADA" in workbook.sheetnames:
            summary_ws = workbook["DESAGREGACIÓN CONSOLIDADA"]
            for row_index, entry in enumerate(apu_sheet_details, start=7):
                sheet_ref = _quote_excel_sheet_name(entry["sheet_name"])
                summary_ws.cell(row=row_index, column=10, value=f'=IFERROR({sheet_ref}!M{entry["vae_total_row"]},0)')
                summary_ws.cell(row=row_index, column=9, value=f'=IFERROR(H{row_index}/H{total_project_row},0)')
                summary_ws.cell(row=row_index, column=11, value=f'=I{row_index}*J{row_index}')

        if template_apu_sheet.title in workbook.sheetnames:
            workbook.remove(template_apu_sheet)

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        safe_stem = re.sub(r"[^A-Za-z0-9._-]+", "_", (article_title or "portal_compras_publicas")).strip("._") or "portal_compras_publicas"
        return buffer.getvalue(), f"{safe_stem}_portal_compras_publicas.xlsx"


public_procurement_portal_import_service = PublicProcurementPortalImportService()
