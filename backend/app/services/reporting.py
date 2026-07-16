import os
import io
import re
import base64
import math
import struct
import urllib.parse
import zlib
import json
from copy import copy
from collections import OrderedDict
from decimal import Decimal, ROUND_HALF_UP
from datetime import date, datetime, time, timezone
from zipfile import ZIP_DEFLATED, ZipFile
from typing import Any, Dict, List, Optional, Tuple
from pathlib import Path
import openpyxl
import httpx
from openpyxl.cell.cell import Cell as OpenPyxlCell, MergedCell
from openpyxl.drawing.image import Image as OpenPyxlImage
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.formula.translate import Translator
from openpyxl.writer.excel import ExcelWriter
from sqlalchemy.orm import selectinload
from app.core.config import settings
from app.core.rounding import round_decimal
from app.core.text_formatting import normalize_sentence_case
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfgen import canvas
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image as ReportLabImage, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from app.api.endpoints.maestros import TIPOS_PROYECTO, CATEGORIAS
from app.services.apu_explosion import collect_apu_exploded_resources
from app.services.apu_resource_readiness import apu_resource_readiness_service

try:
    from PIL import Image as PILImage, ImageDraw
except Exception:  # pragma: no cover - entorno sin Pillow
    PILImage = None
    ImageDraw = None

class ReportingService:
    REPORT_AUTHOR = "GIPROY Network"
    REPORT_EXPORT_RENDER_VERSION = "2026-06-27-cronograma-gantt-draft-report-warnings-v1"
    PROJECT_TYPES_BY_ID = {int(item["id"]): item.get("descripcion") or "" for item in TIPOS_PROYECTO}
    PROJECT_CATEGORIES_BY_ID = {int(item["id"]): item.get("descripcion") or item.get("description") or "" for item in CATEGORIAS}
    INDIRECTOS_CATEGORY_LABELS = {
        "1.1": "Beneficios",
        "1.2": "Personal Tecnico",
        "1.3": "Materiales de Ingenieria",
        "1.4": "Oficina de la Obra",
        "2.1": "Personal de Servicio",
        "2.2": "Transporte",
        "2.3": "Comunicaciones y Fletes",
        "3.1": "Edificios del campamento",
        "3.2": "Caminos, puentes y cerramientos",
        "3.3": "Servicios Basicos",
        "3.4": "Talleres de la Obra",
        "3.5": "Arriendos",
        "4.1": "Entretenimiento",
        "4.2": "Costos medicos, legales y otros",
        "4.3": "Seguros, garantias e impuestos",
        "4.4": "Gastos Varios",
        "5.1": "Material de consumo",
        "5.2": "Reparacion y mantenimiento de equipos",
        "5.3": "Herramientas",
        "6.1": "Utilidad",
        "6.2": "Imprevistos",
        "7.1": "Personalizados",
    }

    def __init__(self, templates_dir: str = None):
        self.templates_dir = templates_dir or str(settings.REPORTS_DIR)
        self._apu_template_metadata_cache: Dict[str, Dict[str, Any]] = {}
        self._workbook_bytes_cache: "OrderedDict[Tuple[Any, ...], bytes]" = OrderedDict()
        self._report_export_cache: "OrderedDict[Tuple[Any, ...], bytes]" = OrderedDict()
        self._map_image_cache: "OrderedDict[Tuple[Any, ...], bytes]" = OrderedDict()
        self._workbook_bytes_cache_limit = 8
        self._report_export_cache_limit = 32
        self._map_image_cache_limit = 16

    def _get_cached_bytes(
        self,
        cache_store: "OrderedDict[Tuple[Any, ...], bytes]",
        cache_key: Tuple[Any, ...],
    ) -> Optional[bytes]:
        payload = cache_store.get(cache_key)
        if payload is None:
            return None
        cache_store.move_to_end(cache_key)
        return payload

    def _set_cached_bytes(
        self,
        cache_store: "OrderedDict[Tuple[Any, ...], bytes]",
        cache_key: Tuple[Any, ...],
        payload: bytes,
        limit: int,
    ) -> None:
        cache_store[cache_key] = payload
        cache_store.move_to_end(cache_key)
        while len(cache_store) > max(int(limit or 0), 1):
            cache_store.popitem(last=False)

    def _number_to_spanish_words(self, value: int) -> str:
        units = {
            0: "cero",
            1: "uno",
            2: "dos",
            3: "tres",
            4: "cuatro",
            5: "cinco",
            6: "seis",
            7: "siete",
            8: "ocho",
            9: "nueve",
            10: "diez",
            11: "once",
            12: "doce",
            13: "trece",
            14: "catorce",
            15: "quince",
            16: "dieciseis",
            17: "diecisiete",
            18: "dieciocho",
            19: "diecinueve",
            20: "veinte",
            21: "veintiuno",
            22: "veintidos",
            23: "veintitres",
            24: "veinticuatro",
            25: "veinticinco",
            26: "veintiseis",
            27: "veintisiete",
            28: "veintiocho",
            29: "veintinueve",
        }
        tens = {
            30: "treinta",
            40: "cuarenta",
            50: "cincuenta",
            60: "sesenta",
            70: "setenta",
            80: "ochenta",
            90: "noventa",
        }
        hundreds = {
            100: "cien",
            200: "doscientos",
            300: "trescientos",
            400: "cuatrocientos",
            500: "quinientos",
            600: "seiscientos",
            700: "setecientos",
            800: "ochocientos",
            900: "novecientos",
        }

        def under_hundred(n: int) -> str:
            if n < 30:
                return units[n]
            ten = (n // 10) * 10
            unit = n % 10
            return tens[ten] if unit == 0 else f"{tens[ten]} y {units[unit]}"

        def under_thousand(n: int) -> str:
            if n < 100:
                return under_hundred(n)
            if n in hundreds:
                return hundreds[n]
            hundred = (n // 100) * 100
            remainder = n % 100
            prefix = "ciento" if hundred == 100 else hundreds[hundred]
            return f"{prefix} {under_hundred(remainder)}"

        if value < 1000:
            return under_thousand(value)
        if value < 1_000_000:
            thousands = value // 1000
            remainder = value % 1000
            if thousands == 1:
                prefix = "mil"
            else:
                prefix = f"{under_thousand(thousands)} mil"
            return prefix if remainder == 0 else f"{prefix} {under_thousand(remainder)}"
        millions = value // 1_000_000
        remainder = value % 1_000_000
        if millions == 1:
            prefix = "un millon"
        else:
            prefix = f"{self._number_to_spanish_words(millions)} millones"
        if remainder == 0:
            return prefix
        return f"{prefix} {self._number_to_spanish_words(remainder)}"

    def _money_to_words_es(self, value: Any, currency_singular: str = "dolar", currency_plural: str = "dolares") -> str:
        amount = Decimal(str(value if value is not None else 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        integer_part = int(amount)
        cents = int((amount - Decimal(integer_part)) * 100)

        integer_words = self._number_to_spanish_words(integer_part)
        currency_word = currency_singular if integer_part == 1 else currency_plural
        cents_words = self._number_to_spanish_words(cents)
        cent_word = "centavo" if cents == 1 else "centavos"

        text = f"Son: {integer_words} con {cents_words} {cent_word} de {currency_word}"
        return text[:1].upper() + text[1:]

    def _save_workbook_buffer(self, wb) -> io.BytesIO:
        wb.properties.creator = self.REPORT_AUTHOR
        wb.properties.lastModifiedBy = self.REPORT_AUTHOR
        wb.properties.lastPrinted = None
        if not wb.properties.title:
            wb.properties.title = "Reporte"
        for ws in wb.worksheets:
            self._normalize_report_sheet_print_setup(ws)
        output = io.BytesIO()
        archive = ZipFile(output, mode="w", compression=ZIP_DEFLATED, compresslevel=1, allowZip64=True)
        try:
            writer = ExcelWriter(wb, archive)
            writer.save()
        finally:
            archive.close()
        output.seek(0)
        return output

    def _buffer_from_cached_bytes(self, payload: bytes) -> io.BytesIO:
        output = io.BytesIO(payload)
        output.seek(0)
        return output

    def _draw_pdf_watermark(
        self,
        pdf_canvas: canvas.Canvas,
        page_width: float,
        page_height: float,
        watermark_text: Optional[str] = None,
    ) -> None:
        text = str(watermark_text or "").strip().upper()
        if not text:
            return

        pdf_canvas.saveState()
        try:
            pdf_canvas.setFillAlpha(0.12)
        except Exception:
            pass
        pdf_canvas.setFillColor(colors.HexColor("#6B7280"))
        font_size = min(max(page_width / 15, 22), 42)
        pdf_canvas.setFont("Helvetica-Bold", font_size)
        pdf_canvas.translate(page_width / 2, page_height / 2)
        pdf_canvas.rotate(34)
        pdf_canvas.drawCentredString(0, 0, text)
        pdf_canvas.restoreState()

        pdf_canvas.saveState()
        try:
            pdf_canvas.setFillAlpha(0.62)
        except Exception:
            pass
        pdf_canvas.setFillColor(colors.HexColor("#6B7280"))
        pdf_canvas.setFont("Helvetica-Bold", 6.5)
        pdf_canvas.drawCentredString(page_width / 2, 6 * mm, text)
        pdf_canvas.restoreState()

    def _clear_sheet_headers_and_footers(self, ws) -> None:
        for header_name in ("oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter"):
            header = getattr(ws, header_name, None)
            if not header:
                continue
            for part_name in ("left", "center", "right"):
                part = getattr(header, part_name, None)
                if part is None:
                    continue
                part.text = ""
                if hasattr(part, "font"):
                    part.font = None
                if hasattr(part, "size"):
                    part.size = None

    def _normalize_report_sheet_print_setup(self, ws) -> None:
        self._clear_sheet_headers_and_footers(ws)
        used_range = ws.calculate_dimension()
        if used_range and used_range != "A1:A1":
            ws.print_area = used_range
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.scale = None
        if ws.sheet_properties.pageSetUpPr is None:
            ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
        else:
            ws.sheet_properties.pageSetUpPr.fitToPage = True

    def _convert_excel_buffer_to_pdf(
        self,
        xlsx_buffer: io.BytesIO,
        timeout_seconds: int = 120,
        watermark_text: Optional[str] = None,
    ) -> io.BytesIO:
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(xlsx_buffer.getvalue()), data_only=True)
        except Exception as exc:
            raise RuntimeError("No fue posible abrir el archivo Excel para convertirlo a PDF por código.") from exc

        output = io.BytesIO()
        pdf = canvas.Canvas(output, pagesize=A4)
        pdf.setAuthor(self.REPORT_AUTHOR)
        pdf.setCreator(self.REPORT_AUTHOR)
        pdf.setTitle("Reporte")

        rendered_any_sheet = False
        for worksheet in workbook.worksheets:
            rendered_any_sheet = self._render_workbook_sheet_to_pdf(pdf, worksheet, watermark_text) or rendered_any_sheet

        if not rendered_any_sheet:
            raise RuntimeError("El archivo Excel no contiene hojas imprimibles para convertir a PDF.")

        pdf.save()
        output.seek(0)
        return output

    def _worksheet_print_bounds(self, ws) -> Tuple[int, int, int, int]:
        print_area = str(getattr(ws, "print_area", "") or "").strip()
        raw_ranges = [print_area] if print_area else [ws.calculate_dimension()]
        boundaries: List[Tuple[int, int, int, int]] = []

        for raw_range in raw_ranges:
            for chunk in str(raw_range).split(","):
                area = chunk.strip()
                if not area:
                    continue
                if "!" in area:
                    area = area.split("!", 1)[1]
                area = area.replace("$", "")
                try:
                    boundaries.append(range_boundaries(area))
                except Exception:
                    continue

        if not boundaries:
            return 1, 1, max(ws.max_row, 1), max(ws.max_column, 1)

        min_col = min(item[0] for item in boundaries)
        min_row = min(item[1] for item in boundaries)
        max_col = max(item[2] for item in boundaries)
        max_row = max(item[3] for item in boundaries)
        return min_col, min_row, max_col, max_row

    def _excel_color_to_reportlab(self, color_obj, fallback: colors.Color = colors.white) -> colors.Color:
        rgb = getattr(color_obj, "rgb", None)
        if not rgb:
            return fallback
        rgb = str(rgb).strip()
        if len(rgb) == 8:
            rgb = rgb[-6:]
        if len(rgb) != 6:
            return fallback
        try:
            return colors.HexColor(f"#{rgb}")
        except Exception:
            return fallback

    def _excel_border_color(self, side) -> colors.Color:
        color_obj = getattr(side, "color", None)
        return self._excel_color_to_reportlab(color_obj, colors.black)

    def _worksheet_column_width_points(self, ws, column_index: int) -> float:
        column_letter = get_column_letter(column_index)
        column_dimension = ws.column_dimensions.get(column_letter)
        if column_dimension and getattr(column_dimension, "hidden", False):
            return 0.0
        width = getattr(column_dimension, "width", None) or 8.43
        return max((float(width) * 7 + 5) * 0.75, 12.0)

    def _worksheet_row_height_points(self, ws, row_index: int) -> float:
        row_dimension = ws.row_dimensions.get(row_index)
        if row_dimension and getattr(row_dimension, "hidden", False):
            return 0.0
        height = getattr(row_dimension, "height", None) or 15
        return max(float(height), 8.0)

    def _format_excel_display_value(self, cell) -> str:
        value = cell.value
        if value in (None, ""):
            return ""

        if isinstance(value, datetime):
            return value.strftime("%d/%m/%Y %H:%M")
        if isinstance(value, date):
            return value.strftime("%d/%m/%Y")
        if isinstance(value, time):
            return value.strftime("%H:%M")

        number_format = str(getattr(cell, "number_format", "") or "")
        if is_date_format(number_format):
            try:
                return value.strftime("%d/%m/%Y")
            except Exception:
                return str(value)

        if isinstance(value, (int, float, Decimal)):
            decimals = 0
            decimal_match = re.search(r"\.([0#]+)", number_format)
            if decimal_match:
                decimals = len(decimal_match.group(1))
            amount = Decimal(str(value))
            if "%" in number_format:
                amount *= Decimal("100")
                quant = Decimal("1") if decimals <= 0 else Decimal("1." + ("0" * decimals))
                rendered = f"{amount.quantize(quant, rounding=ROUND_HALF_UP):,.{decimals}f}"
                return rendered.replace(",", "_").replace(".", ",").replace("_", ".") + "%"
            quant = Decimal("1") if decimals <= 0 else Decimal("1." + ("0" * decimals))
            rendered = f"{amount.quantize(quant, rounding=ROUND_HALF_UP):,.{decimals}f}"
            rendered = rendered.replace(",", "_").replace(".", ",").replace("_", ".")
            if "$" in number_format:
                return f"${rendered}"
            return rendered

        return str(value)

    def _wrap_pdf_text(self, text: str, font_name: str, font_size: float, max_width: float) -> List[str]:
        normalized = str(text or "").replace("\r", "")
        if not normalized:
            return [""]
        wrapped_lines: List[str] = []
        for source_line in normalized.split("\n"):
            words = source_line.split()
            if not words:
                wrapped_lines.append("")
                continue
            current = words[0]
            for word in words[1:]:
                trial = f"{current} {word}".strip()
                if pdfmetrics.stringWidth(trial, font_name, font_size) <= max_width:
                    current = trial
                else:
                    wrapped_lines.append(current)
                    current = word
            wrapped_lines.append(current)
        return wrapped_lines or [normalized]

    def _excel_color_hex(self, color_obj) -> str:
        rgb = getattr(color_obj, "rgb", None)
        if not rgb:
            return ""
        rgb = str(rgb).strip()
        if len(rgb) == 8:
            rgb = rgb[-6:]
        return rgb.upper() if len(rgb) == 6 else ""

    def _is_dark_excel_fill(self, cell) -> bool:
        fill = getattr(cell, "fill", None)
        if getattr(fill, "fill_type", None) != "solid":
            return False
        rgb = self._excel_color_hex(getattr(fill, "fgColor", None))
        if not rgb:
            return False
        try:
            red = int(rgb[0:2], 16)
            green = int(rgb[2:4], 16)
            blue = int(rgb[4:6], 16)
        except ValueError:
            return False
        return ((red * 0.299) + (green * 0.587) + (blue * 0.114)) < 80

    def _worksheet_pdf_section_header_rows(self, ws, row_indexes: List[int], column_indexes: List[int]) -> List[int]:
        known_headers = {
            "ESPECIFICACIONES TÉCNICAS Y CONTRACTUALES",
            "ESPECIFICACIONES TECNICAS Y CONTRACTUALES",
            "ALCANCE DEL PROYECTO",
            "DIMENSIONAMIENTO",
            "INFORMACIÓN FINANCIERA Y CRONOGRAMA",
            "INFORMACION FINANCIERA Y CRONOGRAMA",
            "OBJETIVOS, RESTRICCIONES Y SUPUESTOS",
        }
        header_rows: List[int] = []
        visible_count = max(len(column_indexes), 1)

        for row_idx in row_indexes:
            row_text = " ".join(
                str(ws.cell(row=row_idx, column=col_idx).value or "").strip()
                for col_idx in column_indexes
                if ws.cell(row=row_idx, column=col_idx).value not in (None, "")
            )
            normalized_text = " ".join(row_text.upper().split())
            dark_cells = sum(1 for col_idx in column_indexes if self._is_dark_excel_fill(ws.cell(row=row_idx, column=col_idx)))
            if normalized_text in known_headers or (normalized_text and dark_cells >= max(2, visible_count // 2)):
                header_rows.append(row_idx)

        return header_rows

    def _worksheet_pdf_keep_groups(
        self,
        ws,
        row_indexes: List[int],
        column_indexes: List[int],
        row_heights: Dict[int, float],
        usable_unscaled_height: float,
    ) -> Dict[int, Tuple[int, float]]:
        header_rows = self._worksheet_pdf_section_header_rows(ws, row_indexes, column_indexes)
        if not header_rows:
            return {}

        row_set = set(row_indexes)
        groups: Dict[int, Tuple[int, float]] = {}
        for index, start_row in enumerate(header_rows):
            next_header = header_rows[index + 1] if index + 1 < len(header_rows) else None
            end_row = (next_header - 1) if next_header else start_row
            group_rows = [row_idx for row_idx in row_indexes if start_row <= row_idx <= end_row]
            if not group_rows:
                continue

            if len(group_rows) <= 1:
                continue
            group_height = sum(row_heights.get(row_idx, 0) for row_idx in group_rows)
            if 0 < group_height <= usable_unscaled_height:
                groups[start_row] = (group_rows[-1], group_height)

        return groups

    def _worksheet_pdf_image_range(self, ws, row_idx: int, col_idx: int) -> CellRange:
        cell = ws.cell(row=row_idx, column=col_idx)
        return self._resolve_placeholder_range(ws, cell)

    def _paginate_worksheet_pdf_rows(
        self,
        row_indexes: List[int],
        row_heights: Dict[int, float],
        keep_groups: Dict[int, Tuple[int, float]],
        usable_unscaled_height: float,
    ) -> List[Tuple[int, int, float]]:
        pages: List[Tuple[int, int, float]] = []
        row_pointer = 0
        while row_pointer < len(row_indexes):
            page_start_pointer = row_pointer
            consumed_height = 0.0
            while row_pointer < len(row_indexes):
                candidate_row = row_indexes[row_pointer]
                candidate_height = row_heights[candidate_row]
                keep_group = keep_groups.get(candidate_row)
                if keep_group is not None:
                    _group_end_row, group_height = keep_group
                    if (
                        row_pointer > page_start_pointer
                        and consumed_height + group_height > usable_unscaled_height
                    ):
                        break
                if row_pointer > page_start_pointer and consumed_height + candidate_height > usable_unscaled_height:
                    break
                consumed_height += candidate_height
                row_pointer += 1

            if page_start_pointer == row_pointer:
                row_pointer = page_start_pointer + 1
                consumed_height = row_heights[row_indexes[page_start_pointer]]
            pages.append((row_indexes[page_start_pointer], row_indexes[row_pointer - 1], consumed_height))
        return pages

    def _compact_worksheet_pdf_residual_tail(
        self,
        ws,
        row_indexes: List[int],
        column_indexes: List[int],
        row_heights: Dict[int, float],
        available_height: float,
        scale: float,
        usable_unscaled_height: float,
        keep_groups: Dict[int, Tuple[int, float]],
        page_ranges: List[Tuple[int, int, float]],
    ) -> Tuple[float, float, Dict[int, Tuple[int, float]], List[Tuple[int, int, float]]]:
        if len(page_ranges) <= 2:
            return scale, usable_unscaled_height, keep_groups, page_ranges

        last_page_height = page_ranges[-1][2]
        if last_page_height <= 0 or last_page_height > (usable_unscaled_height * 0.18):
            return scale, usable_unscaled_height, keep_groups, page_ranges

        merge_start_row = page_ranges[-2][0]
        merge_end_row = page_ranges[-1][1]
        merged_tail_height = sum(
            row_heights[row_idx]
            for row_idx in row_indexes
            if merge_start_row <= row_idx <= merge_end_row
        )
        compact_scale = (available_height / max(merged_tail_height, 1.0)) * 0.995
        min_acceptable_scale = scale * 0.86
        if not (min_acceptable_scale <= compact_scale < scale):
            return scale, usable_unscaled_height, keep_groups, page_ranges

        compact_usable_height = available_height / max(compact_scale, 0.01)
        compact_keep_groups = self._worksheet_pdf_keep_groups(
            ws,
            row_indexes,
            column_indexes,
            row_heights,
            compact_usable_height,
        )
        compact_page_ranges = self._paginate_worksheet_pdf_rows(
            row_indexes,
            row_heights,
            compact_keep_groups,
            compact_usable_height,
        )
        if len(compact_page_ranges) >= len(page_ranges):
            return scale, usable_unscaled_height, keep_groups, page_ranges

        return compact_scale, compact_usable_height, compact_keep_groups, compact_page_ranges

    def _render_workbook_sheet_to_pdf(
        self,
        pdf: canvas.Canvas,
        ws,
        watermark_text: Optional[str] = None,
    ) -> bool:
        min_col, min_row, max_col, max_row = self._worksheet_print_bounds(ws)
        column_indexes = [index for index in range(min_col, max_col + 1) if self._worksheet_column_width_points(ws, index) > 0]
        row_indexes = [index for index in range(min_row, max_row + 1) if self._worksheet_row_height_points(ws, index) > 0]
        if not column_indexes or not row_indexes:
            return False

        column_widths = {index: self._worksheet_column_width_points(ws, index) for index in column_indexes}
        row_heights = {index: self._worksheet_row_height_points(ws, index) for index in row_indexes}
        total_width = sum(column_widths.values())
        if total_width <= 0:
            return False

        sheet_landscape = str(getattr(ws.page_setup, "orientation", "") or "").lower() == "landscape"
        pagesize = landscape(A4) if sheet_landscape or total_width > (A4[0] * 0.9) else A4
        page_width, page_height = pagesize
        margin_x = 10 * mm
        margin_y = 10 * mm
        available_width = page_width - (margin_x * 2)
        available_height = page_height - (margin_y * 2)
        scale = min(available_width / total_width, 1.0)
        usable_unscaled_height = available_height / max(scale, 0.01)

        merged_map: Dict[Tuple[int, int], Tuple[int, int]] = {}
        merged_children: set[Tuple[int, int]] = set()
        for merged_range in ws.merged_cells.ranges:
            if merged_range.max_col < min_col or merged_range.min_col > max_col or merged_range.max_row < min_row or merged_range.min_row > max_row:
                continue
            start = (merged_range.min_row, merged_range.min_col)
            merged_map[start] = (merged_range.max_row, merged_range.max_col)
            for row_idx in range(merged_range.min_row, merged_range.max_row + 1):
                for col_idx in range(merged_range.min_col, merged_range.max_col + 1):
                    if (row_idx, col_idx) != start:
                        merged_children.add((row_idx, col_idx))

        for row_idx in row_indexes:
            for col_idx in column_indexes:
                if (row_idx, col_idx) in merged_children:
                    continue
                cell = ws.cell(row=row_idx, column=col_idx)
                rendered_value = self._format_excel_display_value(cell)
                if not rendered_value:
                    continue
                merge_end = merged_map.get((row_idx, col_idx), (row_idx, col_idx))
                span_end_col = merge_end[1]
                visible_cols = [index for index in column_indexes if col_idx <= index <= span_end_col]
                if not visible_cols:
                    continue
                unscaled_cell_width = sum(column_widths[index] for index in visible_cols)
                rendered_cell_width = unscaled_cell_width * scale
                font = getattr(cell, "font", None)
                font_name = "Helvetica-Bold" if getattr(font, "bold", False) else "Helvetica"
                font_size = min(max((float(getattr(font, "sz", 10) or 10) * scale), 5.0), 18.0)
                padding = 2.5
                text_width = max(rendered_cell_width - (padding * 2), 6)
                line_count = max(len(self._wrap_pdf_text(rendered_value, font_name, font_size, text_width)), 1)
                required_height = ((line_count * font_size * 1.15) + (padding * 2)) / max(scale, 0.01)
                if required_height > row_heights[row_idx]:
                    row_heights[row_idx] = required_height

        image_entries: List[Dict[str, Any]] = []
        for image in getattr(ws, "_images", []) or []:
            anchor = getattr(image, "anchor", None)
            anchor_from = getattr(anchor, "_from", None)
            if anchor_from is None:
                continue
            try:
                raw_image = image._data()
            except Exception:
                continue
            image_entries.append(
                {
                    "row": int(getattr(anchor_from, "row", 0)) + 1,
                    "col": int(getattr(anchor_from, "col", 0)) + 1,
                    "width_pt": max(float(getattr(image, "width", 0) or 0) * 0.75, 12.0),
                    "height_pt": max(float(getattr(image, "height", 0) or 0) * 0.75, 12.0),
                    "payload": raw_image,
                }
            )
            image_entries[-1]["range"] = self._worksheet_pdf_image_range(
                ws,
                image_entries[-1]["row"],
                image_entries[-1]["col"],
            )

        keep_groups = self._worksheet_pdf_keep_groups(ws, row_indexes, column_indexes, row_heights, usable_unscaled_height)
        page_ranges = self._paginate_worksheet_pdf_rows(row_indexes, row_heights, keep_groups, usable_unscaled_height)
        scale, usable_unscaled_height, keep_groups, page_ranges = self._compact_worksheet_pdf_residual_tail(
            ws,
            row_indexes,
            column_indexes,
            row_heights,
            available_height,
            scale,
            usable_unscaled_height,
            keep_groups,
            page_ranges,
        )

        for page_start_row, page_end_row, _page_height in page_ranges:
            pdf.setPageSize(pagesize)
            page_rows = [row_idx for row_idx in row_indexes if page_start_row <= row_idx <= page_end_row]

            x_positions: Dict[int, float] = {}
            current_x = margin_x
            for col_idx in column_indexes:
                x_positions[col_idx] = current_x
                current_x += column_widths[col_idx] * scale

            row_top_positions: Dict[int, float] = {}
            current_y = page_height - margin_y
            for row_idx in page_rows:
                row_top_positions[row_idx] = current_y
                current_y -= row_heights[row_idx] * scale

            for row_idx in page_rows:
                for col_idx in column_indexes:
                    if (row_idx, col_idx) in merged_children:
                        continue

                    cell = ws.cell(row=row_idx, column=col_idx)
                    merge_end = merged_map.get((row_idx, col_idx), (row_idx, col_idx))
                    span_end_row = min(merge_end[0], page_rows[-1])
                    span_end_col = merge_end[1]
                    visible_cols = [index for index in column_indexes if col_idx <= index <= span_end_col]
                    visible_rows = [index for index in page_rows if row_idx <= index <= span_end_row]
                    if not visible_cols or not visible_rows:
                        continue

                    cell_x = x_positions[col_idx]
                    cell_width = sum(column_widths[index] for index in visible_cols) * scale
                    cell_height = sum(row_heights[index] for index in visible_rows) * scale
                    cell_top = row_top_positions[row_idx]
                    cell_y = cell_top - cell_height

                    fill = getattr(getattr(cell, "fill", None), "fill_type", None)
                    fill_color = self._excel_color_to_reportlab(getattr(getattr(cell, "fill", None), "fgColor", None), colors.white)
                    if fill == "solid" and fill_color != colors.white:
                        pdf.setFillColor(fill_color)
                        pdf.rect(cell_x, cell_y, cell_width, cell_height, stroke=0, fill=1)

                    border = getattr(cell, "border", None)
                    if border:
                        side_specs = [
                            ("left", cell_x, cell_y, cell_x, cell_y + cell_height),
                            ("right", cell_x + cell_width, cell_y, cell_x + cell_width, cell_y + cell_height),
                            ("top", cell_x, cell_y + cell_height, cell_x + cell_width, cell_y + cell_height),
                            ("bottom", cell_x, cell_y, cell_x + cell_width, cell_y),
                        ]
                        for side_name, x1, y1, x2, y2 in side_specs:
                            side = getattr(border, side_name, None)
                            if side and getattr(side, "style", None):
                                pdf.setStrokeColor(self._excel_border_color(side))
                                pdf.setLineWidth(0.8 if side.style in {"medium", "thick", "double"} else 0.45)
                                pdf.line(x1, y1, x2, y2)

                    rendered_value = self._format_excel_display_value(cell)
                    if not rendered_value:
                        continue

                    font = getattr(cell, "font", None)
                    font_name = "Helvetica-Bold" if getattr(font, "bold", False) else "Helvetica"
                    font_size = min(max((float(getattr(font, "sz", 10) or 10) * scale), 5.0), 18.0)
                    horizontal = str(getattr(getattr(cell, "alignment", None), "horizontal", "") or "").lower()
                    padding = 2.5
                    text_width = max(cell_width - (padding * 2), 6)
                    lines = self._wrap_pdf_text(rendered_value, font_name, font_size, text_width)
                    line_height = font_size * 1.15
                    total_text_height = max(len(lines), 1) * line_height
                    text_start_y = cell_y + cell_height - padding - font_size
                    if total_text_height + (padding * 2) < cell_height:
                        text_start_y = cell_y + ((cell_height + total_text_height) / 2) - font_size + 1

                    text_color = self._excel_color_to_reportlab(getattr(font, "color", None), colors.black)
                    pdf.setFillColor(text_color)
                    pdf.setFont(font_name, font_size)
                    for line_index, line in enumerate(lines):
                        line_y = text_start_y - (line_index * line_height)
                        if line_y < cell_y + padding:
                            break
                        line_width = pdfmetrics.stringWidth(line, font_name, font_size)
                        if horizontal in {"center", "centercontinuous", "distributed"}:
                            line_x = cell_x + max((cell_width - line_width) / 2, padding)
                        elif horizontal == "right":
                            line_x = cell_x + cell_width - line_width - padding
                        else:
                            line_x = cell_x + padding
                        pdf.drawString(line_x, line_y, line)

            first_page_row = page_rows[0]
            last_page_row = page_rows[-1]
            for image_entry in image_entries:
                row_idx = image_entry["row"]
                col_idx = image_entry["col"]
                if row_idx < first_page_row or row_idx > last_page_row or col_idx not in x_positions:
                    continue
                image_range = image_entry.get("range")
                if image_range is None:
                    image_range = CellRange(f"{ws.cell(row=row_idx, column=col_idx).coordinate}:{ws.cell(row=row_idx, column=col_idx).coordinate}")
                visible_image_cols = [
                    index
                    for index in column_indexes
                    if image_range.min_col <= index <= image_range.max_col and index in x_positions
                ]
                visible_image_rows = [
                    index
                    for index in page_rows
                    if image_range.min_row <= index <= image_range.max_row and index in row_top_positions
                ]
                if not visible_image_cols or not visible_image_rows:
                    continue

                box_x = x_positions[visible_image_cols[0]]
                box_top = row_top_positions[visible_image_rows[0]]
                box_width = sum(column_widths[index] for index in visible_image_cols) * scale
                box_height = sum(row_heights[index] for index in visible_image_rows) * scale
                if box_width <= 0 or box_height <= 0:
                    continue

                cell_x = box_x
                image_y = box_top - box_height
                try:
                    pdf.drawImage(
                        ImageReader(io.BytesIO(image_entry["payload"])),
                        cell_x,
                        image_y,
                        width=box_width,
                        height=box_height,
                        preserveAspectRatio=False,
                        mask="auto",
                    )
                except Exception:
                    continue

            self._draw_pdf_watermark(pdf, page_width, page_height, watermark_text)
            pdf.showPage()

        return True

    def _build_presupuesto_bundle_cache_key(
        self,
        presupuesto: Any,
        empresa_id: int,
        template_id: str,
        use_omniclass: bool,
        apus_map: Dict[int, Any],
        official_source_signature: Tuple[Any, ...] | None = None,
    ) -> Tuple[Any, ...]:
        presupuesto_stamp = getattr(presupuesto, "ultima_modificacion", None) or getattr(presupuesto, "fecha_creacion", None)
        apu_stamps = [
            (apu.id, getattr(apu, "ultima_modificacion", None) or getattr(apu, "fecha_creacion", None))
            for apu in apus_map.values()
        ]
        latest_apu_stamp = max((stamp for _, stamp in apu_stamps if stamp is not None), default=None)
        return (
            "presupuesto_bundle",
            self.REPORT_EXPORT_RENDER_VERSION,
            template_id,
            int(empresa_id),
            int(presupuesto.id),
            bool(use_omniclass),
            presupuesto_stamp.isoformat() if presupuesto_stamp else "",
            len(apus_map),
            latest_apu_stamp.isoformat() if latest_apu_stamp else "",
            "official_source",
            *(official_source_signature or ()),
        )

    def _build_presupuesto_official_source_signature(
        self,
        db: Any,
        presupuesto: Any,
        empresa_id: int,
    ) -> Tuple[Any, ...]:
        try:
            from app.services.project_functional_modification import project_functional_modification_service
        except Exception:
            return ()
        _, official_source = project_functional_modification_service.resolve_project_apu_price_overrides(
            db,
            empresa_id=empresa_id,
            proyecto_id=getattr(presupuesto, "proyecto_id", None),
            base_trabajo_id=getattr(getattr(presupuesto, "proyecto", None), "base_trabajo_id", None),
            revision=getattr(presupuesto, "revision", None),
        )
        if not isinstance(official_source, dict):
            return ()
        return (
            official_source.get("source"),
            official_source.get("origin"),
            int(official_source.get("active_modification_id") or 0),
        )

    def _build_stakeholders_export_signature(
        self,
        db: Any,
        proyecto_id: int,
        empresa_id: int,
    ) -> Tuple[Any, ...]:
        from app.models.proyecto import Proyecto
        from app.models.stakeholder import ProyectoStakeholder, Stakeholder

        proyecto = db.query(Proyecto).filter(
            Proyecto.id == proyecto_id,
            Proyecto.empresa_id == empresa_id,
        ).first()
        if proyecto is None:
            return ("missing-project", int(proyecto_id))

        root_code = proyecto.codigo_root or proyecto.codigo or ""
        root_project = db.query(Proyecto).filter(
            Proyecto.empresa_id == empresa_id,
            Proyecto.codigo_root == root_code,
        ).order_by(
            Proyecto.revision.asc(),
            Proyecto.codigo.asc(),
            Proyecto.id.asc(),
        ).first()
        root_project_id = int(root_project.id if root_project is not None else proyecto.id)

        stakeholders = db.query(Stakeholder).filter(
            Stakeholder.proyecto_codigo_root == root_code,
            Stakeholder.empresa_id == empresa_id,
        ).order_by(Stakeholder.id.asc()).all()
        stakeholder_signature = tuple(
            (
                int(stakeholder.id),
                stakeholder.codigo or "",
                stakeholder.nombre or "",
                stakeholder.apellidos or "",
                (
                    getattr(stakeholder, "ultima_modificacion", None)
                    or getattr(stakeholder, "fecha_creacion", None)
                    or ""
                ).isoformat()
                if (
                    getattr(stakeholder, "ultima_modificacion", None)
                    or getattr(stakeholder, "fecha_creacion", None)
                )
                else "",
            )
            for stakeholder in stakeholders
        )

        assignments = db.query(ProyectoStakeholder).options(
            selectinload(ProyectoStakeholder.rol)
        ).filter(
            ProyectoStakeholder.proyecto_id == root_project_id,
        ).order_by(ProyectoStakeholder.stakeholder_id.asc()).all()
        assignment_signature = tuple(
            (
                int(assignment.stakeholder_id),
                int(assignment.rol_id or 0),
                assignment.rol.nombre if assignment.rol else "",
                (
                    getattr(assignment.rol, "ultima_modificacion", None)
                    or getattr(assignment.rol, "fecha_creacion", None)
                    or ""
                ).isoformat()
                if assignment.rol and (
                    getattr(assignment.rol, "ultima_modificacion", None)
                    or getattr(assignment.rol, "fecha_creacion", None)
                )
                else "",
            )
            for assignment in assignments
        )

        return (
            int(root_project_id),
            root_code,
            stakeholder_signature,
            assignment_signature,
        )

    def _build_report_export_cache_key(
        self,
        db: Any,
        report_type: str,
        entity_ids: List[int],
        empresa_id: int,
        template_id: str = "001",
        variant: Optional[str] = None,
        export_format: str = "xlsx",
        filters: Optional[Dict[str, Any]] = None,
        project_id: Optional[int] = None,
        base_trabajo_id: Optional[int] = None,
        revision: Optional[int] = None,
    ) -> Tuple[Any, ...]:
        report_type = str(report_type or "").lower()
        export_format = str(export_format or "xlsx").lower()
        template_id = template_id or "001"
        entity_ids = [int(entity_id) for entity_id in (entity_ids or [])]
        filters_signature = self.normalize_report_filters_signature(filters)

        if report_type == "apu":
            apus_map = self._get_apus_map(db, entity_ids, empresa_id)
            latest_stamp = max(
                (
                    getattr(apu, "ultima_modificacion", None) or getattr(apu, "fecha_creacion", None)
                    for apu in apus_map.values()
                    if apu is not None
                ),
                default=None,
            )
            return (
                "export",
                self.REPORT_EXPORT_RENDER_VERSION,
                report_type,
                export_format,
                template_id,
                variant or "",
                int(empresa_id),
                tuple(entity_ids),
                latest_stamp.isoformat() if latest_stamp else "",
                "project_context",
                int(project_id or 0),
                int(base_trabajo_id or 0),
                int(revision or 0),
                "filters",
                filters_signature,
            )

        if report_type in {"presupuesto", "vae", "polinomica"}:
            presupuesto = self._get_presupuesto(db, entity_ids[0], empresa_id)
            presupuesto_stamp = getattr(presupuesto, "ultima_modificacion", None) or getattr(presupuesto, "fecha_creacion", None)
            latest_apu_stamp = ""
            apu_count = 0
            if report_type == "presupuesto" and variant == "with_apus":
                apus_map = self._get_presupuesto_apus_map(db, presupuesto, empresa_id)
                latest_apu = max(
                    (
                        getattr(apu, "ultima_modificacion", None) or getattr(apu, "fecha_creacion", None)
                        for apu in apus_map.values()
                        if apu is not None
                    ),
                    default=None,
                )
                latest_apu_stamp = latest_apu.isoformat() if latest_apu else ""
                apu_count = len(apus_map)
            return (
                "export",
                self.REPORT_EXPORT_RENDER_VERSION,
                report_type,
                export_format,
                template_id,
                variant or "",
                int(empresa_id),
                int(presupuesto.id),
                presupuesto_stamp.isoformat() if presupuesto_stamp else "",
                apu_count,
                latest_apu_stamp,
                "official_source",
                *self._build_presupuesto_official_source_signature(db, presupuesto, empresa_id),
                "filters",
                filters_signature,
            )

        if report_type == "edt":
            from app.models.proyecto import Proyecto

            proyecto = db.query(Proyecto).filter(Proyecto.id == entity_ids[0], Proyecto.empresa_id == empresa_id).first()
            proyecto_stamp = None
            if proyecto is not None:
                proyecto_stamp = getattr(proyecto, "ultima_modificacion", None) or getattr(proyecto, "fecha_creacion", None)
            return (
                "export",
                self.REPORT_EXPORT_RENDER_VERSION,
                report_type,
                export_format,
                template_id,
                variant or "",
                int(empresa_id),
                int(entity_ids[0]),
                proyecto_stamp.isoformat() if proyecto_stamp else "",
                "filters",
                filters_signature,
            )

        if report_type == "stakeholders":
            return (
                "export",
                self.REPORT_EXPORT_RENDER_VERSION,
                report_type,
                export_format,
                template_id,
                variant or "",
                int(empresa_id),
                int(entity_ids[0]) if entity_ids else 0,
                self._build_stakeholders_export_signature(db, entity_ids[0], empresa_id) if entity_ids else (),
                "filters",
                filters_signature,
            )

        if report_type == "cronograma_valorado" and entity_ids:
            return (
                "export",
                self.REPORT_EXPORT_RENDER_VERSION,
                report_type,
                export_format,
                template_id,
                variant or "",
                int(empresa_id),
                tuple(entity_ids),
                self._gantt_draft_report_signature(db, entity_ids[0], empresa_id),
                "filters",
                filters_signature,
            )

        return (
            "export",
            self.REPORT_EXPORT_RENDER_VERSION,
            report_type,
            export_format,
            template_id,
            variant or "",
            int(empresa_id),
            tuple(entity_ids),
            "filters",
            filters_signature,
        )

    def normalize_report_filters_signature(self, filters: Optional[Dict[str, Any]]) -> str:
        if not filters:
            return ""
        try:
            return json.dumps(filters, sort_keys=True, default=str, ensure_ascii=True, separators=(",", ":"))
        except TypeError:
            return str(sorted(filters.items()))

    def _summarize_gantt_draft_for_reporting(
        self,
        db: Any,
        presupuesto_id: int,
        empresa_id: int,
    ) -> Dict[str, Any]:
        if db is None:
            return {
                "has_pending": False,
                "pending_count": 0,
                "invalidated_count": 0,
                "adjustment_required_count": 0,
                "total_count": 0,
            }
        try:
            from app.models.cronograma_gantt_control import CronogramaGanttDraft
        except Exception:
            return {
                "has_pending": False,
                "pending_count": 0,
                "invalidated_count": 0,
                "adjustment_required_count": 0,
                "total_count": 0,
            }
        draft = (
            db.query(CronogramaGanttDraft)
            .filter(
                CronogramaGanttDraft.presupuesto_id == int(presupuesto_id),
                CronogramaGanttDraft.empresa_id == int(empresa_id),
                CronogramaGanttDraft.status == "draft",
            )
            .order_by(CronogramaGanttDraft.id.desc())
            .first()
        )
        if draft is None:
            return {
                "has_pending": False,
                "pending_count": 0,
                "invalidated_count": 0,
                "adjustment_required_count": 0,
                "total_count": 0,
            }
        intentions = draft.intentions if isinstance(draft.intentions, list) else []
        pending_count = sum(1 for item in intentions if isinstance(item, dict) and item.get("status") == "pending")
        invalidated_count = sum(1 for item in intentions if isinstance(item, dict) and item.get("status") == "invalidated")
        adjustment_required_count = sum(1 for item in intentions if isinstance(item, dict) and item.get("status") == "adjustment_required")
        total_count = pending_count + invalidated_count + adjustment_required_count
        return {
            "has_pending": total_count > 0,
            "draft_id": getattr(draft, "id", None),
            "draft_version": int(getattr(draft, "version", 0) or 0),
            "pending_count": pending_count,
            "invalidated_count": invalidated_count,
            "adjustment_required_count": adjustment_required_count,
            "total_count": total_count,
            "message": (
                f"Existen {total_count} trabajo(s) pendientes en borrador Gantt. "
                "No se incluyen en los calculos oficiales del reporte."
            ) if total_count > 0 else "",
        }

    def _gantt_draft_report_signature(self, db: Any, presupuesto_id: int, empresa_id: int) -> Tuple[Any, ...]:
        summary = self._summarize_gantt_draft_for_reporting(db, presupuesto_id, empresa_id)
        return (
            int(summary.get("draft_id") or 0),
            int(summary.get("draft_version") or 0),
            int(summary.get("pending_count") or 0),
            int(summary.get("invalidated_count") or 0),
            int(summary.get("adjustment_required_count") or 0),
        )

    def _attach_gantt_draft_report_warning(
        self,
        item: Dict[str, Any],
        draft_summary: Dict[str, Any],
    ) -> Dict[str, Any]:
        if not draft_summary.get("has_pending"):
            return item
        warnings = list(item.get("warnings") or [])
        warnings.append({
            "code": "gantt_draft_pending_not_reported",
            "message": draft_summary.get("message") or "Existen trabajos pendientes en borrador Gantt no incluidos en este reporte.",
            "pending_count": int(draft_summary.get("pending_count") or 0),
            "invalidated_count": int(draft_summary.get("invalidated_count") or 0),
            "adjustment_required_count": int(draft_summary.get("adjustment_required_count") or 0),
        })
        item["warnings"] = warnings
        item["gantt_draft_summary"] = draft_summary
        return item

    def get_cached_report_export(self, cache_key: Tuple[Any, ...]) -> Optional[bytes]:
        return self._get_cached_bytes(self._report_export_cache, cache_key)

    def set_cached_report_export(self, cache_key: Tuple[Any, ...], payload: bytes) -> None:
        self._set_cached_bytes(
            self._report_export_cache,
            cache_key,
            payload,
            self._report_export_cache_limit,
        )

    def _money_to_words_body_es(
        self,
        value: Any,
        currency_singular: str = "dolar",
        currency_plural: str = "dolares",
        uppercase: bool = False,
    ) -> str:
        full_text = self._money_to_words_es(value, currency_singular, currency_plural)
        body = full_text[5:] if full_text.startswith("Son: ") else full_text
        return body.upper() if uppercase else body

    def _money_to_words_excel_usd_upper(self, value: Any) -> str:
        amount = Decimal(str(value if value is not None else 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        integer_part = int(amount)
        cents = int((amount - Decimal(integer_part)) * 100)
        integer_words = self._number_to_spanish_words(integer_part).upper()
        return f"{integer_words}  CON {cents:02d}/100 DÓLARES DE LOS ESTADOS UNIDOS DE AMÉRICA"

    def _normalize_report_description(self, value: Any) -> str:
        normalized = normalize_sentence_case(value)
        return normalized or ""

    def _resolve_project_title(self, proyecto: Any = None, fallback: Any = "") -> str:
        if proyecto and getattr(proyecto, "nombre", None):
            return self._normalize_report_description(proyecto.nombre)
        return self._normalize_report_description(fallback)

    def _format_revision_label(self, revision: Any) -> str:
        try:
            revision_number = int(revision or 0)
        except (TypeError, ValueError):
            revision_number = 0
        return f"R{revision_number:03d}"

    def _format_percent_label(self, value: Any, decimals: int = 2) -> str:
        amount = Decimal(str(value if value is not None else 0)).quantize(
            Decimal("1." + ("0" * max(int(decimals or 0), 0))),
            rounding=ROUND_HALF_UP,
        )
        normalized = format(amount.normalize(), "f")
        if "." in normalized:
            normalized = normalized.rstrip("0").rstrip(".")
        return f"{normalized} %"

    def _format_date_label(self, value: Any) -> str:
        if not value:
            return ""
        if hasattr(value, "strftime"):
            return value.strftime("%d/%m/%Y")
        return str(value or "").strip()

    def _resolve_project_type_label(self, detail: Any) -> str:
        raw_id = getattr(detail, "tipo_proyecto_id", None)
        try:
            type_id = int(raw_id) if raw_id is not None else None
        except (TypeError, ValueError):
            type_id = None
        if type_id is None:
            return ""
        return self._normalize_report_description(self.PROJECT_TYPES_BY_ID.get(type_id, ""))

    def _resolve_project_category_label(self, detail: Any) -> str:
        raw_id = getattr(detail, "categoria_id", None)
        try:
            category_id = int(raw_id) if raw_id is not None else None
        except (TypeError, ValueError):
            category_id = None
        if category_id is None:
            return ""
        return self._normalize_report_description(self.PROJECT_CATEGORIES_BY_ID.get(category_id, ""))

    def _format_area_label(self, value: Any) -> str:
        if value in (None, ""):
            return ""
        try:
            return self._format_fixed(float(value), 2)
        except (TypeError, ValueError):
            return str(value or "").strip()

    def _format_money_label(self, value: Any, decimals: int = 2) -> str:
        return f"${self._round_half_up(value or 0, decimals):,.{decimals}f}"

    def _format_money_plain_label(self, value: Any, decimals: int = 2) -> str:
        currency_value = self._round_half_up(value or 0, decimals)
        return f"{currency_value:,.{decimals}f}"

    def _format_geo_reference(self, detail: Any) -> str:
        if not detail:
            return ""
        lat = getattr(detail, "latitud", None)
        lng = getattr(detail, "longitud", None)
        if lat in (None, "") or lng in (None, ""):
            return ""
        try:
            return f"{float(lat):.6f}, {float(lng):.6f}"
        except (TypeError, ValueError):
            return ""

    def _get_project_map_zoom(self, detail: Any) -> int:
        if not detail:
            return 13
        try:
            zoom = int(getattr(detail, "map_zoom", None) or 13)
        except (TypeError, ValueError):
            zoom = 13
        return max(1, min(18, zoom))

    def _build_project_map_image_url(self, detail: Any, width: int = 900, height: int = 520) -> str:
        if not detail:
            return ""
        lat = getattr(detail, "latitud", None)
        lng = getattr(detail, "longitud", None)
        if lat in (None, "") or lng in (None, ""):
            return ""
        try:
            lat_value = f"{float(lat):.6f}"
            lng_value = f"{float(lng):.6f}"
        except (TypeError, ValueError):
            return ""

        params = urllib.parse.urlencode({
            "center": f"{lat_value},{lng_value}",
            "zoom": self._get_project_map_zoom(detail),
            "size": f"{max(int(width), 320)}x{max(int(height), 180)}",
            "maptype": "mapnik",
            "markers": f"{lat_value},{lng_value},lightblue1",
        })
        return f"https://staticmap.openstreetmap.de/staticmap.php?{params}"

    def _build_project_georef_signature(self, detail: Any) -> str:
        if not detail:
            return ""
        lat = getattr(detail, "latitud", None)
        lng = getattr(detail, "longitud", None)
        if lat in (None, "") or lng in (None, ""):
            return ""
        try:
            lat_value = float(lat)
            lng_value = float(lng)
        except (TypeError, ValueError):
            return ""
        return f"{lat_value:.6f}:{lng_value:.6f}:{self._get_project_map_zoom(detail)}"

    def _build_project_georef_map_relative_path(self, detail: Any, empresa_id: int) -> str:
        codigo_root = str(getattr(detail, "codigo_root", None) or "proyecto").strip() or "proyecto"
        safe_root = "".join(char if char.isalnum() or char in ("-", "_") else "_" for char in codigo_root)[:80] or "proyecto"
        return str(Path("uploads") / "proyectos" / str(empresa_id) / safe_root / "georef" / "georreferenciacion.png").replace("\\", "/")

    def _get_project_cached_georef_map_bytes(self, detail: Any) -> Optional[bytes]:
        if not detail:
            return None
        current_signature = self._build_project_georef_signature(detail)
        if not current_signature:
            return None
        if getattr(detail, "georef_map_status", None) != "ready":
            return None
        if str(getattr(detail, "georef_map_signature", "") or "") != current_signature:
            return None
        return self._fetch_local_upload_bytes(str(getattr(detail, "georef_map_url", "") or ""))

    def refresh_project_georef_map_cache(self, db: Any, detail: Any, empresa_id: int) -> bool:
        if not detail:
            return False
        signature = self._build_project_georef_signature(detail)
        if not signature:
            if getattr(detail, "georef_map_status", None) != "missing":
                detail.georef_map_status = "missing"
                detail.georef_map_error = None
                detail.georef_map_signature = None
                db.add(detail)
                db.commit()
                db.refresh(detail)
            return False

        current_signature = str(getattr(detail, "georef_map_signature", "") or "")
        current_status = str(getattr(detail, "georef_map_status", "") or "")
        cached_payload = self._get_project_cached_georef_map_bytes(detail)
        if current_status == "ready" and current_signature == signature and cached_payload:
            return True

        detail.georef_map_status = "pending"
        detail.georef_map_signature = signature
        detail.georef_map_error = None
        db.add(detail)
        db.commit()
        db.refresh(detail)

        try:
            payload = self._build_project_osm_tile_map_bytes(detail, 900, 520)
            if not payload:
                payload = self._fetch_remote_image_bytes(self._build_project_map_image_url(detail, 900, 520))
            if not payload:
                raise RuntimeError("No se pudo generar imagen georreferenciada desde teselas ni static map.")

            relative_path = self._build_project_georef_map_relative_path(detail, empresa_id)
            file_path = Path(relative_path)
            file_path.parent.mkdir(parents=True, exist_ok=True)
            file_path.write_bytes(payload)

            detail.georef_map_url = f"/{relative_path}"
            detail.georef_map_status = "ready"
            detail.georef_map_signature = signature
            detail.georef_map_generated_at = datetime.now(timezone.utc)
            detail.georef_map_error = None
            db.add(detail)
            db.commit()
            db.refresh(detail)
            return True
        except Exception as exc:
            detail.georef_map_status = "pending"
            detail.georef_map_signature = signature
            detail.georef_map_error = str(exc)[:1000]
            db.add(detail)
            db.commit()
            db.refresh(detail)
            return False

    def _project_geo_to_global_pixel(self, lat: float, lng: float, zoom: int, tile_size: int = 256) -> Tuple[float, float]:
        bounded_lat = max(min(float(lat), 85.05112878), -85.05112878)
        bounded_lng = max(min(float(lng), 180.0), -180.0)
        scale = tile_size * (2 ** int(zoom))
        x = (bounded_lng + 180.0) / 360.0 * scale
        sin_lat = math.sin(math.radians(bounded_lat))
        y = (0.5 - math.log((1 + sin_lat) / (1 - sin_lat)) / (4 * math.pi)) * scale
        return x, y

    def _fetch_osm_tile_bytes(self, zoom: int, tile_x: int, tile_y: int) -> Optional[bytes]:
        max_tile = (2 ** int(zoom)) - 1
        if tile_y < 0 or tile_y > max_tile:
            return None
        wrapped_x = int(tile_x) % (2 ** int(zoom))
        cache_key = ("osm-tile", int(zoom), wrapped_x, int(tile_y))
        cached = self._get_cached_bytes(self._map_image_cache, cache_key)
        if cached is not None:
            return cached

        tile_url = f"https://tile.openstreetmap.org/{int(zoom)}/{wrapped_x}/{int(tile_y)}.png"
        try:
            headers = {
                "User-Agent": f"{self.REPORT_AUTHOR}/1.0 reporting-map-renderer",
                "Accept": "image/png,image/*;q=0.8",
            }
            with httpx.Client(timeout=8.0, follow_redirects=True, headers=headers) as client:
                response = client.get(tile_url)
                response.raise_for_status()
                content_type = str(response.headers.get("content-type", "")).lower()
                if not content_type.startswith("image/") or not response.content:
                    return None
                self._set_cached_bytes(
                    self._map_image_cache,
                    cache_key,
                    response.content,
                    self._map_image_cache_limit * 8,
                )
                return response.content
        except Exception:
            return None

    def _build_project_osm_tile_map_bytes(self, detail: Any, width: int = 900, height: int = 520) -> Optional[bytes]:
        if PILImage is None or ImageDraw is None or not detail:
            return None
        lat = getattr(detail, "latitud", None)
        lng = getattr(detail, "longitud", None)
        if lat in (None, "") or lng in (None, ""):
            return None
        try:
            lat_value = float(lat)
            lng_value = float(lng)
        except (TypeError, ValueError):
            return None

        width = max(int(width or 900), 320)
        height = max(int(height or 520), 180)
        zoom = self._get_project_map_zoom(detail)
        cache_key = ("osm-map", round(lat_value, 6), round(lng_value, 6), zoom, width, height)
        cached = self._get_cached_bytes(self._map_image_cache, cache_key)
        if cached is not None:
            return cached

        tile_size = 256
        center_x, center_y = self._project_geo_to_global_pixel(lat_value, lng_value, zoom, tile_size)
        top_left_x = center_x - (width / 2)
        top_left_y = center_y - (height / 2)
        first_tile_x = math.floor(top_left_x / tile_size)
        first_tile_y = math.floor(top_left_y / tile_size)
        last_tile_x = math.floor((top_left_x + width - 1) / tile_size)
        last_tile_y = math.floor((top_left_y + height - 1) / tile_size)

        canvas_image = PILImage.new("RGB", (width, height), (248, 250, 252))
        fetched_tiles = 0
        for tile_y in range(first_tile_y, last_tile_y + 1):
            for tile_x in range(first_tile_x, last_tile_x + 1):
                tile_payload = self._fetch_osm_tile_bytes(zoom, tile_x, tile_y)
                if not tile_payload:
                    continue
                try:
                    tile_image = PILImage.open(io.BytesIO(tile_payload)).convert("RGB")
                except Exception:
                    continue
                paste_x = int((tile_x * tile_size) - top_left_x)
                paste_y = int((tile_y * tile_size) - top_left_y)
                canvas_image.paste(tile_image, (paste_x, paste_y))
                fetched_tiles += 1

        if fetched_tiles == 0:
            return None

        draw = ImageDraw.Draw(canvas_image)
        marker_x = width // 2
        marker_y = height // 2
        orange = (243, 146, 0)
        white = (255, 255, 255)
        blue = (19, 97, 145)
        radius = max(min(width, height) // 34, 12)
        draw.line((marker_x - (radius * 3), marker_y, marker_x - radius, marker_y), fill=blue, width=3)
        draw.line((marker_x + radius, marker_y, marker_x + (radius * 3), marker_y), fill=blue, width=3)
        draw.line((marker_x, marker_y - (radius * 3), marker_x, marker_y - radius), fill=blue, width=3)
        draw.line((marker_x, marker_y + radius, marker_x, marker_y + (radius * 3)), fill=blue, width=3)
        draw.ellipse((marker_x - radius, marker_y - radius, marker_x + radius, marker_y + radius), fill=white, outline=orange, width=4)
        draw.polygon(
            [
                (marker_x, marker_y + radius + 2),
                (marker_x - max(radius // 2, 6), marker_y + radius + 24),
                (marker_x + max(radius // 2, 6), marker_y + radius + 24),
            ],
            fill=orange,
        )

        buffer = io.BytesIO()
        canvas_image.save(buffer, format="PNG", optimize=True)
        payload = buffer.getvalue()
        self._set_cached_bytes(self._map_image_cache, cache_key, payload, self._map_image_cache_limit)
        return payload

    def _write_png_chunk(self, buffer: io.BytesIO, chunk_type: bytes, data: bytes) -> None:
        buffer.write(struct.pack(">I", len(data)))
        buffer.write(chunk_type)
        buffer.write(data)
        buffer.write(struct.pack(">I", zlib.crc32(chunk_type + data) & 0xFFFFFFFF))

    def _build_png_from_rgb_rows(self, width: int, height: int, pixels: List[List[Tuple[int, int, int]]]) -> bytes:
        raw = bytearray()
        for row in pixels:
            raw.append(0)
            for red, green, blue in row:
                raw.extend((red, green, blue))

        buffer = io.BytesIO()
        buffer.write(b"\x89PNG\r\n\x1a\n")
        self._write_png_chunk(buffer, b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        self._write_png_chunk(buffer, b"IDAT", zlib.compress(bytes(raw), 9))
        self._write_png_chunk(buffer, b"IEND", b"")
        return buffer.getvalue()

    def _build_project_map_snapshot_bytes(self, detail: Any, width: int = 900, height: int = 520) -> Optional[bytes]:
        if not detail:
            return None
        lat = getattr(detail, "latitud", None)
        lng = getattr(detail, "longitud", None)
        if lat in (None, "") or lng in (None, ""):
            return None
        try:
            lat_value = float(lat)
            lng_value = float(lng)
        except (TypeError, ValueError):
            return None

        width = max(int(width or 900), 320)
        height = max(int(height or 520), 180)
        bg = (248, 250, 252)
        grid = (226, 232, 240)
        grid_strong = (203, 213, 225)
        orange = (243, 146, 0)
        blue = (19, 97, 145)
        dark = (30, 41, 59)
        pixels = [[bg for _ in range(width)] for _ in range(height)]

        for x in range(0, width, max(width // 12, 48)):
            color = grid_strong if x % max(width // 4, 1) == 0 else grid
            for y in range(height):
                pixels[y][x] = color
        for y in range(0, height, max(height // 8, 40)):
            color = grid_strong if y % max(height // 4, 1) == 0 else grid
            pixels[y] = [color if x % 3 != 0 else pixels[y][x] for x in range(width)]

        center_x = width // 2
        center_y = height // 2
        for x in range(max(center_x - 95, 0), min(center_x + 96, width)):
            for offset in (-1, 0, 1):
                y = center_y + offset
                if 0 <= y < height:
                    pixels[y][x] = blue
        for y in range(max(center_y - 95, 0), min(center_y + 96, height)):
            for offset in (-1, 0, 1):
                x = center_x + offset
                if 0 <= x < width:
                    pixels[y][x] = blue

        radius = max(min(width, height) // 22, 14)
        radius_sq = radius * radius
        for y in range(max(center_y - radius, 0), min(center_y + radius + 1, height)):
            for x in range(max(center_x - radius, 0), min(center_x + radius + 1, width)):
                dx = x - center_x
                dy = y - center_y
                dist = dx * dx + dy * dy
                if dist <= radius_sq:
                    pixels[y][x] = orange if dist >= (radius - 4) * (radius - 4) else (255, 247, 237)
        for y in range(center_y + radius - 2, min(center_y + radius + 24, height)):
            spread = max(1, (y - center_y - radius + 2) // 2)
            for x in range(max(center_x - spread, 0), min(center_x + spread + 1, width)):
                pixels[y][x] = orange

        header_height = min(58, max(36, height // 9))
        for y in range(header_height):
            for x in range(width):
                if x < width:
                    pixels[y][x] = (255, 255, 255)
        for y in range(header_height, min(header_height + 2, height)):
            for x in range(width):
                pixels[y][x] = grid_strong

        # Firma visual mínima sin dependencia de fuentes: barras codifican lat/lng/zoom.
        values = [
            abs(int(lat_value * 1000)) % 100,
            abs(int(lng_value * 1000)) % 100,
            self._get_project_map_zoom(detail) * 5,
        ]
        x_cursor = 24
        for value in values:
            bar_width = max(20, min(180, value + 24))
            for y in range(18, min(36, header_height)):
                for x in range(x_cursor, min(x_cursor + bar_width, width - 24)):
                    pixels[y][x] = dark if y < 24 else orange
            x_cursor += bar_width + 16

        return self._build_png_from_rgb_rows(width, height, pixels)

    def _build_project_map_snapshot_data_url(self, detail: Any, width: int = 900, height: int = 520) -> str:
        payload = self._build_project_map_snapshot_bytes(detail, width, height)
        if not payload:
            return ""
        return f"data:image/png;base64,{base64.b64encode(payload).decode('ascii')}"

    def _get_project_map_image_data_url(self, detail: Any, width: int = 900, height: int = 520) -> str:
        payload = self._get_project_map_image_bytes(detail, width, height)
        if not payload:
            return ""
        return f"data:image/png;base64,{base64.b64encode(payload).decode('ascii')}"

    def _get_project_map_image_bytes(self, detail: Any, width: int = 900, height: int = 520) -> Optional[bytes]:
        cached_payload = self._get_project_cached_georef_map_bytes(detail)
        if cached_payload:
            return cached_payload
        tile_payload = self._build_project_osm_tile_map_bytes(detail, width, height)
        if tile_payload:
            return tile_payload
        remote_payload = self._fetch_remote_image_bytes(self._build_project_map_image_url(detail, width, height))
        if remote_payload:
            return remote_payload
        return self._build_project_map_snapshot_bytes(detail, width, height)

    def _get_project_referential_image_url(self, detail: Any) -> str:
        if not detail:
            return ""
        return str(getattr(detail, "imagen_referencial_url", None) or "").strip()

    def _parse_jsonish_list(self, value: Any) -> List[str]:
        if value in (None, ""):
            return []
        if isinstance(value, (list, tuple)):
            raw_items = value
        else:
            text = str(value or "").strip()
            if not text:
                return []
            raw_items = None
            if text.startswith("[") and text.endswith("]"):
                try:
                    import json
                    decoded = json.loads(text)
                    if isinstance(decoded, list):
                        raw_items = decoded
                except Exception:
                    raw_items = None
            if raw_items is None:
                raw_items = text.splitlines()
        items: List[str] = []
        for item in raw_items:
            normalized = self._normalize_report_description(item).strip()
            if normalized:
                items.append(normalized)
        return items

    def _format_project_list_field(self, value: Any, fallback: str = "") -> str:
        items = self._parse_jsonish_list(value)
        if items:
            return "\n".join(f"- {self._normalize_report_description(item)}" for item in items)
        return self._normalize_report_description(value) or fallback

    def _fetch_local_upload_bytes(self, image_url: str) -> Optional[bytes]:
        normalized_url = str(image_url or "").strip().replace("\\", "/")
        if not normalized_url:
            return None
        if normalized_url.startswith("/uploads/"):
            relative_path = normalized_url.lstrip("/")
        elif normalized_url.startswith("uploads/"):
            relative_path = normalized_url
        else:
            return None

        path_value = Path(relative_path)
        if path_value.is_absolute():
            candidates = [path_value]
        else:
            backend_root = Path(__file__).resolve().parents[2]
            repo_root = backend_root.parent
            candidates = [
                Path.cwd() / path_value,
                backend_root / path_value,
                repo_root / path_value,
            ]
        for file_path in candidates:
            if not file_path.exists() or not file_path.is_file():
                continue
            try:
                return file_path.read_bytes()
            except Exception:
                return None
        return None

    def _fetch_remote_image_bytes(self, image_url: str) -> Optional[bytes]:
        normalized_url = str(image_url or "").strip()
        if not normalized_url:
            return None
        if normalized_url.startswith("data:image/") and ";base64," in normalized_url:
            try:
                _, payload = normalized_url.split(";base64,", 1)
                return base64.b64decode(payload)
            except Exception:
                return None
        cache_key = ("remote-image", normalized_url)
        cached = self._get_cached_bytes(self._map_image_cache, cache_key)
        if cached is not None:
            return cached

        local_payload = self._fetch_local_upload_bytes(normalized_url)
        if local_payload:
            self._set_cached_bytes(
                self._map_image_cache,
                cache_key,
                local_payload,
                self._map_image_cache_limit,
            )
            return local_payload

        try:
            with httpx.Client(timeout=8.0, follow_redirects=True) as client:
                response = client.get(normalized_url)
                response.raise_for_status()
                content_type = str(response.headers.get("content-type", "")).lower()
                if not content_type.startswith("image/"):
                    return None
                payload = response.content
                if not payload:
                    return None
                self._set_cached_bytes(
                    self._map_image_cache,
                    cache_key,
                    payload,
                    self._map_image_cache_limit,
                )
                return payload
        except Exception:
            return None

    def _find_placeholder_cell(self, ws: Any, placeholder: str) -> Optional[Any]:
        placeholder_value = str(placeholder or "").strip()
        if not placeholder_value:
            return None
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and placeholder_value in cell.value:
                    return cell
        return None

    def _resolve_placeholder_range(self, ws: Any, cell: Any) -> CellRange:
        for merged_range in ws.merged_cells.ranges:
            merged = CellRange(str(merged_range))
            if (
                merged.min_row <= cell.row <= merged.max_row
                and merged.min_col <= cell.column <= merged.max_col
            ):
                return merged
        coord = f"{cell.coordinate}:{cell.coordinate}"
        return CellRange(coord)

    def _estimate_range_pixels(self, ws: Any, target_range: CellRange) -> Tuple[int, int]:
        width_units = 0.0
        for column in range(target_range.min_col, target_range.max_col + 1):
            letter = get_column_letter(column)
            width_units += float(ws.column_dimensions[letter].width or 10)
        height_points = 0.0
        for row in range(target_range.min_row, target_range.max_row + 1):
            height_points += float(ws.row_dimensions[row].height or 15)

        width_px = max(int(width_units * 7), 220)
        height_px = max(int(height_points * 96 / 72), 140)
        return width_px, height_px

    def _fit_image_to_placeholder_canvas(self, image_bytes: bytes, width_px: int, height_px: int, padding_px: int = 0) -> bytes:
        if PILImage is None:
            return image_bytes
        try:
            source = PILImage.open(io.BytesIO(image_bytes)).convert("RGBA")
        except Exception:
            return image_bytes

        canvas_width = max(int(width_px or 0), 1)
        canvas_height = max(int(height_px or 0), 1)
        inner_width = max(canvas_width - (padding_px * 2), 1)
        inner_height = max(canvas_height - (padding_px * 2), 1)
        scale = max(inner_width / max(source.width, 1), inner_height / max(source.height, 1))
        target_width = max(int(source.width * scale), 1)
        target_height = max(int(source.height * scale), 1)

        resample = getattr(PILImage, "Resampling", PILImage).LANCZOS
        resized = source.resize((target_width, target_height), resample)
        canvas_image = PILImage.new("RGBA", (canvas_width, canvas_height), (255, 255, 255, 255))
        offset_x = (canvas_width - target_width) // 2
        offset_y = (canvas_height - target_height) // 2
        canvas_image.alpha_composite(resized, (offset_x, offset_y))

        output = io.BytesIO()
        canvas_image.convert("RGB").save(output, format="PNG", optimize=True)
        return output.getvalue()

    def _insert_image_placeholder(self, ws: Any, placeholder: str, image_bytes: Optional[bytes]) -> None:
        cell = self._find_placeholder_cell(ws, placeholder)
        if cell is None:
            return
        cell.value = ""
        if not image_bytes:
            return
        target_range = self._resolve_placeholder_range(ws, cell)
        width_px, height_px = self._estimate_range_pixels(ws, target_range)
        fitted_image_bytes = self._fit_image_to_placeholder_canvas(image_bytes, width_px, height_px)
        image = OpenPyxlImage(io.BytesIO(fitted_image_bytes))
        image.width = width_px
        image.height = height_px
        ws.add_image(image, cell.coordinate)

    def _build_acta_constitucion_preview(self, db: Any, proyecto: Any, empresa_id: int) -> Dict[str, Any]:
        detail = self._get_project_detail(db, proyecto, empresa_id)
        self.refresh_project_georef_map_cache(db, detail, empresa_id)
        money_decimals = self._get_empresa_format_config(db, empresa_id)["money_decimals"]
        start_date = getattr(detail, "fecha_inicio", None) or getattr(proyecto, "fecha_inicio", None)
        end_date = getattr(detail, "fecha_finalizacion", None) or getattr(proyecto, "fecha_fin_estimada", None)
        plazo = getattr(detail, "plazo_ejecucion", None)
        presupuesto_entidad = float(getattr(proyecto, "presupuesto_estimado", 0) or 0)
        ciudad = getattr(detail, "ciudad", None) or ""
        provincia = getattr(detail, "provincia", None) or ""
        pais = getattr(detail, "pais", None) or ""

        fields = [
            {"label": "Código del proyecto", "value": getattr(proyecto, "codigo", None) or getattr(proyecto, "codigo_root", None) or "-"},
            {"label": "Código referencial", "value": getattr(detail, "cod_referencial", None) or "-"},
            {"label": "Dirección del proyecto", "value": getattr(detail, "direccion", None) or "-"},
            {"label": "Ciudad", "value": ciudad or "-"},
            {"label": "Provincia", "value": provincia or "-"},
            {"label": "País", "value": pais or "-"},
            {"label": "Estado del proyecto", "value": getattr(proyecto, "estado", None) or "-"},
            {"label": "Alcance detallado", "value": getattr(detail, "alcance_detallado", None) or "-"},
            {"label": "Ámbito de contratación", "value": getattr(detail, "ambito_contratacion", None) or "-"},
            {"label": "Tipo de contrato", "value": getattr(detail, "tipo_contrato", None) or "-"},
            {"label": "Tipo de proyecto", "value": self._resolve_project_type_label(detail) or "-"},
            {"label": "Categoría del proyecto", "value": self._resolve_project_category_label(detail) or "-"},
            {"label": "Tipo de construcción", "value": getattr(detail, "tipo_construccion", None) or "-"},
            {"label": "Nivel de complejidad", "value": getattr(detail, "nivel_complejidad", None) or "-"},
            {"label": "Tipo de medición", "value": getattr(detail, "tipo_medicion", None) or "-"},
            {"label": "Área de terreno (m²)", "value": self._format_area_label(getattr(detail, "area_terreno", None)) or "-"},
            {"label": "Área de construcción (m²)", "value": self._format_area_label(getattr(detail, "area_construccion", None)) or "-"},
            {"label": "Niveles", "value": getattr(detail, "num_niveles", None) or "-"},
            {"label": "Presupuesto referencial", "value": self._format_money_label(getattr(detail, "presupuesto_referencial", None) or getattr(proyecto, "presupuesto_estimado", None), money_decimals)},
            {"label": "Fuente de financiamiento", "value": getattr(detail, "fuente_financiamiento", None) or "-"},
            {"label": "Cliente", "value": getattr(detail, "cliente_contratante_preliminar", None) or "-"},
            {"label": "Fecha de inicio", "value": self._format_date_label(start_date) or "-"},
            {"label": "Plazo", "value": f"{int(plazo)} días" if plazo not in (None, "") else "-"},
            {"label": "Fecha de finalización", "value": self._format_date_label(end_date) or "-"},
            {"label": "Georreferenciación", "value": self._format_geo_reference(detail) or "-"},
            {"label": "Objetivos clave", "value": self._format_project_list_field(getattr(detail, "objetivos_clave", None), "-")},
            {"label": "Restricciones", "value": self._format_project_list_field(getattr(detail, "restricciones_conocidas", None), "-")},
            {"label": "Supuestos", "value": self._format_project_list_field(getattr(detail, "supuestos_iniciales", None), "-")},
        ]

        return {
            "id": getattr(proyecto, "id", 0) or 0,
            "codigo": getattr(proyecto, "codigo", None) or getattr(proyecto, "codigo_root", None) or "-",
            "descripcion": self._resolve_project_title(proyecto, getattr(proyecto, "descripcion", None)),
            "categoria_base": "Acta de constitución del proyecto",
            "unidad": getattr(detail, "moneda", None) or getattr(proyecto, "moneda", None) or "USD",
            "costo_directo": 0,
            "costo_indirecto": 0,
            "precio_total": presupuesto_entidad,
            "summary_direct_kind": "text",
            "summary_indirect_kind": "text",
            "summary_total_kind": "money",
            "summary_cards": [
                {"label": "Fecha de inicio", "value": self._format_date_label(start_date) or "Sin fecha", "kind": "text"},
                {"label": "Plazo", "value": f"{int(plazo)} días" if plazo not in (None, "") else "Sin plazo", "kind": "text"},
                {"label": "Presupuesto entidad", "value": presupuesto_entidad, "kind": "money", "tone": "total"},
            ],
            "image_referencial_url": self._get_project_referential_image_url(detail),
            "map_image_url": self._get_project_map_image_data_url(detail),
            "fields": fields,
            "lineas": [],
            "metadata_hint": f"Documento base · {self._format_money_label(presupuesto_entidad, money_decimals)}",
        }

    def _build_stakeholders_preview(self, db: Any, proyecto: Any, empresa_id: int) -> Dict[str, Any]:
        from app.repositories.stakeholder import stakeholder_repo

        detail = self._get_project_detail(db, proyecto, empresa_id)
        stakeholders = stakeholder_repo.get_by_project_root(
            db,
            getattr(proyecto, "codigo_root", None) or getattr(proyecto, "codigo", None) or "",
            empresa_id,
            proyecto_id=getattr(proyecto, "id", None),
        )
        role_map = self._build_project_stakeholder_role_map(db, proyecto, empresa_id)

        role_assigned_count = sum(1 for stk in stakeholders if role_map.get(getattr(stk, "id", None)) or getattr(stk, "rol_nombre", None))
        lineas = []
        for stk in stakeholders:
            role_name = role_map.get(getattr(stk, "id", None)) or getattr(stk, "rol_nombre", None) or "Sin rol asignado"
            location_parts = [
                getattr(stk, "ciudad", None),
                getattr(stk, "canton", None),
                getattr(stk, "provincia", None),
                getattr(stk, "pais", None),
            ]
            ubicacion = " / ".join([part.strip() for part in location_parts if isinstance(part, str) and part.strip()])
            contacto = " · ".join([part.strip() for part in [getattr(stk, "movil", None), getattr(stk, "email", None)] if isinstance(part, str) and part.strip()])
            nombre_completo = self._join_report_parts([
                self._normalize_report_description(getattr(stk, "nombre", None)),
                self._normalize_report_description(getattr(stk, "apellidos", None)),
            ], " ")
            profession = self._join_report_parts([
                getattr(stk, "profesion", None),
                getattr(stk, "institucion", None),
            ])
            lineas.append({
                "codigo": getattr(stk, "codigo", None) or "-",
                "nombre_completo": nombre_completo or "-",
                "profesion": self._normalize_report_description(profession) or "-",
                "contacto": contacto or "-",
                "ubicacion": ubicacion or "-",
                "rol": role_name,
            })

        fields = [
            {"label": "Código del proyecto", "value": getattr(proyecto, "codigo", None) or getattr(proyecto, "codigo_root", None) or "-"},
            {"label": "Código referencial", "value": getattr(detail, "cod_referencial", None) or "-"},
            {"label": "Ciudad base", "value": getattr(detail, "ciudad", None) or "-"},
            {"label": "Fecha del reporte", "value": self._format_date_label(getattr(proyecto, "fecha_creacion", None)) or "-"},
        ]

        return {
            "id": getattr(proyecto, "id", 0) or 0,
            "codigo": getattr(proyecto, "codigo", None) or getattr(proyecto, "codigo_root", None) or "-",
            "descripcion": self._resolve_project_title(proyecto, getattr(proyecto, "descripcion", None)),
            "categoria_base": "Equipo del proyecto (Stakeholders)",
            "unidad": "Stakeholders",
            "costo_directo": 0,
            "costo_indirecto": 0,
            "precio_total": len(stakeholders),
            "summary_direct_kind": "text",
            "summary_indirect_kind": "text",
            "summary_total_kind": "text",
            "summary_cards": [
                {"label": "Stakeholders", "value": str(len(stakeholders)), "kind": "text"},
                {"label": "Con rol", "value": str(role_assigned_count), "kind": "text"},
                {"label": "Sin rol asignado", "value": str(max(len(stakeholders) - role_assigned_count, 0)), "kind": "text", "tone": "total"},
            ],
            "fields": fields,
            "lineas": lineas,
            "table_columns": [
                {"key": "codigo", "label": "Código", "width_weight": 0.10},
                {"key": "nombre_completo", "label": "Nombre completo", "width_weight": 0.20},
                {"key": "profesion", "label": "Profesión / Institución", "width_weight": 0.19},
                {"key": "contacto", "label": "Contacto", "width_weight": 0.22},
                {"key": "ubicacion", "label": "Ubicación", "width_weight": 0.17},
                {"key": "rol", "label": "Rol", "width_weight": 0.12},
            ],
            "metadata_hint": "Directorio común del proyecto compartido entre revisiones",
        }

    def _build_project_stakeholder_role_map(self, db: Any, proyecto: Any, empresa_id: int) -> Dict[int, str]:
        from app.models.edo import EdoNode
        from app.models.edt import EdtNode
        from app.models.proyecto import Proyecto
        from app.models.stakeholder import ProyectoStakeholder

        if not proyecto:
            return {}

        root_code = getattr(proyecto, "codigo_root", None) or getattr(proyecto, "codigo", None) or ""
        if not root_code:
            return {}

        root_project = db.query(Proyecto).filter(
            Proyecto.empresa_id == empresa_id,
            Proyecto.codigo_root == root_code,
        ).order_by(
            Proyecto.revision.asc(),
            Proyecto.codigo.asc(),
            Proyecto.id.asc(),
        ).first()
        root_project_id = int(getattr(root_project, "id", None) or getattr(proyecto, "id", 0) or 0)

        role_map: Dict[int, str] = {}
        assignments = db.query(ProyectoStakeholder).options(
            selectinload(ProyectoStakeholder.rol)
        ).filter(
            ProyectoStakeholder.proyecto_id == root_project_id,
        ).all()
        for assignment in assignments:
            stakeholder_id = int(getattr(assignment, "stakeholder_id", 0) or 0)
            role_name = str(getattr(getattr(assignment, "rol", None), "nombre", "") or "").strip()
            if stakeholder_id and role_name:
                role_map[stakeholder_id] = role_name

        for model in (EdoNode, EdtNode):
            nodes = db.query(model).options(selectinload(model.rol)).filter(
                model.proyecto_id == getattr(proyecto, "id", None),
                model.empresa_id == empresa_id,
                model.stakeholder_id.isnot(None),
                model.rol_id.isnot(None),
            ).order_by(model.codigo.asc()).all()
            for node in nodes:
                stakeholder_id = int(getattr(node, "stakeholder_id", 0) or 0)
                role_name = str(getattr(getattr(node, "rol", None), "nombre", "") or "").strip()
                if stakeholder_id and role_name:
                    role_map[stakeholder_id] = role_name

        return role_map

    def generate_acta_constitucion_report(self, db: Any, proyecto_id: int, empresa_id: int, template_id: str = "001") -> io.BytesIO:
        from app.models.proyecto import Proyecto

        proyecto = (
            db.query(Proyecto)
            .filter(Proyecto.id == proyecto_id, Proyecto.empresa_id == empresa_id)
            .first()
        )
        if not proyecto:
            raise ValueError("Proyecto no encontrado")

        detail = self._get_project_detail(db, proyecto, empresa_id)
        filename = "001 - Acta de Constitucion - General.xlsx"
        money_decimals = self._get_empresa_format_config(db, empresa_id)["money_decimals"]
        presupuesto_referencial = getattr(detail, "presupuesto_referencial", None)
        if presupuesto_referencial in (None, ""):
            presupuesto_referencial = getattr(proyecto, "presupuesto_estimado", None)
        fecha_inicio = getattr(detail, "fecha_inicio", None) or getattr(proyecto, "fecha_inicio", None)
        fecha_fin = getattr(detail, "fecha_finalizacion", None) or getattr(proyecto, "fecha_fin_estimada", None)
        data = {
            "#REVISION": self._format_revision_label(getattr(proyecto, "revision", None)),
            "#PRO_TITULO": self._resolve_project_title(proyecto, getattr(proyecto, "descripcion", None)),
            "#CODIGOPROYECTO": getattr(proyecto, "codigo", None) or getattr(proyecto, "codigo_root", None) or "",
            "#CODIGOREFERENCIAL": getattr(detail, "cod_referencial", None) or "",
            "#GEOREFERENCIACION": self._format_geo_reference(detail),
            "#IMAGENPROYECTO": "",
            "#DIRECCIONPROYECTO": getattr(detail, "direccion", None) or "",
            "#TIPOPROYECTO": self._resolve_project_type_label(detail),
            "#CATEGORIA": self._resolve_project_category_label(detail),
            "#ESTADOPROYECTO": getattr(proyecto, "estado", None) or "",
            "#MODALIDADCONTRATACION": getattr(detail, "ambito_contratacion", None) or "",
            "#TIPOCONTRATO": getattr(detail, "tipo_contrato", None) or "",
            "#TIPOCONSTRUCCION": getattr(detail, "tipo_construccion", None) or "",
            "#AMBITO": getattr(detail, "ambito_contratacion", None) or "",
            "#COMPLEJIDAD": getattr(detail, "nivel_complejidad", None) or "",
            "#ALCANCEDETALLADO": self._normalize_report_description(getattr(detail, "alcance_detallado", None)),
            "#TIPOMEDICION": getattr(detail, "tipo_medicion", None) or "",
            "#AREATERRENO": getattr(detail, "area_terreno", None) or 0,
            "#AREACONSTRUCCION": getattr(detail, "area_construccion", None) or 0,
            "#NUMNIVELES": getattr(detail, "num_niveles", None) or "",
            "#PRESUPUESTOREFERENCIAL": self._format_money_plain_label(presupuesto_referencial, money_decimals),
            "#MONEDA": getattr(detail, "moneda", None) or getattr(proyecto, "moneda", None) or "USD",
            "#FUENTEFINANCIAMIENTO": getattr(detail, "fuente_financiamiento", None) or "",
            "#CLIENTE": getattr(detail, "cliente_contratante_preliminar", None) or "",
            "#FECHAINICIO": self._format_date_label(fecha_inicio),
            "#PLAZODIAS": f"{int(detail.plazo_ejecucion)} días" if detail and getattr(detail, "plazo_ejecucion", None) not in (None, "") else "",
            "#FECHAFIN": self._format_date_label(fecha_fin),
            "#PLAZOFIJO": f"{int(detail.plazo_ejecucion)} días" if detail and getattr(detail, "plazo_ejecucion", None) not in (None, "") else "",
            "#OBJETIVOSCLAVE": self._format_project_list_field(getattr(detail, "objetivos_clave", None)),
            "#RESTRICCIONES": self._format_project_list_field(getattr(detail, "restricciones_conocidas", None)),
            "#SUPUESTOS": self._format_project_list_field(getattr(detail, "supuestos_iniciales", None)),
            "#FECHACREACION": self._format_date_label(getattr(proyecto, "fecha_creacion", None)),
        }

        format_config = self._get_empresa_format_config(db, empresa_id)
        formatted_data = self._format_excel_data(
            data,
            {
                "#AREATERRENO": "calc",
                "#AREACONSTRUCCION": "calc",
            },
            format_config["money_decimals"],
            format_config["calc_decimals"],
        )
        template_path = os.path.join(self.templates_dir, filename)
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Plantilla no encontrada: {template_path}")
        wb = openpyxl.load_workbook(template_path)
        ws = wb.active
        self.refresh_project_georef_map_cache(db, detail, empresa_id)
        map_image_bytes = self._get_project_map_image_bytes(detail)
        referential_image_bytes = self._fetch_remote_image_bytes(self._get_project_referential_image_url(detail))
        self._insert_image_placeholder(ws, "#GEOREFERENCIACION", map_image_bytes)
        self._insert_image_placeholder(ws, "#IMAGENPROYECTO", referential_image_bytes or map_image_bytes)
        self._apply_template_to_sheet(ws, formatted_data, [], None)
        self._apply_omniclass_visibility(ws, False)
        self._clear_unresolved_placeholders(ws)
        return self._save_workbook_buffer(wb)

    def _join_report_parts(self, parts: List[Any], separator: str = " · ") -> str:
        normalized_parts: List[str] = []
        for part in parts:
            value = str(part or "").strip()
            if value and value not in normalized_parts:
                normalized_parts.append(value)
        return separator.join(normalized_parts)

    def _estimate_stakeholder_row_height(self, ws, row_index: int, min_col: int = 1, max_col: int = 12) -> float:
        max_lines = 1
        for column_index in range(min_col, max_col + 1):
            cell = ws.cell(row=row_index, column=column_index)
            value = str(cell.value or "").strip()
            if not value:
                continue
            column_letter = get_column_letter(column_index)
            column_width = float(getattr(ws.column_dimensions[column_letter], "width", None) or 10)
            approx_chars_per_line = max(int(column_width * 0.9), 8)
            line_count = 0
            for source_line in value.splitlines() or [value]:
                line_count += max(1, (len(source_line) + approx_chars_per_line - 1) // approx_chars_per_line)
            max_lines = max(max_lines, line_count)
        return min(max(24.0, (max_lines * 11.5) + 8.0), 96.0)

    def _prepare_stakeholders_report_sheet(self, ws, table_start_row: int, table_rows: int) -> None:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = 9
        ws.sheet_view.showGridLines = False
        ws.print_title_rows = "1:9"
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.35
        ws.page_margins.bottom = 0.35

        for column_index in range(1, 13):
            column_letter = get_column_letter(column_index)
            ws.column_dimensions[column_letter].hidden = False

        stakeholder_widths = {
            "A": 6,
            "B": 11,
            "C": 18,
            "D": 18,
            "E": 22,
            "F": 24,
            "G": 16,
            "H": 16,
            "I": 13,
            "J": 17,
            "K": 28,
            "L": 20,
        }
        for column_letter, width in stakeholder_widths.items():
            ws.column_dimensions[column_letter].width = width

        table_end_row = table_start_row + max(int(table_rows or 0), 1) - 1
        for header_row in range(1, table_start_row):
            for column_index in range(1, 13):
                cell = ws.cell(row=header_row, column=column_index)
                if cell.value not in (None, ""):
                    cell.alignment = Alignment(
                        horizontal=getattr(cell.alignment, "horizontal", None) or "center",
                        vertical=getattr(cell.alignment, "vertical", None) or "center",
                        wrap_text=True,
                    )
        for row_index in range(table_start_row, table_end_row + 1):
            ws.row_dimensions[row_index].height = self._estimate_stakeholder_row_height(ws, row_index)
            for column_index in range(1, 13):
                cell = ws.cell(row=row_index, column=column_index)
                cell.alignment = Alignment(
                    horizontal=getattr(cell.alignment, "horizontal", None) or "center",
                    vertical="top",
                    wrap_text=True,
                )
                if cell.font and cell.font.sz and float(cell.font.sz) > 9:
                    source_font = cell.font
                    cell.font = Font(
                        name=source_font.name,
                        sz=9,
                        b=source_font.b,
                        i=source_font.i,
                        vertAlign=source_font.vertAlign,
                        underline=source_font.underline,
                        strike=source_font.strike,
                        color=copy(source_font.color),
                        charset=source_font.charset,
                        family=source_font.family,
                        scheme=source_font.scheme,
                        outline=source_font.outline,
                        shadow=source_font.shadow,
                        condense=source_font.condense,
                        extend=source_font.extend,
                    )

        last_content_row = table_end_row
        for row_index in range(1, ws.max_row + 1):
            if any(ws.cell(row=row_index, column=column_index).value not in (None, "") for column_index in range(1, 13)):
                last_content_row = max(last_content_row, row_index)
        if last_content_row < ws.max_row:
            ws.delete_rows(last_content_row + 1, ws.max_row - last_content_row)

        signature_row = self._normalize_stakeholders_signature_block(ws)
        if signature_row is not None:
            last_content_row = max(last_content_row, signature_row + 1)
        ws.print_area = f"A1:L{last_content_row}"

    def generate_stakeholders_report(self, db: Any, proyecto_id: int, empresa_id: int, template_id: str = "001") -> io.BytesIO:
        from app.models.proyecto import Proyecto
        from app.repositories.stakeholder import stakeholder_repo

        proyecto = (
            db.query(Proyecto)
            .filter(Proyecto.id == proyecto_id, Proyecto.empresa_id == empresa_id)
            .first()
        )
        if not proyecto:
            raise ValueError("Proyecto no encontrado")

        detail = self._get_project_detail(db, proyecto, empresa_id)
        stakeholders = stakeholder_repo.get_by_project_root(
            db,
            getattr(proyecto, "codigo_root", None) or getattr(proyecto, "codigo", None) or "",
            empresa_id,
            proyecto_id=proyecto_id,
        )
        role_map = self._build_project_stakeholder_role_map(db, proyecto, empresa_id)

        filename = "001 - Stakeholders.xlsx"
        data = {
            "#REVISION": self._format_revision_label(getattr(proyecto, "revision", None)),
            "#PRO_TITULO": self._resolve_project_title(proyecto, getattr(proyecto, "descripcion", None)),
            "#CODIGOPROYECTO": getattr(proyecto, "codigo", None) or getattr(proyecto, "codigo_root", None) or "",
            "#CODIGOREFERENCIAL": getattr(detail, "cod_referencial", None) or "",
            "#CIUDAD": getattr(detail, "ciudad", None) or "",
            "#PRO_FECHA": self._format_date_label(getattr(proyecto, "fecha_creacion", None)),
        }
        table_data = []
        for index, stk in enumerate(stakeholders, start=1):
            role_name = role_map.get(getattr(stk, "id", None)) or getattr(stk, "rol_nombre", None) or "Sin rol asignado"
            location = self._join_report_parts([
                getattr(stk, "ciudad", None),
                getattr(stk, "canton", None),
            ], " / ")
            profession = self._join_report_parts([
                getattr(stk, "profesion", None),
                getattr(stk, "institucion", None),
            ])
            table_data.append({
                "#ITEM": index,
                "#CODIGO_STKR": getattr(stk, "codigo", None) or "",
                "#NOMBRES": self._normalize_report_description(getattr(stk, "nombre", None)),
                "#APELLIDOS": self._normalize_report_description(getattr(stk, "apellidos", None)),
                "#PROFESION": self._normalize_report_description(profession),
                "#DIRECCIONSTKR": self._normalize_report_description(getattr(stk, "direccion_detalle", None)),
                "#CIUDADSTKR": location,
                "#PROVINCIASTKR": getattr(stk, "provincia", None) or "",
                "#PAISSTKR": getattr(stk, "pais", None) or "",
                "#TELEFONOSTKR": getattr(stk, "movil", None) or "",
                "#CORREOSTKR": getattr(stk, "email", None) or "",
                "#ROL": role_name,
            })

        template_path = os.path.join(self.templates_dir, filename)
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Plantilla no encontrada: {template_path}")
        wb = openpyxl.load_workbook(template_path)
        ws = wb.active
        self._apply_template_to_sheet(ws, data, table_data, "#ITEM")
        self._apply_omniclass_visibility(ws, False)
        self._prepare_stakeholders_report_sheet(ws, 10, len(table_data))
        return self._save_workbook_buffer(wb)

    def _get_project_reference_code(self, db: Any, presupuesto: Any) -> str:
        from app.models.proyecto_detalle import ProyectoDetalle

        proyecto = getattr(presupuesto, "proyecto", None)
        if not proyecto:
            return ""

        codigo_root = getattr(proyecto, "codigo_root", None)
        empresa_id = getattr(proyecto, "empresa_id", None) or getattr(presupuesto, "empresa_id", None)
        if not codigo_root or not empresa_id:
            return ""

        detail = (
            db.query(ProyectoDetalle)
            .filter(
                ProyectoDetalle.codigo_root == codigo_root,
                ProyectoDetalle.empresa_id == empresa_id,
            )
            .first()
        )
        return getattr(detail, "cod_referencial", "") or ""

    def _get_project_detail(self, db: Any, proyecto: Any, empresa_id: Optional[int] = None) -> Any:
        from app.models.proyecto_detalle import ProyectoDetalle

        if not proyecto:
            return None

        codigo_root = getattr(proyecto, "codigo_root", None)
        resolved_empresa_id = empresa_id or getattr(proyecto, "empresa_id", None)
        if not codigo_root or not resolved_empresa_id:
            return None

        return (
            db.query(ProyectoDetalle)
            .filter(
                ProyectoDetalle.codigo_root == codigo_root,
                ProyectoDetalle.empresa_id == resolved_empresa_id,
            )
            .first()
        )

    def _get_project_location(self, db: Any, proyecto: Any, empresa_id: Optional[int] = None) -> str:
        detail = self._get_project_detail(db, proyecto, empresa_id)
        if not detail:
            return ""

        parts = [
            getattr(detail, "provincia", None),
            getattr(detail, "canton", None),
            getattr(detail, "ciudad", None),
            getattr(detail, "direccion", None),
        ]
        return " · ".join([part.strip() for part in parts if isinstance(part, str) and part.strip()])

    def _get_project_location_city(self, db: Any, proyecto: Any, empresa_id: Optional[int] = None) -> str:
        detail = self._get_project_detail(db, proyecto, empresa_id)
        if not detail:
            return ""
        return (
            str(getattr(detail, "ciudad", "") or "").strip()
            or str(getattr(detail, "canton", "") or "").strip()
            or str(getattr(detail, "provincia", "") or "").strip()
        )

    def _get_budget_line_edt_code(self, line: Any) -> str:
        edt_node = getattr(line, "edt_node", None)
        return (
            getattr(edt_node, "codigo", None)
            or getattr(edt_node, "codigo_completo", None)
            or getattr(line, "codigo_item", None)
            or ""
        )

    def _get_budget_line_visible_item_code(self, line: Any) -> str:
        if getattr(line, "apu_id", None) is None:
            return ""
        return str(getattr(line, "codigo_item", None) or "").strip()

    def _get_budget_line_visible_edt_code(self, line: Any) -> str:
        edt_node = getattr(line, "edt_node", None)
        edt_code = (
            getattr(edt_node, "codigo", None)
            or getattr(edt_node, "codigo_completo", None)
            or ""
        )
        edt_code = str(edt_code or "").strip()
        if edt_code:
            return edt_code

        structural_code = str(getattr(line, "codigo_item", None) or "").strip()
        if getattr(line, "apu_id", None) is None:
            return structural_code

        if "." in structural_code:
            return structural_code.rsplit(".", 1)[0].strip()
        return ""

    def _split_hierarchical_code(self, value: Any) -> Tuple[Any, ...]:
        raw = str(value or "").strip()
        if not raw:
            return (float("inf"),)
        parts: List[Any] = []
        for token in raw.split("."):
            token = token.strip()
            if token.isdigit():
                parts.append(int(token))
            elif token:
                parts.append(token.lower())
        return tuple(parts) if parts else (float("inf"),)

    def _get_budget_line_sort_key(self, line: Any) -> Tuple[Any, ...]:
        edt_code = self._get_budget_line_edt_code(line)
        item_code = getattr(line, "codigo_item", None) or edt_code
        is_structural = getattr(line, "apu_id", None) is None
        return (
            self._split_hierarchical_code(edt_code),
            0 if is_structural else 1,
            self._split_hierarchical_code(item_code),
            getattr(line, "id", 0) or 0,
        )

    def _sort_budget_lines_for_report(self, lines: List[Any]) -> List[Any]:
        return sorted(lines, key=self._get_budget_line_sort_key)

    def _resolve_budget_line_report_item(self, line: Any) -> Any:
        orden = getattr(line, "orden", None)
        if orden is None:
            return ""
        try:
            return int(orden) + 1
        except (TypeError, ValueError):
            return orden

    def _get_budget_line_display_code(self, line: Any) -> str:
        return (getattr(line, "codigo_item", None) or self._get_budget_line_edt_code(line) or "").strip()

    def _get_budget_structural_total(self, line: Any, detail_lines: List[Any]) -> float:
        code = self._get_budget_line_display_code(line)
        if not code:
            return 0.0
        prefix = f"{code}."
        total = 0.0
        for detail in detail_lines:
            if getattr(detail, "apu_id", None) is None:
                continue
            detail_code = self._get_budget_line_display_code(detail)
            if detail_code == code or detail_code.startswith(prefix):
                total += float(getattr(detail, "precio_total", 0) or 0)
        return total

    def _build_presupuesto_table_data(self, presupuesto: Any, template_id: str = "001") -> List[Dict[str, Any]]:
        ordered_lines = self._sort_budget_lines_for_report(list(presupuesto.detalle or []))
        sercop_counter = 0
        table_data: List[Dict[str, Any]] = []

        for line in ordered_lines:
            is_structural = getattr(line, "apu_id", None) is None
            display_code = self._get_budget_line_display_code(line)
            item_visible = self._get_budget_line_visible_item_code(line)
            edt_code_visible = self._get_budget_line_visible_edt_code(line)
            apu = getattr(line, "apu", None)
            apu_code = (getattr(apu, "codigo", None) or "").strip() if apu else ""
            descripcion = str(getattr(line, "descripcion", "") or "").strip()
            structural_total = self._get_budget_structural_total(line, ordered_lines) if is_structural else 0.0

            if template_id == "002":
                item_value = ""
                if not is_structural:
                    sercop_counter += 1
                    item_value = sercop_counter
                table_data.append({
                    "__STRUCTURAL": is_structural,
                    "ITEM": item_value,
                    "ITEM_VISIBLE": item_visible,
                    "CODIGO_EDT": edt_code_visible,
                    "DESCRIPCION": descripcion.upper() if is_structural else self._normalize_report_description(descripcion),
                    "UNIDAD": "" if is_structural else (line.unidad or ""),
                    "CANTIDAD": "" if is_structural else float(line.cantidad) if line.cantidad else 0.0,
                    "PUNITARIO": "" if is_structural else float(line.precio_unitario) if line.precio_unitario else 0.0,
                    "SUBTOTAL": float(structural_total if is_structural else (line.precio_total or 0)),
                })
                continue

            table_data.append({
                "__STRUCTURAL": is_structural,
                "ITEM": display_code,
                "ITEM_VISIBLE": item_visible,
                "CODIGO_EDT": edt_code_visible,
                "CODIGO_APU": "" if is_structural else apu_code,
                "DESCRIPCION": descripcion.upper() if is_structural else self._normalize_report_description(descripcion),
                "UNIDAD": "" if is_structural else (line.unidad or ""),
                "CANTIDAD": "" if is_structural else float(line.cantidad) if line.cantidad else 0.0,
                "PUNITARIO": "" if is_structural else float(line.precio_unitario) if line.precio_unitario else 0.0,
                "SUBTOTAL": float(structural_total if is_structural else (line.precio_total or 0)),
            })

        return table_data

    def _apply_presupuesto_row_typography(self, ws, table_rows: List[Dict[str, Any]], template_id: str) -> None:
        data_start_row = 7 if template_id == "002" else 9
        for offset, row_data in enumerate(table_rows):
            row_idx = data_start_row + offset
            is_structural = bool(row_data.get("__STRUCTURAL"))
            for cell in ws[row_idx]:
                if cell.value in (None, ""):
                    continue
                current_font = copy(cell.font)
                current_font.bold = is_structural
                cell.font = current_font

    def _template_requires_category_engine(self, ws) -> bool:
        category_tokens = (
            "#CODCAT1",
            "#CANTIDAD1",
            "#PRECIO1",
            "#RENDIMIENTO1",
            "#SUBTOTAL1",
            "#PORCENT1",
            "#TOTAL1",
            "#TOTPORCENT1",
            "TOTPORCENT1",
            "#OMNICLASS_COD1",
            "#OMNICLASS_TIT1",
            "@OMNICLASS_TIT1",
        )
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(token in cell.value for token in category_tokens):
                    return True
        return False

    def _get_empresa_format_config(self, db: Any, empresa_id: int) -> Dict[str, int]:
        from app.models.empresa import Empresa

        empresa = db.query(Empresa).filter(Empresa.id == empresa_id).first()
        return {
            "money_decimals": int(getattr(empresa, "decimales_moneda", 2) or 2),
            "calc_decimals": int(getattr(empresa, "decimales_calculos", 4) or 4),
            "use_omniclass": bool(getattr(empresa, "use_omniclass", True)),
        }

    def _resolve_use_omniclass(self, format_config: Dict[str, Any], template_id: Optional[str] = None) -> bool:
        if str(template_id or "") == "002":
            return False
        return bool(format_config.get("use_omniclass", True))

    def _round_half_up(self, value: Any, decimals: int) -> float:
        return float(round_decimal(Decimal(str(value if value is not None else 0)), decimals))

    def _format_fixed(self, value: Any, decimals: int) -> str:
        return f"{self._round_half_up(value, decimals):.{decimals}f}"

    def _format_for_excel(self, value: Any, kind: str, money_decimals: int, calc_decimals: int) -> Any:
        if value is None or value == "":
            return ""
        if kind == "integer":
            return str(int(self._round_half_up(value, 0)))
        if kind == "money":
            return self._format_fixed(value, money_decimals)
        if kind == "percent":
            return self._format_fixed(value, money_decimals)
        if kind == "calc":
            return self._format_fixed(value, calc_decimals)
        return value

    def _format_excel_data(self, data: Dict[str, Any], field_kinds: Dict[str, str], money_decimals: int, calc_decimals: int) -> Dict[str, Any]:
        return {
            key: self._format_for_excel(value, field_kinds.get(key, "raw"), money_decimals, calc_decimals)
            for key, value in data.items()
        }

    def _format_excel_rows(self, rows: List[Dict[str, Any]], field_kinds: Dict[str, str], money_decimals: int, calc_decimals: int) -> List[Dict[str, Any]]:
        formatted_rows = []
        for row in rows:
            formatted_rows.append({
                key: self._format_for_excel(value, field_kinds.get(key, "raw"), money_decimals, calc_decimals)
                for key, value in row.items()
            })
        return formatted_rows

    def _replace_placeholder_token(self, text: str, placeholder: str, value: Any) -> str:
        if not isinstance(text, str) or placeholder not in text:
            return text
        safe_value = "" if value is None else str(value)
        pattern = re.compile(re.escape(placeholder) + r"(?![A-Za-z0-9_])")
        return pattern.sub(safe_value, text)

    def _copy_cell_style(self, source_cell, target_cell) -> None:
        if source_cell.has_style:
            target_cell.font = copy(source_cell.font)
            target_cell.border = copy(source_cell.border)
            target_cell.fill = copy(source_cell.fill)
            target_cell.number_format = copy(source_cell.number_format)
            target_cell.protection = copy(source_cell.protection)
            target_cell.alignment = copy(source_cell.alignment)

    def _ensure_apu_omniclass_columns(self, ws) -> None:
        ws.column_dimensions[get_column_letter(9)].width = max(ws.column_dimensions[get_column_letter(9)].width or 0, 18)
        ws.column_dimensions[get_column_letter(10)].width = max(ws.column_dimensions[get_column_letter(10)].width or 0, 34)

        for category_idx in range(1, 5):
            item_row_idx = self._find_row_with_any_placeholder(ws, [f"#CODCAT{category_idx}"])
            if not item_row_idx:
                continue

            header_row_idx = item_row_idx - 1
            header_code_cell = ws.cell(row=header_row_idx, column=9)
            header_title_cell = ws.cell(row=header_row_idx, column=10)
            item_code_cell = ws.cell(row=item_row_idx, column=9)
            item_title_cell = ws.cell(row=item_row_idx, column=10)

            if not header_code_cell.value:
                header_code_cell.value = "OmniClass Cod"
            if not header_title_cell.value:
                header_title_cell.value = "Clasificación"
            if not item_code_cell.value:
                item_code_cell.value = f"#OMNICLASS_COD{category_idx}"
            if not item_title_cell.value:
                item_title_cell.value = f"#OMNICLASS_TIT{category_idx}"

            style_header_source_row = 10 if ws.cell(row=10, column=9).has_style or ws.cell(row=10, column=10).has_style else header_row_idx
            style_item_source_row = 11 if ws.cell(row=11, column=9).has_style or ws.cell(row=11, column=10).has_style else item_row_idx
            style_header_source_col = 9 if ws.cell(row=style_header_source_row, column=9).has_style else 8
            style_title_source_col = 10 if ws.cell(row=style_header_source_row, column=10).has_style else 8
            style_item_source_col = 9 if ws.cell(row=style_item_source_row, column=9).has_style else 8
            style_item_title_source_col = 10 if ws.cell(row=style_item_source_row, column=10).has_style else 8

            self._copy_cell_style(ws.cell(row=style_header_source_row, column=style_header_source_col), header_code_cell)
            self._copy_cell_style(ws.cell(row=style_header_source_row, column=style_title_source_col), header_title_cell)
            self._copy_cell_style(ws.cell(row=style_item_source_row, column=style_item_source_col), item_code_cell)
            self._copy_cell_style(ws.cell(row=style_item_source_row, column=style_item_title_source_col), item_title_cell)

    def _ensure_apu_omniclass_columns_with_metadata(self, ws, metadata: Dict[str, Any]) -> None:
        ws.column_dimensions[get_column_letter(9)].width = max(ws.column_dimensions[get_column_letter(9)].width or 0, 18)
        ws.column_dimensions[get_column_letter(10)].width = max(ws.column_dimensions[get_column_letter(10)].width or 0, 34)

        for category_idx, category_meta in (metadata.get("categories") or {}).items():
            item_row_idx = category_meta.get("item_row")
            if not item_row_idx:
                continue

            header_row_idx = item_row_idx - 1
            header_code_cell = ws.cell(row=header_row_idx, column=9)
            header_title_cell = ws.cell(row=header_row_idx, column=10)
            item_code_cell = ws.cell(row=item_row_idx, column=9)
            item_title_cell = ws.cell(row=item_row_idx, column=10)

            if not header_code_cell.value:
                header_code_cell.value = "OmniClass Cod"
            if not header_title_cell.value:
                header_title_cell.value = "Clasificación"
            if not item_code_cell.value:
                item_code_cell.value = f"#OMNICLASS_COD{category_idx}"
            if not item_title_cell.value:
                item_title_cell.value = f"#OMNICLASS_TIT{category_idx}"

            header_style_source = ws.cell(row=header_row_idx, column=8 if ws.max_column >= 8 else header_title_cell.column)
            item_style_source = ws.cell(row=item_row_idx, column=8 if ws.max_column >= 8 else item_title_cell.column)
            self._copy_cell_style(header_style_source, header_code_cell)
            self._copy_cell_style(header_style_source, header_title_cell)
            self._copy_cell_style(item_style_source, item_code_cell)
            self._copy_cell_style(item_style_source, item_title_cell)

    def _apply_omniclass_visibility(self, ws, use_omniclass: bool) -> None:
        if use_omniclass:
            return

        columns_to_hide = set()
        token_markers = ("OMNICLASS", "OMNICLASS_COD", "OMNICLASS_TIT")
        literal_markers = {"OmniClass Cod", "Clasificación"}

        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                value = cell.value.strip()
                upper_value = value.upper()
                if any(token in upper_value for token in token_markers) or value in literal_markers:
                    columns_to_hide.add(cell.column)

        for column_index in columns_to_hide:
            ws.column_dimensions[get_column_letter(column_index)].hidden = True

    def _shift_merged_ranges_below_row(self, ws, start_row: int, delta_rows: int) -> None:
        if not delta_rows:
            return
        ranges_to_shift = [CellRange(str(rng)) for rng in ws.merged_cells.ranges if rng.min_row >= start_row]
        for merged_range in ranges_to_shift:
            try:
                ws.unmerge_cells(str(merged_range))
            except KeyError:
                if str(merged_range) in ws.merged_cells:
                    ws.merged_cells.remove(str(merged_range))
        for merged_range in ranges_to_shift:
            ws.merge_cells(
                start_row=merged_range.min_row + delta_rows,
                start_column=merged_range.min_col,
                end_row=merged_range.max_row + delta_rows,
                end_column=merged_range.max_col,
            )

    def _apply_template_to_sheet(self, ws, data: Dict[str, Any], table_data: List[Dict[str, Any]] = None, table_row_marker: str = None) -> None:
        if table_data and table_row_marker:
            marker_row_idx = -1
            marker_cols = {}

            for row in ws.iter_rows():
                found = False
                for cell in row:
                    if isinstance(cell.value, str) and table_row_marker in cell.value:
                        marker_row_idx = cell.row
                        found = True
                        break
                if found:
                    for cell in ws[marker_row_idx]:
                        if isinstance(cell.value, str) and cell.value.startswith('#'):
                            marker_cols[cell.value] = cell.column
                    break

            if marker_row_idx != -1:
                footer_placeholders = [
                    key for key in data.keys()
                    if isinstance(key, str) and key.startswith('#') and key not in marker_cols
                ]
                footer_row_idx = self._find_next_placeholder_row(ws, marker_row_idx + 1, footer_placeholders)
                if footer_row_idx and footer_row_idx > marker_row_idx + 1:
                    self._remove_rows_between(ws, marker_row_idx + 1, footer_row_idx - 1)

                num_rows = len(table_data)
                if num_rows > 1:
                    self._shift_merged_ranges_below_row(ws, marker_row_idx + 1, num_rows - 1)
                    ws.insert_rows(marker_row_idx + 1, num_rows - 1)
                    source_row = ws[marker_row_idx]
                    for i in range(1, num_rows):
                        target_row_idx = marker_row_idx + i
                        for cell in source_row:
                            new_cell = ws.cell(row=target_row_idx, column=cell.column)
                            if cell.has_style:
                                from copy import copy
                                new_cell.font = copy(cell.font)
                                new_cell.border = copy(cell.border)
                                new_cell.fill = copy(cell.fill)
                                new_cell.number_format = copy(cell.number_format)
                                new_cell.protection = copy(cell.protection)
                                new_cell.alignment = copy(cell.alignment)
                        if marker_row_idx in ws.row_dimensions:
                            ws.row_dimensions[target_row_idx].height = ws.row_dimensions[marker_row_idx].height
                        self._clone_single_row_merged_ranges(ws, marker_row_idx, target_row_idx)

                for i, row_data in enumerate(table_data):
                    current_row_idx = marker_row_idx + i
                    for placeholder, col_idx in marker_cols.items():
                        clean_key = placeholder.replace('#', '')
                        import re
                        clean_key = re.sub(r'\d+$', '', clean_key)
                        if placeholder in row_data:
                            val = row_data.get(placeholder)
                        elif clean_key in row_data:
                            val = row_data.get(clean_key)
                        elif placeholder.replace('#', '') in row_data:
                            val = row_data.get(placeholder.replace('#', ''))
                        else:
                            val = ""
                        ws.cell(row=current_row_idx, column=col_idx).value = val

        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    for key, val in data.items():
                        if cell.value == key:
                            cell.value = val
                            break
                        cell.value = self._replace_placeholder_token(cell.value, key, val)

    def _copy_sheet_row(self, ws, source_row_idx: int, target_row_idx: int) -> None:
        source_row = ws[source_row_idx]
        for cell in source_row:
            new_cell = ws.cell(row=target_row_idx, column=cell.column)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                try:
                    new_cell.value = Translator(cell.value, origin=cell.coordinate).translate_formula(new_cell.coordinate)
                except Exception:
                    new_cell.value = cell.value
            else:
                new_cell.value = cell.value
            if cell.has_style:
                new_cell.font = copy(cell.font)
                new_cell.border = copy(cell.border)
                new_cell.fill = copy(cell.fill)
                new_cell.number_format = copy(cell.number_format)
                new_cell.protection = copy(cell.protection)
                new_cell.alignment = copy(cell.alignment)
        if source_row_idx in ws.row_dimensions:
            ws.row_dimensions[target_row_idx].height = ws.row_dimensions[source_row_idx].height
        self._clone_single_row_merged_ranges(ws, source_row_idx, target_row_idx)

    def _capture_row_blueprint(self, ws, row_idx: int) -> Dict[str, Any]:
        cells = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cells.append({
                "col": col_idx,
                "value": cell.value,
                "font": copy(cell.font) if cell.has_style else None,
                "border": copy(cell.border) if cell.has_style else None,
                "fill": copy(cell.fill) if cell.has_style else None,
                "number_format": copy(cell.number_format) if cell.has_style else None,
                "protection": copy(cell.protection) if cell.has_style else None,
                "alignment": copy(cell.alignment) if cell.has_style else None,
            })
        merges = [
            (merged_range.min_col, merged_range.max_col)
            for merged_range in ws.merged_cells.ranges
            if merged_range.min_row == row_idx and merged_range.max_row == row_idx
        ]
        return {
            "height": ws.row_dimensions[row_idx].height,
            "cells": cells,
            "merges": merges,
        }

    def _clear_single_row_merges(self, ws, row_idx: int) -> None:
        ranges_to_remove = [
            str(merged_range)
            for merged_range in ws.merged_cells.ranges
            if merged_range.min_row == row_idx and merged_range.max_row == row_idx
        ]
        for range_ref in ranges_to_remove:
            try:
                ws.unmerge_cells(range_ref)
            except KeyError:
                if range_ref in ws.merged_cells:
                    ws.merged_cells.remove(range_ref)

    def _restore_row_blueprint(self, ws, row_idx: int, blueprint: Dict[str, Any]) -> None:
        self._clear_single_row_merges(ws, row_idx)
        for cell_meta in blueprint.get("cells", []):
            self._set_sheet_value_safe(ws, row_idx, cell_meta["col"], cell_meta.get("value"))
            target_cell = ws.cell(row=row_idx, column=cell_meta["col"])
            if isinstance(target_cell, MergedCell):
                continue
            if cell_meta.get("font") is not None:
                target_cell.font = copy(cell_meta["font"])
                target_cell.border = copy(cell_meta["border"])
                target_cell.fill = copy(cell_meta["fill"])
                target_cell.number_format = copy(cell_meta["number_format"])
                target_cell.protection = copy(cell_meta["protection"])
                target_cell.alignment = copy(cell_meta["alignment"])
        ws.row_dimensions[row_idx].height = blueprint.get("height")
        for min_col, max_col in blueprint.get("merges", []):
            ws.merge_cells(start_row=row_idx, start_column=min_col, end_row=row_idx, end_column=max_col)

    def _clone_single_row_merged_ranges(self, ws, source_row_idx: int, target_row_idx: int) -> None:
        source_ranges = [
            CellRange(str(merged_range))
            for merged_range in ws.merged_cells.ranges
            if merged_range.min_row == source_row_idx and merged_range.max_row == source_row_idx
        ]
        for merged_range in source_ranges:
            target_range = CellRange(
                min_col=merged_range.min_col,
                max_col=merged_range.max_col,
                min_row=target_row_idx,
                max_row=target_row_idx,
            )
            target_range_str = str(target_range)
            if target_range_str in ws.merged_cells:
                continue
            ws.merge_cells(target_range_str)

    def _replace_placeholders_in_row(self, ws, row_idx: int, replacements: Dict[str, Any]) -> None:
        for cell in ws[row_idx]:
            if not isinstance(cell.value, str):
                continue
            direct_placeholder = next((placeholder for placeholder in replacements if cell.value == placeholder), None)
            if direct_placeholder is not None:
                cell.value = replacements.get(direct_placeholder, "")
                continue
            new_value = cell.value
            for placeholder, value in replacements.items():
                new_value = self._replace_placeholder_token(new_value, placeholder, value)
            cell.value = new_value

    def _find_row_with_any_placeholder(self, ws, placeholders: List[str]) -> Optional[int]:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(placeholder in cell.value for placeholder in placeholders):
                    return cell.row
        return None

    def _find_next_placeholder_row(self, ws, start_row: int, placeholders: List[str]) -> Optional[int]:
        if not placeholders:
            return None
        for row in ws.iter_rows(min_row=max(1, start_row)):
            for cell in row:
                if isinstance(cell.value, str) and any(placeholder in cell.value for placeholder in placeholders):
                    return cell.row
        return None

    def _remove_rows_between(self, ws, start_row: int, end_row: int) -> int:
        if end_row < start_row:
            return 0
        count = end_row - start_row + 1
        deleted_end = start_row + count - 1
        ranges_to_process = [CellRange(str(rng)) for rng in ws.merged_cells.ranges if rng.max_row >= start_row]
        for merged_range in ranges_to_process:
            try:
                ws.unmerge_cells(str(merged_range))
            except KeyError:
                if str(merged_range) in ws.merged_cells:
                    ws.merged_cells.remove(str(merged_range))
        ws.delete_rows(start_row, count)
        for merged_range in ranges_to_process:
            if merged_range.min_row > deleted_end:
                ws.merge_cells(
                    start_row=merged_range.min_row - count,
                    start_column=merged_range.min_col,
                    end_row=merged_range.max_row - count,
                    end_column=merged_range.max_col,
                )
            elif merged_range.min_row < start_row and merged_range.max_row > deleted_end:
                ws.merge_cells(
                    start_row=merged_range.min_row,
                    start_column=merged_range.min_col,
                    end_row=merged_range.max_row - count,
                    end_column=merged_range.max_col,
                )
        return count

    def _clear_unresolved_placeholders(self, ws) -> None:
        token_pattern = re.compile(r"[#@][A-Za-z0-9_%]+")
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and ("#" in cell.value or "@" in cell.value):
                    cell.value = token_pattern.sub("", cell.value).strip()

    def _resolve_apu_line_category(self, linea: Any) -> int:
        if getattr(linea, "apu_hijo_id", None):
            return 2
        recurso = getattr(linea, "recurso", None)
        codigo = getattr(recurso, "codigo", "") if recurso else ""
        try:
            return int(str(codigo).strip().split("-")[0])
        except (ValueError, IndexError):
            return 1

    def _build_apu_category_payload(self, apu: Any, money_decimals: int = 2, calc_decimals: int = 4, use_omniclass: bool = True) -> Dict[int, Dict[str, Any]]:
        costo_directo = float(apu.costo_directo or 0.0)
        categories: Dict[int, Dict[str, Any]] = {
            1: {"items": [], "total": 0.0, "percent": 0.0},
            2: {"items": [], "total": 0.0, "percent": 0.0},
            3: {"items": [], "total": 0.0, "percent": 0.0},
            4: {"items": [], "total": 0.0, "percent": 0.0},
        }

        for linea in sorted(apu.lineas, key=lambda item: ((item.orden or 0), item.id or 0)):
            recurso = getattr(linea, "recurso", None)
            apu_hijo = getattr(linea, "apu_hijo", None)
            category_id = self._resolve_apu_line_category(linea)
            if category_id not in categories:
                continue

            codigo = ""
            descripcion = ""
            unidad = ""
            precio = 0.0
            omniclass_cod = ""
            omniclass_tit = ""

            if recurso:
                codigo = self._resolve_apu_report_code(recurso)
                descripcion = self._normalize_report_description(recurso.descripcion)
                unidad = recurso.unidad.descripcion if (recurso.unidad and hasattr(recurso.unidad, 'descripcion')) else ""
                precio = float(recurso.precio or 0.0)
                omniclass_cod = recurso.omniclass_codigo or "" if use_omniclass else ""
                omniclass_tit = recurso.omniclass_titulo or "" if use_omniclass else ""
            elif apu_hijo:
                codigo = self._resolve_apu_report_code(apu_hijo)
                descripcion = self._normalize_report_description(apu_hijo.descripcion)
                unidad = apu_hijo.unidad or ""
                precio = float(apu_hijo.precio_unitario_total or 0.0)
                omniclass_cod = apu_hijo.omniclass_codigo or "" if use_omniclass else ""
                omniclass_tit = apu_hijo.omniclass_titulo or "" if use_omniclass else ""

            cantidad = float(linea.cantidad or 0.0)
            rendimiento = float(linea.rendimiento or 0.0) if linea.rendimiento is not None else ""
            subtotal = float(linea.subtotal or 0.0)
            percent = (subtotal / costo_directo) if costo_directo > 0 else 0.0

            categories[category_id]["items"].append({
                "CODCAT": codigo,
                "DESCRIPCION": descripcion,
                "UNIDAD": unidad,
                "CANTIDAD": self._round_half_up(cantidad, calc_decimals),
                "PRECIO": self._round_half_up(precio, money_decimals),
                "RENDIMIENTO": "" if rendimiento == "" else self._round_half_up(rendimiento, calc_decimals),
                "SUBTOTAL": self._round_half_up(subtotal, money_decimals),
                "PORCENT": self._round_half_up(percent, 8),
                "OMNICLASS_COD": omniclass_cod,
                "OMNICLASS_TIT": omniclass_tit,
            })
            categories[category_id]["total"] += subtotal

        for category_id, category in categories.items():
            category["total"] = self._round_half_up(category["total"], money_decimals)
            category["percent"] = self._round_half_up((float(category["total"]) / costo_directo) if costo_directo > 0 else 0.0, 8)

        return categories

    def _apply_apu_category_template(self, ws, apu: Any, data: Dict[str, Any], money_decimals: int = 2, calc_decimals: int = 4, use_omniclass: bool = True) -> None:
        self._apply_apu_category_template_with_metadata(
            ws,
            apu,
            data,
            self._extract_apu_template_metadata(ws),
            money_decimals,
            calc_decimals,
            use_omniclass,
        )

    def _apply_apu_category_template_with_metadata(
        self,
        ws,
        apu: Any,
        data: Dict[str, Any],
        metadata: Dict[str, Any],
        money_decimals: int = 2,
        calc_decimals: int = 4,
        use_omniclass: bool = True,
    ) -> None:
        layout_kind = metadata.get("layout_kind") or self._resolve_apu_layout_kind(ws)
        if use_omniclass:
            self._ensure_apu_omniclass_columns_with_metadata(ws, metadata)
        for cell_meta in metadata.get("data_cells", []):
            row_idx = cell_meta["row"]
            col_idx = cell_meta["col"]
            raw_value = cell_meta["value"]
            new_value = raw_value
            for key, val in data.items():
                new_value = self._replace_placeholder_token(new_value, key, val)
            self._set_sheet_value(ws, row_idx, col_idx, new_value)

        categories = self._build_apu_category_payload(apu, money_decimals, calc_decimals, use_omniclass)
        row_offset = 0
        summary_targets: List[Tuple[int, Dict[str, Any], int, str]] = []

        category_order = self._resolve_apu_template_category_order(metadata)
        for category_idx in category_order:
            category_meta = metadata.get("categories", {}).get(category_idx)
            if not category_meta:
                continue

            item_row_idx = category_meta["item_row"] + row_offset
            summary_row_idx = (category_meta["summary_row"] + row_offset) if category_meta.get("summary_row") else None
            gap_rows = max(0, int(category_meta.get("gap_rows") or 0))

            if summary_row_idx and gap_rows:
                removed_rows = self._remove_rows_between(ws, item_row_idx + 1, summary_row_idx - 1)
                summary_row_idx -= removed_rows
                row_offset -= removed_rows

            category = categories.get(category_idx) or {"items": [], "total": 0.0, "percent": 0.0}
            items = category["items"] or [{}]
            extra_rows = max(0, len(items) - 1)

            if extra_rows:
                self._shift_merged_ranges_below_row(ws, item_row_idx + 1, extra_rows)
                ws.insert_rows(item_row_idx + 1, extra_rows)
                item_blueprint = category_meta.get("item_blueprint")
                for offset in range(1, extra_rows + 1):
                    if item_blueprint:
                        self._restore_row_blueprint(ws, item_row_idx + offset, item_blueprint)
                    else:
                        self._copy_sheet_row(ws, item_row_idx, item_row_idx + offset)

            adjusted_summary_row_idx = summary_row_idx + extra_rows if summary_row_idx and summary_row_idx > item_row_idx else summary_row_idx
            self._restore_apu_category_from_metadata(
                ws,
                category_meta,
                item_row_idx,
                adjusted_summary_row_idx,
            )
            if not adjusted_summary_row_idx and layout_kind == "sercop":
                adjusted_summary_row_idx = self._find_row_containing_text(
                    ws,
                    self._resolve_apu_sercop_summary_label(category_idx),
                )

            for offset, item in enumerate(items):
                current_row_idx = item_row_idx + offset
                replacements = {
                    f"#CODCAT{category_idx}": item.get("CODCAT", ""),
                    f"#DESCRIPCION{category_idx}": item.get("DESCRIPCION", ""),
                    f"#UNIDAD{category_idx}": item.get("UNIDAD", ""),
                    f"#CANTIDAD{category_idx}": item.get("CANTIDAD", ""),
                    f"#PRECIO{category_idx}": item.get("PRECIO", ""),
                    f"#RENDIMIENTO{category_idx}": item.get("RENDIMIENTO", ""),
                    f"#SUBTOTAL{category_idx}": item.get("SUBTOTAL", ""),
                    f"#PORCENT{category_idx}": item.get("PORCENT", ""),
                    f"#OMNICLASS_COD{category_idx}": item.get("OMNICLASS_COD", ""),
                    f"#OMNICLASS_TIT{category_idx}": item.get("OMNICLASS_TIT", ""),
                    f"@OMNICLASS_TIT{category_idx}": item.get("OMNICLASS_TIT", ""),
                }
                self._replace_placeholders_in_row(ws, current_row_idx, replacements)
                self._ensure_apu_item_row_layout(ws, current_row_idx, item.get("DESCRIPCION", ""), layout_kind)

            if adjusted_summary_row_idx:
                self._replace_placeholders_in_row(ws, adjusted_summary_row_idx, {
                    f"#TOTAL{category_idx}": category["total"],
                    f"#TOTPORCENT{category_idx}": category["percent"],
                    f"TOTPORCENT{category_idx}": category["percent"],
                })
                summary_targets.append((category_idx, category, adjusted_summary_row_idx, layout_kind))
            row_offset += extra_rows

        self._sync_apu_footer_values(ws, data, metadata=metadata, row_offset=row_offset, money_decimals=money_decimals)
        for category_idx, category, summary_row_idx, summary_layout_kind in summary_targets:
            self._apply_apu_category_summary_value(
                ws,
                category_idx,
                category,
                summary_row_idx,
                summary_layout_kind,
            )
        if layout_kind == "sercop":
            self._apply_apu_sercop_section_title_fills(ws)
        self._clear_unresolved_placeholders(ws)

    def _resolve_apu_template_category_order(self, metadata: Dict[str, Any]) -> List[int]:
        categories = metadata.get("categories") or {}
        return [
            category_idx
            for category_idx, _category_meta in sorted(
                categories.items(),
                key=lambda item: (int((item[1] or {}).get("item_row") or 0), int(item[0])),
            )
        ]

    def _find_row_containing_text(self, ws, needle: str) -> Optional[int]:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and needle in cell.value:
                    return cell.row
        return None

    def _set_sheet_value_safe(self, ws, row: int, column: int, value: Any) -> None:
        target_range = None
        for merged_range in ws.merged_cells.ranges:
            if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= column <= merged_range.max_col:
                target_range = merged_range
                break
        if target_range:
            try:
                ws.cell(target_range.min_row, target_range.min_col).value = value
            except AttributeError:
                ws._cells[(target_range.min_row, target_range.min_col)] = OpenPyxlCell(
                    ws,
                    row=target_range.min_row,
                    column=target_range.min_col,
                    value=value,
                )
        else:
            try:
                ws.cell(row, column).value = value
            except AttributeError:
                ws._cells[(row, column)] = OpenPyxlCell(ws, row=row, column=column, value=value)

    def _set_sheet_value(self, ws, row: int, column: int, value: Any) -> None:
        try:
            ws.cell(row, column).value = value
        except AttributeError:
            self._set_sheet_value_safe(ws, row, column, value)

    def _resolve_apu_layout_kind(self, ws) -> str:
        if ws["A3"].value == "Rubro:":
            return "sercop"
        return "general"

    def _get_apu_category_blueprints(self, layout_kind: str, category_idx: int) -> Tuple[Dict[int, Any], Dict[int, Any]]:
        if layout_kind == "sercop":
            header_map = {
                1: {1: "Descripción", 2: "Cantidad", 3: "Tarifa", 4: "Costo hora", 5: "Rendimiento", 6: "Costo"},
                2: {1: "Descripción", 3: "Unidad", 4: "Cantidad", 5: "Precio unitario", 6: "Costo"},
                3: {1: "Descripción", 3: "Unidad", 4: "Cantidad", 5: "Tarifa", 6: "Costo"},
                4: {1: "Descripción", 2: "Cantidad", 3: "Jornal/hr", 4: "Costo hora", 5: "Rendimiento", 6: "Costo"},
            }
            item_map = {
                1: {1: "#DESCRIPCION1", 2: "#CANTIDAD1", 3: "#PRECIO1", 5: "#RENDIMIENTO1", 6: "#SUBTOTAL1"},
                2: {1: "#DESCRIPCION2", 3: "#UNIDAD2", 4: "#CANTIDAD2", 5: "#PRECIO2", 6: "#SUBTOTAL2"},
                3: {1: "#DESCRIPCION3", 3: "#UNIDAD3", 4: "#CANTIDAD3", 5: "#PRECIO3", 6: "#SUBTOTAL3"},
                4: {1: "#DESCRIPCION4", 2: "#CANTIDAD4", 3: "#PRECIO4", 5: "#RENDIMIENTO4", 6: "#SUBTOTAL4"},
            }
            return header_map.get(category_idx, {}), item_map.get(category_idx, {})

        header_map = {
            1: {1: "Código", 2: "Descripción", 3: "Unidad", 4: "Cantidad", 5: "Precio", 6: "Rendim.", 7: "Total", 8: "%"},
            2: {1: "Código", 2: "Descripción", 3: "Unidad", 4: "Cantidad", 5: "Precio", 7: "Total", 8: "%"},
            3: {1: "Código", 2: "Descripción", 3: "Unidad", 4: "Cantidad", 5: "Tarifa/U", 6: "Distancia", 7: "Total", 8: "%"},
            4: {1: "Código", 2: "Descripción", 4: "Número", 5: "S.R.H.", 6: "Rendim.", 7: "Total", 8: "%"},
        }
        item_map = {
            1: {1: "#CODCAT1", 2: "#DESCRIPCION1", 3: "#UNIDAD1", 4: "#CANTIDAD1", 5: "#PRECIO1", 6: "#RENDIMIENTO1", 7: "#SUBTOTAL1", 8: "#PORCENT1"},
            2: {1: "#CODCAT2", 2: "#DESCRIPCION2", 3: "#UNIDAD2", 4: "#CANTIDAD2", 5: "#PRECIO2", 7: "#SUBTOTAL2", 8: "#PORCENT2"},
            3: {1: "#CODCAT3", 2: "#DESCRIPCION3", 3: "#UNIDAD3", 4: "#CANTIDAD3", 5: "#PRECIO3", 6: "#RENDIMIENTO3", 7: "#SUBTOTAL3", 8: "#PORCENT3"},
            4: {1: "#CODCAT4", 2: "#DESCRIPCION4", 4: "#CANTIDAD4", 5: "#PRECIO4", 6: "#RENDIMIENTO4", 7: "#SUBTOTAL4", 8: "#PORCENT4"},
        }
        return header_map.get(category_idx, {}), item_map.get(category_idx, {})

    def _restore_apu_category_row_blueprints(self, ws, category_idx: int, item_row_idx: int) -> None:
        layout_kind = self._resolve_apu_layout_kind(ws)
        header_values, item_values = self._get_apu_category_blueprints(layout_kind, category_idx)
        header_row_idx = item_row_idx - 1
        for column, value in header_values.items():
            self._set_sheet_value(ws, header_row_idx, column, value)
        for column, value in item_values.items():
            self._set_sheet_value(ws, item_row_idx, column, value)

    def _restore_apu_category_from_metadata(self, ws, category_meta: Dict[str, Any], item_row_idx: int, summary_row_idx: Optional[int]) -> None:
        pre_title_blueprint = category_meta.get("pre_title_blueprint")
        title_blueprint = category_meta.get("title_blueprint")
        header_blueprint = category_meta.get("header_blueprint")
        item_blueprint = category_meta.get("item_blueprint")
        summary_blueprint = category_meta.get("summary_blueprint")
        if pre_title_blueprint:
            self._restore_row_blueprint(ws, item_row_idx - 3, pre_title_blueprint)
        if title_blueprint:
            self._restore_row_blueprint(ws, item_row_idx - 2, title_blueprint)
        if header_blueprint:
            self._restore_row_blueprint(ws, item_row_idx - 1, header_blueprint)
        if item_blueprint:
            self._restore_row_blueprint(ws, item_row_idx, item_blueprint)
        if summary_row_idx and summary_blueprint:
            self._restore_row_blueprint(ws, summary_row_idx, summary_blueprint)

    def _apply_apu_sercop_section_title_fills(self, ws) -> None:
        fill = PatternFill(fill_type="solid", fgColor="FFEAF3F8")
        for label in ("EQUIPOS", "MANO DE OBRA", "MATERIALES", "TRANSPORTE"):
            row_idx = self._find_row_containing_text(ws, label)
            if not row_idx:
                continue
            for col_idx in range(1, 7):
                ws.cell(row=row_idx, column=col_idx).fill = copy(fill)

    def _resolve_apu_sercop_summary_label(self, category_idx: int) -> str:
        return {
            1: "SUBTOTAL M",
            4: "SUBTOTAL N",
            2: "SUBTOTAL O",
            3: "SUBTOTAL P",
        }.get(category_idx, "SUBTOTAL")

    def _apply_apu_category_summary_value(
        self,
        ws,
        category_idx: int,
        category: Dict[str, Any],
        summary_row_idx: Optional[int],
        layout_kind: str,
    ) -> None:
        if not summary_row_idx:
            return
        total_value = category.get("total", 0)
        if layout_kind == "sercop":
            self._set_sheet_value_safe(ws, summary_row_idx, 1, self._resolve_apu_sercop_summary_label(category_idx))
            self._set_sheet_value(ws, summary_row_idx, 6, total_value)
            return

    def _ensure_apu_item_row_layout(self, ws, row_idx: int, description: Any, layout_kind: str) -> None:
        desc_col = 1 if layout_kind == "sercop" else 2
        desc_cell = ws.cell(row=row_idx, column=desc_col)
        current_alignment = copy(desc_cell.alignment) if desc_cell.alignment else Alignment()
        current_alignment.wrap_text = True
        current_alignment.vertical = current_alignment.vertical or "center"
        desc_cell.alignment = current_alignment

        text = str(description or "").strip()
        if not text:
            return

        chars_per_line = 34 if layout_kind == "sercop" else 28
        line_count = max(1, text.count("\n") + 1)
        estimated_lines = max(line_count, (len(text) // chars_per_line) + (1 if len(text) % chars_per_line else 0))

        if estimated_lines <= 1:
            target_height = 15
        elif estimated_lines == 2:
            target_height = 24
        else:
            target_height = min(14 * estimated_lines, 56)

        current_height = ws.row_dimensions[row_idx].height or 15
        if target_height > current_height:
            ws.row_dimensions[row_idx].height = target_height

    def _sync_apu_footer_values(
        self,
        ws,
        data: Dict[str, Any],
        metadata: Optional[Dict[str, Any]] = None,
        row_offset: int = 0,
        money_decimals: int = 2,
    ) -> None:
        footer_rows = (metadata or {}).get("footer_rows", {})
        footer_block = (metadata or {}).get("footer_block") or {}
        layout_kind = (metadata or {}).get("layout_kind") or self._resolve_apu_layout_kind(ws)
        footer_start = footer_block.get("start_row")
        footer_rows_blueprints = footer_block.get("rows") or []
        if footer_start and footer_rows_blueprints:
            target_start = footer_start + row_offset
            for idx, blueprint in enumerate(footer_rows_blueprints):
                self._restore_row_blueprint(ws, target_start + idx, blueprint)

        if layout_kind != "sercop":
            direct_row = footer_rows.get("direct_row")
            if direct_row is not None:
                direct_row += row_offset
            else:
                direct_row = self._find_row_containing_text(ws, "Costo Directo Total")
            if direct_row:
                self._set_sheet_value_safe(ws, direct_row, 1, "Costo Directo Total: ")
                self._set_sheet_value(ws, direct_row, 8, data.get("#COSTODIRECTO", ""))

            indirect_title_row = footer_rows.get("indirect_title_row")
            if indirect_title_row is not None:
                indirect_title_row += row_offset
            else:
                indirect_title_row = self._find_row_containing_text(ws, "COSTOS INDIRECTOS")
            if indirect_title_row:
                self._set_sheet_value_safe(ws, indirect_title_row, 1, "COSTOS INDIRECTOS")
                indirect_row = indirect_title_row + 1
                if ws.max_column >= 8:
                    self._set_sheet_value(ws, indirect_row, 1, self._format_percent_label(data.get("#%INDIRECTO", 0), money_decimals))
                    self._set_sheet_value(ws, indirect_row, 8, data.get("#COSTOINDIRECTO", ""))

            total_row = footer_rows.get("total_row")
            if total_row is not None:
                total_row += row_offset
            else:
                total_row = self._find_row_containing_text(ws, "Precio Unitario Total")
            if total_row:
                self._set_sheet_value_safe(
                    ws,
                    total_row,
                    1,
                    "Precio Unitario Total .................................................................................................",
                )
                self._set_sheet_value(ws, total_row, 8, data.get("#TTOTAL", ""))

            words_row = footer_rows.get("words_row")
            if words_row is not None:
                words_row += row_offset
            else:
                words_row = self._find_row_containing_text(ws, "Son:")
            if words_row:
                self._set_sheet_value_safe(ws, words_row, 1, "Son:")
                self._set_sheet_value(ws, words_row, 2, data.get("#TEXTOTOTAL", ""))

        if layout_kind == "sercop":
            sercop_direct_row = footer_rows.get("sercop_direct_row")
            if sercop_direct_row is not None:
                sercop_direct_row += row_offset
            else:
                sercop_direct_row = self._find_row_containing_text(ws, "TOTAL COSTO DIRECTO (M+N+O+P)")
            if sercop_direct_row:
                self._set_sheet_value(ws, sercop_direct_row, 6, data.get("#COSTODIRECTO", ""))

            sercop_indirect_row = footer_rows.get("sercop_indirect_row")
            if sercop_indirect_row is not None:
                sercop_indirect_row += row_offset
            else:
                sercop_indirect_row = self._find_row_containing_text(ws, "INDIRECTOS")
            if sercop_indirect_row:
                self._set_sheet_value(ws, sercop_indirect_row, 5, self._format_percent_label(data.get("#%INDIRECTO", 0), money_decimals))
                self._set_sheet_value(ws, sercop_indirect_row, 6, data.get("#COSTOINDIRECTO", ""))

            sercop_total_row = footer_rows.get("sercop_total_row")
            if sercop_total_row is not None:
                sercop_total_row += row_offset
            else:
                sercop_total_row = self._find_row_containing_text(ws, "COSTO TOTAL DEL RUBRO")
            if sercop_total_row:
                self._set_sheet_value(ws, sercop_total_row, 6, data.get("#TTOTAL", ""))

            sercop_offer_row = footer_rows.get("sercop_offer_row")
            if sercop_offer_row is not None:
                sercop_offer_row += row_offset
            else:
                sercop_offer_row = self._find_row_containing_text(ws, "VALOR OFERTADO")
            if sercop_offer_row:
                self._set_sheet_value(ws, sercop_offer_row, 6, data.get("#TTOTAL", ""))

    def _extract_apu_template_metadata(self, ws) -> Dict[str, Any]:
        data_tokens = {
            "#CODIGO_APU",
            "#DESCRIPCION",
            "#CATEGORIA_BASE",
            "#UNIDAD",
            "#TTOTAL",
            "#COSTODIRECTO",
            "#COSTOINDIRECTO",
            "#%INDIRECTO",
            "#TOTALO",
            "#TEXTOTOTAL",
        }
        metadata = {
            "layout_kind": self._resolve_apu_layout_kind(ws),
            "data_cells": [],
            "categories": {},
            "footer_rows": {},
            "footer_block": {},
        }

        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                if any(token in cell.value for token in data_tokens):
                    metadata["data_cells"].append({
                        "row": cell.row,
                        "col": cell.column,
                        "value": cell.value,
                    })

        for category_idx in range(1, 5):
            item_placeholders = [
                f"#CODCAT{category_idx}",
                f"#DESCRIPCION{category_idx}",
                f"#UNIDAD{category_idx}",
                f"#CANTIDAD{category_idx}",
                f"#PRECIO{category_idx}",
                f"#RENDIMIENTO{category_idx}",
                f"#SUBTOTAL{category_idx}",
                f"#PORCENT{category_idx}",
                f"#OMNICLASS_COD{category_idx}",
                f"#OMNICLASS_TIT{category_idx}",
                f"@OMNICLASS_TIT{category_idx}",
            ]
            summary_placeholders = [
                f"#TOTAL{category_idx}",
                f"#TOTPORCENT{category_idx}",
                f"TOTPORCENT{category_idx}",
            ]
            item_row = self._find_row_with_any_placeholder(ws, item_placeholders)
            summary_row = self._find_row_with_any_placeholder(ws, summary_placeholders)
            if item_row:
                metadata["categories"][category_idx] = {
                    "pre_title_row": item_row - 3 if item_row - 3 >= 1 else None,
                    "title_row": item_row - 2 if item_row - 2 >= 1 else None,
                    "header_row": item_row - 1,
                    "item_row": item_row,
                    "summary_row": summary_row,
                    "gap_rows": max(0, (summary_row or item_row) - item_row - 1),
                    "pre_title_blueprint": self._capture_row_blueprint(ws, item_row - 3) if item_row - 3 >= 1 else None,
                    "title_blueprint": self._capture_row_blueprint(ws, item_row - 2) if item_row - 2 >= 1 else None,
                    "header_blueprint": self._capture_row_blueprint(ws, item_row - 1),
                    "item_blueprint": self._capture_row_blueprint(ws, item_row),
                    "summary_blueprint": self._capture_row_blueprint(ws, summary_row) if summary_row else None,
                }

        footer_needles = {
            "direct_row": "Costo Directo Total",
            "indirect_title_row": "COSTOS INDIRECTOS",
            "total_row": "Precio Unitario Total",
            "words_row": "Son:",
            "sercop_direct_row": "TOTAL COSTO DIRECTO (M+N+O+P)",
            "sercop_indirect_row": "INDIRECTOS",
            "sercop_total_row": "COSTO TOTAL DEL RUBRO",
            "sercop_offer_row": "VALOR OFERTADO",
        }
        for key, needle in footer_needles.items():
            metadata["footer_rows"][key] = self._find_row_containing_text(ws, needle)

        footer_row_candidates = [row for row in metadata["footer_rows"].values() if row]
        if footer_row_candidates:
            footer_start = max(1, min(footer_row_candidates) - 1)
            footer_end = max(footer_row_candidates)
            metadata["footer_block"] = {
                "start_row": footer_start,
                "rows": [self._capture_row_blueprint(ws, row_idx) for row_idx in range(footer_start, footer_end + 1)],
            }

        return metadata

    def _get_apu_template_metadata(self, template_id: str) -> Dict[str, Any]:
        cached = self._apu_template_metadata_cache.get(template_id)
        if cached:
            return cached
        filename = self._resolve_apu_template_filename(template_id)
        template_path = os.path.join(self.templates_dir, filename)
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Plantilla no encontrada: {template_path}")
        wb = openpyxl.load_workbook(template_path)
        metadata = self._extract_apu_template_metadata(wb.active)
        self._apu_template_metadata_cache[template_id] = metadata
        return metadata

    def _sanitize_sheet_title(self, value: str, fallback: str) -> str:
        invalid = set(r'[]:*?/\\')
        cleaned = ''.join(ch for ch in (value or "") if ch not in invalid).strip()
        cleaned = cleaned[:31] if cleaned else fallback
        return cleaned or fallback

    def _build_generated_by_giproy_label(self, report_label: str) -> str:
        normalized_label = str(report_label or "").strip().upper()
        report_name = f"Reporte {normalized_label}" if normalized_label else "Reporte documental"
        return f"{report_name} generado por GIPROY®"

    def _build_hierarchy_document_identity_table(
        self,
        doc_width: float,
        project_name: str,
        project_code: str,
        revision: str,
        reference_code: str,
        label_style: ParagraphStyle,
        body_style: ParagraphStyle,
        text_renderer,
    ) -> Table:
        def cell_label(value: str) -> Paragraph:
            return Paragraph(text_renderer(value), label_style)

        def cell_value(value: str) -> Paragraph:
            return Paragraph(text_renderer(value), body_style)

        table = Table(
            [
                [
                    cell_label("Cod del Proyecto"),
                    cell_value(project_code),
                    cell_label("Rev"),
                    cell_value(revision),
                    cell_label("Cod Referencial"),
                    cell_value(reference_code),
                ],
                [
                    cell_label("Nombre del Proyecto"),
                    cell_value(project_name),
                    "",
                    "",
                    "",
                    "",
                ],
            ],
            colWidths=[
                doc_width * 0.18,
                doc_width * 0.18,
                doc_width * 0.07,
                doc_width * 0.07,
                doc_width * 0.12,
                doc_width * 0.38,
            ],
        )
        table.setStyle(TableStyle([
            ("SPAN", (1, 1), (-1, 1)),
            ("BACKGROUND", (0, 0), (-1, -1), colors.white),
            ("BACKGROUND", (0, 0), (0, 1), colors.HexColor("#F8FAFC")),
            ("BACKGROUND", (2, 0), (2, 0), colors.HexColor("#F8FAFC")),
            ("BACKGROUND", (4, 0), (4, 0), colors.HexColor("#F8FAFC")),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D4D4D8")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]))
        return table

    def _build_hierarchy_document_pdf(
        self,
        preview: Dict[str, Any],
        watermark_text: Optional[str] = None,
    ) -> io.BytesIO:
        report_type = str(preview.get("report_type") or "").lower()
        item = (preview.get("items") or [{}])[0]
        is_edt = report_type == "edt"
        title = "Estructura de Descomposición del Trabajo" if is_edt else "Estructura de Descomposición de la Organización"
        acronym = "EDT" if is_edt else "EDO"
        subtitle = (
            "Sección documental para incorporar al informe técnico del proyecto."
            if is_edt
            else "Sección documental para incorporar al informe organizacional del proyecto."
        )

        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=A4,
            leftMargin=18 * mm,
            rightMargin=18 * mm,
            topMargin=17 * mm,
            bottomMargin=16 * mm,
            title=f"{title} - {acronym}",
        )
        doc.author = self.REPORT_AUTHOR
        doc.creator = self.REPORT_AUTHOR

        styles = getSampleStyleSheet()
        eyebrow_style = ParagraphStyle(
            "HierarchyDocEyebrow",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#F39200"),
            spaceAfter=5,
        )
        title_style = ParagraphStyle(
            "HierarchyDocTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#111827"),
            spaceAfter=5,
        )
        subtitle_style = ParagraphStyle(
            "HierarchyDocSubtitle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#4B5563"),
            spaceAfter=8,
        )
        section_style = ParagraphStyle(
            "HierarchyDocSection",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            textColor=colors.HexColor("#111827"),
            spaceBefore=5,
            spaceAfter=5,
        )
        body_style = ParagraphStyle(
            "HierarchyDocBody",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#374151"),
        )
        small_style = ParagraphStyle(
            "HierarchyDocSmall",
            parent=body_style,
            fontSize=7.4,
            leading=9,
            textColor=colors.HexColor("#6B7280"),
        )
        label_style = ParagraphStyle(
            "HierarchyDocLabel",
            parent=small_style,
            fontName="Helvetica-Bold",
            textColor=colors.HexColor("#6B7280"),
        )
        card_title_style = ParagraphStyle(
            "HierarchyDocCardTitle",
            parent=body_style,
            fontName="Helvetica-Bold",
            fontSize=8.8,
            leading=11,
            textColor=colors.HexColor("#111827"),
        )
        card_meta_style = ParagraphStyle(
            "HierarchyDocCardMeta",
            parent=small_style,
            fontName="Helvetica-Bold",
            textColor=colors.HexColor("#136191"),
        )

        def text(value: Any) -> str:
            raw = "" if value is None else str(value)
            return (
                raw.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
                .replace('"', "&quot;")
            )

        def clean(value: Any, fallback: str = "-") -> str:
            rendered = str(value or "").strip()
            return rendered or fallback

        field_map = {
            str(field.get("label") or ""): clean(field.get("value"))
            for field in (item.get("fields") or [])
        }
        project_name = field_map.get("Nombre del Proyecto") or clean(item.get("descripcion"))
        revision = field_map.get("Revisión") or "-"
        project_code = field_map.get("Código del Proyecto") or clean(item.get("codigo"))
        reference_code = field_map.get("Código Referencial") or "-"

        summary_cards = item.get("summary_cards") or []
        summary_row = [
            [
                Paragraph(text(card.get("label") or "-").upper(), label_style),
                Paragraph(text(clean(card.get("value"), "0")), card_title_style),
            ]
            for card in summary_cards
        ]
        summary_table = Table(
            summary_row,
            colWidths=[doc.width / max(len(summary_row), 1)] * max(len(summary_row), 1),
        ) if summary_row else None
        if summary_table:
            summary_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D4D4D8")),
                ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#E4E4E7")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ]))

        meta_table = self._build_hierarchy_document_identity_table(
            doc.width,
            project_name,
            project_code,
            revision,
            reference_code,
            label_style,
            body_style,
            text,
        )

        narrative = (
            "La EDT organiza el alcance del proyecto en cuentas y paquetes de trabajo para facilitar "
            "la lectura documental, la asignación de responsables y la trazabilidad de entregables."
            if is_edt
            else "La EDO identifica hitos organizacionales, roles y responsables asociados al proyecto "
            "para facilitar la lectura documental de responsabilidades y coordinación operativa."
        )

        story = [
            Paragraph(text(self._build_generated_by_giproy_label(acronym)), eyebrow_style),
            Paragraph(text(title), title_style),
            Paragraph(text(subtitle), subtitle_style),
            meta_table,
            Spacer(1, 5 * mm),
        ]
        if summary_table:
            story.extend([summary_table, Spacer(1, 5 * mm)])
        story.extend([
            Paragraph("Propósito de la sección", section_style),
            Paragraph(text(narrative), body_style),
            Spacer(1, 4 * mm),
            Paragraph("Detalle documental", section_style),
        ])

        for line_index, line in enumerate(item.get("lineas") or [], start=1):
            level_code = clean(line.get("codigo_edt") if is_edt else line.get("item"), str(line_index))
            level = max(0, min(str(level_code).count("."), 4))
            left_padding = 6 + (level * 7)
            if is_edt:
                main_title = clean(line.get("descripcion_cuenta_paquete"))
                meta_bits = [
                    f"Código EDT: {clean(line.get('codigo_edt'))}",
                    f"Responsable: {clean(line.get('nombre_responsable'))}",
                ]
                body_bits = [
                    ("Definición", clean(line.get("definicion"), "Sin definición registrada.")),
                ]
                if line.get("subtotal_cuenta") not in (None, "", 0):
                    body_bits.append(("Valor referencial", f"${self._round_half_up(line.get('subtotal_cuenta'), 2):,.2f}"))
            else:
                main_title = clean(line.get("nombre_responsable"))
                meta_bits = [
                    f"Ítem: {clean(line.get('item'), str(line_index))}",
                    f"Rol asignado: {clean(line.get('rol_asignado'))}",
                ]
                body_bits = [
                    ("Actividades clave", clean(line.get("actividades_clave"), "Sin actividades clave registradas.")),
                ]

            card_rows = [
                [Paragraph(text(main_title), card_title_style)],
                [Paragraph(text(" · ".join(meta_bits)), card_meta_style)],
            ]
            for label, value in body_bits:
                card_rows.append([Paragraph(f"<b>{text(label)}:</b> {text(value)}", body_style)])

            card = Table(card_rows, colWidths=[doc.width - left_padding])
            card.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFFFFF")),
                ("BOX", (0, 0), (-1, -1), 0.45, colors.HexColor("#E4E4E7")),
                ("LEFTPADDING", (0, 0), (-1, -1), left_padding),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(KeepTogether([card, Spacer(1, 2.2 * mm)]))

        story.extend([
            Spacer(1, 4 * mm),
            Paragraph("Control documental", section_style),
            Paragraph(
                text("Este reporte fue generado por GiProy para integrarse como anexo o sección dentro del informe documental del proyecto."),
                small_style,
            ),
        ])

        def apply_pdf_metadata(pdf_canvas, _doc):
            pdf_canvas.setAuthor(self.REPORT_AUTHOR)
            pdf_canvas.setCreator(self.REPORT_AUTHOR)
            pdf_canvas.setTitle(f"{title} - {project_name}")
            pdf_canvas.saveState()
            pdf_canvas.setFont("Helvetica", 7)
            pdf_canvas.setFillColor(colors.HexColor("#6B7280"))
            footer = f"{acronym} · {project_code} · pág. {_doc.page}"
            pdf_canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, footer)
            pdf_canvas.restoreState()
            self._draw_pdf_watermark(pdf_canvas, A4[0], A4[1], watermark_text)

        doc.build(story, onFirstPage=apply_pdf_metadata, onLaterPages=apply_pdf_metadata)
        output.seek(0)
        return output

    def _build_pdf_report(
        self,
        title: str,
        items: List[Dict[str, Any]],
        money_decimals: int = 2,
        calc_decimals: int = 4,
        watermark_text: Optional[str] = None,
    ) -> io.BytesIO:
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=A4,
            leftMargin=12 * mm,
            rightMargin=12 * mm,
            topMargin=12 * mm,
            bottomMargin=12 * mm,
            title=title or "Reporte",
        )
        doc.author = self.REPORT_AUTHOR
        doc.creator = self.REPORT_AUTHOR

        def apply_pdf_metadata(canvas, _doc):
            canvas.setAuthor(self.REPORT_AUTHOR)
            canvas.setCreator(self.REPORT_AUTHOR)
            canvas.setTitle(title or "Reporte")
            self._draw_pdf_watermark(canvas, A4[0], A4[1], watermark_text)

        styles = getSampleStyleSheet()
        eyebrow_style = ParagraphStyle(
            "ReportEyebrow",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#F39200"),
            spaceAfter=4,
        )
        title_style = ParagraphStyle(
            "ReportTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=19,
            textColor=colors.HexColor("#111827"),
            spaceAfter=4,
        )
        meta_style = ParagraphStyle(
            "ReportMeta",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#6B7280"),
            spaceAfter=10,
        )
        summary_label_style = ParagraphStyle(
            "ReportSummaryLabel",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=7,
            leading=9,
            textColor=colors.HexColor("#6B7280"),
        )
        summary_value_style = ParagraphStyle(
            "ReportSummaryValue",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=13,
            textColor=colors.HexColor("#111827"),
        )
        summary_total_value_style = ParagraphStyle(
            "ReportSummaryTotalValue",
            parent=summary_value_style,
            textColor=colors.HexColor("#F39200"),
        )
        cell_style = ParagraphStyle(
            "ReportCell",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=7,
            leading=9,
            textColor=colors.HexColor("#111827"),
        )
        cell_bold_style = ParagraphStyle(
            "ReportCellBold",
            parent=cell_style,
            fontName="Helvetica-Bold",
        )
        cell_indented_style = ParagraphStyle(
            "ReportCellIndented",
            parent=cell_style,
            leftIndent=10,
        )

        def money(value: Any) -> str:
            return f"${self._round_half_up(value or 0, money_decimals):,.{money_decimals}f}"

        def calc(value: Any) -> str:
            if value in (None, ""):
                return "-"
            return self._format_fixed(value, calc_decimals)

        def integer(value: Any) -> str:
            if value in (None, ""):
                return "-"
            return str(int(self._round_half_up(value, 0)))

        def render_value(value: Any, kind: str) -> str:
            if value in (None, ""):
                return "-"
            if kind == "money":
                return money(value)
            if kind == "percent":
                return self._format_fixed(value, money_decimals)
            if kind == "integer":
                return integer(value)
            if kind == "calc":
                return calc(value)
            return text(value)

        def text(value: Any) -> str:
            raw = "" if value is None else str(value)
            return (
                raw.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
                .replace('"', "&quot;")
            )

        story = []
        for index, item in enumerate(items):
            if index > 0:
                story.append(PageBreak())

            header_block = [
                Paragraph(text(title or "Reporte"), eyebrow_style),
                Paragraph(f"{text(item.get('codigo') or '-') } - {text(item.get('descripcion') or '-')}", title_style),
            ]
            meta_parts = [
                f"Categoría base: {text(item.get('categoria_base') or '-')}",
                f"Unidad: {text(item.get('unidad') or '-')}",
            ]
            metadata_hint = item.get("metadata_hint")
            if metadata_hint:
                meta_parts.append(text(metadata_hint))
            else:
                meta_parts.extend([
                    f"Costo Directo: {money(item.get('costo_directo'))}",
                    f"Costo Indirecto: {money(item.get('costo_indirecto'))}",
                    f"Total: {money(item.get('precio_total'))}",
                ])
            header_block.append(Paragraph(" | ".join(meta_parts), meta_style))

            summary_cards = item.get("summary_cards") or [
                {"label": "DIRECTO", "value": item.get("costo_directo"), "kind": item.get("summary_direct_kind", "money")},
                {"label": "INDIRECTO", "value": item.get("costo_indirecto"), "kind": item.get("summary_indirect_kind", "money")},
                {"label": "TOTAL", "value": item.get("precio_total"), "kind": item.get("summary_total_kind", "money"), "tone": "total"},
            ]
            summary_data = [
                [Paragraph(text(card.get("label") or ""), summary_label_style) for card in summary_cards],
                [
                    Paragraph(
                        render_value(card.get("value"), card.get("kind", "text")),
                        summary_total_value_style if card.get("tone") == "total" else summary_value_style,
                    )
                    for card in summary_cards
                ],
            ]
            summary_table = Table(
                summary_data,
                colWidths=[doc.width / max(len(summary_cards), 1)] * max(len(summary_cards), 1),
            )
            summary_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#F8FAFC")),
                ("BACKGROUND", (len(summary_cards) - 1, 0), (len(summary_cards) - 1, 1), colors.HexColor("#FFF7ED")),
                ("BOX", (0, 0), (-1, 1), 0.6, colors.HexColor("#D4D4D8")),
                ("INNERGRID", (0, 0), (-1, 1), 0.4, colors.HexColor("#E4E4E7")),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]))
            header_block.append(summary_table)
            header_block.append(Spacer(1, 6 * mm))
            story.append(KeepTogether(header_block))

            fields = item.get("fields") or []
            if fields:
                field_rows = [
                    [
                        Paragraph(text(field.get("label") or ""), cell_bold_style),
                        Paragraph(text(field.get("value") or "-"), cell_style),
                    ]
                    for field in fields
                ]
                fields_table = Table(field_rows, colWidths=[doc.width * 0.32, doc.width * 0.68])
                fields_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F8FAFC")),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#E4E4E7")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ]))
                story.append(KeepTogether([fields_table, Spacer(1, 6 * mm)]))

            visual_blocks = []
            referential_image_url = item.get("image_referencial_url")
            referential_image_bytes = self._fetch_remote_image_bytes(referential_image_url) if referential_image_url else None
            if referential_image_bytes:
                visual_blocks.append(("Imagen referencial", referential_image_bytes))
            map_image_url = item.get("map_image_url")
            map_image_bytes = self._fetch_remote_image_bytes(map_image_url) if map_image_url else None
            if map_image_bytes:
                visual_blocks.append(("Georreferenciación", map_image_bytes))
            for block_label, block_bytes in visual_blocks:
                try:
                    image = ReportLabImage(io.BytesIO(block_bytes), width=doc.width, height=min(90 * mm, doc.width * 0.55))
                    image.hAlign = "CENTER"
                    story.append(KeepTogether([
                        Paragraph(text(block_label), meta_style),
                        Spacer(1, 2 * mm),
                        image,
                        Spacer(1, 6 * mm),
                    ]))
                except Exception:
                    pass

            lineas = item.get("lineas", [])
            if lineas:
                table_columns = item.get("table_columns") or [
                    {"key": "codigo", "label": "Código"},
                    {"key": "descripcion", "label": "Descripción"},
                    {"key": "unidad", "label": "Unidad"},
                    {"key": "cantidad", "label": "Cant.", "kind": "calc"},
                    {"key": "precio", "label": "Precio", "kind": "money"},
                    {"key": "rendimiento", "label": "Rend.", "kind": "calc"},
                    {"key": "subtotal", "label": "Subtotal", "kind": "money"},
                ]
                table_rows = [[
                    Paragraph(text((column.get("label") or "-").upper()), cell_bold_style)
                    for column in table_columns
                ]]
                for linea in lineas:
                    is_structural = bool(linea.get("is_structural"))
                    row_cells = []
                    for column in table_columns:
                        column_key = column.get("key")
                        raw_value = linea.get(column_key)
                        kind = column.get("kind")
                        if column_key == "descripcion":
                            descripcion = text(raw_value or "-")
                            paragraph = Paragraph(
                                descripcion,
                                cell_bold_style if is_structural else cell_indented_style,
                            )
                            row_cells.append(paragraph)
                        elif column_key == "unidad":
                            row_cells.append(Paragraph(text(raw_value or "-"), cell_style))
                        else:
                            if kind == "money":
                                rendered_value = render_value(raw_value, "money")
                            elif kind == "integer":
                                rendered_value = render_value(raw_value, "integer")
                            elif kind == "calc":
                                rendered_value = render_value(raw_value, "calc")
                            else:
                                rendered_value = text(raw_value or "-")
                            row_cells.append(Paragraph(rendered_value, cell_style))
                    table_rows.append(row_cells)

                total_width = doc.width
                width_weights = [
                    float(column.get("width_weight") or 0)
                    for column in table_columns
                ]
                if any(weight > 0 for weight in width_weights):
                    total_weight = sum(weight for weight in width_weights if weight > 0) or 1
                    fallback_weight = total_weight / max(len(table_columns), 1)
                    normalized_weights = [
                        weight if weight > 0 else fallback_weight
                        for weight in width_weights
                    ]
                    normalized_total = sum(normalized_weights) or 1
                    col_widths = [
                        total_width * (weight / normalized_total)
                        for weight in normalized_weights
                    ]
                elif len(table_columns) <= 4:
                    col_widths = [total_width / len(table_columns)] * len(table_columns)
                else:
                    description_index = next((idx for idx, column in enumerate(table_columns) if column.get("key") == "descripcion"), None)
                    remaining = len(table_columns) - (1 if description_index is not None else 0)
                    narrow_width = min(21 * mm, (total_width * 0.62) / max(remaining, 1))
                    description_width = max(total_width - (narrow_width * max(remaining, 1)), 55 * mm)
                    col_widths = []
                    for idx, _column in enumerate(table_columns):
                        col_widths.append(description_width if idx == description_index else narrow_width)
                report_table = Table(
                    table_rows,
                    repeatRows=1,
                    colWidths=col_widths,
                )
                report_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F4F4F5")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#52525B")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("LEADING", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#E4E4E7")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAFAFA")]),
                ]))
                story.append(report_table)

        doc.build(story, onFirstPage=apply_pdf_metadata, onLaterPages=apply_pdf_metadata)
        output.seek(0)
        return output

    def fill_template(
        self,
        template_name: str,
        data: Dict[str, Any],
        table_data: List[Dict[str, Any]] = None,
        table_row_marker: str = None,
        use_omniclass: bool = True,
    ) -> io.BytesIO:
        """
        Llena una plantilla de Excel con datos simples y opcionalmente una tabla.
        - template_name: nombre del archivo (.xlsx)
        - data: mapeo de #PLACEHOLDER -> valor para campos simples
        - table_data: lista de diccionarios para filas de tabla
        - table_row_marker: un placeholder único que identifica la fila de la tabla (ej: '#DESCRIPCION')
        """
        template_path = os.path.join(self.templates_dir, template_name)
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Plantilla no encontrada: {template_path}")

        wb = openpyxl.load_workbook(template_path)
        ws = wb.active
        if self._template_requires_category_engine(ws):
            raise ValueError(
                "La plantilla Excel seleccionada requiere motor categorizado por bloques y no puede "
                "rellenarse con el motor genérico de tabla plana."
            )
        self._apply_template_to_sheet(ws, data, table_data, table_row_marker)
        self._apply_omniclass_visibility(ws, use_omniclass)

        # Guardar en buffer
        return self._save_workbook_buffer(wb)

    def _normalize_report_signature_block(self, ws) -> Optional[int]:
        signature_text = "FIRMA DEL SPONSOR O\nPERSONA QUE AUTORIZA"
        signature_row = None
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "FIRMA DEL" in cell.value.upper():
                    signature_row = cell.row
                    break
            if signature_row is not None:
                break
        if signature_row is None:
            return None

        target_range = CellRange(
            min_row=signature_row,
            max_row=signature_row + 1,
            min_col=1,
            max_col=5,
        )
        for merged_range in list(ws.merged_cells.ranges):
            if CellRange(str(merged_range)).isdisjoint(target_range):
                continue
            try:
                ws.unmerge_cells(str(merged_range))
            except KeyError:
                if str(merged_range) in ws.merged_cells:
                    ws.merged_cells.remove(str(merged_range))

        for row_idx in range(signature_row, signature_row + 2):
            ws.row_dimensions[row_idx].height = 12
            for col_idx in range(1, 6):
                cell = ws.cell(row=row_idx, column=col_idx)
                if isinstance(cell, MergedCell):
                    continue
                cell.value = None

        signature_cell = ws.cell(row=signature_row, column=1)
        signature_cell.value = signature_text
        signature_cell.font = Font(name="Arial", size=8, bold=True)
        signature_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.merge_cells(
            start_row=signature_row,
            start_column=1,
            end_row=signature_row + 1,
            end_column=5,
        )
        return signature_row

    def _normalize_edo_signature_block(self, ws) -> None:
        self._normalize_report_signature_block(ws)

    def _normalize_edt_signature_block(self, ws) -> None:
        self._normalize_report_signature_block(ws)

    def _normalize_stakeholders_signature_block(self, ws) -> Optional[int]:
        return self._normalize_report_signature_block(ws)

    def _find_header_row_with_labels(self, ws, labels: List[str]) -> Optional[int]:
        expected = {str(label or "").strip().lower() for label in labels if label}
        if not expected:
            return None
        for row in ws.iter_rows():
            values = {
                str(cell.value or "").strip().lower()
                for cell in row
                if cell.value not in (None, "")
            }
            if expected.issubset(values):
                return row[0].row
        return None

    def _normalize_edo_document_header(self, ws) -> None:
        self._normalize_hierarchy_excel_document_header(ws)

    def _normalize_edt_document_header(self, ws) -> None:
        self._normalize_hierarchy_excel_document_header(ws)

    def _normalize_hierarchy_excel_document_header(self, ws) -> None:
        def first_non_empty(*values: Any, rejected: Optional[set[str]] = None) -> str:
            rejected_values = {str(value or "").strip().lower() for value in (rejected or set())}
            for value in values:
                if value not in (None, ""):
                    rendered = str(value).strip()
                    if rendered.lower() in rejected_values:
                        continue
                    return rendered
            return ""

        project_name = first_non_empty(
            ws["D5"].value,
            ws["C5"].value,
            ws["A5"].value,
            ws["B5"].value,
            rejected={"Nombre del Proyecto:"},
        )
        project_code = first_non_empty(
            ws["C4"].value,
            ws["C6"].value,
            ws["D6"].value,
            ws["E6"].value,
            rejected={"Cod del Proyecto:", "Código del Proyecto:"},
        )
        revision = first_non_empty(
            ws["E4"].value,
            ws["D4"].value,
            ws["F4"].value,
            ws["G4"].value,
            rejected={"Revisión:", "Revision:"},
        )
        reference_code = first_non_empty(
            ws["G4"].value,
            ws["D6"].value,
            ws["H4"].value,
            ws["G6"].value,
            ws["E6"].value,
            rejected={"Cod Referencial:", "Código Referencial:"},
        )

        header_range = CellRange("A4:H8")
        for merged_range in list(ws.merged_cells.ranges):
            if CellRange(str(merged_range)).isdisjoint(header_range):
                continue
            ws.unmerge_cells(str(merged_range))

        thin_side = Side(style="thin", color="000000")
        header_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        label_font = Font(name="Arial", size=10, bold=True)
        value_font = Font(name="Arial", size=10, bold=False)
        label_alignment = Alignment(horizontal="center", vertical="center", wrap_text=False, shrink_to_fit=True)
        value_alignment = Alignment(horizontal="center", vertical="center", wrap_text=False, shrink_to_fit=True)

        for row_idx in range(4, 9):
            ws.row_dimensions[row_idx].hidden = False
            ws.row_dimensions[row_idx].height = 20 if row_idx in (4, 5) else 4
            for col_idx in range(1, 9):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value = None
                cell.border = header_border if row_idx in (4, 5) else Border()
                cell.fill = PatternFill(fill_type=None)
                cell.alignment = value_alignment
                cell.font = value_font

        for row_idx in range(6, 9):
            ws.row_dimensions[row_idx].hidden = True

        ws.merge_cells("A4:B4")
        ws.merge_cells("G4:H4")
        ws.merge_cells("A5:C5")
        ws.merge_cells("D5:H5")

        header_cells = [
            ("A4", "Cod del Proyecto:", label_font, label_alignment),
            ("C4", project_code, value_font, value_alignment),
            ("D4", "Rev:", label_font, label_alignment),
            ("E4", revision, value_font, value_alignment),
            ("F4", "Cod Referencial:", label_font, label_alignment),
            ("G4", reference_code, value_font, value_alignment),
            ("A5", "Nombre del Proyecto:", label_font, label_alignment),
            ("D5", project_name, value_font, value_alignment),
        ]
        for coordinate, value, font, alignment in header_cells:
            cell = ws[coordinate]
            cell.value = value
            cell.font = font
            cell.alignment = alignment
            cell.border = header_border

    def _normalize_edo_document_columns(self, ws) -> None:
        ws.column_dimensions["A"].width = 5
        ws.column_dimensions["B"].width = 12.5
        ws.column_dimensions["C"].width = max(float(ws.column_dimensions["C"].width or 0), 26)
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions["E"].width = 8
        ws.column_dimensions["F"].width = 12
        ws.column_dimensions["G"].width = 34
        ws.column_dimensions["H"].width = 36
        header_row = self._find_header_row_with_labels(ws, ["Ítem", "Rol Asignado"])
        if header_row is None:
            return
        for row in ws.iter_rows(min_row=header_row, max_row=ws.max_row, min_col=1, max_col=2):
            for cell in row:
                cell.alignment = copy(cell.alignment)
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical=getattr(cell.alignment, "vertical", None) or "center",
                    wrap_text=getattr(cell.alignment, "wrap_text", None),
                )

    def _normalize_edt_document_columns(self, ws, report_type: str) -> None:
        variant = str(report_type or "listado").lower()
        if variant == "listado":
            ws.column_dimensions["A"].width = 5
            ws.column_dimensions["B"].width = 8
            ws.column_dimensions["C"].width = 56
            ws.column_dimensions["D"].width = 12
            ws.column_dimensions["E"].width = 8
            ws.column_dimensions["F"].width = 12
            ws.column_dimensions["G"].width = 34
            ws.column_dimensions["H"].width = 36
            self._normalize_edt_document_header(ws)
            header_row = self._find_header_row_with_labels(ws, ["Ítem", "Cod. EDT"])
            self._normalize_edt_listado_table_width(ws, header_row)
            table_columns = (1, 2)
        elif variant == "diccionario":
            ws.column_dimensions["A"].width = 10
            ws.column_dimensions["B"].width = 52
            ws.column_dimensions["C"].width = 26
            ws.column_dimensions["D"].width = 12
            ws.column_dimensions["E"].width = 8
            ws.column_dimensions["F"].width = 12
            ws.column_dimensions["G"].width = 34
            ws.column_dimensions["H"].width = 36
            self._normalize_edt_document_header(ws)
            header_row = self._find_header_row_with_labels(ws, ["Cod. EDT", "Descripción Cuenta / Paquete de Trabajo"])
            table_columns = (1,)
        elif variant == "valorada":
            ws.column_dimensions["A"].width = 10
            ws.column_dimensions["B"].width = 52
            ws.column_dimensions["C"].width = 26
            ws.column_dimensions["D"].width = 12
            ws.column_dimensions["E"].width = 8
            ws.column_dimensions["F"].width = 12
            ws.column_dimensions["G"].width = 34
            ws.column_dimensions["H"].width = 36
            self._normalize_edt_document_header(ws)
            header_row = self._find_header_row_with_labels(ws, ["Cod. EDT", "Descripción Cuenta / Paquete de Trabajo"])
            table_columns = (1,)
        else:
            return
        if header_row is None:
            return
        for column_index in table_columns:
            for row in ws.iter_rows(min_row=header_row, max_row=ws.max_row, min_col=column_index, max_col=column_index):
                for cell in row:
                    cell.alignment = copy(cell.alignment)
                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical=getattr(cell.alignment, "vertical", None) or "center",
                        wrap_text=getattr(cell.alignment, "wrap_text", None),
                    )

    def _normalize_edt_listado_table_width(self, ws, header_row: Optional[int]) -> None:
        if header_row is None:
            return

        for merged_range in list(ws.merged_cells.ranges):
            if (
                merged_range.min_row >= header_row
                and merged_range.max_row == merged_range.min_row
                and merged_range.min_col == 3
                and merged_range.max_col <= 8
            ):
                try:
                    ws.unmerge_cells(str(merged_range))
                except KeyError:
                    if str(merged_range) in ws.merged_cells:
                        ws.merged_cells.remove(str(merged_range))

        for row_idx in range(header_row, ws.max_row + 1):
            first_value = str(ws.cell(row=row_idx, column=1).value or "").strip().lower()
            if first_value in {"#ciudad", "cuenca"} or "firma del" in first_value:
                break
            has_table_value = any(ws.cell(row=row_idx, column=col_idx).value not in (None, "") for col_idx in (1, 2, 3))
            if not has_table_value:
                continue
            ws.merge_cells(start_row=row_idx, start_column=3, end_row=row_idx, end_column=8)

    def _render_edt_listado_document_blocks(self, ws, table_data: List[Dict[str, Any]]) -> None:
        header_row = self._find_header_row_with_labels(ws, ["Ítem", "Cod. EDT"])
        if header_row is None or not table_data:
            return

        def clear_row_merges(row_idx: int) -> None:
            row_range = CellRange(min_row=row_idx, max_row=row_idx, min_col=1, max_col=8)
            for merged_range in list(ws.merged_cells.ranges):
                if CellRange(str(merged_range)).isdisjoint(row_range):
                    continue
                try:
                    ws.unmerge_cells(str(merged_range))
                except KeyError:
                    if str(merged_range) in ws.merged_cells:
                        ws.merged_cells.remove(str(merged_range))

        thin_side = Side(style="thin", color="000000")
        block_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        title_font = Font(name="Arial", size=10, bold=True)
        meta_font = Font(name="Arial", size=9, bold=True, color="136191")
        body_font = Font(name="Arial", size=9, bold=False)
        section_font = Font(name="Arial", size=10, bold=True)
        white_fill = PatternFill(fill_type=None)
        header_fill = PatternFill(fill_type="solid", fgColor="E9EEF3")
        account_fill = PatternFill(fill_type="solid", fgColor="EEF6FB")

        document_area = CellRange(min_row=header_row, max_row=ws.max_row, min_col=1, max_col=8)
        for merged_range in list(ws.merged_cells.ranges):
            if not CellRange(str(merged_range)).isdisjoint(document_area):
                try:
                    ws.unmerge_cells(str(merged_range))
                except KeyError:
                    if str(merged_range) in ws.merged_cells:
                        ws.merged_cells.remove(str(merged_range))

        for row_idx in range(1, header_row + 1):
            if ws.row_dimensions[row_idx].hidden:
                continue
            for col_idx in range(1, 9):
                ws.cell(row=row_idx, column=col_idx).fill = header_fill

        clear_row_merges(header_row)
        section_cell = ws.cell(row=header_row, column=1)
        section_cell.value = "Detalle documental"
        section_cell.font = section_font
        section_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
        for col_idx in range(1, 9):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.border = block_border
            cell.fill = header_fill
        ws.merge_cells(start_row=header_row, start_column=1, end_row=header_row, end_column=8)
        ws.row_dimensions[header_row].height = 22

        original_rows = len(table_data)
        spacer_row = header_row + 1
        data_start_row = header_row + 2
        extra_rows = (original_rows * 2) + 1
        if extra_rows:
            insert_at = header_row + original_rows + 1
            self._shift_merged_ranges_below_row(ws, insert_at, extra_rows)
            ws.insert_rows(insert_at, extra_rows)

        clear_row_merges(spacer_row)
        for col_idx in range(1, 9):
            cell = ws.cell(row=spacer_row, column=col_idx)
            cell.value = None
            cell.border = Border()
            cell.fill = white_fill
        ws.row_dimensions[spacer_row].height = 10

        for index, row_data in enumerate(table_data):
            block_row = data_start_row + (index * 3)
            codigo = str(row_data.get("CODIGOEDT") or "-").strip() or "-"
            descripcion = str(row_data.get("DESCRIPCION_CUENTA_PAQUETE") or "-").strip() or "-"
            responsable = str(row_data.get("NOMBRE_RESPONSABLE") or "-").strip() or "-"
            definicion = str(row_data.get("DEFINICION") or "").strip() or "Sin definición registrada."
            level = max(0, min(codigo.count("."), 4))

            block_values = [
                (f"Código EDT: {codigo} · {descripcion}", title_font, account_fill, 22, False),
                (f"Responsable: {responsable}", meta_font, white_fill, 20, False),
                (f"Definición: {definicion}", body_font, white_fill, 30, True),
            ]
            for offset, (value, font, fill, height, wrap_text) in enumerate(block_values):
                row_idx = block_row + offset
                clear_row_merges(row_idx)
                for col_idx in range(1, 9):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.value = None
                    cell.border = block_border
                    cell.fill = fill
                    cell.font = font
                    cell.alignment = Alignment(
                        horizontal="left",
                        vertical="center",
                        wrap_text=wrap_text,
                        indent=level,
                    )
                anchor = ws.cell(row=row_idx, column=1)
                anchor.value = value
                anchor.font = font
                anchor.alignment = Alignment(
                    horizontal="left",
                    vertical="center",
                    wrap_text=wrap_text,
                    indent=level,
                )
                ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=8)
                ws.row_dimensions[row_idx].height = height

    def _resolve_apu_template_filename(self, template_id: str) -> str:
        filename = f"{template_id} - Analisis - APUS - General.xlsx"
        if template_id == "002":
            filename = "002 - Analisis - APUS - SERCOP.xlsx"
        elif template_id == "003":
            filename = "001 - Analisis - APUS - General.xlsx"
        return filename

    def _get_apu(self, db: Any, apu_id: int, empresa_id: int) -> Any:
        from app.models.apu import APU, APULinea
        from app.models.recurso import Recurso

        apu = (
            db.query(APU)
            .options(
                selectinload(APU.subcategoria_item),
                selectinload(APU.lineas)
                .selectinload(APULinea.recurso)
                .selectinload(Recurso.unidad),
                selectinload(APU.lineas).selectinload(APULinea.apu_hijo),
            )
            .filter(APU.id == apu_id, APU.empresa_id == empresa_id)
            .first()
        )
        if not apu:
            raise ValueError(f"APU no encontrado: {apu_id}")
        return apu

    def _get_apus_map(self, db: Any, apu_ids: List[int], empresa_id: int) -> Dict[int, Any]:
        from app.models.apu import APU, APULinea
        from app.models.recurso import Recurso

        if not apu_ids:
            return {}

        apus = (
            db.query(APU)
            .options(
                selectinload(APU.subcategoria_item),
                selectinload(APU.lineas)
                .selectinload(APULinea.recurso)
                .selectinload(Recurso.unidad),
                selectinload(APU.lineas).selectinload(APULinea.apu_hijo),
            )
            .filter(APU.empresa_id == empresa_id, APU.id.in_(apu_ids))
            .all()
        )
        return {apu.id: apu for apu in apus}

    def _get_presupuesto(self, db: Any, presupuesto_id: int, empresa_id: int) -> Any:
        from app.models.edt import EdtNode
        from app.models.presupuesto import Presupuesto, PresupuestoDetalle
        from app.services.presupuesto import refresh_presupuesto_prices
        presupuesto = (
            db.query(Presupuesto)
            .options(
                selectinload(Presupuesto.indirectos),
                selectinload(Presupuesto.empresa),
                selectinload(Presupuesto.proyecto),
                selectinload(Presupuesto.detalle).selectinload(PresupuestoDetalle.apu),
                selectinload(Presupuesto.detalle).selectinload(PresupuestoDetalle.edt_node),
            )
            .filter(Presupuesto.id == presupuesto_id, Presupuesto.empresa_id == empresa_id)
            .first()
        )
        if not presupuesto:
            raise ValueError(f"Presupuesto no encontrado: {presupuesto_id}")
        refresh_presupuesto_prices(db, presupuesto.id)
        db.refresh(presupuesto)
        return (
            db.query(Presupuesto)
            .options(
                selectinload(Presupuesto.indirectos),
                selectinload(Presupuesto.empresa),
                selectinload(Presupuesto.proyecto),
                selectinload(Presupuesto.detalle).selectinload(PresupuestoDetalle.apu),
                selectinload(Presupuesto.detalle).selectinload(PresupuestoDetalle.edt_node),
            )
            .filter(Presupuesto.id == presupuesto_id, Presupuesto.empresa_id == empresa_id)
            .first()
        )

    def _get_presupuesto_distinct_apu_ids(self, presupuesto: Any) -> List[int]:
        seen = set()
        apu_ids: List[int] = []
        for line in self._sort_budget_lines_for_report(list(presupuesto.detalle or [])):
            apu_id = getattr(line, "apu_id", None)
            if not apu_id or apu_id in seen:
                continue
            seen.add(apu_id)
            apu_ids.append(apu_id)
        return apu_ids

    def _get_presupuesto_apus_map(self, db: Any, presupuesto: Any, empresa_id: int) -> Dict[int, Any]:
        apu_ids = self._get_presupuesto_distinct_apu_ids(presupuesto)
        preloaded_map = {}
        missing_ids: List[int] = []

        for apu_id in apu_ids:
            matching_line = next((line for line in presupuesto.detalle if getattr(line, "apu_id", None) == apu_id), None)
            apu = getattr(matching_line, "apu", None) if matching_line else None
            if apu is None:
                missing_ids.append(apu_id)
                continue
            preloaded_map[apu_id] = apu

        if missing_ids:
            preloaded_map.update(self._get_apus_map(db, missing_ids, empresa_id))

        return {apu_id: preloaded_map[apu_id] for apu_id in apu_ids if apu_id in preloaded_map}

    def _resolve_apu_base_category(self, apu: Any) -> str:
        subcategoria = getattr(apu, "subcategoria_item", None)
        descripcion = getattr(subcategoria, "descripcion", None) if subcategoria else None
        return descripcion or "-"

    def _resolve_apu_line_category_label(self, linea: Any) -> str:
        category_id = self._resolve_apu_line_category(linea)
        return {
            1: "Equipos y Herramientas",
            2: "Materiales",
            3: "Transporte",
            4: "Mano de Obra",
        }.get(category_id, "-")

    def _resolve_apu_report_code(self, entity: Any) -> str:
        # Los reportes APU deben exponer siempre el código interno real del sistema.
        return str(getattr(entity, "codigo", None) or "").strip()

    def _clone_sheet(self, source_ws, target_ws) -> None:
        for row in source_ws.iter_rows():
            for cell in row:
                target_cell = target_ws.cell(row=cell.row, column=cell.column)
                target_cell.value = cell.value
                if cell.has_style:
                    target_cell.font = copy(cell.font)
                    target_cell.border = copy(cell.border)
                    target_cell.fill = copy(cell.fill)
                    target_cell.number_format = copy(cell.number_format)
                    target_cell.protection = copy(cell.protection)
                    target_cell.alignment = copy(cell.alignment)
        for col_letter, dim in source_ws.column_dimensions.items():
            target_ws.column_dimensions[col_letter].width = dim.width
            target_ws.column_dimensions[col_letter].hidden = dim.hidden
        for row_idx, dim in source_ws.row_dimensions.items():
            target_ws.row_dimensions[row_idx].height = dim.height
            target_ws.row_dimensions[row_idx].hidden = dim.hidden
        for merged_range in source_ws.merged_cells.ranges:
            target_ws.merge_cells(str(merged_range))
        if source_ws.sheet_view:
            target_ws.sheet_view.zoomScale = source_ws.sheet_view.zoomScale
            target_ws.sheet_view.showGridLines = source_ws.sheet_view.showGridLines
        target_ws.freeze_panes = source_ws.freeze_panes

    def _get_formula(self, db: Any, presupuesto_id: int) -> Any:
        from app.models.polinomica import FormulaPolinomica
        formula = db.query(FormulaPolinomica).filter(FormulaPolinomica.presupuesto_id == presupuesto_id).first()
        if not formula:
            raise ValueError("Fórmula polinómica no encontrada para este presupuesto")
        return formula

    def _build_apu_report_payload(self, apu: Any, money_decimals: int = 2, calc_decimals: int = 4, use_omniclass: bool = True) -> Dict[str, Any]:
        categoria_base = self._resolve_apu_base_category(apu)
        data = {
            "#CODIGO_APU": self._resolve_apu_report_code(apu),
            "#DESCRIPCION": self._normalize_report_description(apu.descripcion),
            "#CATEGORIA_BASE": categoria_base,
            "#UNIDAD": apu.unidad,
            "#TTOTAL": self._round_half_up(float(apu.precio_unitario_total or 0.0), money_decimals),
            "#COSTODIRECTO": self._round_half_up(float(apu.costo_directo or 0.0), money_decimals),
            "#COSTOINDIRECTO": self._round_half_up(float(apu.costo_indirecto or 0.0), money_decimals),
            "#%INDIRECTO": self._round_half_up(((float(apu.costo_indirecto or 0) / float(apu.costo_directo or 1)) * 100.0) if float(apu.costo_directo or 0) > 0 else 0.0, money_decimals),
            "#TOTALO": self._round_half_up(float(apu.precio_unitario_total or 0.0), money_decimals),
        }

        table_data = []
        categorized_preview_lines: Dict[int, List[Dict[str, Any]]] = {1: [], 2: [], 3: [], 4: []}
        for linea in sorted(apu.lineas, key=lambda item: ((item.orden or 0), item.id or 0)):
            codigo = ""
            desc = ""
            unidad = ""
            precio = 0.0
            recurso = getattr(linea, 'recurso', None)
            apu_hijo = getattr(linea, 'apu_hijo', None)

            if recurso:
                codigo = self._resolve_apu_report_code(recurso)
                desc = self._normalize_report_description(recurso.descripcion)
                unidad = recurso.unidad.descripcion if (recurso.unidad and hasattr(recurso.unidad, 'descripcion')) else ""
                precio = float(recurso.precio or 0.0)
                omniclass_cod = recurso.omniclass_codigo or "" if use_omniclass else ""
                omniclass_tit = recurso.omniclass_titulo or "" if use_omniclass else ""
            elif apu_hijo:
                codigo = self._resolve_apu_report_code(apu_hijo)
                desc = self._normalize_report_description(apu_hijo.descripcion)
                unidad = apu_hijo.unidad or ""
                precio = float(apu_hijo.precio_unitario_total or 0.0)
                omniclass_cod = apu_hijo.omniclass_codigo or "" if use_omniclass else ""
                omniclass_tit = apu_hijo.omniclass_titulo or "" if use_omniclass else ""
            else:
                desc = "--- ITEM NO VINCULADO ---"
                omniclass_cod = ""
                omniclass_tit = ""

            quantity = float(linea.cantidad or 0)
            rendimiento = float(linea.rendimiento or 1.0)
            subtotal = float(linea.subtotal or 0.0)
            percent = float((subtotal / float(apu.costo_directo) * 100) if (apu.costo_directo and float(apu.costo_directo) > 0) else 0)

            table_data.append({
                "ITEM": codigo,
                "DESCRIPCION": desc,
                "UNIDAD": unidad,
                "CANTIDAD": quantity,
                "PRECIO": precio,
                "RENDIMIENTO": rendimiento,
                "SUBTOTAL": subtotal,
                "PORCENT": percent,
                "OMNICLASS_COD": omniclass_cod,
                "OMNICLASS_TIT": omniclass_tit
            })
            category_id = self._resolve_apu_line_category(linea)
            categorized_preview_lines.setdefault(category_id, []).append({
                "codigo": codigo,
                "descripcion": desc,
                "unidad": unidad,
                "categoria_principal": self._resolve_apu_line_category_label(linea),
                "cantidad": quantity,
                "precio": precio,
                "rendimiento": rendimiento,
                "subtotal": subtotal,
                "is_structural": False,
                "cantidad_kind": "calc",
                "precio_kind": "money",
                "rendimiento_kind": "calc",
                "subtotal_kind": "money",
            })

        category_labels = {
            1: "Equipos y Herramientas",
            2: "Materiales",
            3: "Transporte",
            4: "Mano de Obra",
        }
        category_payload = self._build_apu_category_payload(apu, money_decimals, calc_decimals, use_omniclass)
        preview_lines = []
        for category_id in (1, 2, 3, 4):
            category_label = category_labels.get(category_id, "-")
            category_lines = categorized_preview_lines.get(category_id) or []
            if not category_lines:
                continue
            preview_lines.append({
                "codigo": "",
                "descripcion": category_label,
                "unidad": "",
                "cantidad": "",
                "precio": "",
                "rendimiento": "",
                "subtotal": "",
                "is_structural": True,
                "cantidad_kind": "raw",
                "precio_kind": "raw",
                "rendimiento_kind": "raw",
                "subtotal_kind": "raw",
            })
            preview_lines.extend(category_lines)
            category_summary = category_payload.get(category_id) or {}
            preview_lines.append({
                "codigo": "",
                "descripcion": f"Subtotal {category_label}",
                "unidad": "",
                "categoria_principal": category_label,
                "cantidad": "",
                "precio": "",
                "rendimiento": "",
                "subtotal": float(category_summary.get("total") or 0),
                "is_structural": True,
                "cantidad_kind": "raw",
                "precio_kind": "raw",
                "rendimiento_kind": "raw",
                "subtotal_kind": "money",
            })

        return {
            "data": data,
            "table_data": table_data,
            "preview": {
                "id": apu.id,
                "codigo": apu.codigo,
                "descripcion": self._normalize_report_description(apu.descripcion),
                "unidad": apu.unidad,
                "categoria_base": categoria_base,
                "costo_directo": float(apu.costo_directo or 0),
                "costo_indirecto": float(apu.costo_indirecto or 0),
                "precio_total": float(apu.precio_unitario_total or 0),
                "is_structural": False,
                "summary_direct_kind": "money",
                "summary_indirect_kind": "money",
                "summary_total_kind": "money",
                "lineas": preview_lines,
            }
        }

    def _build_apu_template_data(self, apu: Any, money_decimals: int = 2) -> Dict[str, Any]:
        categoria_base = self._resolve_apu_base_category(apu)
        return {
            "#CODIGO_APU": self._resolve_apu_report_code(apu),
            "#DESCRIPCION": self._normalize_report_description(apu.descripcion),
            "#CATEGORIA_BASE": categoria_base,
            "#UNIDAD": apu.unidad,
            "#TTOTAL": self._round_half_up(float(apu.precio_unitario_total or 0.0), money_decimals),
            "#COSTODIRECTO": self._round_half_up(float(apu.costo_directo or 0.0), money_decimals),
            "#COSTOINDIRECTO": self._round_half_up(float(apu.costo_indirecto or 0.0), money_decimals),
            "#%INDIRECTO": self._round_half_up(
                ((float(apu.costo_indirecto or 0) / float(apu.costo_directo or 1)) * 100.0)
                if float(apu.costo_directo or 0) > 0
                else 0.0,
                money_decimals,
            ),
            "#TOTALO": self._round_half_up(float(apu.precio_unitario_total or 0.0), money_decimals),
            "#TEXTOTOTAL": self._money_to_words_excel_usd_upper(float(apu.precio_unitario_total or 0.0)),
        }

    def _build_presupuesto_preview(self, presupuesto: Any) -> Dict[str, Any]:
        lineas = []
        for line in self._sort_budget_lines_for_report(list(presupuesto.detalle or [])):
            lineas.append({
                "item_visible": self._get_budget_line_visible_item_code(line),
                "edt_code_visible": self._get_budget_line_visible_edt_code(line),
                "descripcion": self._normalize_report_description(line.descripcion),
                "unidad": line.unidad or "",
                "cantidad": float(line.cantidad or 0),
                "precio": float(line.precio_unitario or 0),
                "rendimiento": "",
                "subtotal": float(line.precio_total or 0),
                "is_structural": getattr(line, "apu_id", None) is None,
                "cantidad_kind": "calc",
                "precio_kind": "money",
                "rendimiento_kind": "calc",
                "subtotal_kind": "money",
            })

        return {
            "id": presupuesto.id,
            "codigo": presupuesto.codigo or f"PRES-{presupuesto.id}",
            "descripcion": self._normalize_report_description(presupuesto.descripcion),
            "unidad": "Presupuesto",
            "costo_directo": float(presupuesto.subtotal or 0),
            "costo_indirecto": float(presupuesto.indirectos_total or 0),
            "precio_total": float(presupuesto.total or 0),
            "is_structural": True,
            "summary_direct_kind": "money",
            "summary_indirect_kind": "money",
            "summary_total_kind": "money",
            "table_columns": [
                { "key": "item_visible", "label": "Item" },
                { "key": "edt_code_visible", "label": "Cod EDT" },
                { "key": "descripcion", "label": "Descripción" },
                { "key": "unidad", "label": "Unidad" },
                { "key": "cantidad", "label": "Cant.", "kind": "calc" },
                { "key": "precio", "label": "Precio", "kind": "money" },
                { "key": "rendimiento", "label": "Rend." },
                { "key": "subtotal", "label": "Subtotal", "kind": "money" },
            ],
            "lineas": lineas,
        }

    def _resolve_indirecto_category_label(self, categoria_codigo: Any) -> str:
        code = str(categoria_codigo or "").strip()
        label = self.INDIRECTOS_CATEGORY_LABELS.get(code)
        return f"{code} · {label}" if label else (code or "Sin categoria")

    def _resolve_indirecto_type_label(self, indirecto: Any) -> str:
        if bool(getattr(indirecto, "custom", False)):
            return "Personalizado"
        if bool(getattr(indirecto, "fijo", False)):
            return "Fijo"
        if bool(getattr(indirecto, "usuario", False)):
            return "Usuario"
        return "Catalogo"

    def _resolve_presupuesto_indirectos_totals(self, presupuesto: Any, money_decimals: Optional[int] = None) -> Dict[str, Decimal]:
        money_decimals = int(money_decimals if money_decimals is not None else (getattr(presupuesto, "dec_moneda", 2) or 2))
        operative_without_tax = Decimal(str(getattr(presupuesto, "subtotal", 0) or 0))
        indirectos_porcentaje = Decimal(str(getattr(presupuesto, "indirectos_porcentaje", 0) or 0))
        iva_porcentaje = Decimal(str(getattr(presupuesto, "iva_aplicado", 0) or 0))
        indirectos_factor = Decimal("1") + (indirectos_porcentaje / Decimal("100"))

        if indirectos_factor > 0:
            base_calculo_raw = operative_without_tax / indirectos_factor
            base_calculo = round_decimal(base_calculo_raw, money_decimals)
        else:
            base_calculo_raw = operative_without_tax
            base_calculo = round_decimal(operative_without_tax, money_decimals)

        indirectos_total = round_decimal(base_calculo_raw * (indirectos_porcentaje / Decimal("100")), money_decimals)
        base_con_indirectos = base_calculo + indirectos_total
        iva_monto = round_decimal(base_con_indirectos * (iva_porcentaje / Decimal("100")), money_decimals)
        total_presupuesto = base_con_indirectos + iva_monto

        return {
            "base_calculo": base_calculo,
            "base_calculo_raw": base_calculo_raw,
            "indirectos_porcentaje": indirectos_porcentaje,
            "indirectos_total": indirectos_total,
            "base_con_indirectos": base_con_indirectos,
            "iva_porcentaje": iva_porcentaje,
            "iva_monto": iva_monto,
            "total_presupuesto": total_presupuesto,
        }

    def _build_presupuesto_indirectos_rows(self, presupuesto: Any) -> List[Dict[str, Any]]:
        money_decimals = int(getattr(presupuesto, "dec_moneda", 2) or 2)
        totals = self._resolve_presupuesto_indirectos_totals(presupuesto, money_decimals)
        base_calculo = totals["base_calculo"]
        base_calculo_raw = totals["base_calculo_raw"]
        configured_indirectos = list(getattr(presupuesto, "indirectos", None) or [])
        configured_indirectos.sort(
            key=lambda item: (
                str(getattr(item, "categoria_codigo", "") or ""),
                0 if bool(getattr(item, "fijo", False)) else 1,
                str(getattr(item, "nombre", "") or "").lower(),
            )
        )

        rows: List[Dict[str, Any]] = []
        current_category = None
        category_percent = Decimal("0")
        category_cost = Decimal("0")
        item_index = 1

        def append_category_total() -> None:
            nonlocal category_percent, category_cost, current_category
            if current_category is None:
                return
            rows.append({
                "item": "",
                "categoria": f"Subtotal {current_category}",
                "codigo": "",
                "concepto": "",
                "descripcion": "",
                "tipo": "",
                "porcentaje": self._round_half_up(category_percent, 4),
                "base_calculo": "",
                "costo_aplicado": self._round_half_up(category_cost, money_decimals),
                "observaciones": "",
                "is_structural": True,
                "ITEM": "",
                "CAT_INDIRECTOS": f"Subtotal {current_category}",
                "CODIGO_INDIRECTO": "",
                "CUENTA_INDIRECTOS": "",
                "TIPO_INDIRECTO": "",
                "PORCENT_ASIG": self._round_half_up(category_percent, 4),
                "BASE_CALCULO": "",
                "COSTO_APLICADO": self._round_half_up(category_cost, money_decimals),
                "OBSERVACIONES": "",
            })

        for indirecto in configured_indirectos:
            category = self._resolve_indirecto_category_label(getattr(indirecto, "categoria_codigo", None))
            if current_category is not None and category != current_category:
                append_category_total()
                category_percent = Decimal("0")
                category_cost = Decimal("0")
            current_category = category

            porcentaje = Decimal(str(getattr(indirecto, "porcentaje", 0) or 0))
            costo_aplicado = round_decimal(base_calculo_raw * (porcentaje / Decimal("100")), money_decimals)
            category_percent += porcentaje
            category_cost += costo_aplicado
            rows.append({
                "item": item_index,
                "categoria": category,
                "codigo": getattr(indirecto, "concepto_codigo", "") or "",
                "concepto": self._normalize_report_description(getattr(indirecto, "nombre", "") or ""),
                "descripcion": self._normalize_report_description(getattr(indirecto, "nombre", "") or ""),
                "tipo": self._resolve_indirecto_type_label(indirecto),
                "porcentaje": self._round_half_up(porcentaje, 4),
                "base_calculo": self._round_half_up(base_calculo, money_decimals),
                "costo_aplicado": self._round_half_up(costo_aplicado, money_decimals),
                "observaciones": self._normalize_report_description(getattr(indirecto, "observaciones", "") or ""),
                "is_structural": False,
                "ITEM": item_index,
                "CAT_INDIRECTOS": category,
                "CODIGO_INDIRECTO": getattr(indirecto, "concepto_codigo", "") or "",
                "CUENTA_INDIRECTOS": self._normalize_report_description(getattr(indirecto, "nombre", "") or ""),
                "TIPO_INDIRECTO": self._resolve_indirecto_type_label(indirecto),
                "PORCENT_ASIG": self._round_half_up(porcentaje, 4),
                "BASE_CALCULO": self._round_half_up(base_calculo, money_decimals),
                "COSTO_APLICADO": self._round_half_up(costo_aplicado, money_decimals),
                "OBSERVACIONES": self._normalize_report_description(getattr(indirecto, "observaciones", "") or ""),
            })
            item_index += 1

        append_category_total()
        if not rows:
            rows.append({
                "item": "",
                "categoria": "Sin indirectos configurados",
                "codigo": "",
                "concepto": "No existen cuentas de indirectos configuradas para este presupuesto.",
                "descripcion": "No existen cuentas de indirectos configuradas para este presupuesto.",
                "tipo": "",
                "porcentaje": 0,
                "base_calculo": self._round_half_up(base_calculo, money_decimals),
                "costo_aplicado": 0,
                "observaciones": "",
                "is_structural": True,
                "ITEM": "",
                "CAT_INDIRECTOS": "Sin indirectos configurados",
                "CODIGO_INDIRECTO": "",
                "CUENTA_INDIRECTOS": "No existen cuentas de indirectos configuradas para este presupuesto.",
                "TIPO_INDIRECTO": "",
                "PORCENT_ASIG": 0,
                "BASE_CALCULO": self._round_half_up(base_calculo, money_decimals),
                "COSTO_APLICADO": 0,
                "OBSERVACIONES": "",
            })
        return rows

    def _build_presupuesto_indirectos_preview(self, presupuesto: Any) -> Dict[str, Any]:
        money_decimals = int(getattr(presupuesto, "dec_moneda", 2) or 2)
        totals = self._resolve_presupuesto_indirectos_totals(presupuesto, money_decimals)
        base_calculo = totals["base_calculo"]
        indirectos_porcentaje = totals["indirectos_porcentaje"]
        indirectos_total = totals["indirectos_total"]
        base_con_indirectos = totals["base_con_indirectos"]
        iva_porcentaje = totals["iva_porcentaje"]
        iva_monto = totals["iva_monto"]
        total_presupuesto = totals["total_presupuesto"]
        rows = self._build_presupuesto_indirectos_rows(presupuesto)
        return {
            "id": presupuesto.id,
            "codigo": presupuesto.codigo or f"PRES-{presupuesto.id}",
            "descripcion": self._normalize_report_description(
                f"{self._resolve_project_title(presupuesto.proyecto, presupuesto.descripcion)} · Indirectos"
            ),
            "unidad": "Presupuesto",
            "metadata_hint": "Listado de indirectos configurados y costo aplicado sobre el total sin indirectos.",
            "costo_directo": float(base_calculo),
            "costo_indirecto": float(indirectos_total),
            "precio_total": float(total_presupuesto),
            "is_structural": True,
            "summary_layout": "single_line",
            "summary_cards": [
                {"label": "Total sin indirectos", "value": float(base_calculo), "kind": "money"},
                {"label": "% indirectos", "value": f"{self._format_fixed(indirectos_porcentaje, 4)}%"},
                {"label": "Costo indirectos", "value": float(indirectos_total), "kind": "money"},
                {"label": "Total presupuesto", "value": float(total_presupuesto), "kind": "money", "tone": "total"},
            ],
            "fields": [
                {"label": "Código del Proyecto", "value": getattr(presupuesto.proyecto, "codigo", "") if presupuesto.proyecto else ""},
                {"label": "Revisión", "value": self._format_revision_label(getattr(presupuesto, "revision", None))},
                {"label": "Código Referencial", "value": ""},
                {"label": "IVA aplicado", "value": f"{self._format_fixed(iva_porcentaje, 2)}%"},
                {"label": "Base con indirectos", "value": self._format_money_plain_label(base_con_indirectos, money_decimals)},
                {"label": "IVA calculado", "value": self._format_money_plain_label(iva_monto, money_decimals)},
            ],
            "table_columns": [
                {"key": "item", "label": "Item", "kind": "integer"},
                {"key": "categoria", "label": "Categoría"},
                {"key": "codigo", "label": "Código"},
                {"key": "descripcion", "label": "Concepto indirecto"},
                {"key": "tipo", "label": "Tipo"},
                {"key": "porcentaje", "label": "% aplicado", "kind": "calc"},
                {"key": "base_calculo", "label": "Base cálculo", "kind": "money"},
                {"key": "costo_aplicado", "label": "Costo aplicado", "kind": "money"},
                {"key": "observaciones", "label": "Observaciones"},
            ],
            "lineas": rows,
        }

    def _resolve_hierarchy_node_type(self, node: Any) -> str:
        raw_type = getattr(getattr(node, "tipo_nodo", None), "value", getattr(node, "tipo_nodo", "")) or ""
        return str(raw_type).replace("_", " ").strip().title()

    def _resolve_hierarchy_node_level(self, node: Any) -> int:
        code = str(getattr(node, "codigo", "") or "").strip()
        if not code:
            return 0
        return len([part for part in code.split(".") if part])

    def _resolve_hierarchy_node_role(self, node: Any) -> str:
        role = getattr(node, "rol", None)
        return str(getattr(role, "nombre", "") or "").strip()

    def _resolve_hierarchy_node_stakeholder_name(self, node: Any) -> str:
        stakeholder = getattr(node, "stakeholder", None)
        parts = [
            str(getattr(stakeholder, "nombre", "") or "").strip(),
            str(getattr(stakeholder, "apellidos", "") or "").strip(),
        ]
        return " ".join([part for part in parts if part]).strip()

    def _resolve_hierarchy_node_description(self, node: Any) -> str:
        node_type = getattr(getattr(node, "tipo_nodo", None), "value", getattr(node, "tipo_nodo", "")) or ""
        if str(node_type) == "STAKEHOLDER":
            return self._resolve_hierarchy_node_stakeholder_name(node) or "Responsable sin nombre"
        return (
            str(getattr(node, "nombre", "") or "").strip()
            or str(getattr(node, "descripcion", "") or "").strip()
            or str(getattr(node, "definicion", "") or "").strip()
        )

    def _is_hierarchy_stakeholder_node(self, node: Any) -> bool:
        node_type = getattr(getattr(node, "tipo_nodo", None), "value", getattr(node, "tipo_nodo", "")) or ""
        return str(node_type) == "STAKEHOLDER"

    def _resolve_hierarchy_stakeholder_parent_id(
        self,
        node: Any,
        structural_by_id: Dict[int, Any],
        structural_id_by_code: Dict[str, int],
    ) -> Optional[int]:
        parent_id = getattr(node, "parent_id", None)
        if parent_id in structural_by_id:
            return parent_id

        code = str(getattr(node, "codigo", "") or "").strip()
        candidates: List[str] = []
        if ".R" in code:
            candidates.append(code.split(".R", 1)[0])
        if "." in code:
            candidates.append(".".join(code.split(".")[:-1]))
        for candidate in candidates:
            candidate = candidate.strip()
            if candidate in structural_id_by_code:
                return structural_id_by_code[candidate]
        return None

    def _build_hierarchy_responsibility_context(self, nodes: List[Any]) -> Tuple[List[Any], Dict[int, Dict[str, str]]]:
        structural_nodes = [
            node for node in nodes
            if not self._is_hierarchy_stakeholder_node(node)
        ]
        structural_by_id = {
            getattr(node, "id", None): node
            for node in structural_nodes
            if getattr(node, "id", None) is not None
        }
        structural_id_by_code = {
            str(getattr(node, "codigo", "") or "").strip(): getattr(node, "id", None)
            for node in structural_nodes
            if getattr(node, "id", None) is not None and str(getattr(node, "codigo", "") or "").strip()
        }
        grouped: Dict[int, List[Any]] = {}
        for node in nodes:
            if not self._is_hierarchy_stakeholder_node(node):
                continue
            parent_id = self._resolve_hierarchy_stakeholder_parent_id(node, structural_by_id, structural_id_by_code)
            if parent_id is None:
                continue
            grouped.setdefault(parent_id, []).append(node)

        summaries: Dict[int, Dict[str, str]] = {}
        for parent_id, stakeholder_nodes in grouped.items():
            names = [
                self._resolve_hierarchy_node_stakeholder_name(node)
                for node in stakeholder_nodes
            ]
            roles = [
                self._resolve_hierarchy_node_role(node)
                for node in stakeholder_nodes
            ]
            codes = [
                str(getattr(getattr(node, "stakeholder", None), "codigo", "") or "").strip()
                for node in stakeholder_nodes
            ]
            activities = [
                str(getattr(node, "actividades_claves", "") or "").strip()
                for node in stakeholder_nodes
            ]
            summaries[parent_id] = {
                "nombre_responsable": self._join_report_parts(names, " / "),
                "rol_asignado": self._join_report_parts(roles, " / "),
                "codigo_stkr": self._join_report_parts(codes, " / "),
                "actividades_clave": self._join_report_parts(activities, " / "),
            }
        return structural_nodes, summaries

    def _build_edt_preview(self, db: Any, proyecto_id: int, empresa_id: int, report_type: str) -> Dict[str, Any]:
        from app.models.edt import EdtNode
        from app.models.proyecto import Proyecto

        proyecto = db.query(Proyecto).filter(Proyecto.id == proyecto_id, Proyecto.empresa_id == empresa_id).first()
        if not proyecto:
            raise ValueError("Proyecto no encontrado")

        project_detail = self._get_project_detail(db, proyecto, empresa_id)
        nodos = db.query(EdtNode).filter(EdtNode.proyecto_id == proyecto_id).order_by(EdtNode.codigo).all()
        variant = (report_type or "listado").lower()
        structural_nodes, responsible_by_parent_id = self._build_hierarchy_responsibility_context(nodos)
        lineas = []
        for index, nodo in enumerate(structural_nodes, start=1):
            responsible = responsible_by_parent_id.get(getattr(nodo, "id", None), {})
            lineas.append({
                "item": index,
                "codigo_edt": getattr(nodo, "codigo", "") or "",
                "descripcion_cuenta_paquete": self._normalize_report_description(
                    self._resolve_hierarchy_node_description(nodo)
                ),
                "nombre_responsable": responsible.get("nombre_responsable", ""),
                "codigo_stkr": responsible.get("codigo_stkr", ""),
                "definicion": getattr(nodo, "definicion", "") or getattr(nodo, "actividades_claves", "") or "",
                "subtotal_cuenta": 0,
                "is_structural": True,
            })

        columns_by_variant = {
            "diccionario": [
                {"key": "codigo_edt", "label": "Cod. EDT"},
                {"key": "descripcion_cuenta_paquete", "label": "Descripción Cuenta / Paquete de Trabajo"},
                {"key": "nombre_responsable", "label": "Nombre del Responsable"},
                {"key": "codigo_stkr", "label": "Código"},
                {"key": "definicion", "label": "Definición"},
            ],
            "valorada": [
                {"key": "codigo_edt", "label": "Cod. EDT"},
                {"key": "descripcion_cuenta_paquete", "label": "Descripción Cuenta / Paquete de Trabajo"},
                {"key": "subtotal_cuenta", "label": "P. Total", "kind": "money"},
            ],
            "listado": [
                {"key": "item", "label": "Ítem", "kind": "integer"},
                {"key": "codigo_edt", "label": "Cod. EDT"},
                {"key": "descripcion_cuenta_paquete", "label": "Descripción Cuenta / Paquete de Trabajo"},
            ],
        }
        variant_label = {
            "diccionario": "Diccionario",
            "valorada": "Valorada",
            "listado": "Listado",
        }.get(variant, "Listado")
        cuentas_count = len([
            n for n in nodos
            if getattr(getattr(n, "tipo_nodo", None), "value", getattr(n, "tipo_nodo", "")) == "CUENTA_PAQUETE"
        ])
        responsables_count = len([
            n for n in nodos
            if getattr(getattr(n, "tipo_nodo", None), "value", getattr(n, "tipo_nodo", "")) == "STAKEHOLDER"
        ])

        return {
            "id": proyecto.id,
            "codigo": proyecto.codigo or f"PRO-{proyecto.id}",
            "descripcion": self._normalize_report_description(f"{proyecto.nombre} · EDT {variant_label}"),
            "categoria_base": "Estructura de Descomposición del Trabajo",
            "unidad": f"EDT {variant_label}",
            "costo_directo": 0,
            "costo_indirecto": 0,
            "precio_total": len(structural_nodes),
            "is_structural": True,
            "summary_direct_kind": "integer",
            "summary_indirect_kind": "integer",
            "summary_total_kind": "integer",
            "summary_cards": [
                { "label": "Nodos", "value": len(nodos), "kind": "integer" },
                { "label": "Cuentas", "value": cuentas_count, "kind": "integer" },
                { "label": "Responsables", "value": responsables_count, "kind": "integer", "tone": "total" },
            ],
            "preview_layout": "edt_document",
            "document_title": f"Estructura de Descomposición del Trabajo - EDT ({variant_label})",
            "fields": [
                {"label": "Nombre del Proyecto", "value": self._resolve_project_title(proyecto)},
                {"label": "Revisión", "value": self._format_revision_label(getattr(proyecto, "revision", None))},
                {"label": "Código del Proyecto", "value": proyecto.codigo or ""},
                {"label": "Código Referencial", "value": getattr(project_detail, "cod_referencial", "") or ""},
                {"label": "Oferente", "value": proyecto.empresa.nombre if proyecto.empresa else ""},
                {"label": "Ubicación", "value": self._get_project_location(db, proyecto, empresa_id)},
                {"label": "Ciudad", "value": self._get_project_location_city(db, proyecto, empresa_id)},
                {"label": "Fecha", "value": proyecto.fecha_creacion.strftime("%d/%m/%Y") if proyecto.fecha_creacion else ""},
            ],
            "table_columns": columns_by_variant.get(variant, columns_by_variant["listado"]),
            "lineas": lineas,
            "metadata_hint": f"Documento EDT · {variant_label}",
        }

    def _build_edo_preview(self, db: Any, proyecto_id: int, empresa_id: int) -> Dict[str, Any]:
        from app.models.edo import EdoNode
        from app.models.proyecto import Proyecto

        proyecto = db.query(Proyecto).filter(Proyecto.id == proyecto_id, Proyecto.empresa_id == empresa_id).first()
        if not proyecto:
            raise ValueError("Proyecto no encontrado")

        project_detail = self._get_project_detail(db, proyecto, empresa_id)
        nodos = db.query(EdoNode).filter(EdoNode.proyecto_id == proyecto_id).order_by(EdoNode.codigo).all()
        structural_nodes, responsible_by_parent_id = self._build_hierarchy_responsibility_context(nodos)
        lineas = []
        for index, nodo in enumerate(structural_nodes, start=1):
            responsible = responsible_by_parent_id.get(getattr(nodo, "id", None), {})
            lineas.append({
                "item": index,
                "rol_asignado": responsible.get("rol_asignado", "") or "Hito",
                "nombre_responsable": responsible.get("nombre_responsable", "") or self._normalize_report_description(
                    self._resolve_hierarchy_node_description(nodo)
                ),
                "codigo_stkr": responsible.get("codigo_stkr", ""),
                "actividades_clave": responsible.get("actividades_clave", "") or getattr(nodo, "actividades_claves", "") or "",
                "is_structural": True,
            })

        hitos_count = len([
            n for n in nodos
            if getattr(getattr(n, "tipo_nodo", None), "value", getattr(n, "tipo_nodo", "")) == "HITO"
        ])
        responsables_count = len([
            n for n in nodos
            if getattr(getattr(n, "tipo_nodo", None), "value", getattr(n, "tipo_nodo", "")) == "STAKEHOLDER"
        ])

        return {
            "id": proyecto.id,
            "codigo": proyecto.codigo or f"PRO-{proyecto.id}",
            "descripcion": self._normalize_report_description(f"{proyecto.nombre} · EDO"),
            "categoria_base": "Estructura de Descomposición de la Organización",
            "unidad": "EDO",
            "costo_directo": 0,
            "costo_indirecto": 0,
            "precio_total": len(structural_nodes),
            "is_structural": True,
            "summary_direct_kind": "integer",
            "summary_indirect_kind": "integer",
            "summary_total_kind": "integer",
            "summary_cards": [
                { "label": "Nodos", "value": len(nodos), "kind": "integer" },
                { "label": "Hitos", "value": hitos_count, "kind": "integer" },
                { "label": "Responsables", "value": responsables_count, "kind": "integer", "tone": "total" },
            ],
            "preview_layout": "edo_document",
            "fields": [
                {"label": "Nombre del Proyecto", "value": self._resolve_project_title(proyecto)},
                {"label": "Revisión", "value": self._format_revision_label(getattr(proyecto, "revision", None))},
                {"label": "Código del Proyecto", "value": proyecto.codigo or ""},
                {"label": "Código Referencial", "value": getattr(project_detail, "cod_referencial", "") or ""},
                {"label": "Ciudad", "value": self._get_project_location_city(db, proyecto, empresa_id)},
                {"label": "Fecha", "value": proyecto.fecha_creacion.strftime("%d/%m/%Y") if proyecto.fecha_creacion else ""},
            ],
            "table_columns": [
                {"key": "item", "label": "Ítem", "kind": "integer"},
                {"key": "rol_asignado", "label": "Rol Asignado"},
                {"key": "nombre_responsable", "label": "Nombre del Responsable"},
                {"key": "actividades_clave", "label": "Actividades clave"},
            ],
            "lineas": lineas,
            "metadata_hint": "Documento EDO · Hitos y responsabilidades del proyecto",
        }

    def _build_vae_preview(self, db: Any, presupuesto: Any) -> Dict[str, Any]:
        project_total = float(presupuesto.total or 0)
        lineas = []
        project_vae_total = 0.0

        for line in self._sort_budget_lines_for_report(list(presupuesto.detalle or [])):
            vae_rubro = 0.0
            if line.apu:
                costo_directo = float(line.apu.costo_directo or 1)
                for linea_apu in sorted(line.apu.lineas, key=lambda item: ((item.orden or 0), item.id or 0)):
                    subtotal = float(linea_apu.subtotal or 0)
                    cpc_porc = 0.0
                    if linea_apu.recurso and linea_apu.recurso.cpc:
                        cpc_porc = float(linea_apu.recurso.cpc.porcentaje or 0)
                    vae_rubro += (subtotal / costo_directo) * (cpc_porc / 100)
            peso_relativo = float(line.precio_total or 0) / project_total if project_total > 0 else 0
            vae_ponderado = vae_rubro * peso_relativo
            project_vae_total += vae_ponderado
            lineas.append({
                "codigo": str(self._resolve_budget_line_report_item(line) or ""),
                "descripcion": self._normalize_report_description(line.descripcion),
                "unidad": line.unidad or "",
                "cantidad": float(line.cantidad or 0),
                "precio": float(line.precio_unitario or 0),
                "rendimiento": float(round_decimal(vae_rubro, 4)),
                "subtotal": float(round_decimal(vae_ponderado, 4)),
                "cantidad_kind": "calc",
                "precio_kind": "money",
                "rendimiento_kind": "calc",
                "subtotal_kind": "calc",
            })

        return {
            "id": presupuesto.id,
            "codigo": presupuesto.codigo or f"VAE-{presupuesto.id}",
            "descripcion": self._normalize_report_description(f"{presupuesto.descripcion} · VAE"),
            "unidad": "VAE",
            "costo_directo": project_total,
            "costo_indirecto": 0,
            "precio_total": project_vae_total,
            "summary_direct_kind": "money",
            "summary_indirect_kind": "money",
            "summary_total_kind": "calc",
            "lineas": lineas,
        }

    def _build_polinomica_preview(self, presupuesto: Any, formula: Any) -> Dict[str, Any]:
        base_calculo = float(formula.costo_directo_total or presupuesto.subtotal or 0)
        lineas = [{
            "codigo": m.simbolo,
            "descripcion": self._normalize_report_description(m.descripcion or m.indice_inec.descripcion if m.indice_inec else ""),
            "unidad": m.indice_inec.codigo if m.indice_inec else "",
            "cantidad": float(m.coeficiente or 0),
            "precio": base_calculo,
            "rendimiento": "",
            "subtotal": float(m.coeficiente or 0) * base_calculo,
            "cantidad_kind": "calc",
            "precio_kind": "money",
            "rendimiento_kind": "calc",
            "subtotal_kind": "money",
        } for m in sorted(formula.monomios, key=lambda x: x.simbolo)]

        return {
            "id": presupuesto.id,
            "codigo": presupuesto.codigo or f"POL-{presupuesto.id}",
            "descripcion": self._normalize_report_description(f"{presupuesto.descripcion} · Fórmula Polinómica"),
            "unidad": "Polinómica",
            "costo_directo": float(presupuesto.subtotal or 0),
            "costo_indirecto": float(presupuesto.indirectos_total or 0),
            "precio_total": float((presupuesto.subtotal or 0) + (presupuesto.indirectos_total or 0)),
            "summary_direct_kind": "money",
            "summary_indirect_kind": "money",
            "summary_total_kind": "money",
            "lineas": lineas,
        }

    def _apply_official_apu_report_overlay(
        self,
        db: Any,
        apu: Any,
        empresa_id: int,
        *,
        project_id: Optional[int] = None,
        base_trabajo_id: Optional[int] = None,
        revision: Optional[int] = None,
    ) -> Any:
        if not project_id or not apu:
            return apu
        from app.services.project_functional_modification import project_functional_modification_service

        overrides, _official_source = project_functional_modification_service.resolve_project_apu_price_overrides(
            db,
            empresa_id=empresa_id,
            proyecto_id=project_id,
            base_trabajo_id=base_trabajo_id or getattr(apu, "base_trabajo_id", None),
            revision=revision if revision is not None else getattr(apu, "revision", None),
        )
        override = overrides.get(int(getattr(apu, "id", 0) or 0))
        if override is None:
            return apu
        report_apu = copy(apu)
        report_apu.precio_unitario_total = override
        return report_apu

    def generate_apu_report(
        self,
        db: Any,
        apu_id: int,
        empresa_id: int,
        template_id: str = "001",
        *,
        project_id: Optional[int] = None,
        base_trabajo_id: Optional[int] = None,
        revision: Optional[int] = None,
    ) -> io.BytesIO:
        apu = self._get_apu(db, apu_id, empresa_id)
        apu = self._apply_official_apu_report_overlay(
            db,
            apu,
            empresa_id,
            project_id=project_id,
            base_trabajo_id=base_trabajo_id,
            revision=revision,
        )
        format_config = self._get_empresa_format_config(db, empresa_id)
        use_omniclass = self._resolve_use_omniclass(format_config, template_id)
        template_data = self._build_apu_template_data(apu, format_config["money_decimals"])
        template_metadata = self._get_apu_template_metadata(template_id)
        filename = self._resolve_apu_template_filename(template_id)
        template_path = os.path.join(self.templates_dir, filename)
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Plantilla no encontrada: {template_path}")
        wb = openpyxl.load_workbook(template_path)
        ws = wb.active
        ws.title = self._sanitize_sheet_title(apu.codigo, "APU")
        self._apply_apu_category_template_with_metadata(
            ws,
            apu,
            template_data,
            template_metadata,
            format_config["money_decimals"],
            format_config["calc_decimals"],
            use_omniclass,
        )
        self._apply_omniclass_visibility(ws, use_omniclass)
        return self._save_workbook_buffer(wb)

    def generate_apu_report_bundle(
        self,
        db: Any,
        apu_ids: List[int],
        empresa_id: int,
        template_id: str = "001",
        *,
        project_id: Optional[int] = None,
        base_trabajo_id: Optional[int] = None,
        revision: Optional[int] = None,
    ) -> io.BytesIO:
        if not apu_ids:
            raise ValueError("No se recibieron APUs para exportar.")
        format_config = self._get_empresa_format_config(db, empresa_id)
        use_omniclass = self._resolve_use_omniclass(format_config, template_id)
        apus_map = self._get_apus_map(db, apu_ids, empresa_id)
        template_metadata = self._get_apu_template_metadata(template_id)

        filename = self._resolve_apu_template_filename(template_id)
        template_path = os.path.join(self.templates_dir, filename)
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Plantilla no encontrada: {template_path}")

        wb = openpyxl.load_workbook(template_path)
        template_ws = wb.active
        generated_sheets = []

        for idx, apu_id in enumerate(apu_ids, start=1):
            apu = apus_map.get(apu_id)
            if not apu:
                raise ValueError(f"APU no encontrado: {apu_id}")
            apu = self._apply_official_apu_report_overlay(
                db,
                apu,
                empresa_id,
                project_id=project_id,
                base_trabajo_id=base_trabajo_id,
                revision=revision,
            )
            template_data = self._build_apu_template_data(apu, format_config["money_decimals"])
            ws = wb.copy_worksheet(template_ws)
            ws.title = self._sanitize_sheet_title(apu.codigo, f"APU {idx}")
            self._apply_apu_category_template_with_metadata(
                ws,
                apu,
                template_data,
                template_metadata,
                format_config["money_decimals"],
                format_config["calc_decimals"],
                use_omniclass,
            )
            self._apply_omniclass_visibility(ws, use_omniclass)
            generated_sheets.append(ws)

        wb.remove(template_ws)
        if generated_sheets:
            wb.active = 0

        if not use_omniclass:
            for ws in wb.worksheets:
                self._apply_omniclass_visibility(ws, False)

        return self._save_workbook_buffer(wb)

    def generate_presupuesto_report_bundle(self, db: Any, presupuesto_id: int, empresa_id: int, template_id: str = "001") -> io.BytesIO:
        presupuesto = self._get_presupuesto(db, presupuesto_id, empresa_id)
        apus_map = self._get_presupuesto_apus_map(db, presupuesto, empresa_id)
        apu_ids = list(apus_map.keys())
        if not apu_ids:
            return self._generate_presupuesto_report_from_entity(db, presupuesto, empresa_id, template_id)

        format_config = self._get_empresa_format_config(db, empresa_id)
        use_omniclass = self._resolve_use_omniclass(format_config, template_id)
        cache_key = self._build_presupuesto_bundle_cache_key(
            presupuesto,
            empresa_id,
            template_id,
            use_omniclass,
            apus_map,
            official_source_signature=self._build_presupuesto_official_source_signature(db, presupuesto, empresa_id),
        )
        cached_payload = self._get_cached_bytes(self._workbook_bytes_cache, cache_key)
        if cached_payload is not None:
            return self._buffer_from_cached_bytes(cached_payload)

        presupuesto_buffer = self._generate_presupuesto_report_from_entity(db, presupuesto, empresa_id, template_id)
        wb = openpyxl.load_workbook(presupuesto_buffer)

        template_metadata = self._get_apu_template_metadata(template_id)
        apu_template_filename = self._resolve_apu_template_filename(template_id)
        apu_template_path = os.path.join(self.templates_dir, apu_template_filename)
        if not os.path.exists(apu_template_path):
            raise FileNotFoundError(f"Plantilla no encontrada: {apu_template_path}")

        apu_template_wb = openpyxl.load_workbook(apu_template_path)
        apu_template_ws = apu_template_wb.active
        seed_ws = wb.create_sheet(title="_apu_seed_template")
        self._clone_sheet(apu_template_ws, seed_ws)

        for idx, apu_id in enumerate(apu_ids, start=1):
            apu = apus_map.get(apu_id)
            if not apu:
                raise ValueError(f"APU no encontrado: {apu_id}")
            apu = self._apply_official_apu_report_overlay(
                db,
                apu,
                empresa_id,
                project_id=getattr(presupuesto, "proyecto_id", None),
                base_trabajo_id=getattr(getattr(presupuesto, "proyecto", None), "base_trabajo_id", None),
                revision=getattr(presupuesto, "revision", None),
            )
            template_data = self._build_apu_template_data(apu, format_config["money_decimals"])
            ws = wb.copy_worksheet(seed_ws)
            ws.title = self._sanitize_sheet_title(apu.codigo, f"APU {idx}")
            self._apply_apu_category_template_with_metadata(
                ws,
                apu,
                template_data,
                template_metadata,
                format_config["money_decimals"],
                format_config["calc_decimals"],
                use_omniclass,
            )
            self._apply_omniclass_visibility(ws, use_omniclass)

        wb.remove(seed_ws)

        if not use_omniclass:
            for ws in wb.worksheets:
                self._apply_omniclass_visibility(ws, False)

        output = self._save_workbook_buffer(wb)
        self._set_cached_bytes(
            self._workbook_bytes_cache,
            cache_key,
            output.getvalue(),
            self._workbook_bytes_cache_limit,
        )
        output.seek(0)
        return output

    def preview_report(
        self,
        db: Any,
        report_type: str,
        entity_ids: List[int],
        empresa_id: int,
        template_id: str = "001",
        variant: Optional[str] = None,
        filters: Optional[Dict[str, Any]] = None,
        project_id: Optional[int] = None,
        base_trabajo_id: Optional[int] = None,
        revision: Optional[int] = None,
    ) -> Dict[str, Any]:
        items = []
        title = "Reporte"

        if report_type == "apu":
            title = "Reporte de APUs"
            format_config = self._get_empresa_format_config(db, empresa_id)
            use_omniclass = self._resolve_use_omniclass(format_config, template_id)
            apus_map = self._get_apus_map(db, entity_ids, empresa_id)
            for apu_id in entity_ids:
                apu = apus_map.get(apu_id)
                if not apu:
                    raise ValueError(f"APU no encontrado: {apu_id}")
                apu = self._apply_official_apu_report_overlay(
                    db,
                    apu,
                    empresa_id,
                    project_id=project_id,
                    base_trabajo_id=base_trabajo_id,
                    revision=revision,
                )
                items.append(
                    self._build_apu_report_payload(
                        apu,
                        format_config["money_decimals"],
                        format_config["calc_decimals"],
                        use_omniclass,
                    )["preview"]
                )
        elif report_type == "presupuesto":
            title = "Reporte de Presupuesto"
            presupuesto = self._get_presupuesto(db, entity_ids[0], empresa_id)
            if variant == "indirectos":
                title = "Reporte de Indirectos del Presupuesto"
                items.append(self._build_presupuesto_indirectos_preview(presupuesto))
            else:
                items.append(self._build_presupuesto_preview(presupuesto))
            if variant == "with_apus":
                title = "Reporte de Presupuesto + APUs"
                format_config = self._get_empresa_format_config(db, empresa_id)
                use_omniclass = self._resolve_use_omniclass(format_config, template_id)
                apus_map = self._get_presupuesto_apus_map(db, presupuesto, empresa_id)
                for apu_id in self._get_presupuesto_distinct_apu_ids(presupuesto):
                    apu = apus_map.get(apu_id)
                    if not apu:
                        raise ValueError(f"APU no encontrado: {apu_id}")
                    apu = self._apply_official_apu_report_overlay(
                        db,
                        apu,
                        empresa_id,
                        project_id=getattr(presupuesto, "proyecto_id", None),
                        base_trabajo_id=getattr(getattr(presupuesto, "proyecto", None), "base_trabajo_id", None),
                        revision=getattr(presupuesto, "revision", None),
                    )
                    items.append(
                        self._build_apu_report_payload(
                            apu,
                            format_config["money_decimals"],
                            format_config["calc_decimals"],
                            use_omniclass,
                        )["preview"]
                    )
        elif report_type == "cronograma_valorado":
            report_variant = str(variant or "valorado").lower()
            gantt_draft_summary = self._summarize_gantt_draft_for_reporting(db, entity_ids[0], empresa_id)
            if report_variant in {"cash_flow", "flujo_caja", "caja"}:
                title = "Reporte de Flujo de Caja"
                items.append(self._build_cronograma_cash_flow_preview(db, entity_ids[0], empresa_id))
            elif report_variant == "pareto":
                title = "Reporte de Pareto Temporal"
                items.append(self._build_cronograma_pareto_preview(db, entity_ids[0], empresa_id))
            elif report_variant == "gantt":
                title = "Reporte de Cronograma Gantt"
                items.append(self._build_cronograma_gantt_preview(db, entity_ids[0], empresa_id))
            elif report_variant == "integrado":
                title = "Reporte Integrado Gantt / Valorado / Caja"
                items.append(self._build_cronograma_gantt_preview(db, entity_ids[0], empresa_id))
                items.append(self._build_cronograma_valorado_preview(db, entity_ids[0], empresa_id))
                items.append(self._build_cronograma_cash_flow_preview(db, entity_ids[0], empresa_id))
            elif self._is_cronograma_resource_usage_variant(report_variant):
                title = "Reporte de Uso de Recursos por Rango" if self._is_cronograma_resource_usage_range_variant(report_variant) else "Reporte de Uso de Recursos"
                items.append(self._build_cronograma_resource_usage_payload(db, entity_ids[0], empresa_id, filters if self._is_cronograma_resource_usage_range_variant(report_variant) else None))
            else:
                title = "Reporte de Cronograma Valorado"
                items.append(self._build_cronograma_valorado_preview(db, entity_ids[0], empresa_id))
            items = [
                self._attach_gantt_draft_report_warning(item, gantt_draft_summary)
                for item in items
            ]
        elif report_type == "edt":
            report_variant = variant or "listado"
            title = f"Reporte EDT · {report_variant}"
            items.append(self._build_edt_preview(db, entity_ids[0], empresa_id, report_variant))
        elif report_type == "edo":
            title = "Reporte EDO"
            items.append(self._build_edo_preview(db, entity_ids[0], empresa_id))
        elif report_type == "vae":
            title = "Reporte VAE"
            presupuesto = self._get_presupuesto(db, entity_ids[0], empresa_id)
            items.append(self._build_vae_preview(db, presupuesto))
        elif report_type == "polinomica":
            title = "Reporte Fórmula Polinómica"
            presupuesto = self._get_presupuesto(db, entity_ids[0], empresa_id)
            formula = self._get_formula(db, presupuesto.id)
            items.append(self._build_polinomica_preview(presupuesto, formula))
        elif report_type == "acta_constitucion":
            from app.models.proyecto import Proyecto

            title = "Acta de Constitución del Proyecto"
            proyecto = (
                db.query(Proyecto)
                .filter(Proyecto.id == entity_ids[0], Proyecto.empresa_id == empresa_id)
                .first()
            )
            if not proyecto:
                raise ValueError(f"Proyecto no encontrado: {entity_ids[0]}")
            items.append(self._build_acta_constitucion_preview(db, proyecto, empresa_id))
        elif report_type == "stakeholders":
            from app.models.proyecto import Proyecto

            title = "Equipo del Proyecto (Stakeholders)"
            proyecto = (
                db.query(Proyecto)
                .filter(Proyecto.id == entity_ids[0], Proyecto.empresa_id == empresa_id)
                .first()
            )
            if not proyecto:
                raise ValueError(f"Proyecto no encontrado: {entity_ids[0]}")
            items.append(self._build_stakeholders_preview(db, proyecto, empresa_id))
        else:
            raise ValueError(f"Vista previa no soportada para: {report_type}")

        return {
            "report_type": report_type,
            "template_id": template_id,
            "variant": variant,
            "filters": filters,
            "title": title,
            "selection_count": len(items),
            "items": items,
        }

    def generate_preview_pdf(
        self,
        db: Any,
        report_type: str,
        entity_ids: List[int],
        empresa_id: int,
        template_id: str = "001",
        variant: Optional[str] = None,
        filters: Optional[Dict[str, Any]] = None,
        watermark_text: Optional[str] = None,
        project_id: Optional[int] = None,
        base_trabajo_id: Optional[int] = None,
        revision: Optional[int] = None,
    ) -> io.BytesIO:
        if report_type == "cronograma_valorado" and self._is_cronograma_resource_usage_variant(variant):
            resource_filters = filters if self._is_cronograma_resource_usage_range_variant(variant) else None
            gantt_draft_summary = self._summarize_gantt_draft_for_reporting(db, entity_ids[0], empresa_id)
            resource_payload = self._attach_gantt_draft_report_warning(
                self._build_cronograma_resource_usage_payload(db, entity_ids[0], empresa_id, resource_filters),
                gantt_draft_summary,
            )
            return self._build_cronograma_resource_usage_executive_pdf(
                db,
                entity_ids[0],
                empresa_id,
                resource_filters,
                watermark_text,
                payload_override=resource_payload,
            )
        preview = self.preview_report(
            db,
            report_type,
            entity_ids,
            empresa_id,
            template_id,
            variant,
            filters,
            project_id=project_id,
            base_trabajo_id=base_trabajo_id,
            revision=revision,
        )
        if report_type in {"edo", "edt"}:
            return self._build_hierarchy_document_pdf(preview, watermark_text)
        format_config = self._get_empresa_format_config(db, empresa_id)
        return self._build_pdf_report(
            preview.get("title") or "Reporte",
            preview.get("items") or [],
            format_config["money_decimals"],
            format_config["calc_decimals"],
            watermark_text,
        )

    def generate_edo_report(self, db: Any, proyecto_id: int, empresa_id: int) -> io.BytesIO:
        from app.models.edo import EdoNode
        from app.models.proyecto import Proyecto

        proyecto = db.query(Proyecto).filter(Proyecto.id == proyecto_id, Proyecto.empresa_id == empresa_id).first()
        if not proyecto:
            raise ValueError("Proyecto no encontrado")

        format_config = self._get_empresa_format_config(db, empresa_id)
        project_detail = self._get_project_detail(db, proyecto, empresa_id)
        data = {
            "#REVISION": self._format_revision_label(getattr(proyecto, "revision", None)),
            "#PRO_TITULO": self._resolve_project_title(proyecto),
            "#CODIGOPROYECTO": getattr(proyecto, "codigo", "") or "",
            "#CODIGOREFERENCIAL": getattr(project_detail, "cod_referencial", "") or "",
            "#CIUDAD": self._get_project_location_city(db, proyecto, empresa_id),
            "#PRO_FECHA": proyecto.fecha_creacion.strftime("%d/%m/%Y") if proyecto.fecha_creacion else "",
        }

        nodos = db.query(EdoNode).filter(EdoNode.proyecto_id == proyecto_id).order_by(EdoNode.codigo).all()
        structural_nodes, responsible_by_parent_id = self._build_hierarchy_responsibility_context(nodos)
        table_data = []
        for index, nodo in enumerate(structural_nodes, start=1):
            responsible = responsible_by_parent_id.get(getattr(nodo, "id", None), {})
            table_data.append({
                "ITEM": index,
                "ROL": responsible.get("rol_asignado", "") or self._resolve_hierarchy_node_type(nodo),
                "NOMBRE_RESPONSABLE": responsible.get("nombre_responsable", "") or self._resolve_hierarchy_node_description(nodo),
                "CODIGO_STKR": responsible.get("codigo_stkr", ""),
                "ACTIVIDADES_CLAVE": responsible.get("actividades_clave", "") or getattr(nodo, "actividades_claves", "") or "",
            })

        formatted_rows = self._format_excel_rows(
            table_data,
            {"ITEM": "integer"},
            format_config["money_decimals"],
            format_config["calc_decimals"],
        )
        report_buffer = self.fill_template(
            "001 - EDO.xlsx",
            data,
            formatted_rows,
            table_row_marker="#ITEM",
            use_omniclass=format_config["use_omniclass"],
        )
        workbook = openpyxl.load_workbook(report_buffer)
        self._normalize_edo_document_header(workbook.active)
        self._normalize_edo_signature_block(workbook.active)
        self._normalize_edo_document_columns(workbook.active)
        return self._save_workbook_buffer(workbook)

    def generate_edt_report(self, db: Any, proyecto_id: int, empresa_id: int, report_type: str = "listado") -> io.BytesIO:
        """
        Genera reportes de EDT: listado, diccionario o valorada.
        """
        from app.models.edt import EdtNode
        from app.models.proyecto import Proyecto
        
        proyecto = db.query(Proyecto).filter(Proyecto.id == proyecto_id, Proyecto.empresa_id == empresa_id).first()
        if not proyecto:
            raise ValueError("Proyecto no encontrado")

        # Determinar plantilla y datos
        filename = ""
        table_row_marker = "#CODIGOEDT"
        
        if report_type == "listado":
            filename = "001 - EDT - listado.xlsx"
        elif report_type == "diccionario":
            filename = "001 - EDT - Diccionario.xlsx"
        elif report_type == "valorada":
            filename = "001 - EDT - Valorada.xlsx"
        else:
            raise ValueError(f"Tipo de reporte EDT no válido: {report_type}")

        format_config = self._get_empresa_format_config(db, empresa_id)

        data = {
            "#REVISION": self._format_revision_label(getattr(proyecto, "revision", None)),
            "#PRO_TITULO": self._resolve_project_title(proyecto),
            "#CODIGOPROYECTO": proyecto.codigo or "",
            "#CODIGOREFERENCIAL": getattr(self._get_project_detail(db, proyecto, empresa_id), "cod_referencial", "") or "",
            "#PRO_OFERENTE": proyecto.empresa.nombre if proyecto.empresa else "",
            "#PRO_UBICACION": self._get_project_location(db, proyecto, empresa_id),
            "#CIUDAD": self._get_project_location_city(db, proyecto, empresa_id),
            "#PRO_FECHA": proyecto.fecha_creacion.strftime("%d/%m/%Y") if proyecto.fecha_creacion else ""
        }

        # Obtener nodos EDT
        nodos = db.query(EdtNode).filter(EdtNode.proyecto_id == proyecto_id).order_by(EdtNode.codigo).all()
        structural_nodes, responsible_by_parent_id = self._build_hierarchy_responsibility_context(nodos)
        
        table_data = []
        for index, nodo in enumerate(structural_nodes, start=1):
            responsible = responsible_by_parent_id.get(getattr(nodo, "id", None), {})
            row = {
                "ITEM": index,
                "CODIGOEDT": nodo.codigo or "",
                "DESCRIPCION_CUENTA_PAQUETE": self._normalize_report_description(self._resolve_hierarchy_node_description(nodo)),
                "NOMBRE_RESPONSABLE": responsible.get("nombre_responsable", ""),
                "CODIGO_STKR": responsible.get("codigo_stkr", ""),
                "DEFINICION": getattr(nodo, "definicion", "") or getattr(nodo, "actividades_claves", "") or "",
                "SUBTOTAL_CUENTA": 0,
            }
            table_data.append(row)

        if report_type == "valorada":
            data["#TEXTOTOTAL"] = ""

        formatted_data = self._format_excel_data(
            data,
            {
                "#SUBTOTAL_CUENTA": "money",
            },
            format_config["money_decimals"],
            format_config["calc_decimals"],
        )
        formatted_rows = self._format_excel_rows(
            table_data,
            {
                "ITEM": "integer",
                "SUBTOTAL_CUENTA": "money",
            },
            format_config["money_decimals"],
            format_config["calc_decimals"],
        )

        report_buffer = self.fill_template(
            filename,
            formatted_data,
            formatted_rows,
            table_row_marker=table_row_marker,
            use_omniclass=format_config["use_omniclass"],
        )
        if report_type in {"listado", "diccionario", "valorada"}:
            workbook = openpyxl.load_workbook(report_buffer)
            self._normalize_edt_document_columns(workbook.active, report_type)
            if report_type == "listado":
                self._render_edt_listado_document_blocks(workbook.active, table_data)
            self._normalize_edt_signature_block(workbook.active)
            return self._save_workbook_buffer(workbook)
        return report_buffer

    def _generate_presupuesto_report_from_entity(self, db: Any, presupuesto: Any, empresa_id: int, template_id: str = "001") -> io.BytesIO:
        # 001 - Presupuesto.xlsx
        filename = "001 - Presupuesto.xlsx"
        if template_id == "002":
            filename = "002 - Presupuesto - SERCOP.xlsx"

        format_config = self._get_empresa_format_config(db, empresa_id)
        use_omniclass = self._resolve_use_omniclass(format_config, template_id)
        reference_code = self._get_project_reference_code(db, presupuesto)

        subtotal_presupuesto = float(presupuesto.subtotal or 0)
        iva_porcentaje = float(presupuesto.iva_aplicado or 0)
        iva_reporte = subtotal_presupuesto * (iva_porcentaje / 100.0)
        total_reporte = subtotal_presupuesto + iva_reporte

        data = {
            "#PRO_TITULO": self._resolve_project_title(presupuesto.proyecto, presupuesto.descripcion),
            "#CODIGOPROYECTO": presupuesto.proyecto.codigo if presupuesto.proyecto else "",
            "#CODIGOREFERENCIAL": reference_code,
            "#REVISION": self._format_revision_label(presupuesto.revision),
            "#PRO_OFERENTE": presupuesto.empresa.nombre if presupuesto.empresa else "",
            "#PRO_UBICACION": self._get_project_location(db, presupuesto.proyecto, empresa_id),
            "#PRO_FECHA": presupuesto.fecha_creacion.strftime("%d/%m/%Y") if presupuesto.fecha_creacion else "",
            "#%IVA": f"{self._format_fixed(float(presupuesto.iva_aplicado or 0), format_config['money_decimals'])}%",
            "#SUBTOTALSINIVA": subtotal_presupuesto,
            "#IVA": iva_reporte,
            "#TTOTAL": total_reporte,
            "#TEXTOTOTAL": self._money_to_words_body_es(total_reporte)
        }

        table_data = self._build_presupuesto_table_data(presupuesto, template_id)

        formatted_data = self._format_excel_data(
            data,
            {
                "#SUBTOTALSINIVA": "money",
                "#IVA": "money",
                "#TTOTAL": "money",
            },
            format_config["money_decimals"],
            format_config["calc_decimals"],
        )
        formatted_rows = self._format_excel_rows(
            table_data,
            {
                "CANTIDAD": "calc",
                "PUNITARIO": "money",
                "SUBTOTAL": "money",
            },
            format_config["money_decimals"],
            format_config["calc_decimals"],
        )
        output = self.fill_template(
            filename,
            formatted_data,
            formatted_rows,
            table_row_marker="#DESCRIPCION",
            use_omniclass=use_omniclass,
        )
        output.seek(0)
        wb = openpyxl.load_workbook(output)
        ws = wb.active
        self._apply_presupuesto_row_typography(ws, table_data, template_id)
        if not use_omniclass:
            for current_ws in wb.worksheets:
                self._apply_omniclass_visibility(current_ws, False)
        return self._save_workbook_buffer(wb)

    def generate_presupuesto_report(self, db: Any, presupuesto_id: int, empresa_id: int, template_id: str = "001") -> io.BytesIO:
        presupuesto = self._get_presupuesto(db, presupuesto_id, empresa_id)
        return self._generate_presupuesto_report_from_entity(db, presupuesto, empresa_id, template_id)

    def generate_presupuesto_indirectos_report(self, db: Any, presupuesto_id: int, empresa_id: int, template_id: str = "001") -> io.BytesIO:
        presupuesto = self._get_presupuesto(db, presupuesto_id, empresa_id)
        filename = "001 - Porcentajes de Indirectos.xlsx"
        format_config = self._get_empresa_format_config(db, empresa_id)
        money_decimals = int(format_config.get("money_decimals", getattr(presupuesto, "dec_moneda", 2) or 2))
        rows = self._build_presupuesto_indirectos_rows(presupuesto)
        configured_rows = [row for row in rows if not row.get("is_structural")]
        totals = self._resolve_presupuesto_indirectos_totals(presupuesto, money_decimals)
        base_calculo = totals["base_calculo"]
        indirectos_porcentaje = totals["indirectos_porcentaje"]
        indirectos_total = totals["indirectos_total"]
        base_con_indirectos = totals["base_con_indirectos"]
        iva_porcentaje = totals["iva_porcentaje"]
        iva_monto = totals["iva_monto"]
        total_presupuesto = totals["total_presupuesto"]
        reference_code = self._get_project_reference_code(db, presupuesto)
        data = {
            "#PRO_TITULO": self._resolve_project_title(presupuesto.proyecto, presupuesto.descripcion),
            "#CODIGOPROYECTO": presupuesto.proyecto.codigo if presupuesto.proyecto else "",
            "#CODIGOREFERENCIAL": reference_code,
            "#REVISION": self._format_revision_label(getattr(presupuesto, "revision", None)),
            "#PRO_OFERENTE": presupuesto.empresa.nombre if presupuesto.empresa else "",
            "#PRO_UBICACION": self._get_project_location(db, presupuesto.proyecto, empresa_id),
            "#PRO_FECHA": date.today().strftime("%d/%m/%Y"),
            "#CIUDAD": self._get_project_location(db, presupuesto.proyecto, empresa_id),
            "#TOTAL_DIRECTO": self._round_half_up(base_calculo, money_decimals),
            "#TOTAL_PORCENT_ASIG": self._round_half_up(indirectos_porcentaje, 4),
            "#TOTAL_INDIRECTOS": self._round_half_up(indirectos_total, money_decimals),
            "#SUBTOTAL_CON_INDIRECTOS": self._round_half_up(base_con_indirectos, money_decimals),
            "#IVA_PORCENTAJE": self._round_half_up(iva_porcentaje, 2),
            "#IVA_MONTO": self._round_half_up(iva_monto, money_decimals),
            "#TOTAL_PRESUPUESTO": self._round_half_up(total_presupuesto, money_decimals),
            "#TOTAL_CUENTAS": len(configured_rows),
        }
        return self.fill_template(
            filename,
            data,
            rows,
            table_row_marker="#ITEM",
            use_omniclass=False,
        )

    def _get_cronograma_valorado_payload(self, db: Any, presupuesto_id: int, empresa_id: int) -> Tuple[Any, Any]:
        from app.api.endpoints.cronogramas import _get_or_create_cronograma, _serialize_cronograma

        presupuesto = self._get_presupuesto(db, presupuesto_id, empresa_id)
        cronograma = _get_or_create_cronograma(db, presupuesto)
        return presupuesto, _serialize_cronograma(db, presupuesto, cronograma)

    def _format_cronograma_period_values(self, values: List[Any], decimals: int = 2, suffix: str = "") -> str:
        return " | ".join(f"{self._format_fixed(value, decimals)}{suffix}" for value in values)

    def _is_cronograma_resource_usage_variant(self, variant: Optional[str]) -> bool:
        return str(variant or "").strip().lower() in {
            "resources",
            "resource_usage",
            "uso_recursos",
            "uso_de_recursos",
            "resources_range",
            "resource_usage_range",
            "uso_recursos_rango",
        }

    def _is_cronograma_resource_usage_range_variant(self, variant: Optional[str]) -> bool:
        return str(variant or "").strip().lower() in {
            "resources_range",
            "resource_usage_range",
            "uso_recursos_rango",
        }

    def _decimal_value(self, value: Any, default: str = "0") -> Decimal:
        if value in (None, ""):
            return Decimal(default)
        try:
            return Decimal(str(value))
        except Exception:
            return Decimal(default)

    def _normalize_report_datetime(self, value: Any, end_of_day: bool = False) -> Optional[datetime]:
        if value in (None, ""):
            return None
        parsed: Optional[datetime] = None
        if isinstance(value, datetime):
            parsed = value
        elif isinstance(value, date):
            parsed = datetime.combine(value, time.max if end_of_day else time.min)
        elif isinstance(value, str):
            raw_value = value.strip()
            if not raw_value:
                return None
            if raw_value.endswith("Z"):
                raw_value = f"{raw_value[:-1]}+00:00"
            try:
                parsed = datetime.fromisoformat(raw_value)
            except ValueError:
                try:
                    parsed_date = date.fromisoformat(raw_value[:10])
                    parsed = datetime.combine(parsed_date, time.max if end_of_day else time.min)
                except ValueError:
                    return None
            if "T" not in raw_value and len(raw_value) <= 10:
                parsed = datetime.combine(parsed.date(), time.max if end_of_day else time.min)
        if parsed is None:
            return None
        if parsed.tzinfo is not None:
            parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
        return parsed.replace(tzinfo=None)

    def _format_report_date_filter(self, value: Optional[datetime]) -> Optional[str]:
        return value.date().isoformat() if value else None

    def _format_report_date_label(self, value: Optional[datetime]) -> str:
        return value.strftime("%d/%m/%Y") if value else "-"

    def _resolve_resource_usage_range_context(
        self,
        periods: List[Any],
        filters: Optional[Dict[str, Any]],
    ) -> Dict[str, Any]:
        filters = filters or {}
        period_bounds = [
            (
                self._normalize_report_datetime(getattr(period, "starts_at", None) if not isinstance(period, dict) else period.get("starts_at")),
                self._normalize_report_datetime(getattr(period, "ends_at", None) if not isinstance(period, dict) else period.get("ends_at"), end_of_day=True),
            )
            for period in periods
        ]
        available_starts = [start for start, _ in period_bounds if start is not None]
        available_ends = [end for _, end in period_bounds if end is not None]
        default_start = min(available_starts) if available_starts else None
        default_end = max(available_ends) if available_ends else None

        start_value = filters.get("date_start") or filters.get("fecha_inicio") or filters.get("start_date")
        end_value = filters.get("date_end") or filters.get("fecha_fin") or filters.get("end_date")
        start_at = self._normalize_report_datetime(start_value) or default_start
        end_at = self._normalize_report_datetime(end_value, end_of_day=True) or default_end

        if start_at and end_at and start_at > end_at:
            raise ValueError("La fecha de inicio del reporte de recursos no puede ser posterior a la fecha final.")

        period_factors: List[Decimal] = []
        selected_period_indexes: List[int] = []
        range_active = bool(filters)
        for index, (period_start, period_end) in enumerate(period_bounds):
            factor = Decimal("1")
            if range_active:
                factor = Decimal("0")
                if period_start is not None and period_end is not None and start_at is not None and end_at is not None:
                    total_seconds = Decimal(str(max((period_end - period_start).total_seconds(), 0)))
                    if total_seconds > 0:
                        overlap_start = max(period_start, start_at)
                        overlap_end = min(period_end, end_at)
                        overlap_seconds = Decimal(str(max((overlap_end - overlap_start).total_seconds(), 0)))
                        factor = min(Decimal("1"), max(Decimal("0"), overlap_seconds / total_seconds))
            period_factors.append(factor)
            if not range_active or factor > 0:
                selected_period_indexes.append(index)

        return {
            "range_active": range_active,
            "start_at": start_at,
            "end_at": end_at,
            "period_factors": period_factors,
            "selected_period_indexes": selected_period_indexes,
            "filters": {
                "date_start": self._format_report_date_filter(start_at),
                "date_end": self._format_report_date_filter(end_at),
            },
        }


    def _get_apu_for_resource_usage(self, db: Any, apu_id: int, empresa_id: int) -> Any:
        from app.models.apu import APU, APULinea
        from app.models.recurso import Recurso

        apu = (
            db.query(APU)
            .options(
                selectinload(APU.subcategoria_item),
                selectinload(APU.lineas)
                .selectinload(APULinea.recurso)
                .selectinload(Recurso.unidad),
                selectinload(APU.lineas)
                .selectinload(APULinea.recurso)
                .selectinload(Recurso.subcategoria_item),
                selectinload(APU.lineas).selectinload(APULinea.apu_hijo),
            )
            .filter(APU.id == apu_id, APU.empresa_id == empresa_id)
            .first()
        )
        if not apu:
            raise ValueError(f"APU no encontrado: {apu_id}")
        return apu

    def _resolve_resource_usage_category(self, recurso: Any, linea: Any = None) -> Tuple[int, str]:
        category_id = None
        subcategoria = getattr(recurso, "subcategoria_item", None) if recurso else None
        if subcategoria is not None:
            category_id = getattr(subcategoria, "subcategoria_codigo", None)
        if category_id is None and recurso is not None:
            try:
                category_id = int(str(getattr(recurso, "codigo", "") or "").split("-")[0])
            except (ValueError, IndexError):
                category_id = None
        if category_id is None and linea is not None:
            category_id = self._resolve_apu_line_category(linea)
        try:
            category_id = int(category_id or 1)
        except (TypeError, ValueError):
            category_id = 1
        labels = {
            1: "1. Equipo/Herramientas",
            2: "2. Materiales",
            3: "3. Transporte",
            4: "4. Mano de Obra",
        }
        return category_id, labels.get(category_id, f"{category_id}. Recursos")

    def _resolve_resource_usage_subcategory(self, recurso: Any) -> str:
        subcategoria = getattr(recurso, "subcategoria_item", None) if recurso else None
        if not subcategoria:
            return "-"
        code = str(getattr(subcategoria, "codigo", "") or "").strip()
        description = str(getattr(subcategoria, "descripcion", "") or "").strip()
        if code and description:
            return f"{code} - {description}"
        return description or code or "-"

    def _collect_cronograma_resource_usage_from_apu(
        self,
        db: Any,
        apu: Any,
        empresa_id: int,
        factor: Decimal,
        accumulator: Dict[int, Dict[str, Any]],
        visited: Optional[set] = None,
    ) -> None:
        labels = {
            1: "1. Equipo/Herramientas",
            2: "2. Materiales",
            3: "3. Transporte",
            4: "4. Mano de Obra",
        }
        exploded = collect_apu_exploded_resources(
            apu,
            inherited_factor=factor,
            accumulator={},
            active_path=visited,
            child_loader=(
                (lambda child_id: self._get_apu_for_resource_usage(db, int(child_id), empresa_id))
                if db is not None
                else None
            ),
            category_labels=labels,
        )
        for item in exploded.values():
            recurso_id = int(item.get("recurso_id") or 0)
            if recurso_id <= 0:
                continue
            category_id = int(item.get("categoria_id") or 1)
            entry = accumulator.setdefault(recurso_id, {
                "recurso_id": recurso_id,
                "categoria_id": category_id,
                "categoria": str(item.get("categoria") or labels.get(category_id, f"{category_id}. Recursos")),
                "subcategoria": str(item.get("subcategoria") or "-"),
                "recurso": self._normalize_report_description(str(item.get("descripcion") or "")),
                "codigo": str(item.get("codigo") or "").strip(),
                "unidad": str(item.get("unidad") or "").strip(),
                "precio_unitario": self._decimal_value(item.get("precio_unitario")),
                "cantidad_base": Decimal("0"),
            })
            entry["cantidad_base"] += self._decimal_value(item.get("cantidad"))

    def _build_cronograma_resource_usage_payload(
        self,
        db: Any,
        presupuesto_id: int,
        empresa_id: int,
        filters: Optional[Dict[str, Any]] = None,
    ) -> Dict[str, Any]:
        presupuesto, cronograma = self._get_cronograma_valorado_payload(db, presupuesto_id, empresa_id)
        apu_resource_readiness_service.ensure_budget_ready(db, presupuesto, empresa_id)
        periods = list(getattr(cronograma, "periods", []) or [])
        period_count = len(periods)
        range_context = self._resolve_resource_usage_range_context(periods, filters)
        selected_period_indexes = list(range_context["selected_period_indexes"])
        selected_periods = [periods[index] for index in selected_period_indexes]
        selected_period_position = {
            original_index: selected_index
            for selected_index, original_index in enumerate(selected_period_indexes)
        }
        period_factors = list(range_context["period_factors"])
        selected_period_count = len(selected_periods)
        budget_lines_by_id = {
            int(getattr(line, "id", 0)): line
            for line in list(getattr(presupuesto, "detalle", []) or [])
            if getattr(line, "id", None) is not None
        }
        apu_cache: Dict[int, Any] = {}
        resources: Dict[int, Dict[str, Any]] = {}

        for row in list(getattr(cronograma, "rows", []) or []):
            source_line = budget_lines_by_id.get(int(getattr(row, "linea_id", 0) or 0))
            apu = getattr(source_line, "apu", None) if source_line is not None else None
            apu_id = getattr(row, "apu_id", None) or getattr(source_line, "apu_id", None)
            if apu is None and apu_id:
                apu_id_int = int(apu_id)
                apu = apu_cache.get(apu_id_int)
                if apu is None:
                    apu = self._get_apu_for_resource_usage(db, apu_id_int, empresa_id)
                    apu_cache[apu_id_int] = apu
            if apu is None:
                continue

            row_resources: Dict[int, Dict[str, Any]] = {}
            row_quantity = self._decimal_value(getattr(row, "cantidad", None) or getattr(source_line, "cantidad", None))
            self._collect_cronograma_resource_usage_from_apu(db, apu, empresa_id, row_quantity, row_resources)

            distribution = list(getattr(row, "distribution", []) or [])
            if period_count:
                distribution = (distribution + [0] * period_count)[:period_count]
            for resource_id, row_resource in row_resources.items():
                target = resources.setdefault(resource_id, {
                    **{key: row_resource[key] for key in (
                        "recurso_id",
                        "categoria_id",
                        "categoria",
                        "subcategoria",
                        "recurso",
                        "codigo",
                        "unidad",
                        "precio_unitario",
                    )},
                    "period_quantities": [Decimal("0") for _ in range(selected_period_count)],
                    "period_costs": [Decimal("0") for _ in range(selected_period_count)],
                    "cantidad_total": Decimal("0"),
                    "costo_total": Decimal("0"),
                })
                quantity_base = self._decimal_value(row_resource.get("cantidad_base"))
                price = self._decimal_value(row_resource.get("precio_unitario"))
                total_quantity = Decimal("0")
                total_cost = Decimal("0")
                if period_count:
                    for index, pct in enumerate(distribution):
                        period_quantity = (
                            quantity_base
                            * self._decimal_value(pct)
                            / Decimal("100")
                            * (period_factors[index] if index < len(period_factors) else Decimal("1"))
                        )
                        period_cost = period_quantity * price
                        selected_index = selected_period_position.get(index)
                        if selected_index is not None:
                            target["period_quantities"][selected_index] += period_quantity
                            target["period_costs"][selected_index] += period_cost
                        total_quantity += period_quantity
                        total_cost += period_cost
                else:
                    total_quantity = quantity_base
                    total_cost = quantity_base * price
                target["cantidad_total"] += total_quantity
                target["costo_total"] += total_cost

        lineas = [
            {
                **resource,
                "precio_unitario": float(resource["precio_unitario"]),
                "period_quantities": [float(value) for value in resource["period_quantities"]],
                "period_costs": [float(value) for value in resource["period_costs"]],
                "cantidad_total": float(resource["cantidad_total"]),
                "costo_total": float(resource["costo_total"]),
            }
            for resource in resources.values()
            if resource["cantidad_total"] > 0
        ]
        lineas.sort(key=lambda item: (
            int(item.get("categoria_id") or 0),
            str(item.get("subcategoria") or ""),
            str(item.get("recurso") or ""),
            str(item.get("codigo") or ""),
        ))

        period_columns = [
            {"key": f"periodo_{index}_costo", "label": getattr(period, "label", f"P{index + 1}"), "kind": "money", "width_weight": 0.75}
            for index, period in enumerate(selected_periods[:8])
        ]
        preview_lines = []
        for line in lineas:
            preview_line = {
                "categoria": line["categoria"],
                "subcategoria": line["subcategoria"],
                "recurso": line["recurso"],
                "codigo": line["codigo"],
                "unidad": line["unidad"],
                "cantidad_total": line["cantidad_total"],
                "costo_total": line["costo_total"],
            }
            for index in range(min(len(selected_periods), 8)):
                preview_line[f"periodo_{index}_costo"] = line["period_costs"][index]
            preview_lines.append(preview_line)

        period_note = ""
        if len(selected_periods) > 8:
            period_note = " · vista GiProy/PDF muestra los primeros 8 periodos; Excel contiene todos"
        range_note = ""
        if range_context["range_active"]:
            range_note = (
                f" · rango {self._format_report_date_label(range_context['start_at'])}"
                f" - {self._format_report_date_label(range_context['end_at'])}"
            )
        report_label = "Uso de recursos por rango" if range_context["range_active"] else "Uso de recursos"

        return {
            "id": getattr(presupuesto, "id", None),
            "preview_layout": "resource_usage",
            "filters": range_context["filters"] if range_context["range_active"] else None,
            "date_range": range_context["filters"],
            "codigo": getattr(getattr(presupuesto, "proyecto", None), "codigo", None) or getattr(getattr(presupuesto, "proyecto", None), "codigo_root", None) or "",
            "descripcion": self._resolve_project_title(getattr(presupuesto, "proyecto", None), getattr(presupuesto, "descripcion", None)),
            "unidad": getattr(cronograma, "moneda", None) or "USD",
            "metadata_hint": (
                f"{report_label} · {getattr(cronograma, 'period_type', '')} · {len(selected_periods)} periodo(s)"
                f"{range_note}"
                f" · modo {getattr(cronograma, 'distribution_mode', '')}{period_note}"
            ),
            "summary_cards": [
                {"label": "Recursos", "value": len(lineas), "kind": "integer"},
                {"label": "Periodos", "value": len(selected_periods), "kind": "integer"},
                *([
                    {"label": "Inicio", "value": self._format_report_date_label(range_context["start_at"]), "kind": "raw"},
                    {"label": "Fin", "value": self._format_report_date_label(range_context["end_at"]), "kind": "raw"},
                ] if range_context["range_active"] else []),
                {"label": "Costo directo", "value": sum(float(line["costo_total"]) for line in lineas), "kind": "money", "tone": "total"},
            ],
            "table_columns": [
                {"key": "categoria", "label": "Categoria", "width_weight": 1.0},
                {"key": "subcategoria", "label": "Subcategoria", "width_weight": 1.0},
                {"key": "recurso", "label": "Recurso", "width_weight": 2.0},
                {"key": "unidad", "label": "Unid.", "width_weight": 0.55},
                {"key": "cantidad_total", "label": "Cant. total", "kind": "calc", "width_weight": 0.8},
                {"key": "costo_total", "label": "Costo total", "kind": "money", "width_weight": 0.9},
                *period_columns,
            ],
            "lineas": preview_lines,
            "resource_usage_rows": lineas,
            "periods": [
                {
                    "label": getattr(period, "label", f"P{index + 1}"),
                    "starts_at": getattr(period, "starts_at", None),
                    "ends_at": getattr(period, "ends_at", None),
                    "overlap_factor": float(period_factors[selected_period_indexes[index]]) if index < len(selected_period_indexes) else 1.0,
                }
                for index, period in enumerate(selected_periods)
            ],
        }

    def _summarize_cronograma_manual_schedule_pending(self, cronograma: Any) -> Dict[str, Any]:
        rows = list(getattr(cronograma, "rows", []) or [])
        pending_rows = [row for row in rows if bool(getattr(row, "requires_manual_schedule", False))]
        pending_amount = sum(float(getattr(row, "precio_total", 0) or 0) for row in pending_rows)
        total_amount = sum(float(getattr(row, "precio_total", 0) or 0) for row in rows)
        scheduled_amount = max(0.0, total_amount - pending_amount)
        pending_pct = (pending_amount / total_amount * 100.0) if total_amount > 0 else 0.0
        scheduled_pct = (scheduled_amount / total_amount * 100.0) if total_amount > 0 else 0.0
        return {
            "count": len(pending_rows),
            "pending_amount": pending_amount,
            "scheduled_amount": scheduled_amount,
            "total_amount": total_amount,
            "pending_pct": round(pending_pct, 4),
            "scheduled_pct": round(scheduled_pct, 4),
            "has_pending": bool(pending_rows),
        }

    def _summarize_gantt_operational_reporting(self, schedule: Any, cronograma: Any = None) -> Dict[str, Any]:
        rows = list(getattr(schedule, "rows", []) or [])
        valuado_rows = {
            str(getattr(row, "linea_id", "") or "").strip(): row
            for row in list(getattr(cronograma, "rows", []) or [])
            if str(getattr(row, "linea_id", "") or "").strip()
        }
        valuado_periods = list(getattr(cronograma, "periods", []) or [])
        subbar_rows = 0
        subbar_count = 0
        draft_subbar_count = 0
        manual_temporal_rows = 0
        interparent_merge_rows = 0
        interparent_merge_count = 0
        conflict_rows = 0
        conflict_periods = 0

        for row in rows:
            metadata = dict(getattr(row, "metadata", None) or {})
            operational = dict(metadata.get("gantt_operational") or {})
            raw_subbars = list(metadata.get("gantt_subbars") or [])
            budget_line_id = str(
                getattr(row, "presupuesto_linea_id", None)
                or getattr(row, "linea_id", None)
                or getattr(row, "budget_line_id", None)
                or operational.get("budget_line_id")
                or ""
            ).strip()
            if operational.get("has_subbars"):
                subbar_rows += 1
            subbar_count += int(operational.get("subbar_count") or 0)
            draft_subbar_count += sum(
                1 for segment in raw_subbars
                if str(segment.get("status", "") or "").strip().lower() == "draft_session"
            )
            row_interparent_merge_count = sum(
                1 for segment in raw_subbars
                if str(segment.get("source", "") or "").strip().lower() == "gantt_interparent_merge"
            )
            interparent_merge_count += row_interparent_merge_count
            if row_interparent_merge_count > 0:
                interparent_merge_rows += 1
            if operational.get("has_manual_temporal_window"):
                manual_temporal_rows += 1

            valuado_row = valuado_rows.get(budget_line_id)
            if not valuado_row or not raw_subbars or not valuado_periods:
                continue

            distribution = list(getattr(valuado_row, "distribution", []) or [])
            line_total = float(getattr(valuado_row, "precio_total", 0) or 0)
            period_buckets: Dict[str, Dict[str, float]] = {}
            for segment in raw_subbars:
                period_id = str(segment.get("period_id") or segment.get("parent_period_id") or "").strip()
                if not period_id:
                    continue
                current = period_buckets.get(period_id) or {"actual_percent": 0.0, "actual_amount": 0.0}
                current["actual_percent"] += float(segment.get("percent") or 0)
                current["actual_amount"] += float(segment.get("amount") or 0)
                period_buckets[period_id] = current

            row_conflicts = 0
            for index, period in enumerate(valuado_periods):
                period_id = str(getattr(period, "id", None) or f"P{index + 1}")
                actual = period_buckets.get(period_id) or {"actual_percent": 0.0, "actual_amount": 0.0}
                allowed_percent = float(distribution[index] or 0) if index < len(distribution) else 0.0
                allowed_amount = float(line_total * (allowed_percent / 100.0)) if line_total > 0 else 0.0
                if actual["actual_percent"] - allowed_percent > 0.0001 or actual["actual_amount"] - allowed_amount > 0.01:
                    row_conflicts += 1
            if row_conflicts > 0:
                conflict_rows += 1
                conflict_periods += row_conflicts

        has_operational_state = bool(subbar_rows or manual_temporal_rows or draft_subbar_count or conflict_rows)

        return {
            "subbar_rows": subbar_rows,
            "subbar_count": subbar_count,
            "draft_subbar_count": draft_subbar_count,
            "manual_temporal_rows": manual_temporal_rows,
            "interparent_merge_rows": interparent_merge_rows,
            "interparent_merge_count": interparent_merge_count,
            "conflict_rows": conflict_rows,
            "conflict_periods": conflict_periods,
            "has_operational_state": has_operational_state,
        }

    def _build_cronograma_valorado_preview(self, db: Any, presupuesto_id: int, empresa_id: int) -> Dict[str, Any]:
        presupuesto, cronograma = self._get_cronograma_valorado_payload(db, presupuesto_id, empresa_id)
        pending_summary = self._summarize_cronograma_manual_schedule_pending(cronograma)
        budget_lines_by_id = {
            int(getattr(line, "id", 0)): line
            for line in list(presupuesto.detalle or [])
            if getattr(line, "id", None) is not None
        }
        period_columns = [
            { "key": f"periodo_{index}", "label": period.label, "kind": "money" }
            for index, period in enumerate(cronograma.periods)
        ]
        lineas = []
        for row in cronograma.rows:
            source_line = budget_lines_by_id.get(int(getattr(row, "linea_id", 0) or 0))
            line = {
                "item_visible": str(getattr(row, "codigo_item", None) or "").strip(),
                "edt_code_visible": self._get_budget_line_visible_edt_code(source_line) if source_line else "",
                "descripcion": self._normalize_report_description(row.descripcion),
                "unidad": row.unidad or "",
                "cantidad": row.cantidad,
                "precio_unitario": row.precio_unitario,
                "precio_total": row.precio_total,
            }
            for index, pct in enumerate(row.distribution):
                line[f"periodo_{index}"] = row.precio_total * (pct / 100)
            lineas.append(line)

        return {
            "id": presupuesto.id,
            "codigo": getattr(presupuesto.proyecto, "codigo", None) or getattr(presupuesto.proyecto, "codigo_root", None) or "",
            "descripcion": self._resolve_project_title(presupuesto.proyecto, presupuesto.descripcion),
            "unidad": cronograma.moneda,
            "metadata_hint": (
                f"Cronograma valorado · {cronograma.period_type} · {len(cronograma.periods)} periodo(s)"
                + (" · temporalidad parcial" if pending_summary["has_pending"] else "")
            ),
            "summary_cards": [
                { "label": "Presupuesto", "value": sum(row.precio_total for row in cronograma.rows), "kind": "money", "tone": "total" },
                { "label": "Periodos", "value": len(cronograma.periods), "kind": "calc" },
                { "label": "Modo", "value": cronograma.distribution_mode, "kind": "raw" },
                *([
                    { "label": "Pendiente manual", "value": pending_summary["pending_amount"], "kind": "money" },
                    { "label": "Cobertura valorada", "value": pending_summary["scheduled_pct"], "kind": "calc" },
                ] if pending_summary["has_pending"] else []),
            ],
            "table_columns": [
                { "key": "item_visible", "label": "Item" },
                { "key": "edt_code_visible", "label": "Cod EDT" },
                { "key": "descripcion", "label": "Descripción" },
                { "key": "unidad", "label": "Unid." },
                { "key": "cantidad", "label": "Cant.", "kind": "calc" },
                { "key": "precio_unitario", "label": "P. Unit.", "kind": "money" },
                { "key": "precio_total", "label": "Subtotal", "kind": "money" },
                *period_columns,
            ],
            "lineas": lineas,
        }

    def _build_cronograma_cash_flow_preview(self, db: Any, presupuesto_id: int, empresa_id: int) -> Dict[str, Any]:
        presupuesto, cronograma = self._get_cronograma_valorado_payload(db, presupuesto_id, empresa_id)
        pending_summary = self._summarize_cronograma_manual_schedule_pending(cronograma)
        lineas = [
            {
                "periodo": point.label,
                "fecha_inicio": point.starts_at.strftime("%d/%m/%Y %H:%M"),
                "fecha_fin": point.ends_at.strftime("%d/%m/%Y %H:%M"),
                "horas_periodo": point.work_hours,
                "costo_periodo": point.cost,
                "equipos": (point.category_costs or {}).get("Equipos", 0),
                "mano_obra": (point.category_costs or {}).get("Mano de obra", 0),
                "transporte": (point.category_costs or {}).get("Transporte", 0),
                "categoria_dominante": point.dominant_category or "",
                "costo_acumulado": point.cumulative_cost,
                "porcentaje_periodo": point.cost_pct,
                "porcentaje_acumulado": point.cumulative_pct,
            }
            for point in cronograma.cash_flow
        ]

        return {
            "id": presupuesto.id,
            "codigo": getattr(presupuesto.proyecto, "codigo", None) or getattr(presupuesto.proyecto, "codigo_root", None) or "",
            "descripcion": self._resolve_project_title(presupuesto.proyecto, presupuesto.descripcion),
            "unidad": cronograma.moneda,
            "metadata_hint": (
                f"Flujo de caja · {cronograma.period_type} · {len(cronograma.cash_flow)} periodo(s)"
                + (" · cobertura parcial" if pending_summary["has_pending"] else "")
            ),
            "summary_cards": [
                { "label": "Costo total", "value": cronograma.cash_flow[-1].cumulative_cost if cronograma.cash_flow else 0, "kind": "money", "tone": "total" },
                { "label": "Periodos", "value": len(cronograma.cash_flow), "kind": "calc" },
                { "label": "Horas efectivas", "value": sum(float(point.work_hours or 0) for point in cronograma.cash_flow), "kind": "calc" },
                { "label": "Categoría líder", "value": max(
                    (
                        (
                            label,
                            sum(float((point.category_costs or {}).get(label, 0) or 0) for point in cronograma.cash_flow)
                        )
                        for label in {"Equipos", "Mano de obra", "Transporte"}
                    ),
                    key=lambda item: item[1],
                    default=("Sin categoría", 0),
                )[0], "kind": "raw" },
                { "label": "Origen", "value": "Gantt parcial" if cronograma.distribution_mode == "gantt" and pending_summary["has_pending"] else ("Gantt" if cronograma.distribution_mode == "gantt" else "Valorado"), "kind": "raw" },
                *([
                    { "label": "Pendiente manual", "value": pending_summary["pending_amount"], "kind": "money" },
                    { "label": "Cobertura valorada", "value": pending_summary["scheduled_pct"], "kind": "calc" },
                ] if pending_summary["has_pending"] else []),
            ],
            "table_columns": [
                { "key": "periodo", "label": "Periodo" },
                { "key": "fecha_inicio", "label": "Inicio" },
                { "key": "fecha_fin", "label": "Fin" },
                { "key": "horas_periodo", "label": "Horas", "kind": "calc" },
                { "key": "costo_periodo", "label": "Costo", "kind": "money" },
                { "key": "equipos", "label": "Equipos", "kind": "money" },
                { "key": "mano_obra", "label": "Mano de obra", "kind": "money" },
                { "key": "transporte", "label": "Transporte", "kind": "money" },
                { "key": "categoria_dominante", "label": "Dominante" },
                { "key": "costo_acumulado", "label": "Costo acum.", "kind": "money" },
                { "key": "porcentaje_periodo", "label": "% periodo", "kind": "calc" },
                { "key": "porcentaje_acumulado", "label": "% acum.", "kind": "calc" },
            ],
            "lineas": lineas,
        }

    def _build_cronograma_gantt_preview(self, db: Any, presupuesto_id: int, empresa_id: int) -> Dict[str, Any]:
        from app.services.cronograma_trabajo import cronograma_trabajo_service

        presupuesto = self._get_presupuesto(db, presupuesto_id, empresa_id)
        schedule = cronograma_trabajo_service.get_schedule(
            db,
            presupuesto_id=presupuesto.id,
            proyecto_id=presupuesto.proyecto_id,
            empresa_id=empresa_id,
        )
        _, cronograma = self._get_cronograma_valorado_payload(db, presupuesto_id, empresa_id)
        operational_summary = self._summarize_gantt_operational_reporting(schedule, cronograma)
        budget_lines_by_id = {
            int(getattr(line, "id", 0)): line
            for line in list(presupuesto.detalle or [])
            if getattr(line, "id", None) is not None
        }
        lineas = []
        for row in schedule.rows or []:
            dependencies = getattr(row, "dependencies", []) or []
            dependency_labels = [
                f"{dependency.source_id}:{dependency.type}{'+' if dependency.lag_days and dependency.lag_days > 0 else ''}{dependency.lag_days or 0:g}d"
                for dependency in dependencies
            ]
            source_line = budget_lines_by_id.get(int(getattr(row, "presupuesto_linea_id", 0) or 0))
            lineas.append({
                "item_visible": str(getattr(row, "codigo_item", None) or "").strip(),
                "edt_code_visible": self._get_budget_line_visible_edt_code(source_line) if source_line else "",
                "descripcion": self._normalize_report_description(row.descripcion),
                "inicio": row.start_date.strftime("%d/%m/%Y") if row.start_date else "",
                "fin": row.end_date.strftime("%d/%m/%Y") if row.end_date else "",
                "duracion": row.dias_calendario or row.dias_utiles or 0,
                "avance": row.progress_pct or 0,
                "predecesoras": ", ".join(str(item) for item in (row.predecessors or [])),
                "dependencias": " | ".join(dependency_labels),
            })

        return {
            "id": presupuesto.id,
            "codigo": getattr(presupuesto.proyecto, "codigo", None) or getattr(presupuesto.proyecto, "codigo_root", None) or "",
            "descripcion": self._resolve_project_title(presupuesto.proyecto, presupuesto.descripcion),
            "unidad": "días",
            "metadata_hint": (
                f"Gantt · {len(lineas)} tarea(s) · {schedule.config.dias_laborables_semana if schedule.config else 5:g} días laborables/semana"
                + (" · reconciliación operativa" if operational_summary["has_operational_state"] else "")
            ),
            "summary_cards": [
                { "label": "Tareas", "value": len(lineas), "kind": "calc" },
                { "label": "Trabajo útil", "value": schedule.summary.trabajo_total if schedule.summary else 0, "kind": "calc" },
                { "label": "Cuadrilla", "value": schedule.summary.cuadrilla_total if schedule.summary else 0, "kind": "calc" },
                *([
                    { "label": "Líneas con subbarras", "value": operational_summary["subbar_rows"], "kind": "calc" },
                    { "label": "Tramos operativos", "value": operational_summary["subbar_count"], "kind": "calc" },
                    { "label": "Tramos en borrador", "value": operational_summary["draft_subbar_count"], "kind": "calc" },
                    { "label": "Fusiones intertramo", "value": operational_summary["interparent_merge_count"], "kind": "calc" },
                    { "label": "Ventanas manuales", "value": operational_summary["manual_temporal_rows"], "kind": "calc" },
                    { "label": "Conflictos valorado", "value": operational_summary["conflict_periods"], "kind": "calc" },
                ] if operational_summary["has_operational_state"] else []),
            ],
            "table_columns": [
                { "key": "item_visible", "label": "Item" },
                { "key": "edt_code_visible", "label": "Cod EDT" },
                { "key": "descripcion", "label": "Descripción" },
                { "key": "inicio", "label": "Inicio" },
                { "key": "fin", "label": "Fin" },
                { "key": "duracion", "label": "Dur.", "kind": "calc" },
                { "key": "avance", "label": "% Av.", "kind": "calc" },
                { "key": "predecesoras", "label": "Predec." },
                { "key": "dependencias", "label": "Dependencias" },
            ],
            "lineas": lineas,
        }

    def _build_cronograma_pareto_preview(self, db: Any, presupuesto_id: int, empresa_id: int) -> Dict[str, Any]:
        from app.api.endpoints.cronogramas_trabajo import _build_cronograma_trabajo_pareto_response
        from app.services.cronograma_trabajo import cronograma_trabajo_service

        presupuesto = self._get_presupuesto(db, presupuesto_id, empresa_id)
        schedule = cronograma_trabajo_service.get_schedule(
            db,
            presupuesto_id=presupuesto.id,
            proyecto_id=presupuesto.proyecto_id,
            empresa_id=empresa_id,
        )
        price_map = {
            int(linea.id): float(linea.precio_total or 0.0)
            for linea in (presupuesto.detalle or [])
        }
        pareto = _build_cronograma_trabajo_pareto_response(
            presupuesto_id=presupuesto.id,
            rows=schedule.rows,
            price_map=price_map,
            view="integrated",
            top=20,
        )
        lineas = []
        for item in pareto.items:
            lineas.append({
                "ranking": item.ranking,
                "codigo": item.codigo or "",
                "descripcion": self._normalize_report_description(item.descripcion),
                "costo": item.cost_value,
                "duracion_dias": item.duration_days,
                "horas_efectivas": item.work_hours,
                "impacto": item.porcentaje,
                "impacto_acumulado": item.porcentaje_acumulado,
            })

        return {
            "id": presupuesto.id,
            "codigo": getattr(presupuesto.proyecto, "codigo", None) or getattr(presupuesto.proyecto, "codigo_root", None) or "",
            "descripcion": self._resolve_project_title(presupuesto.proyecto, presupuesto.descripcion),
            "unidad": presupuesto.moneda or "USD",
            "metadata_hint": f"Pareto temporal · {pareto.view} · top {pareto.visible_items}",
            "summary_cards": [
                {"label": "Items", "value": pareto.visible_items, "kind": "calc"},
                {"label": "Acumulado", "value": pareto.visible_acumulado, "kind": "calc"},
                {"label": "Modo", "value": "integrado", "kind": "raw"},
            ],
            "table_columns": [
                {"key": "ranking", "label": "Rank", "kind": "calc"},
                {"key": "codigo", "label": "Item"},
                {"key": "descripcion", "label": "Descripción"},
                {"key": "costo", "label": "Costo", "kind": "money"},
                {"key": "duracion_dias", "label": "Dur. días", "kind": "calc"},
                {"key": "horas_efectivas", "label": "Horas", "kind": "calc"},
                {"key": "impacto", "label": "% impacto", "kind": "calc"},
                {"key": "impacto_acumulado", "label": "% acum.", "kind": "calc"},
            ],
            "lineas": lineas,
        }

    def _build_simple_report_workbook(self, title: str, items: List[Dict[str, Any]], money_decimals: int, calc_decimals: int) -> io.BytesIO:
        wb = openpyxl.Workbook()
        default_ws = wb.active
        wb.remove(default_ws)
        for item_index, item in enumerate(items, start=1):
            sheet_title = (item.get("metadata_hint") or title or f"Reporte {item_index}").split("·")[0].strip()[:31] or f"Reporte {item_index}"
            ws = wb.create_sheet(sheet_title)
            ws.cell(row=1, column=1, value=title)
            ws.cell(row=1, column=1).font = Font(bold=True, size=14)
            ws.cell(row=2, column=1, value=item.get("descripcion") or "")
            ws.cell(row=3, column=1, value=item.get("metadata_hint") or "")
            warnings = item.get("warnings") if isinstance(item.get("warnings"), list) else []
            start_row = 5
            if warnings:
                for warning_index, warning in enumerate(warnings, start=1):
                    message = warning.get("message") if isinstance(warning, dict) else str(warning or "")
                    ws.cell(row=3 + warning_index, column=1, value=f"Advertencia: {message}")
                    ws.cell(row=3 + warning_index, column=1).font = Font(bold=True, color="9A3412")
                start_row = 5 + len(warnings)
            columns = item.get("table_columns") or []
            for column_index, column in enumerate(columns, start=1):
                cell = ws.cell(row=start_row, column=column_index, value=column.get("label") or column.get("key"))
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
            for row_index, line in enumerate(item.get("lineas") or [], start=start_row + 1):
                for column_index, column in enumerate(columns, start=1):
                    value = line.get(column.get("key"))
                    kind = column.get("kind")
                    if kind == "money":
                        value = self._round_half_up(value, money_decimals)
                    elif kind == "calc":
                        value = self._round_half_up(value, calc_decimals)
                    ws.cell(row=row_index, column=column_index, value=value)
            for column_index, column in enumerate(columns, start=1):
                width = max(12, min(42, len(str(column.get("label") or column.get("key"))) + 6))
                ws.column_dimensions[get_column_letter(column_index)].width = width
            ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
        wb.properties.title = title
        return self._save_workbook_buffer(wb)

    def _build_cronograma_resource_usage_workbook(
        self,
        db: Any,
        presupuesto_id: int,
        empresa_id: int,
        filters: Optional[Dict[str, Any]] = None,
        payload_override: Optional[Dict[str, Any]] = None,
    ) -> io.BytesIO:
        format_config = self._get_empresa_format_config(db, empresa_id)
        payload = payload_override or self._build_cronograma_resource_usage_payload(db, presupuesto_id, empresa_id, filters)
        template_path = os.path.join(self.templates_dir, "001 - Uso de Recursos Cronograma.xlsx")
        if os.path.exists(template_path):
            wb = openpyxl.load_workbook(template_path)
            ws = wb.active
            for row in ws.iter_rows():
                for cell in row:
                    cell.value = None
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
        is_range_report = bool(payload.get("filters"))
        report_title = "Cronograma de Uso de Recursos por Rango" if is_range_report else "Cronograma de Uso de Recursos"
        ws.title = "Recursos por rango" if is_range_report else "Uso de recursos"
        wb.properties.title = report_title

        rows = list(payload.get("resource_usage_rows") or [])
        periods = list(payload.get("periods") or [])
        money_decimals = format_config["money_decimals"]
        calc_decimals = format_config["calc_decimals"]

        title_fill = PatternFill("solid", fgColor="111827")
        header_fill = PatternFill("solid", fgColor="F39200")
        group_fill = PatternFill("solid", fgColor="E5E7EB")
        subtotal_fill = PatternFill("solid", fgColor="F8FAFC")
        thin_border = Border(bottom=Side(style="thin", color="D4D4D8"))

        ws.cell(row=1, column=1, value=payload.get("descripcion") or report_title)
        ws.cell(row=1, column=1).font = Font(bold=True, size=14, color="FFFFFF")
        ws.cell(row=1, column=1).fill = title_fill
        ws.cell(row=2, column=1, value=report_title)
        ws.cell(row=2, column=1).font = Font(bold=True, color="F39200")
        ws.cell(row=3, column=1, value=payload.get("metadata_hint") or "")
        warnings = payload.get("warnings") if isinstance(payload.get("warnings"), list) else []
        if warnings:
            first_warning = warnings[0] if isinstance(warnings[0], dict) else {"message": str(warnings[0] or "")}
            warning_message = first_warning.get("message") or "Existen trabajos pendientes no incluidos en este reporte."
            ws.cell(row=4, column=1, value=f"Advertencia: {warning_message}")
            ws.cell(row=4, column=1).font = Font(bold=True, color="9A3412")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(7, 5 + len(periods) * 2 + 2))

        base_headers = ["Categoria", "Subcategoria", "Recurso", "Codigo", "Unidad"]
        for index, label in enumerate(base_headers, start=1):
            cell = ws.cell(row=5, column=index, value=label)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        current_col = 6
        for index, period in enumerate(periods, start=1):
            start_col = current_col
            ws.cell(row=5, column=start_col, value=getattr(period, "label", None) or period.get("label") or f"Periodo {index}")
            ws.merge_cells(start_row=5, start_column=start_col, end_row=5, end_column=start_col + 1)
            for column in (start_col, start_col + 1):
                cell = ws.cell(row=5, column=column)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=6, column=start_col, value="Cantidad")
            ws.cell(row=6, column=start_col + 1, value="Costo Directo")
            current_col += 2
        total_quantity_col = current_col
        total_cost_col = current_col + 1
        ws.cell(row=5, column=total_quantity_col, value="Total General")
        ws.merge_cells(start_row=5, start_column=total_quantity_col, end_row=5, end_column=total_cost_col)
        ws.cell(row=6, column=total_quantity_col, value="Cantidad")
        ws.cell(row=6, column=total_cost_col, value="Costo Directo")
        for column in range(1, total_cost_col + 1):
            for row_index in (5, 6):
                cell = ws.cell(row=row_index, column=column)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border

        row_index = 7
        grand_quantities = [0.0 for _ in periods]
        grand_costs = [0.0 for _ in periods]
        grand_quantity = 0.0
        grand_cost = 0.0
        current_category = None
        current_subcategory = None

        for resource in rows:
            category = resource.get("categoria") or "-"
            subcategory = resource.get("subcategoria") or "-"
            if category != current_category:
                current_category = category
                current_subcategory = None
                ws.cell(row=row_index, column=1, value=category)
                ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=total_cost_col)
                ws.cell(row=row_index, column=1).font = Font(bold=True)
                ws.cell(row=row_index, column=1).fill = group_fill
                row_index += 1
            if subcategory != current_subcategory:
                current_subcategory = subcategory
                ws.cell(row=row_index, column=2, value=subcategory)
                ws.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=total_cost_col)
                ws.cell(row=row_index, column=2).font = Font(bold=True)
                ws.cell(row=row_index, column=2).fill = subtotal_fill
                row_index += 1

            ws.cell(row=row_index, column=1, value=category)
            ws.cell(row=row_index, column=2, value=subcategory)
            ws.cell(row=row_index, column=3, value=resource.get("recurso") or "")
            ws.cell(row=row_index, column=4, value=resource.get("codigo") or "")
            ws.cell(row=row_index, column=5, value=resource.get("unidad") or "")
            current_col = 6
            period_quantities = list(resource.get("period_quantities") or [])
            period_costs = list(resource.get("period_costs") or [])
            for period_index in range(len(periods)):
                quantity = float(period_quantities[period_index] if period_index < len(period_quantities) else 0)
                cost = float(period_costs[period_index] if period_index < len(period_costs) else 0)
                ws.cell(row=row_index, column=current_col, value=self._round_half_up(quantity, calc_decimals))
                ws.cell(row=row_index, column=current_col + 1, value=self._round_half_up(cost, money_decimals))
                grand_quantities[period_index] += quantity
                grand_costs[period_index] += cost
                current_col += 2
            total_quantity = float(resource.get("cantidad_total") or 0)
            total_cost = float(resource.get("costo_total") or 0)
            ws.cell(row=row_index, column=total_quantity_col, value=self._round_half_up(total_quantity, calc_decimals))
            ws.cell(row=row_index, column=total_cost_col, value=self._round_half_up(total_cost, money_decimals))
            grand_quantity += total_quantity
            grand_cost += total_cost
            row_index += 1

        ws.cell(row=row_index, column=1, value="Total General")
        ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=5)
        ws.cell(row=row_index, column=1).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=row_index, column=1).fill = title_fill
        current_col = 6
        for period_index in range(len(periods)):
            ws.cell(row=row_index, column=current_col, value=self._round_half_up(grand_quantities[period_index], calc_decimals))
            ws.cell(row=row_index, column=current_col + 1, value=self._round_half_up(grand_costs[period_index], money_decimals))
            current_col += 2
        ws.cell(row=row_index, column=total_quantity_col, value=self._round_half_up(grand_quantity, calc_decimals))
        ws.cell(row=row_index, column=total_cost_col, value=self._round_half_up(grand_cost, money_decimals))
        for column in range(1, total_cost_col + 1):
            cell = ws.cell(row=row_index, column=column)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = title_fill

        money_format = f'#,##0.{"0" * money_decimals}'
        calc_format = f'#,##0.{"0" * calc_decimals}'
        for row in ws.iter_rows(min_row=7, max_row=row_index, min_col=6, max_col=total_cost_col):
            for cell in row:
                cell.number_format = money_format if (cell.column - 6) % 2 == 1 else calc_format
                cell.alignment = Alignment(horizontal="right")
        for column in range(1, total_cost_col + 1):
            ws.column_dimensions[get_column_letter(column)].width = 16
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["C"].width = 48
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 12
        ws.freeze_panes = "F7"
        ws.auto_filter.ref = f"A6:{get_column_letter(total_cost_col)}{row_index}"
        return self._save_workbook_buffer(wb)

    def _build_cronograma_resource_usage_executive_pdf(
        self,
        db: Any,
        presupuesto_id: int,
        empresa_id: int,
        filters: Optional[Dict[str, Any]] = None,
        watermark_text: Optional[str] = None,
        payload_override: Optional[Dict[str, Any]] = None,
    ) -> io.BytesIO:
        format_config = self._get_empresa_format_config(db, empresa_id)
        money_decimals = format_config["money_decimals"]
        calc_decimals = format_config["calc_decimals"]
        payload = payload_override or self._build_cronograma_resource_usage_payload(db, presupuesto_id, empresa_id, filters)
        rows = list(payload.get("resource_usage_rows") or [])
        periods = list(payload.get("periods") or [])
        is_range_report = bool(payload.get("filters"))
        report_title = "Reporte de Uso de Recursos por Rango" if is_range_report else "Reporte de Uso de Recursos"

        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=A4,
            leftMargin=11 * mm,
            rightMargin=11 * mm,
            topMargin=11 * mm,
            bottomMargin=11 * mm,
            title=report_title,
        )
        doc.author = self.REPORT_AUTHOR
        doc.creator = self.REPORT_AUTHOR

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "ResourceUsageTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=17,
            leading=20,
            textColor=colors.HexColor("#111827"),
            spaceAfter=4,
        )
        section_style = ParagraphStyle(
            "ResourceUsageSection",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            textColor=colors.HexColor("#111827"),
            spaceBefore=8,
            spaceAfter=5,
        )
        text_style = ParagraphStyle(
            "ResourceUsageText",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#374151"),
        )
        small_style = ParagraphStyle(
            "ResourceUsageSmall",
            parent=text_style,
            fontSize=7,
            leading=9,
            textColor=colors.HexColor("#6B7280"),
        )
        label_style = ParagraphStyle(
            "ResourceUsageLabel",
            parent=small_style,
            fontName="Helvetica-Bold",
            textColor=colors.HexColor("#6B7280"),
        )
        strong_style = ParagraphStyle(
            "ResourceUsageStrong",
            parent=text_style,
            fontName="Helvetica-Bold",
            textColor=colors.HexColor("#111827"),
        )
        category_style = ParagraphStyle(
            "ResourceUsageCategory",
            parent=text_style,
            fontName="Helvetica-Bold",
            fontSize=9,
            leading=11,
            textColor=colors.HexColor("#111827"),
        )
        subcategory_style = ParagraphStyle(
            "ResourceUsageSubcategory",
            parent=text_style,
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#136191"),
        )

        def pdf_text(value: Any) -> str:
            raw = "" if value is None else str(value)
            return (
                raw.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
                .replace('"', "&quot;")
            )

        def money(value: Any) -> str:
            return f"${self._round_half_up(value or 0, money_decimals):,.{money_decimals}f}"

        def calc(value: Any, decimals: Optional[int] = None) -> str:
            return self._format_fixed(value or 0, calc_decimals if decimals is None else decimals)

        def apply_pdf_metadata(canvas, _doc):
            canvas.setAuthor(self.REPORT_AUTHOR)
            canvas.setCreator(self.REPORT_AUTHOR)
            canvas.setTitle("Reporte de Uso de Recursos")
            self._draw_pdf_watermark(canvas, A4[0], A4[1], watermark_text)

        total_cost = sum(float(row.get("costo_total") or 0) for row in rows)
        total_quantity = sum(float(row.get("cantidad_total") or 0) for row in rows)
        category_map: "OrderedDict[str, Dict[str, Any]]" = OrderedDict()
        for row in rows:
            category = str(row.get("categoria") or "Recursos")
            current = category_map.setdefault(category, {
                "rows": [],
                "cost": 0.0,
                "quantity": 0.0,
                "subcategories": OrderedDict(),
            })
            current["rows"].append(row)
            current["cost"] += float(row.get("costo_total") or 0)
            current["quantity"] += float(row.get("cantidad_total") or 0)
            subcategory = str(row.get("subcategoria") or "-")
            subcurrent = current["subcategories"].setdefault(subcategory, {
                "rows": [],
                "cost": 0.0,
                "quantity": 0.0,
            })
            subcurrent["rows"].append(row)
            subcurrent["cost"] += float(row.get("costo_total") or 0)
            subcurrent["quantity"] += float(row.get("cantidad_total") or 0)

        categories = sorted(category_map.items(), key=lambda item: str(item[0]))
        top_resources = sorted(rows, key=lambda row: float(row.get("costo_total") or 0), reverse=True)[:12]
        top_10_cost = sum(float(row.get("costo_total") or 0) for row in top_resources[:10])
        dominant_category = max(categories, key=lambda item: item[1]["cost"], default=("-", {"cost": 0.0}))
        period_summaries = []
        for period_index, period in enumerate(periods):
            cost = sum(float(row.get("period_costs", [])[period_index] or 0) for row in rows if period_index < len(row.get("period_costs", [])))
            quantity = sum(float(row.get("period_quantities", [])[period_index] or 0) for row in rows if period_index < len(row.get("period_quantities", [])))
            period_summaries.append({
                "label": period.get("label") if isinstance(period, dict) else getattr(period, "label", f"P{period_index + 1}"),
                "cost": cost,
                "quantity": quantity,
            })
        peak_period = max(period_summaries, key=lambda item: item["cost"], default={"label": "-", "cost": 0.0, "quantity": 0.0})
        max_period_cost = max([point["cost"] for point in period_summaries] or [1.0])

        story = []
        story.append(Paragraph(report_title, title_style))
        story.append(Paragraph(pdf_text(payload.get("descripcion") or "-"), strong_style))
        story.append(Paragraph(pdf_text(payload.get("metadata_hint") or ""), small_style))
        warnings = payload.get("warnings") if isinstance(payload.get("warnings"), list) else []
        for warning in warnings:
            warning_message = warning.get("message") if isinstance(warning, dict) else str(warning or "")
            if warning_message:
                story.append(Paragraph(f"Advertencia: {pdf_text(warning_message)}", small_style))
        story.append(Spacer(1, 4 * mm))

        kpi_data = [
            [
                Paragraph("RECURSOS UNICOS", label_style),
                Paragraph("PERIODOS", label_style),
                Paragraph("COSTO DIRECTO", label_style),
                Paragraph("PERIODO PICO", label_style),
                Paragraph("CATEGORIA DOMINANTE", label_style),
            ],
            [
                Paragraph(str(len(rows)), strong_style),
                Paragraph(str(len(periods)), strong_style),
                Paragraph(money(total_cost), strong_style),
                Paragraph(f"{pdf_text(peak_period['label'])}<br/>{money(peak_period['cost'])}", strong_style),
                Paragraph(f"{pdf_text(dominant_category[0])}<br/>{calc((dominant_category[1]['cost'] / total_cost * 100) if total_cost else 0, 2)}%", strong_style),
            ],
        ]
        kpi_table = Table(kpi_data, colWidths=[doc.width / 5] * 5)
        kpi_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#FFF7ED")),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#E4E4E7")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(kpi_table)
        story.append(Spacer(1, 5 * mm))

        executive_notes = [
            f"El reporte consolida {len(rows)} recursos finales despues de explotar los APUs anidados.",
            f"La categoria de mayor peso es {dominant_category[0]}, con {money(dominant_category[1]['cost'])}.",
            f"El periodo de mayor demanda es {peak_period['label']}, con {money(peak_period['cost'])}.",
            f"Los 10 recursos gobernantes concentran {calc((top_10_cost / total_cost * 100) if total_cost else 0, 2)}% del costo directo de recursos.",
        ]
        story.append(Paragraph("Lectura ejecutiva", section_style))
        story.append(Paragraph("<br/>".join(f"- {pdf_text(note)}" for note in executive_notes), text_style))

        story.append(Paragraph("Distribucion por categoria", section_style))
        category_rows = [[
            Paragraph("Categoria", label_style),
            Paragraph("Recursos", label_style),
            Paragraph("Cantidad", label_style),
            Paragraph("Costo directo", label_style),
            Paragraph("Peso", label_style),
        ]]
        for category, data in categories:
            category_rows.append([
                Paragraph(pdf_text(category), text_style),
                Paragraph(str(len(data["rows"])), text_style),
                Paragraph(calc(data["quantity"]), text_style),
                Paragraph(money(data["cost"]), text_style),
                Paragraph(f"{calc((data['cost'] / total_cost * 100) if total_cost else 0, 2)}%", text_style),
            ])
        category_table = Table(category_rows, colWidths=[doc.width * 0.34, doc.width * 0.13, doc.width * 0.17, doc.width * 0.2, doc.width * 0.16])
        category_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#111827")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#FFFFFF")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E4E4E7")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(category_table)

        story.append(Paragraph("Demanda por periodos", section_style))
        period_rows = []
        for period in period_summaries:
            bar_width = int(((period["cost"] / max_period_cost) if max_period_cost else 0) * 28)
            bar = "|" * max(bar_width, 1)
            period_rows.append([
                Paragraph(pdf_text(period["label"]), strong_style),
                Paragraph(f'<font color="#136191">{bar}</font>', text_style),
                Paragraph(money(period["cost"]), text_style),
                Paragraph(f"{calc((period['cost'] / total_cost * 100) if total_cost else 0, 2)}%", text_style),
            ])
        period_table = Table(period_rows, colWidths=[doc.width * 0.14, doc.width * 0.48, doc.width * 0.22, doc.width * 0.16])
        period_table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E4E4E7")),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFFFFF")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(period_table)

        story.append(Paragraph("Recursos gobernantes", section_style))
        top_rows = [[
            Paragraph("#", label_style),
            Paragraph("Recurso", label_style),
            Paragraph("Categoria", label_style),
            Paragraph("Cantidad", label_style),
            Paragraph("Costo", label_style),
            Paragraph("Peso", label_style),
        ]]
        for index, row in enumerate(top_resources, start=1):
            top_rows.append([
                Paragraph(str(index), text_style),
                Paragraph(f"{pdf_text(row.get('recurso') or '-')}<br/><font color='#6B7280'>{pdf_text(row.get('codigo') or '')} · {pdf_text(row.get('unidad') or '')}</font>", text_style),
                Paragraph(pdf_text(row.get("categoria") or "-"), text_style),
                Paragraph(calc(row.get("cantidad_total")), text_style),
                Paragraph(money(row.get("costo_total")), text_style),
                Paragraph(f"{calc((float(row.get('costo_total') or 0) / total_cost * 100) if total_cost else 0, 2)}%", text_style),
            ])
        top_table = Table(top_rows, colWidths=[doc.width * 0.06, doc.width * 0.36, doc.width * 0.2, doc.width * 0.13, doc.width * 0.15, doc.width * 0.1])
        top_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F4F4F5")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E4E4E7")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAFAFA")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(top_table)

        story.append(PageBreak())
        story.append(Paragraph("Anexo ejecutivo de recursos", title_style))
        story.append(Paragraph(
            "Listado categorizado y subcategorizado de recursos finales. Los periodos se presentan como lectura compacta por recurso; la matriz operativa completa se conserva en Excel.",
            text_style,
        ))

        for category, data in categories:
            story.append(Spacer(1, 3 * mm))
            story.append(Paragraph(
                f"{pdf_text(category)} · {len(data['rows'])} recurso(s) · {money(data['cost'])}",
                category_style,
            ))
            sorted_subcategories = sorted(data["subcategories"].items(), key=lambda item: str(item[0]))
            for subcategory, subdata in sorted_subcategories:
                story.append(Paragraph(
                    f"{pdf_text(subcategory)} · {len(subdata['rows'])} recurso(s) · {money(subdata['cost'])}",
                    subcategory_style,
                ))
                for row in sorted(subdata["rows"], key=lambda item: str(item.get("recurso") or "")):
                    period_parts = []
                    quantities = list(row.get("period_quantities") or [])
                    costs = list(row.get("period_costs") or [])
                    for period_index, period in enumerate(periods):
                        quantity = float(quantities[period_index] if period_index < len(quantities) else 0)
                        cost = float(costs[period_index] if period_index < len(costs) else 0)
                        label = period.get("label") if isinstance(period, dict) else getattr(period, "label", f"P{period_index + 1}")
                        period_parts.append(f"<b>{pdf_text(label)}:</b> {calc(quantity)} / {money(cost)}")
                    period_text = " · ".join(period_parts) if period_parts else "Sin demanda periodica"
                    resource_block = Table(
                        [[
                            Paragraph(
                                f"<b>{pdf_text(row.get('recurso') or '-')}</b><br/>"
                                f"<font color='#6B7280'>{pdf_text(row.get('codigo') or '')} · Unidad: {pdf_text(row.get('unidad') or '-')} · "
                                f"Cantidad total: {calc(row.get('cantidad_total'))} · Costo total: {money(row.get('costo_total'))}</font><br/>"
                                f"<font color='#374151'>{period_text}</font>",
                                text_style,
                            )
                        ]],
                        colWidths=[doc.width],
                    )
                    resource_block.setStyle(TableStyle([
                        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFFFFF")),
                        ("BOX", (0, 0), (-1, -1), 0.25, colors.HexColor("#E4E4E7")),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ]))
                    story.append(resource_block)
                    story.append(Spacer(1, 1.2 * mm))

        doc.build(story, onFirstPage=apply_pdf_metadata, onLaterPages=apply_pdf_metadata)
        output.seek(0)
        return output

    def generate_cronograma_valorado_report(
        self,
        db: Any,
        presupuesto_id: int,
        empresa_id: int,
        template_id: str = "001",
        variant: Optional[str] = None,
        filters: Optional[Dict[str, Any]] = None,
    ) -> io.BytesIO:
        report_variant = str(variant or "valorado").lower()
        gantt_draft_summary = self._summarize_gantt_draft_for_reporting(db, presupuesto_id, empresa_id)
        if self._is_cronograma_resource_usage_variant(report_variant):
            resource_filters = filters if self._is_cronograma_resource_usage_range_variant(report_variant) else None
            resource_payload = self._attach_gantt_draft_report_warning(
                self._build_cronograma_resource_usage_payload(db, presupuesto_id, empresa_id, resource_filters),
                gantt_draft_summary,
            )
            return self._build_cronograma_resource_usage_workbook(
                db,
                presupuesto_id,
                empresa_id,
                resource_filters,
                payload_override=resource_payload,
            )
        if report_variant in {"cash_flow", "flujo_caja", "caja"}:
            format_config = self._get_empresa_format_config(db, empresa_id)
            return self._build_simple_report_workbook(
                "Flujo de Caja",
                [
                    self._attach_gantt_draft_report_warning(
                        self._build_cronograma_cash_flow_preview(db, presupuesto_id, empresa_id),
                        gantt_draft_summary,
                    )
                ],
                format_config["money_decimals"],
                format_config["calc_decimals"],
            )
        if report_variant == "pareto":
            format_config = self._get_empresa_format_config(db, empresa_id)
            return self._build_simple_report_workbook(
                "Pareto Temporal",
                [
                    self._attach_gantt_draft_report_warning(
                        self._build_cronograma_pareto_preview(db, presupuesto_id, empresa_id),
                        gantt_draft_summary,
                    )
                ],
                format_config["money_decimals"],
                format_config["calc_decimals"],
            )
        if report_variant == "gantt":
            format_config = self._get_empresa_format_config(db, empresa_id)
            return self._build_simple_report_workbook(
                "Cronograma Gantt",
                [
                    self._attach_gantt_draft_report_warning(
                        self._build_cronograma_gantt_preview(db, presupuesto_id, empresa_id),
                        gantt_draft_summary,
                    )
                ],
                format_config["money_decimals"],
                format_config["calc_decimals"],
            )
        if report_variant == "integrado":
            format_config = self._get_empresa_format_config(db, empresa_id)
            return self._build_simple_report_workbook(
                "Cronograma Integrado",
                [
                    self._attach_gantt_draft_report_warning(
                        self._build_cronograma_gantt_preview(db, presupuesto_id, empresa_id),
                        gantt_draft_summary,
                    ),
                    self._attach_gantt_draft_report_warning(
                        self._build_cronograma_valorado_preview(db, presupuesto_id, empresa_id),
                        gantt_draft_summary,
                    ),
                    self._attach_gantt_draft_report_warning(
                        self._build_cronograma_cash_flow_preview(db, presupuesto_id, empresa_id),
                        gantt_draft_summary,
                    ),
                ],
                format_config["money_decimals"],
                format_config["calc_decimals"],
            )

        presupuesto, cronograma = self._get_cronograma_valorado_payload(db, presupuesto_id, empresa_id)
        filename = "002 - Cronograma Valorado SERCOP.xlsx" if template_id == "002" else "001 - Cronograma Valorado General.xlsx"
        format_config = self._get_empresa_format_config(db, empresa_id)
        reference_code = self._get_project_reference_code(db, presupuesto)
        total_presupuesto = sum(row.precio_total for row in cronograma.rows)
        period_labels = [period.label for period in cronograma.periods]

        data = {
            "#PRO_TITULO": self._resolve_project_title(presupuesto.proyecto, presupuesto.descripcion),
            "#CODIGOPROYECTO": presupuesto.proyecto.codigo if presupuesto.proyecto else "",
            "#CODIGOREFERENCIAL": reference_code,
            "#REVISION": self._format_revision_label(presupuesto.revision),
            "#PRO_OFERENTE": presupuesto.empresa.nombre if presupuesto.empresa else "",
            "#PRO_UBICACION": self._get_project_location(db, presupuesto.proyecto, empresa_id),
            "#PRO_FECHA_INICIO": cronograma.periods[0].starts_at.strftime("%d/%m/%Y") if cronograma.periods else "",
            "#PRO_FECHA_FIN": cronograma.periods[-1].ends_at.strftime("%d/%m/%Y") if cronograma.periods else "",
            "#PRO_PLAZO": len(cronograma.periods),
            "#PRO_TIPO_PERIODO": cronograma.period_type,
            "#PRO_NUM_PERIODO": " | ".join(period_labels),
            "#TOTAL_PRESUPUESTO": total_presupuesto,
            "#TOTAL_PARCIAL": self._format_cronograma_period_values(cronograma.footer.inversion_parcial, format_config["money_decimals"]),
            "#PORCENT_PARCIAL": self._format_cronograma_period_values(cronograma.footer.avance_parcial_pct, format_config["calc_decimals"], "%"),
            "#TOTAL_ACUM": self._format_cronograma_period_values(cronograma.footer.inversion_acumulada, format_config["money_decimals"]),
            "#PORCENT_ACUM": self._format_cronograma_period_values(cronograma.footer.avance_acumulado_pct, format_config["calc_decimals"], "%"),
            "#CIUDAD": self._get_project_location(db, presupuesto.proyecto, empresa_id),
            "#PRO_FECHA": date.today().strftime("%d/%m/%Y"),
        }

        table_data = []
        budget_lines_by_id = {
            int(getattr(line, "id", 0)): line
            for line in list(presupuesto.detalle or [])
            if getattr(line, "id", None) is not None
        }
        for index, row in enumerate(cronograma.rows, start=1):
            period_amounts = [row.precio_total * (pct / 100) for pct in row.distribution]
            source_line = budget_lines_by_id.get(int(getattr(row, "linea_id", 0) or 0))
            table_data.append({
                "ITEM": row.codigo_item or index,
                "ITEM_VISIBLE": str(getattr(row, "codigo_item", None) or "").strip(),
                "CODIGO_EDT": self._get_budget_line_visible_edt_code(source_line) if source_line else "",
                "CODIGO_APU": row.apu_id or "",
                "DESCRIPCION": self._normalize_report_description(row.descripcion),
                "UNIDAD": row.unidad or "",
                "CANTIDAD": row.cantidad,
                "PUNITARIO": row.precio_unitario,
                "SUBTOTAL": row.precio_total,
                "DATOPERIODO": self._format_cronograma_period_values(period_amounts, format_config["money_decimals"]),
                "DATOS_PERIODO": self._format_cronograma_period_values(period_amounts, format_config["money_decimals"]),
            })

        formatted_data = self._format_excel_data(
            data,
            {
                "#TOTAL_PRESUPUESTO": "money",
            },
            format_config["money_decimals"],
            format_config["calc_decimals"],
        )
        formatted_rows = self._format_excel_rows(
            table_data,
            {
                "CANTIDAD": "calc",
                "PUNITARIO": "money",
                "SUBTOTAL": "money",
            },
            format_config["money_decimals"],
            format_config["calc_decimals"],
        )
        return self.fill_template(
            filename,
            formatted_data,
            formatted_rows,
            table_row_marker="#DESCRIPCION",
            use_omniclass=False,
        )
    
    def generate_vae_report(self, db: Any, presupuesto_id: int, empresa_id: int, template_id: str = "001") -> io.BytesIO:
        from app.models.presupuesto import Presupuesto
        from app.models.recurso import Recurso
        from app.models.apu import APU
        
        presupuesto = db.query(Presupuesto).filter(Presupuesto.id == presupuesto_id, Presupuesto.empresa_id == empresa_id).first()
        if not presupuesto:
            raise ValueError("Presupuesto no encontrado")

        filename = "001 - VAE Proyecto - General.xlsx"
        if template_id == "002":
            filename = "002 - VAE Proyecto - SERCOP.xlsx"

        format_config = self._get_empresa_format_config(db, empresa_id)
        use_omniclass = self._resolve_use_omniclass(format_config, template_id)
        reference_code = self._get_project_reference_code(db, presupuesto)

        project_total = float(presupuesto.total or 0)
        
        data = {
            "#PRO_TITULO": self._resolve_project_title(presupuesto.proyecto, presupuesto.descripcion),
            "#CODIGOPROYECTO": presupuesto.proyecto.codigo if presupuesto.proyecto else "",
            "#CODIGOREFERENCIAL": reference_code,
            "#REVISION": self._format_revision_label(presupuesto.revision),
            "#PRO_FECHA": presupuesto.fecha_creacion.strftime("%d/%m/%Y") if presupuesto.fecha_creacion else "",
            "#%IVA": f"{self._format_fixed(float(presupuesto.iva_aplicado or 0), format_config['money_decimals'])}%",
            "#IVA": float(presupuesto.impuestos or 0),
            "#TTOTAL": project_total
        }

        table_data = []
        project_vae_total = 0.0
        
        for line in self._sort_budget_lines_for_report(list(presupuesto.detalle or [])):
            vae_rubro = 0.0
            if line.apu:
                costo_directo = float(line.apu.costo_directo or 1)
                for linea_apu in sorted(line.apu.lineas, key=lambda item: ((item.orden or 0), item.id or 0)):
                    subtotal = float(linea_apu.subtotal or 0)
                    cpc_porc = 0.0
                    if linea_apu.recurso and linea_apu.recurso.cpc:
                        cpc_porc = float(linea_apu.recurso.cpc.porcentaje or 0)
                    vae_rubro += (subtotal / costo_directo) * (cpc_porc / 100)
            
            peso_relativo = float(line.precio_total) / project_total if project_total > 0 else 0
            vae_ponderado = vae_rubro * peso_relativo
            project_vae_total += vae_ponderado

            table_data.append({
                "ITEM": self._resolve_budget_line_report_item(line),
                "DESCRIPCION": self._normalize_report_description(line.descripcion),
                "UNIDAD": line.unidad or "",
                "CANTIDAD": float(line.cantidad) if line.cantidad else 0.0,
                "PUNITARIO": float(line.precio_unitario) if line.precio_unitario else 0.0,
                "PGLOBAL": float(line.precio_total) if line.precio_total else 0.0,
                "PESORELAT": peso_relativo,
                "VAERUBRO": vae_rubro,
                "VAEPOND": vae_ponderado
            })

        data["#VAE_TOTAL"] = project_vae_total
        formatted_data = self._format_excel_data(
            data,
            {
                "#IVA": "money",
                "#TTOTAL": "money",
                "#VAE_TOTAL": "calc",
            },
            format_config["money_decimals"],
            format_config["calc_decimals"],
        )
        formatted_rows = self._format_excel_rows(
            table_data,
            {
                "CANTIDAD": "calc",
                "PUNITARIO": "money",
                "PGLOBAL": "money",
                "PESORELAT": "percent",
                "VAERUBRO": "calc",
                "VAEPOND": "calc",
            },
            format_config["money_decimals"],
            format_config["calc_decimals"],
        )
        return self.fill_template(
            filename,
            formatted_data,
            formatted_rows,
            table_row_marker="#DESCRIPCION",
            use_omniclass=use_omniclass,
        )

    def generate_polinomica_report(self, db: Any, presupuesto_id: int, empresa_id: int, template_id: str = "001") -> io.BytesIO:
        from app.models.presupuesto import Presupuesto
        from app.models.polinomica import FormulaPolinomica
        
        presupuesto = db.query(Presupuesto).filter(Presupuesto.id == presupuesto_id, Presupuesto.empresa_id == empresa_id).first()
        if not presupuesto:
            raise ValueError("Presupuesto no encontrado")

        formula = db.query(FormulaPolinomica).filter(FormulaPolinomica.presupuesto_id == presupuesto_id).first()
        if not formula:
            raise ValueError("Fórmula polinómica no encontrada para este presupuesto")

        filename = "001 - Formula Polinomica.xlsx"
        if template_id == "002":
            filename = "002 - Formula Polinómica - sin desglose de equipo.xlsx"

        format_config = self._get_empresa_format_config(db, empresa_id)
        use_omniclass = self._resolve_use_omniclass(format_config, template_id)
        reference_code = self._get_project_reference_code(db, presupuesto)

        data = {
            "#PRO_TITULO": self._resolve_project_title(presupuesto.proyecto, presupuesto.descripcion),
            "#CODIGOPROYECTO": presupuesto.proyecto.codigo if presupuesto.proyecto else "",
            "#CODIGOREFERENCIAL": reference_code,
            "#REVISION": self._format_revision_label(presupuesto.revision),
            "#PRO_FECHA": presupuesto.fecha_creacion.strftime("%d/%m/%Y") if presupuesto.fecha_creacion else "",
            "#COEF_FIJO": float(formula.coeficiente_fijo or 0),
            "#FORMULA_POLINOMICA": ""
        }

        monomios_str = []
        table_data = []
        for m in sorted(formula.monomios, key=lambda x: x.simbolo):
            prefix = "+ " if monomios_str else ""
            monomios_str.append(f"{prefix}{self._format_fixed(float(m.coeficiente or 0), format_config['calc_decimals'])} ({m.simbolo})")
            table_data.append({
                "SIMBOLO": m.simbolo,
                "DESCRIPCION": self._normalize_report_description(m.descripcion or m.indice_inec.descripcion if m.indice_inec else ""),
                "COEFICIENTE": float(m.coeficiente),
                "COD_INEC": m.indice_inec.codigo if m.indice_inec else ""
            })
        
        data["#FORMULA_POLINOMICA"] = " ".join(monomios_str)
        formatted_data = self._format_excel_data(
            data,
            {
                "#COEF_FIJO": "calc",
            },
            format_config["money_decimals"],
            format_config["calc_decimals"],
        )
        formatted_rows = self._format_excel_rows(
            table_data,
            {
                "COEFICIENTE": "calc",
            },
            format_config["money_decimals"],
            format_config["calc_decimals"],
        )

        return self.fill_template(
            filename,
            formatted_data,
            formatted_rows,
            table_row_marker="#DESCRIPCION",
            use_omniclass=use_omniclass,
        )

reporting_service = ReportingService()
