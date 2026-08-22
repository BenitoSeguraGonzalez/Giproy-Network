from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
import pytest

from app.api.endpoints.cronogramas import (
    _build_cash_flow_from_footer,
    _build_gantt_distribution_for_periods,
    _build_periods,
)
from app.schemas.cronograma import CronogramaFooter


REPO_ROOT = Path(__file__).resolve().parents[3]
REFERENCE_DIR = REPO_ROOT / "docs" / "adicionales" / "gantt"
OLE_HEADER = bytes.fromhex("d0cf11e0a1b11ae1")
ZIP_HEADER = b"PK"


def _reference_path(name: str) -> Path:
    path = REFERENCE_DIR / name
    if not path.exists():
        pytest.skip(f"No existe la referencia local {path}")
    return path


def _cash_flow_task_row(task_name: str) -> tuple:
    workbook = load_workbook(
        _reference_path("Flujo de caja de project.xlsx"),
        data_only=True,
        read_only=True,
    )
    worksheet = workbook["Uso de tareas"]
    normalized = task_name.casefold()
    for row in worksheet.iter_rows(values_only=True):
        # The reference sheet contains summary rows with the same labels in
        # Tarea 3/Tarea 4; leaf rows only populate Tarea 4.
        if row[3] is None and isinstance(row[4], str) and row[4].casefold() == normalized:
            return row
    for row in worksheet.iter_rows(values_only=True):
        if any(isinstance(cell, str) and cell.casefold() == normalized for cell in row[:5]):
            return row
    raise AssertionError(f"No se encontro la tarea de referencia {task_name!r}")


def _period_values(row: tuple, first_column: int = 5, step: int = 2, period_count: int = 4) -> list[float]:
    return [float(row[first_column + index * step] or 0) for index in range(period_count)]


def test_reference_files_are_classified_without_direct_mpp_import():
    valorado = _reference_path("Cronograma valorado.xls")
    trabajo = _reference_path("Cronograma de trabajo.mpp")
    cash_flow = _reference_path("Flujo de caja de project.xlsx")

    assert valorado.read_bytes().startswith(OLE_HEADER)
    assert trabajo.read_bytes().startswith(OLE_HEADER)
    assert cash_flow.read_bytes().startswith(ZIP_HEADER)
    assert b"MSProject" in trabajo.read_bytes()
    assert not trabajo.read_bytes().lstrip().startswith(b"<?xml")


def test_project_cash_flow_reference_uses_monthly_cost_and_cumulative_pairs():
    workbook = load_workbook(
        _reference_path("Flujo de caja de project.xlsx"),
        data_only=True,
        read_only=True,
    )
    worksheet = workbook["Uso de tareas"]
    header_months = next(worksheet.iter_rows(min_row=5, max_row=5, values_only=True))
    header_fields = next(worksheet.iter_rows(min_row=6, max_row=6, values_only=True))

    assert [header_months[index] for index in (5, 7, 9, 11)] == [
        "Agosto",
        "Septiembre",
        "Octubre",
        "Noviembre",
    ]
    assert [header_fields[index] for index in (5, 7, 9, 11)] == ["Costo"] * 4
    assert [header_fields[index] for index in (6, 8, 10, 12)] == ["Costo acumulado"] * 4

    replanteo = _cash_flow_task_row("Replanteo y nivelación")
    assert _period_values(replanteo) == [1.0, 0.0, 0.0, 0.0]
    assert _period_values(replanteo, first_column=6) == [replanteo[14]] * 4

    excavacion = _cash_flow_task_row("Excavación cimientos")
    assert sum(_period_values(excavacion)) == pytest.approx(1.0)
    assert _period_values(excavacion, first_column=6)[-1] == pytest.approx(excavacion[14])


def test_reference_equivalent_flow_maps_gantt_periodization_to_cash_flow():
    periods = _build_periods(
        datetime(2014, 8, 4, tzinfo=timezone.utc),
        datetime(2014, 12, 1, 23, 59, 59, tzinfo=timezone.utc),
        "mensual",
        plazo_days=120,
    )
    replanteo = _cash_flow_task_row("Replanteo y nivelación")
    reference_ratios = _period_values(replanteo)
    reference_total = float(replanteo[14])

    distribution = _build_gantt_distribution_for_periods(
        periods,
        datetime(2014, 8, 4, tzinfo=timezone.utc),
        datetime(2014, 8, 30, tzinfo=timezone.utc),
    )
    mapped_ratios = [*reference_ratios, 0.0]
    footer = CronogramaFooter(
        inversion_parcial=[reference_total * ratio for ratio in mapped_ratios],
        inversion_acumulada=[reference_total for _ in mapped_ratios],
        avance_parcial_pct=[ratio * 100 for ratio in mapped_ratios],
        avance_acumulado_pct=[100.0 for _ in mapped_ratios],
    )
    cash_flow = _build_cash_flow_from_footer(periods, footer)

    assert len(periods) == 5
    assert [period.starts_at.date().isoformat() for period in periods] == [
        "2014-08-04",
        "2014-09-01",
        "2014-10-01",
        "2014-11-01",
        "2014-12-01",
    ]
    assert [round(value / 100, 6) for value in distribution] == [
        round(value, 6) for value in [*reference_ratios, 0.0]
    ]
    assert [point.cost for point in cash_flow] == pytest.approx([reference_total, 0, 0, 0, 0])
    assert cash_flow[-1].cumulative_cost == pytest.approx(reference_total)
