import io
from datetime import datetime

import pdfplumber
from openpyxl import load_workbook
from fastapi import HTTPException
from types import SimpleNamespace

from app.api.endpoints import reporting as reporting_endpoint
from app.schemas.reporting import ReportExportRequest
from app.services.reporting import reporting_service


def _preview_payload(label: str):
    return {
        "descripcion": label,
        "metadata_hint": label,
        "summary_cards": [],
        "table_columns": [
            {"key": "periodo", "label": "Periodo"},
            {"key": "costo", "label": "Costo", "kind": "money"},
        ],
        "lineas": [
            {"periodo": "P1", "costo": 100.0},
        ],
    }


def _export_payload(report_format: str) -> ReportExportRequest:
    return ReportExportRequest(
        report_type="cronograma_valorado",
        entity_ids=[1],
        format=report_format,
    )


def test_reporting_export_policy_blocks_xlsx_when_excel_exports_disabled(monkeypatch):
    monkeypatch.setattr(
        reporting_endpoint.commercial_capabilities_service,
        "resolve_company_capabilities",
        lambda *args: {
            "capabilities": {"excel_exports": False},
            "restrictions": {"requires_watermark": True},
        },
    )

    try:
        reporting_endpoint._ensure_reporting_export_allowed(_export_payload("xlsx"), None, 1)
    except HTTPException as exc:
        assert exc.status_code == 403
        assert "Excel" in exc.detail
    else:
        raise AssertionError("La politica de reporting debe bloquear XLSX sin capacidad Excel.")


def test_reporting_export_policy_blocks_pdf_excel_when_excel_exports_disabled(monkeypatch):
    monkeypatch.setattr(
        reporting_endpoint.commercial_capabilities_service,
        "resolve_company_capabilities",
        lambda *args: {
            "capabilities": {"excel_exports": False},
            "restrictions": {"requires_watermark": True},
        },
    )

    try:
        reporting_endpoint._ensure_reporting_export_allowed(_export_payload("pdf_excel"), None, 1)
    except HTTPException as exc:
        assert exc.status_code == 403
        assert "Excel" in exc.detail
    else:
        raise AssertionError("La politica de reporting debe bloquear PDF+Excel sin capacidad Excel.")


def test_reporting_export_policy_allows_pdf_when_only_excel_exports_disabled(monkeypatch):
    monkeypatch.setattr(
        reporting_endpoint.commercial_capabilities_service,
        "resolve_company_capabilities",
        lambda *args: {
            "capabilities": {"excel_exports": False},
            "restrictions": {"requires_watermark": True},
        },
    )

    commercial_state = reporting_endpoint._ensure_reporting_export_allowed(_export_payload("pdf"), None, 1)

    assert commercial_state["restrictions"]["requires_watermark"] is True


def test_reporting_watermark_text_uses_non_commercial_restriction():
    watermark_text = reporting_endpoint._build_reporting_watermark_text({
        "restrictions": {
            "requires_watermark": True,
            "non_commercial": True,
        },
    })

    assert watermark_text == "GIPROY - USO NO COMERCIAL"


def test_pdf_report_renders_server_side_watermark_text():
    watermark_text = "GIPROY - USO NO COMERCIAL"
    pdf_buffer = reporting_service._build_pdf_report(
        "Reporte QA",
        [{"codigo": "QA-001", "descripcion": "Reporte de prueba", "summary_cards": []}],
        watermark_text=watermark_text,
    )

    with pdfplumber.open(io.BytesIO(pdf_buffer.getvalue())) as pdf:
        extracted_text = "\n".join(page.extract_text() or "" for page in pdf.pages)

    assert "USO NO COMERCIAL" in extracted_text


def test_cronograma_report_preview_variants_reuse_common_reporting(monkeypatch):
    monkeypatch.setattr(reporting_service, "_build_cronograma_gantt_preview", lambda *args: _preview_payload("Gantt QA"))
    monkeypatch.setattr(reporting_service, "_build_cronograma_valorado_preview", lambda *args: _preview_payload("Valorado QA"))
    monkeypatch.setattr(reporting_service, "_build_cronograma_cash_flow_preview", lambda *args: _preview_payload("Caja QA"))
    monkeypatch.setattr(reporting_service, "_build_cronograma_pareto_preview", lambda *args: _preview_payload("Pareto QA"))
    monkeypatch.setattr(reporting_service, "_build_cronograma_resource_usage_payload", lambda *args: _preview_payload("Recursos QA"))

    cash_preview = reporting_service.preview_report(None, "cronograma_valorado", [1], 1, variant="cash_flow")
    gantt_preview = reporting_service.preview_report(None, "cronograma_valorado", [1], 1, variant="gantt")
    integrated_preview = reporting_service.preview_report(None, "cronograma_valorado", [1], 1, variant="integrado")
    pareto_preview = reporting_service.preview_report(None, "cronograma_valorado", [1], 1, variant="pareto")
    resources_preview = reporting_service.preview_report(None, "cronograma_valorado", [1], 1, variant="resources")

    assert cash_preview["title"] == "Reporte de Flujo de Caja"
    assert cash_preview["items"][0]["descripcion"] == "Caja QA"
    assert gantt_preview["title"] == "Reporte de Cronograma Gantt"
    assert gantt_preview["items"][0]["descripcion"] == "Gantt QA"
    assert pareto_preview["title"] == "Reporte de Pareto Temporal"
    assert pareto_preview["items"][0]["descripcion"] == "Pareto QA"
    assert resources_preview["title"] == "Reporte de Uso de Recursos"
    assert resources_preview["items"][0]["descripcion"] == "Recursos QA"
    assert integrated_preview["title"] == "Reporte Integrado Gantt / Valorado / Caja"
    assert [item["descripcion"] for item in integrated_preview["items"]] == [
        "Gantt QA",
        "Valorado QA",
        "Caja QA",
    ]


def test_cronograma_report_preview_warns_pending_gantt_draft_without_changing_values(monkeypatch):
    monkeypatch.setattr(reporting_service, "_build_cronograma_gantt_preview", lambda *args: _preview_payload("Gantt QA"))
    monkeypatch.setattr(
        reporting_service,
        "_summarize_gantt_draft_for_reporting",
        lambda *args: {
            "has_pending": True,
            "draft_id": 17,
            "draft_version": 4,
            "pending_count": 1,
            "invalidated_count": 1,
            "adjustment_required_count": 1,
            "total_count": 3,
            "message": "Existen 3 trabajo(s) pendientes en borrador Gantt. No se incluyen en los calculos oficiales del reporte.",
        },
    )

    preview = reporting_service.preview_report(None, "cronograma_valorado", [1], 1, variant="gantt")
    item = preview["items"][0]

    assert item["lineas"][0]["costo"] == 100.0
    assert item["warnings"][0]["code"] == "gantt_draft_pending_not_reported"
    assert item["warnings"][0]["pending_count"] == 1
    assert item["warnings"][0]["invalidated_count"] == 1
    assert item["warnings"][0]["adjustment_required_count"] == 1
    assert item["gantt_draft_summary"]["total_count"] == 3


def test_cronograma_export_cache_key_tracks_pending_gantt_draft_signature(monkeypatch):
    monkeypatch.setattr(
        reporting_service,
        "_gantt_draft_report_signature",
        lambda *args: (17, 1, 1, 0, 0),
    )
    first_key = reporting_service._build_report_export_cache_key(
        None,
        "cronograma_valorado",
        [1],
        1,
        variant="gantt",
        filters={"date_start": "2026-01-01"},
    )
    monkeypatch.setattr(
        reporting_service,
        "_gantt_draft_report_signature",
        lambda *args: (17, 2, 1, 1, 0),
    )
    second_key = reporting_service._build_report_export_cache_key(
        None,
        "cronograma_valorado",
        [1],
        1,
        variant="gantt",
        filters={"date_start": "2026-01-01"},
    )

    assert first_key != second_key


def test_cronograma_report_exports_warn_pending_gantt_draft(monkeypatch):
    monkeypatch.setattr(reporting_service, "_get_empresa_format_config", lambda *args: {"money_decimals": 2, "calc_decimals": 4})
    monkeypatch.setattr(reporting_service, "_build_cronograma_gantt_preview", lambda *args: _preview_payload("Gantt QA"))
    monkeypatch.setattr(
        reporting_service,
        "_summarize_gantt_draft_for_reporting",
        lambda *args: {
            "has_pending": True,
            "pending_count": 2,
            "invalidated_count": 0,
            "adjustment_required_count": 0,
            "total_count": 2,
            "message": "Existen 2 trabajo(s) pendientes en borrador Gantt. No se incluyen en los calculos oficiales del reporte.",
        },
    )

    workbook_buffer = reporting_service.generate_cronograma_valorado_report(
        None,
        presupuesto_id=1,
        empresa_id=1,
        variant="gantt",
    )
    workbook = load_workbook(workbook_buffer, data_only=True, read_only=True)

    assert workbook["Gantt QA"]["A4"].value.startswith("Advertencia: Existen 2 trabajo(s) pendientes")
    assert workbook["Gantt QA"]["A7"].value == "P1"
    assert workbook["Gantt QA"]["B7"].value == 100


def test_cronograma_integrated_xlsx_contains_gantt_valorado_and_cash_flow_sheets(monkeypatch):
    monkeypatch.setattr(reporting_service, "_get_empresa_format_config", lambda *args: {"money_decimals": 2, "calc_decimals": 4})
    monkeypatch.setattr(reporting_service, "_build_cronograma_gantt_preview", lambda *args: _preview_payload("Gantt QA"))
    monkeypatch.setattr(reporting_service, "_build_cronograma_valorado_preview", lambda *args: _preview_payload("Valorado QA"))
    monkeypatch.setattr(reporting_service, "_build_cronograma_cash_flow_preview", lambda *args: _preview_payload("Caja QA"))

    workbook_buffer = reporting_service.generate_cronograma_valorado_report(
        None,
        presupuesto_id=1,
        empresa_id=1,
        variant="integrado",
    )
    workbook = load_workbook(workbook_buffer, data_only=True, read_only=True)

    assert workbook.sheetnames == ["Gantt QA", "Valorado QA", "Caja QA"]
    assert workbook["Caja QA"]["A6"].value == "P1"
    assert workbook["Caja QA"]["B6"].value == 100


def test_cronograma_pareto_xlsx_uses_common_report_workbook(monkeypatch):
    monkeypatch.setattr(reporting_service, "_get_empresa_format_config", lambda *args: {"money_decimals": 2, "calc_decimals": 4})
    monkeypatch.setattr(reporting_service, "_build_cronograma_pareto_preview", lambda *args: _preview_payload("Pareto QA"))

    workbook_buffer = reporting_service.generate_cronograma_valorado_report(
        None,
        presupuesto_id=1,
        empresa_id=1,
        variant="pareto",
    )
    workbook = load_workbook(workbook_buffer, data_only=True, read_only=True)

    assert workbook.sheetnames == ["Pareto QA"]
    assert workbook["Pareto QA"]["A6"].value == "P1"
    assert workbook["Pareto QA"]["B6"].value == 100


def test_cronograma_resource_usage_explodes_nested_apus_and_keeps_all_periods(monkeypatch):
    unidad_m3 = SimpleNamespace(simbolo="m3", descripcion="metro cubico")
    unidad_u = SimpleNamespace(simbolo="u", descripcion="unidad")
    subcat_materiales = SimpleNamespace(subcategoria_codigo=2, codigo="2-001", descripcion="Materiales petreos")

    cement = SimpleNamespace(
        id=101,
        codigo="2-001-001",
        descripcion="Cemento portland",
        precio="10",
        unidad=unidad_m3,
        subcategoria_item=subcat_materiales,
    )
    sand = SimpleNamespace(
        id=102,
        codigo="2-001-002",
        descripcion="Arena fina",
        precio="5",
        unidad=unidad_u,
        subcategoria_item=subcat_materiales,
    )
    child_apu = SimpleNamespace(
        id=202,
        lineas=[
            SimpleNamespace(id=3, orden=1, cantidad="3", rendimiento="1", recurso=cement, recurso_id=cement.id),
            SimpleNamespace(id=4, orden=2, cantidad="4", rendimiento="1", recurso=sand, recurso_id=sand.id),
        ],
    )
    parent_apu = SimpleNamespace(
        id=201,
        lineas=[
            SimpleNamespace(id=1, orden=1, cantidad="2", rendimiento="1", recurso=cement, recurso_id=cement.id),
            SimpleNamespace(id=2, orden=2, cantidad="0.5", rendimiento="1", apu_hijo=child_apu, apu_hijo_id=child_apu.id),
        ],
    )
    presupuesto = SimpleNamespace(
        id=77,
        descripcion="Presupuesto QA",
        proyecto=SimpleNamespace(codigo="QA-001", codigo_root="QA-001"),
        detalle=[
            SimpleNamespace(id=10, apu_id=parent_apu.id, apu=parent_apu, cantidad="10"),
        ],
    )
    cronograma = SimpleNamespace(
        period_type="mensual",
        distribution_mode="gantt",
        moneda="USD",
        periods=[
            SimpleNamespace(label="P1"),
            SimpleNamespace(label="P2"),
            SimpleNamespace(label="P3"),
        ],
        rows=[
            SimpleNamespace(linea_id=10, apu_id=parent_apu.id, cantidad="10", distribution=[50, 0, 50]),
        ],
    )

    monkeypatch.setattr(reporting_service, "_get_cronograma_valorado_payload", lambda *args: (presupuesto, cronograma))
    monkeypatch.setattr(reporting_service, "_resolve_project_title", lambda *args: "Proyecto QA")

    payload = reporting_service._build_cronograma_resource_usage_payload(None, presupuesto.id, 3)
    rows_by_resource = {row["recurso"]: row for row in payload["resource_usage_rows"]}

    assert payload["summary_cards"][0]["value"] == 2
    assert len(payload["periods"]) == 3
    assert set(rows_by_resource) == {"Cemento portland", "Arena fina"}
    assert rows_by_resource["Cemento portland"]["cantidad_total"] == 35.0
    assert rows_by_resource["Cemento portland"]["costo_total"] == 350.0
    assert rows_by_resource["Cemento portland"]["period_quantities"] == [17.5, 0.0, 17.5]
    assert rows_by_resource["Cemento portland"]["period_costs"] == [175.0, 0.0, 175.0]
    assert rows_by_resource["Arena fina"]["cantidad_total"] == 20.0
    assert rows_by_resource["Arena fina"]["period_quantities"] == [10.0, 0.0, 10.0]
    assert rows_by_resource["Arena fina"]["period_costs"] == [50.0, 0.0, 50.0]


def test_cronograma_resource_usage_reads_materialized_parent_and_child_apu_lines_after_gantt_apply(monkeypatch):
    unidad_h = SimpleNamespace(simbolo="h", descripcion="hora")
    subcat_mano_obra = SimpleNamespace(subcategoria_codigo=4, codigo="4-001", descripcion="Albanileria")
    peon = SimpleNamespace(
        id=102,
        codigo="4-001-001",
        descripcion="Peon",
        precio="4",
        unidad=unidad_h,
        subcategoria_item=subcat_mano_obra,
    )
    child_apu = SimpleNamespace(
        id=202,
        lineas=[
            SimpleNamespace(id=3, orden=1, cantidad="2", rendimiento="0.75", recurso=peon, recurso_id=peon.id),
        ],
    )
    parent_apu = SimpleNamespace(
        id=201,
        lineas=[
            SimpleNamespace(id=1, orden=1, cantidad="1", rendimiento="1", recurso=peon, recurso_id=peon.id),
            SimpleNamespace(id=2, orden=2, cantidad="0.62", rendimiento="1", apu_hijo=child_apu, apu_hijo_id=child_apu.id),
        ],
    )
    presupuesto = SimpleNamespace(
        id=77,
        descripcion="Presupuesto QA",
        proyecto=SimpleNamespace(codigo="QA-001", codigo_root="QA-001"),
        detalle=[
            SimpleNamespace(id=10, apu_id=parent_apu.id, apu=parent_apu, cantidad="10"),
        ],
    )
    cronograma = SimpleNamespace(
        period_type="mensual",
        distribution_mode="gantt",
        moneda="USD",
        periods=[SimpleNamespace(label="P1"), SimpleNamespace(label="P2")],
        rows=[SimpleNamespace(linea_id=10, apu_id=parent_apu.id, cantidad="10", distribution=[25, 75])],
    )

    monkeypatch.setattr(reporting_service, "_get_cronograma_valorado_payload", lambda *args: (presupuesto, cronograma))
    monkeypatch.setattr(reporting_service, "_resolve_project_title", lambda *args: "Proyecto QA")

    payload = reporting_service._build_cronograma_resource_usage_payload(None, presupuesto.id, 3)
    rows_by_resource = {row["recurso"]: row for row in payload["resource_usage_rows"]}

    assert set(rows_by_resource) == {"Peon"}
    assert rows_by_resource["Peon"]["cantidad_total"] == 12.0
    assert rows_by_resource["Peon"]["costo_total"] == 48.0
    assert rows_by_resource["Peon"]["period_quantities"] == [3.0, 9.0]
    assert rows_by_resource["Peon"]["period_costs"] == [12.0, 36.0]


def test_cronograma_resource_usage_blocks_apu_without_reachable_resources(monkeypatch):
    empty_apu = SimpleNamespace(id=201, codigo="5-001", descripcion="APU sin recursos", lineas=[])
    presupuesto = SimpleNamespace(
        id=77,
        descripcion="Presupuesto QA",
        proyecto=SimpleNamespace(codigo="QA-001", codigo_root="QA-001"),
        detalle=[SimpleNamespace(id=10, apu_id=empty_apu.id, apu=empty_apu, cantidad="10")],
    )
    cronograma = SimpleNamespace(
        period_type="mensual",
        distribution_mode="gantt",
        moneda="USD",
        periods=[SimpleNamespace(label="P1")],
        rows=[SimpleNamespace(linea_id=10, apu_id=empty_apu.id, cantidad="10", distribution=[100])],
    )

    monkeypatch.setattr(reporting_service, "_get_cronograma_valorado_payload", lambda *args: (presupuesto, cronograma))

    try:
        reporting_service._build_cronograma_resource_usage_payload(None, presupuesto.id, 3)
    except HTTPException as exc:
        assert exc.status_code == 422
        assert exc.detail["code"] == "apu_resources_incomplete"
        assert exc.detail["issues"][0]["apu_id"] == empty_apu.id
    else:
        raise AssertionError("El reporte de recursos debe bloquear APUs sin recursos alcanzables.")


def test_cronograma_resource_usage_range_prorates_by_period_overlap(monkeypatch):
    unidad_u = SimpleNamespace(simbolo="u", descripcion="unidad")
    resource = SimpleNamespace(
        id=501,
        codigo="4-001-001",
        descripcion="Cuadrilla albañil",
        precio="20",
        unidad=unidad_u,
        subcategoria_item=SimpleNamespace(subcategoria_codigo=4, codigo="4-001", descripcion="Mano de obra"),
    )
    apu = SimpleNamespace(
        id=601,
        lineas=[
            SimpleNamespace(id=1, orden=1, cantidad="1", rendimiento="1", recurso=resource, recurso_id=resource.id),
        ],
    )
    presupuesto = SimpleNamespace(
        id=88,
        descripcion="Presupuesto QA",
        proyecto=SimpleNamespace(codigo="QA-002", codigo_root="QA-002"),
        detalle=[SimpleNamespace(id=20, apu_id=apu.id, apu=apu, cantidad="100")],
    )
    cronograma = SimpleNamespace(
        period_type="mensual",
        distribution_mode="gantt",
        moneda="USD",
        periods=[
            SimpleNamespace(label="P1", starts_at=datetime(2026, 1, 1), ends_at=datetime(2026, 1, 31)),
            SimpleNamespace(label="P2", starts_at=datetime(2026, 2, 1), ends_at=datetime(2026, 2, 28)),
        ],
        rows=[SimpleNamespace(linea_id=20, apu_id=apu.id, cantidad="100", distribution=[50, 50])],
    )

    monkeypatch.setattr(reporting_service, "_get_cronograma_valorado_payload", lambda *args: (presupuesto, cronograma))
    monkeypatch.setattr(reporting_service, "_resolve_project_title", lambda *args: "Proyecto QA")

    payload = reporting_service._build_cronograma_resource_usage_payload(
        None,
        presupuesto.id,
        3,
        {"date_start": "2026-01-01T00:00:00", "date_end": "2026-01-16T00:00:00"},
    )
    row = payload["resource_usage_rows"][0]

    assert payload["filters"] == {"date_start": "2026-01-01", "date_end": "2026-01-16"}
    assert payload["summary_cards"][1]["value"] == 1
    assert len(payload["periods"]) == 1
    assert row["period_quantities"] == [25.0]
    assert row["period_costs"] == [500.0]
    assert row["cantidad_total"] == 25.0
    assert row["costo_total"] == 500.0


def test_cronograma_resource_usage_range_workbook_and_pdf_use_range_title(monkeypatch):
    monkeypatch.setattr(reporting_service, "_get_empresa_format_config", lambda *args: {"money_decimals": 2, "calc_decimals": 4})
    monkeypatch.setattr(reporting_service, "_build_cronograma_resource_usage_payload", lambda *args: {
        "descripcion": "Proyecto QA",
        "metadata_hint": "Uso de recursos por rango · mensual · 1 periodo(s) · rango 01/01/2026 - 16/01/2026",
        "filters": {"date_start": "2026-01-01", "date_end": "2026-01-16"},
        "resource_usage_rows": [
            {
                "categoria": "4. Mano de Obra",
                "subcategoria": "4-001 - Mano de obra",
                "recurso": "Cuadrilla albañil",
                "codigo": "4-001-001",
                "unidad": "u",
                "period_quantities": [25.0],
                "period_costs": [500.0],
                "cantidad_total": 25.0,
                "costo_total": 500.0,
            }
        ],
        "periods": [{"label": "P1"}],
    })

    workbook_buffer = reporting_service.generate_cronograma_valorado_report(
        None,
        presupuesto_id=1,
        empresa_id=1,
        variant="resources_range",
        filters={"date_start": "2026-01-01", "date_end": "2026-01-16"},
    )
    workbook = load_workbook(workbook_buffer, data_only=True, read_only=True)
    sheet = workbook["Recursos por rango"]
    assert sheet["A2"].value == "Cronograma de Uso de Recursos por Rango"
    assert sheet["F9"].value == 25
    assert sheet["G9"].value == 500

    pdf_buffer = reporting_service.generate_preview_pdf(
        None,
        "cronograma_valorado",
        [1],
        1,
        variant="resources_range",
        filters={"date_start": "2026-01-01", "date_end": "2026-01-16"},
    )
    with pdfplumber.open(io.BytesIO(pdf_buffer.getvalue())) as pdf:
        extracted_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Uso de Recursos por Rango" in extracted_text
    assert "Cuadrilla albañil" in extracted_text


def test_cronograma_resource_usage_exports_warn_pending_gantt_draft(monkeypatch):
    monkeypatch.setattr(reporting_service, "_get_empresa_format_config", lambda *args: {"money_decimals": 2, "calc_decimals": 4})
    monkeypatch.setattr(reporting_service, "_build_cronograma_resource_usage_payload", lambda *args: {
        "descripcion": "Proyecto QA",
        "metadata_hint": "Uso de recursos QA",
        "resource_usage_rows": [
            {
                "categoria": "2. Materiales",
                "subcategoria": "2-001 - Materiales petreos",
                "recurso": "Cemento portland",
                "codigo": "2-001-001",
                "unidad": "m3",
                "period_quantities": [1.5],
                "period_costs": [15.0],
                "cantidad_total": 1.5,
                "costo_total": 15.0,
            }
        ],
        "periods": [{"label": "P1"}],
    })
    monkeypatch.setattr(
        reporting_service,
        "_summarize_gantt_draft_for_reporting",
        lambda *args: {
            "has_pending": True,
            "pending_count": 1,
            "invalidated_count": 0,
            "adjustment_required_count": 0,
            "total_count": 1,
            "message": "Existen 1 trabajo(s) pendientes en borrador Gantt. No se incluyen en los calculos oficiales del reporte.",
        },
    )

    workbook_buffer = reporting_service.generate_cronograma_valorado_report(
        None,
        presupuesto_id=1,
        empresa_id=1,
        variant="resources",
    )
    workbook = load_workbook(workbook_buffer, data_only=True, read_only=True)
    sheet = workbook["Uso de recursos"]

    assert sheet["A4"].value.startswith("Advertencia: Existen 1 trabajo(s) pendientes")
    assert sheet["F5"].value == "P1"
    assert sheet["F9"].value == 1.5

    pdf_buffer = reporting_service.generate_preview_pdf(
        None,
        "cronograma_valorado",
        [1],
        1,
        variant="resources",
    )
    with pdfplumber.open(io.BytesIO(pdf_buffer.getvalue())) as pdf:
        extracted_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Advertencia: Existen 1 trabajo(s) pendientes" in extracted_text


def test_cronograma_resource_usage_workbook_keeps_zero_period_columns(monkeypatch):
    monkeypatch.setattr(reporting_service, "_get_empresa_format_config", lambda *args: {"money_decimals": 2, "calc_decimals": 4})
    monkeypatch.setattr(reporting_service, "_build_cronograma_resource_usage_payload", lambda *args: {
        "descripcion": "Proyecto QA",
        "metadata_hint": "Uso de recursos QA",
        "resource_usage_rows": [
            {
                "categoria": "2. Materiales",
                "subcategoria": "2-001 - Materiales petreos",
                "recurso": "Cemento portland",
                "codigo": "2-001-001",
                "unidad": "m3",
                "period_quantities": [1.5, 0.0, 2.0],
                "period_costs": [15.0, 0.0, 20.0],
                "cantidad_total": 3.5,
                "costo_total": 35.0,
            }
        ],
        "periods": [
            {"label": "P1"},
            {"label": "P2"},
            {"label": "P3"},
        ],
    })

    workbook_buffer = reporting_service.generate_cronograma_valorado_report(
        None,
        presupuesto_id=1,
        empresa_id=1,
        variant="resources",
    )
    workbook = load_workbook(workbook_buffer, data_only=True, read_only=True)
    sheet = workbook["Uso de recursos"]

    assert sheet["F5"].value == "P1"
    assert sheet["H5"].value == "P2"
    assert sheet["J5"].value == "P3"
    assert sheet["F9"].value == 1.5
    assert sheet["G9"].value == 15
    assert sheet["H9"].value == 0
    assert sheet["I9"].value == 0
    assert sheet["J9"].value == 2
    assert sheet["K9"].value == 20
    assert sheet["L9"].value == 3.5
    assert sheet["M9"].value == 35


def test_cash_flow_preview_includes_hours_column(monkeypatch):
    class _Point:
        label = "P1"
        starts_at = __import__("datetime").datetime(2026, 3, 24, 8, 0, 0)
        ends_at = __import__("datetime").datetime(2026, 3, 31, 16, 0, 0)
        work_hours = 64.0
        cost = 100.0
        cumulative_cost = 100.0
        cost_pct = 25.0
        cumulative_pct = 25.0
        category_costs = {"Equipos": 60.0, "Mano de obra": 25.0, "Transporte": 15.0}
        dominant_category = "Equipos"

    class _Cronograma:
        cash_flow = [_Point()]
        period_type = "semanal"
        distribution_mode = "gantt"
        moneda = "USD"

    class _Proyecto:
        codigo = "PRO-1"
        codigo_root = "PRO-1"

    class _Presupuesto:
        id = 1
        descripcion = "Presupuesto QA"
        proyecto = _Proyecto()

    monkeypatch.setattr(reporting_service, "_get_cronograma_valorado_payload", lambda *args: (_Presupuesto(), _Cronograma()))
    monkeypatch.setattr(reporting_service, "_resolve_project_title", lambda *args: "Proyecto QA")

    preview = reporting_service._build_cronograma_cash_flow_preview(None, 1, 1)

    assert preview["summary_cards"][2]["label"] == "Horas efectivas"
    assert preview["summary_cards"][2]["value"] == 64.0
    assert preview["summary_cards"][3]["label"] == "Categoría líder"
    assert preview["summary_cards"][3]["value"] == "Equipos"
    assert preview["table_columns"][3]["key"] == "horas_periodo"
    assert preview["table_columns"][5]["key"] == "equipos"
    assert preview["lineas"][0]["equipos"] == 60.0
    assert preview["lineas"][0]["fecha_inicio"] == "24/03/2026 08:00"


def test_cronograma_valorado_preview_reports_pending_manual_coverage(monkeypatch):
    class _Row:
        def __init__(self, linea_id, codigo_item, descripcion, precio_total, distribution, requires_manual_schedule=False):
            self.linea_id = linea_id
            self.codigo_item = codigo_item
            self.descripcion = descripcion
            self.unidad = "u"
            self.cantidad = 1.0
            self.precio_unitario = precio_total
            self.precio_total = precio_total
            self.distribution = distribution
            self.requires_manual_schedule = requires_manual_schedule

    class _Period:
        def __init__(self, label):
            self.label = label

    class _Cronograma:
        rows = [
            _Row(1, "1.1.1", "Línea programada", 100.0, [100.0, 0.0], requires_manual_schedule=False),
            _Row(2, "1.1.2", "Línea pendiente", 50.0, [0.0, 0.0], requires_manual_schedule=True),
        ]
        periods = [_Period("P1"), _Period("P2")]
        period_type = "mensual"
        distribution_mode = "gantt"
        moneda = "USD"
        cash_flow = []

    class _Proyecto:
        codigo = "PRO-1"
        codigo_root = "PRO-1"

    class _LineaPresupuesto:
        def __init__(self, line_id):
            self.id = line_id
            self.edt_id = None
            self.codigo_item = f"COD-{line_id}"

    class _Presupuesto:
        id = 1
        descripcion = "Presupuesto QA"
        proyecto = _Proyecto()
        detalle = [_LineaPresupuesto(1), _LineaPresupuesto(2)]

    monkeypatch.setattr(reporting_service, "_get_cronograma_valorado_payload", lambda *args: (_Presupuesto(), _Cronograma()))
    monkeypatch.setattr(reporting_service, "_resolve_project_title", lambda *args: "Proyecto QA")
    monkeypatch.setattr(reporting_service, "_get_budget_line_visible_edt_code", lambda *args: "1.1")

    preview = reporting_service._build_cronograma_valorado_preview(None, 1, 1)

    assert "temporalidad parcial" in preview["metadata_hint"]
    assert preview["summary_cards"][3]["label"] == "Pendiente manual"
    assert preview["summary_cards"][3]["value"] == 50.0
    assert preview["summary_cards"][4]["label"] == "Cobertura valorada"
    assert round(float(preview["summary_cards"][4]["value"]), 4) == 66.6667


def test_cronograma_cash_flow_preview_reports_partial_origin_when_manual_pending(monkeypatch):
    class _Point:
        label = "P1"
        starts_at = __import__("datetime").datetime(2026, 3, 24, 8, 0, 0)
        ends_at = __import__("datetime").datetime(2026, 3, 31, 16, 0, 0)
        work_hours = 64.0
        cost = 100.0
        cumulative_cost = 100.0
        cost_pct = 66.6667
        cumulative_pct = 66.6667
        category_costs = {"Equipos": 60.0, "Mano de obra": 25.0, "Transporte": 15.0}
        dominant_category = "Equipos"

    class _Row:
        def __init__(self, precio_total, requires_manual_schedule=False):
            self.precio_total = precio_total
            self.requires_manual_schedule = requires_manual_schedule

    class _Cronograma:
        cash_flow = [_Point()]
        period_type = "semanal"
        distribution_mode = "gantt"
        moneda = "USD"
        rows = [_Row(100.0, False), _Row(50.0, True)]

    class _Proyecto:
        codigo = "PRO-1"
        codigo_root = "PRO-1"

    class _Presupuesto:
        id = 1
        descripcion = "Presupuesto QA"
        proyecto = _Proyecto()

    monkeypatch.setattr(reporting_service, "_get_cronograma_valorado_payload", lambda *args: (_Presupuesto(), _Cronograma()))
    monkeypatch.setattr(reporting_service, "_resolve_project_title", lambda *args: "Proyecto QA")

    preview = reporting_service._build_cronograma_cash_flow_preview(None, 1, 1)

    assert "cobertura parcial" in preview["metadata_hint"]
    assert preview["summary_cards"][4]["label"] == "Origen"
    assert preview["summary_cards"][4]["value"] == "Gantt parcial"
    assert preview["summary_cards"][5]["label"] == "Pendiente manual"
    assert preview["summary_cards"][5]["value"] == 50.0


def test_cronograma_gantt_preview_reports_operational_reconciliation(monkeypatch):
    class _Summary:
        trabajo_total = 12.0
        cuadrilla_total = 3.0

    class _Config:
        dias_laborables_semana = 5

    class _Row:
        def __init__(self):
            self.presupuesto_linea_id = 1
            self.codigo_item = "1.1.1"
            self.descripcion = "Replanteo y nivelación"
            self.start_date = __import__("datetime").datetime(2026, 4, 1)
            self.end_date = __import__("datetime").datetime(2026, 4, 10)
            self.dias_calendario = 10
            self.dias_utiles = 8
            self.progress_pct = 25.0
            self.predecessors = []
            self.dependencies = []
            self.metadata = {
                "gantt_operational": {
                    "budget_line_id": "1",
                    "has_subbars": True,
                    "subbar_count": 2,
                    "has_manual_temporal_window": True,
                },
                "gantt_subbars": [
                    {"period_id": "P1", "status": "draft_session", "percent": 6.0, "amount": 60.0},
                    {"period_id": "P1", "status": "accepted_session", "source": "gantt_interparent_merge", "percent": 5.0, "amount": 50.0},
                ],
            }

    class _Schedule:
        rows = [_Row()]
        summary = _Summary()
        config = _Config()

    class _ValoradoRow:
        linea_id = 1
        distribution = [10.0, 0.0]
        precio_total = 100.0

    class _Period:
        def __init__(self, period_id):
            self.id = period_id

    class _Cronograma:
        rows = [_ValoradoRow()]
        periods = [_Period("P1"), _Period("P2")]

    class _Proyecto:
        codigo = "PRO-1"
        codigo_root = "PRO-1"
        id = 7

    class _LineaPresupuesto:
        id = 1

    class _Presupuesto:
        id = 1
        descripcion = "Presupuesto QA"
        proyecto = _Proyecto()
        proyecto_id = 7
        detalle = [_LineaPresupuesto()]

    monkeypatch.setattr(reporting_service, "_get_presupuesto", lambda *args: _Presupuesto())
    monkeypatch.setattr(reporting_service, "_resolve_project_title", lambda *args: "Proyecto QA")
    monkeypatch.setattr(reporting_service, "_get_budget_line_visible_edt_code", lambda *args: "1.1")
    monkeypatch.setattr(reporting_service, "_get_cronograma_valorado_payload", lambda *args: (_Presupuesto(), _Cronograma()))

    from app.services.cronograma_trabajo import cronograma_trabajo_service

    monkeypatch.setattr(cronograma_trabajo_service, "get_schedule", lambda *args, **kwargs: _Schedule())

    preview = reporting_service._build_cronograma_gantt_preview(None, 1, 1)

    assert "reconciliación operativa" in preview["metadata_hint"]
    assert preview["summary_cards"][3]["label"] == "Líneas con subbarras"
    assert preview["summary_cards"][3]["value"] == 1
    assert preview["summary_cards"][4]["label"] == "Tramos operativos"
    assert preview["summary_cards"][4]["value"] == 2
    assert preview["summary_cards"][5]["label"] == "Tramos en borrador"
    assert preview["summary_cards"][5]["value"] == 1
    assert preview["summary_cards"][6]["label"] == "Fusiones intertramo"
    assert preview["summary_cards"][6]["value"] == 1
    assert preview["summary_cards"][7]["label"] == "Ventanas manuales"
    assert preview["summary_cards"][7]["value"] == 1
    assert preview["summary_cards"][8]["label"] == "Conflictos valorado"
    assert preview["summary_cards"][8]["value"] == 1


def test_summarize_gantt_operational_reporting_ignores_manual_window_without_conflict():
    class _ScheduleRow:
        def __init__(self):
            self.presupuesto_linea_id = 2
            self.metadata = {
                "gantt_operational": {
                    "budget_line_id": "2",
                    "has_subbars": False,
                    "subbar_count": 0,
                    "has_manual_temporal_window": True,
                },
                "manual_temporal_window": {
                    "starts_at": "2026-04-01T08:00:00",
                    "ends_at": "2026-04-03T17:00:00",
                    "source": "manual_temporal_material_only",
                },
            }

    class _Schedule:
        rows = [_ScheduleRow()]

    class _ValoradoRow:
        linea_id = 2
        distribution = [0.0]
        precio_total = 50.0

    class _Period:
        id = "P1"

    class _Cronograma:
        rows = [_ValoradoRow()]
        periods = [_Period()]

    summary = reporting_service._summarize_gantt_operational_reporting(_Schedule(), _Cronograma())

    assert summary["manual_temporal_rows"] == 1
    assert summary["subbar_rows"] == 0
    assert summary["draft_subbar_count"] == 0
    assert summary["interparent_merge_rows"] == 0
    assert summary["interparent_merge_count"] == 0
    assert summary["conflict_rows"] == 0
    assert summary["conflict_periods"] == 0
    assert summary["has_operational_state"] is True


def test_summarize_gantt_operational_reporting_counts_interparent_merge():
    class _ScheduleRow:
        def __init__(self):
            self.presupuesto_linea_id = 3
            self.metadata = {
                "gantt_operational": {
                    "budget_line_id": "3",
                    "has_subbars": True,
                    "subbar_count": 2,
                    "has_manual_temporal_window": False,
                },
                "gantt_subbars": [
                    {"period_id": "P1", "status": "accepted_session", "source": "gantt_interparent_merge", "percent": 6.0, "amount": 60.0},
                    {"period_id": "P2", "status": "draft_session", "source": "gantt_subbar_split", "percent": 4.0, "amount": 40.0},
                ],
            }

    class _Schedule:
        rows = [_ScheduleRow()]

    class _ValoradoRow:
        linea_id = 3
        distribution = [0.0, 10.0]
        precio_total = 100.0

    class _Period:
        def __init__(self, period_id):
            self.id = period_id

    class _Cronograma:
        rows = [_ValoradoRow()]
        periods = [_Period("P1"), _Period("P2")]

    summary = reporting_service._summarize_gantt_operational_reporting(_Schedule(), _Cronograma())

    assert summary["subbar_rows"] == 1
    assert summary["subbar_count"] == 2
    assert summary["draft_subbar_count"] == 1
    assert summary["interparent_merge_rows"] == 1
    assert summary["interparent_merge_count"] == 1
    assert summary["conflict_rows"] == 1
    assert summary["conflict_periods"] == 2
