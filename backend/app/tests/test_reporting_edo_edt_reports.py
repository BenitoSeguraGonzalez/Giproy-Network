import openpyxl
from pathlib import Path
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

from app.models.edo import EdoNode, TipoNodo
from app.models.edt import EdtNode, TipoNodoEdt
from app.models.proyecto import Proyecto
from app.models.stakeholder import Rol, Stakeholder
from app.services.reporting import ReportingService


REPORT_TEMPLATES_DIR = Path(__file__).resolve().parents[3] / "docs" / "reportes"


def _create_project(db, sample_empresa):
    project = Proyecto(
        nombre="Proyecto Reportes EDO EDT",
        codigo="PRO-EDO-EDT",
        codigo_root="PRO-EDO-EDT",
        revision=2,
        empresa_id=sample_empresa.id,
    )
    db.add(project)
    db.commit()
    db.refresh(project)
    return project


def _create_stakeholder_role(db, sample_empresa, project):
    stakeholder = Stakeholder(
        codigo="STK-0001",
        nombre="Santiago",
        apellidos="Bermeo",
        email="santiago@example.com",
        proyecto_codigo_root=project.codigo_root,
        empresa_id=sample_empresa.id,
    )
    role = Rol(
        codigo="ROL-0001",
        nombre="Gerente de Proyecto",
        empresa_id=sample_empresa.id,
    )
    db.add_all([stakeholder, role])
    db.commit()
    db.refresh(stakeholder)
    db.refresh(role)
    return stakeholder, role


def _sheet_values(buffer):
    workbook = openpyxl.load_workbook(buffer, data_only=False)
    worksheet = workbook.active
    values = []
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value not in (None, ""):
                values.append(str(cell.value))
    return values


def _assert_no_overlapping_merged_ranges(worksheet):
    ranges = list(worksheet.merged_cells.ranges)
    for index, current in enumerate(ranges):
        for candidate in ranges[index + 1:]:
            assert current.isdisjoint(candidate), f"Merged ranges overlap: {current} / {candidate}"


def test_generated_by_giproy_label_is_reusable_for_document_reports():
    service = ReportingService()

    assert service._build_generated_by_giproy_label("edo") == "Reporte EDO generado por GIPROY®"
    assert service._build_generated_by_giproy_label("EDT") == "Reporte EDT generado por GIPROY®"
    assert service._build_generated_by_giproy_label("") == "Reporte documental generado por GIPROY®"


def test_edo_report_generates_from_template_without_unresolved_placeholders(db, sample_empresa):
    project = _create_project(db, sample_empresa)
    stakeholder, role = _create_stakeholder_role(db, sample_empresa, project)
    db.add_all([
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.HITO,
            orden=1,
            codigo="1",
            nombre="Dirección del proyecto",
        ),
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.HITO,
            orden=2,
            codigo="2",
            nombre="Control de calidad",
        ),
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.STAKEHOLDER,
            orden=3,
            codigo="1.R1",
            stakeholder_id=stakeholder.id,
            rol_id=role.id,
            actividades_claves="Coordinar decisiones principales",
        ),
    ])
    db.commit()

    values = _sheet_values(ReportingService().generate_edo_report(db, project.id, sample_empresa.id))

    assert "PRO-EDO-EDT" in values
    assert "Gerente de Proyecto" in values
    assert "Santiago Bermeo" in values
    assert "Coordinar decisiones principales" in values
    assert not any(value.startswith("#") for value in values)


def test_edo_generated_workbook_keeps_signature_in_two_line_wide_block(db, sample_empresa):
    project = _create_project(db, sample_empresa)
    stakeholder, role = _create_stakeholder_role(db, sample_empresa, project)
    nodes = [
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.HITO,
            orden=index,
            codigo=str(index),
            nombre=f"Hito {index}",
        )
        for index in range(1, 10)
    ]
    nodes.append(
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.STAKEHOLDER,
            orden=10,
            codigo="10.R1",
            stakeholder_id=stakeholder.id,
            rol_id=role.id,
        )
    )
    db.add_all(nodes)
    db.commit()

    workbook = openpyxl.load_workbook(ReportingService().generate_edo_report(db, project.id, sample_empresa.id))
    worksheet = workbook.active
    signature_cell = next(
        cell
        for row in worksheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and "FIRMA DEL" in cell.value
    )
    signature_range = f"A{signature_cell.row}:E{signature_cell.row + 1}"

    assert signature_cell.value == "FIRMA DEL SPONSOR O\nPERSONA QUE AUTORIZA"
    assert signature_cell.font.sz <= 8
    assert signature_cell.alignment.wrap_text is True
    assert signature_range in [str(merged_range) for merged_range in worksheet.merged_cells.ranges]


def test_edo_preview_pdf_contract_matches_edo_tree_semantics(db, sample_empresa):
    project = _create_project(db, sample_empresa)
    stakeholder, role = _create_stakeholder_role(db, sample_empresa, project)
    db.add_all([
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.HITO,
            orden=1,
            codigo="1",
            nombre="Dirección del proyecto",
        ),
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.HITO,
            orden=2,
            codigo="2",
            nombre="Control de calidad",
        ),
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.STAKEHOLDER,
            orden=3,
            codigo="1.R1",
            stakeholder_id=stakeholder.id,
            rol_id=role.id,
            actividades_claves="Coordinar decisiones principales",
        ),
    ])
    db.commit()

    preview = ReportingService().preview_report(db, "edo", [project.id], sample_empresa.id)
    item = preview["items"][0]

    assert item["metadata_hint"] == "Documento EDO · Hitos y responsabilidades del proyecto"
    assert item["preview_layout"] == "edo_document"
    assert item["summary_cards"] == [
        {"label": "Nodos", "value": 3, "kind": "integer"},
        {"label": "Hitos", "value": 2, "kind": "integer"},
        {"label": "Responsables", "value": 1, "kind": "integer", "tone": "total"},
    ]
    assert [column["key"] for column in item["table_columns"]] == [
        "item",
        "rol_asignado",
        "nombre_responsable",
        "actividades_clave",
    ]
    assert {"cantidad", "precio", "rendimiento", "subtotal"}.isdisjoint(
        {column["key"] for column in item["table_columns"]}
    )
    assert item["lineas"][0]["item"] == 1
    assert item["lineas"][0]["rol_asignado"] == "Gerente de Proyecto"
    assert item["lineas"][0]["nombre_responsable"] == "Santiago Bermeo"
    assert item["lineas"][0]["actividades_clave"] == "Coordinar decisiones principales"
    assert item["lineas"][1]["item"] == 2
    assert item["lineas"][1]["rol_asignado"] == "Hito"
    assert item["lineas"][1]["nombre_responsable"] == "Control de calidad"
    assert len(item["lineas"]) == 2


def test_edo_pdf_direct_generates_document_pdf(db, sample_empresa):
    project = _create_project(db, sample_empresa)
    stakeholder, role = _create_stakeholder_role(db, sample_empresa, project)
    db.add_all([
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.HITO,
            orden=1,
            codigo="1",
            nombre="Dirección del proyecto",
        ),
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.HITO,
            orden=2,
            codigo="2",
            nombre="Control de calidad",
        ),
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.STAKEHOLDER,
            orden=3,
            codigo="1.R1",
            stakeholder_id=stakeholder.id,
            rol_id=role.id,
            actividades_claves="Coordinar decisiones principales",
        ),
    ])
    db.commit()

    payload = ReportingService().generate_preview_pdf(db, "edo", [project.id], sample_empresa.id).getvalue()

    assert payload.startswith(b"%PDF")
    assert len(payload) > 1500
    assert ReportingService.REPORT_EXPORT_RENDER_VERSION.endswith("edt-listado-document-blocks-v4")


def test_edo_template_signature_block_is_two_compact_lines():
    workbook = openpyxl.load_workbook(REPORT_TEMPLATES_DIR / "001 - EDO.xlsx")
    worksheet = workbook.active

    assert worksheet["A19"].value == "FIRMA DEL SPONSOR O\nPERSONA QUE AUTORIZA"
    assert worksheet["A19"].font.sz <= 8
    assert worksheet["A19"].alignment.wrap_text is True
    assert "A19:E20" in [str(merged_range) for merged_range in worksheet.merged_cells.ranges]


def test_edo_template_header_uses_two_row_document_identity():
    workbook = openpyxl.load_workbook(REPORT_TEMPLATES_DIR / "001 - EDO.xlsx")
    worksheet = workbook.active
    merged_ranges = {str(merged_range) for merged_range in worksheet.merged_cells.ranges}

    assert worksheet["A4"].value == "Cod del Proyecto:"
    assert worksheet["C4"].value == "#CODIGOPROYECTO"
    assert worksheet["D4"].value == "Rev:"
    assert worksheet["E4"].value == "#REVISION"
    assert worksheet["F4"].value == "Cod Referencial:"
    assert worksheet["G4"].value == "#CODIGOREFERENCIAL"
    assert worksheet["A5"].value == "Nombre del Proyecto:"
    assert worksheet["D5"].value == "#PRO_TITULO"
    assert {"A4:B4", "G4:H4", "A5:C5", "D5:H5"}.issubset(merged_ranges)
    assert not {"A5:D8", "E6:F8", "G6:H8"}.intersection(merged_ranges)
    assert worksheet["A4"].alignment.wrap_text is not True
    assert worksheet["A5"].alignment.wrap_text is not True
    assert worksheet.row_dimensions[6].hidden is True
    assert worksheet.row_dimensions[7].hidden is True
    assert worksheet.row_dimensions[8].hidden is True
    assert worksheet.column_dimensions["D"].width == 12
    assert worksheet.column_dimensions["E"].width == 8
    assert worksheet.column_dimensions["G"].width == 34
    assert worksheet.column_dimensions["H"].width == 36


def test_edo_generated_workbook_keeps_project_identity_without_header_overflow(db, sample_empresa):
    project = _create_project(db, sample_empresa)
    db.add(
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.HITO,
            orden=1,
            codigo="1",
            nombre="Dirección del proyecto",
        )
    )
    db.commit()

    workbook = openpyxl.load_workbook(ReportingService().generate_edo_report(db, project.id, sample_empresa.id))
    worksheet = workbook.active
    merged_ranges = {str(merged_range) for merged_range in worksheet.merged_cells.ranges}

    assert worksheet["A4"].value == "Cod del Proyecto:"
    assert worksheet["C4"].value == "PRO-EDO-EDT"
    assert worksheet["D4"].value == "Rev:"
    assert worksheet["E4"].value == "R002"
    assert worksheet["F4"].value == "Cod Referencial:"
    assert worksheet["A5"].value == "Nombre del Proyecto:"
    assert worksheet["D5"].value == "Proyecto reportes edo edt"
    assert {"A4:B4", "G4:H4", "A5:C5", "D5:H5"}.issubset(merged_ranges)
    assert "A5:D8" not in merged_ranges
    assert worksheet.row_dimensions[6].hidden is True
    assert worksheet["C4"].alignment.shrink_to_fit is True
    assert worksheet["A5"].alignment.wrap_text is not True
    assert worksheet["D5"].alignment.wrap_text is not True


def test_edo_direct_pdf_identity_table_uses_requested_two_row_layout():
    styles = getSampleStyleSheet()
    label_style = ParagraphStyle("TestLabel", parent=styles["Normal"])
    body_style = ParagraphStyle("TestBody", parent=styles["Normal"])

    table = ReportingService()._build_hierarchy_document_identity_table(
        520,
        "Proyecto prueba compartir 1",
        "SantiagoBermeo-2026-001",
        "R000",
        "REF-001",
        label_style,
        body_style,
        lambda value: str(value),
    )

    assert table._cellvalues[0][0].getPlainText() == "Cod del Proyecto"
    assert table._cellvalues[0][1].getPlainText() == "SantiagoBermeo-2026-001"
    assert table._cellvalues[0][2].getPlainText() == "Rev"
    assert table._cellvalues[0][3].getPlainText() == "R000"
    assert table._cellvalues[0][4].getPlainText() == "Cod Referencial"
    assert table._cellvalues[0][5].getPlainText() == "REF-001"
    assert table._cellvalues[1][0].getPlainText() == "Nombre del Proyecto"
    assert table._cellvalues[1][1].getPlainText() == "Proyecto prueba compartir 1"
    assert ("SPAN", (1, 1), (-1, 1)) in table._spanCmds


def test_edo_generated_workbook_centers_item_and_role_columns(db, sample_empresa):
    project = _create_project(db, sample_empresa)
    stakeholder, role = _create_stakeholder_role(db, sample_empresa, project)
    db.add_all([
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.HITO,
            orden=1,
            codigo="1",
            nombre="Dirección del proyecto",
        ),
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.STAKEHOLDER,
            orden=2,
            codigo="1.R1",
            stakeholder_id=stakeholder.id,
            rol_id=role.id,
        ),
    ])
    db.commit()

    workbook = openpyxl.load_workbook(ReportingService().generate_edo_report(db, project.id, sample_empresa.id))
    worksheet = workbook.active

    assert worksheet.column_dimensions["A"].width == 5
    assert worksheet.column_dimensions["B"].width == 12.5
    assert worksheet.column_dimensions["D"].width == 12
    assert worksheet.column_dimensions["G"].width == 34
    assert worksheet.column_dimensions["H"].width == 36
    assert worksheet["A10"].alignment.horizontal == "center"
    assert worksheet["B10"].alignment.horizontal == "center"


def test_edo_generated_workbook_preserves_activities_merged_cells_on_inserted_rows(db, sample_empresa):
    project = _create_project(db, sample_empresa)
    stakeholder, role = _create_stakeholder_role(db, sample_empresa, project)
    db.add_all([
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.HITO,
            orden=1,
            codigo="1",
            nombre="Dirección del proyecto",
        ),
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.HITO,
            orden=2,
            codigo="2",
            nombre="Control de calidad",
        ),
        EdoNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodo.STAKEHOLDER,
            orden=3,
            codigo="1.R1",
            stakeholder_id=stakeholder.id,
            rol_id=role.id,
            actividades_claves=(
                "El Gerente del proyecto gestionará todas las actividades necesarias "
                "para que los diferentes ítems constructivos puedan realizarse."
            ),
        ),
    ])
    db.commit()

    workbook = openpyxl.load_workbook(ReportingService().generate_edo_report(db, project.id, sample_empresa.id))
    worksheet = workbook.active

    merged_ranges = {str(merged_range) for merged_range in worksheet.merged_cells.ranges}
    assert "E10:H10" in merged_ranges
    assert "E11:H11" in merged_ranges
    assert worksheet["E11"].alignment.wrap_text is True
    assert worksheet["E11"].alignment.horizontal == "center"


def test_edt_report_variants_match_template_placeholders(db, sample_empresa):
    project = _create_project(db, sample_empresa)
    stakeholder, role = _create_stakeholder_role(db, sample_empresa, project)
    db.add_all([
        EdtNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodoEdt.CUENTA_PAQUETE,
            orden=1,
            codigo="1",
            nombre="Diseño",
            definicion="Paquete de diseño integral",
        ),
        EdtNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodoEdt.STAKEHOLDER,
            orden=2,
            codigo="1.R1",
            stakeholder_id=stakeholder.id,
            rol_id=role.id,
            actividades_claves="Validar entregables",
        ),
    ])
    db.commit()

    service = ReportingService()
    for variant in ("listado", "diccionario", "valorada"):
        values = _sheet_values(service.generate_edt_report(db, project.id, sample_empresa.id, variant))
        assert "PRO-EDO-EDT" in values
        assert any("Diseño" in value for value in values)
        assert not any(value.startswith("#") for value in values)

    dictionary_values = _sheet_values(service.generate_edt_report(db, project.id, sample_empresa.id, "diccionario"))
    assert "Santiago Bermeo" in dictionary_values
    assert "STK-0001" in dictionary_values


def test_edt_preview_contract_matches_listado_template_semantics(db, sample_empresa):
    project = _create_project(db, sample_empresa)
    stakeholder, role = _create_stakeholder_role(db, sample_empresa, project)
    db.add_all([
        EdtNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodoEdt.CUENTA_PAQUETE,
            orden=1,
            codigo="1",
            nombre="Diseño",
            definicion="Paquete de diseño integral",
        ),
        EdtNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodoEdt.STAKEHOLDER,
            orden=2,
            codigo="1.R1",
            stakeholder_id=stakeholder.id,
            rol_id=role.id,
            actividades_claves="Validar entregables",
        ),
    ])
    db.commit()

    preview = ReportingService().preview_report(db, "edt", [project.id], sample_empresa.id, variant="listado")
    item = preview["items"][0]

    assert item["metadata_hint"] == "Documento EDT · Listado"
    assert item["preview_layout"] == "edt_document"
    assert item["summary_cards"] == [
        {"label": "Nodos", "value": 2, "kind": "integer"},
        {"label": "Cuentas", "value": 1, "kind": "integer"},
        {"label": "Responsables", "value": 1, "kind": "integer", "tone": "total"},
    ]
    assert [column["key"] for column in item["table_columns"]] == [
        "item",
        "codigo_edt",
        "descripcion_cuenta_paquete",
    ]
    assert {"cantidad", "precio", "rendimiento", "subtotal"}.isdisjoint(
        {column["key"] for column in item["table_columns"]}
    )
    assert item["lineas"][0]["item"] == 1
    assert item["lineas"][0]["codigo_edt"] == "1"
    assert item["lineas"][0]["descripcion_cuenta_paquete"] == "Diseño"
    assert item["lineas"][0]["nombre_responsable"] == "Santiago Bermeo"
    assert len(item["lineas"]) == 1


def test_edt_pdf_direct_generates_document_pdf(db, sample_empresa):
    project = _create_project(db, sample_empresa)
    stakeholder, role = _create_stakeholder_role(db, sample_empresa, project)
    db.add_all([
        EdtNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodoEdt.CUENTA_PAQUETE,
            orden=1,
            codigo="1",
            nombre="Diseño",
            definicion="Paquete de diseño integral",
        ),
        EdtNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodoEdt.STAKEHOLDER,
            orden=2,
            codigo="1.R1",
            stakeholder_id=stakeholder.id,
            rol_id=role.id,
            actividades_claves="Validar entregables",
        ),
    ])
    db.commit()

    payload = ReportingService().generate_preview_pdf(db, "edt", [project.id], sample_empresa.id, variant="listado").getvalue()

    assert payload.startswith(b"%PDF")
    assert len(payload) > 1500


def test_edt_preview_contract_matches_diccionario_template_semantics(db, sample_empresa):
    project = _create_project(db, sample_empresa)
    stakeholder, role = _create_stakeholder_role(db, sample_empresa, project)
    db.add_all([
        EdtNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodoEdt.CUENTA_PAQUETE,
            orden=1,
            codigo="1",
            nombre="Diseño",
            definicion="Paquete de diseño integral",
        ),
        EdtNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodoEdt.STAKEHOLDER,
            orden=2,
            codigo="1.R1",
            stakeholder_id=stakeholder.id,
            rol_id=role.id,
            actividades_claves="Validar entregables",
        )
    ])
    db.commit()

    preview = ReportingService().preview_report(db, "edt", [project.id], sample_empresa.id, variant="diccionario")
    item = preview["items"][0]

    assert item["metadata_hint"] == "Documento EDT · Diccionario"
    assert [column["key"] for column in item["table_columns"]] == [
        "codigo_edt",
        "descripcion_cuenta_paquete",
        "nombre_responsable",
        "codigo_stkr",
        "definicion",
    ]
    assert item["lineas"][0]["codigo_edt"] == "1"
    assert item["lineas"][0]["descripcion_cuenta_paquete"] == "Diseño"
    assert item["lineas"][0]["nombre_responsable"] == "Santiago Bermeo"
    assert item["lineas"][0]["codigo_stkr"] == "STK-0001"
    assert len(item["lineas"]) == 1
    assert item["lineas"][0]["codigo_stkr"] == "STK-0001"
    assert item["lineas"][0]["definicion"] == "Paquete de diseño integral"


def test_edt_generated_workbook_keeps_signature_in_two_line_wide_block(db, sample_empresa):
    project = _create_project(db, sample_empresa)
    nodes = [
        EdtNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodoEdt.CUENTA_PAQUETE,
            orden=index,
            codigo=str(index),
            nombre=f"Paquete {index}",
        )
        for index in range(1, 10)
    ]
    db.add_all(nodes)
    db.commit()

    for variant in ("listado", "diccionario"):
        workbook = openpyxl.load_workbook(ReportingService().generate_edt_report(db, project.id, sample_empresa.id, variant))
        worksheet = workbook.active
        signature_cell = next(
            cell
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and "FIRMA DEL" in cell.value
        )
        signature_range = f"A{signature_cell.row}:E{signature_cell.row + 1}"

        assert signature_cell.value == "FIRMA DEL SPONSOR O\nPERSONA QUE AUTORIZA"
        assert signature_cell.font.sz <= 8
        assert signature_cell.alignment.wrap_text is True
        assert signature_range in [str(merged_range) for merged_range in worksheet.merged_cells.ranges]


def test_edt_generated_workbook_compacts_item_and_code_columns(db, sample_empresa):
    project = _create_project(db, sample_empresa)
    stakeholder, role = _create_stakeholder_role(db, sample_empresa, project)
    db.add(
        EdtNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodoEdt.CUENTA_PAQUETE,
            orden=1,
            codigo="1.1",
            nombre="Paquete documental",
            definicion="Alcance documental de prueba",
        )
    )
    db.flush()
    db.add(
        EdtNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodoEdt.STAKEHOLDER,
            orden=2,
            codigo="1.1.R1",
            stakeholder_id=stakeholder.id,
            rol_id=role.id,
        )
    )
    db.commit()

    workbook = openpyxl.load_workbook(ReportingService().generate_edt_report(db, project.id, sample_empresa.id, "listado"))
    worksheet = workbook.active
    merged_ranges = {str(merged_range) for merged_range in worksheet.merged_cells.ranges}
    _assert_no_overlapping_merged_ranges(worksheet)

    assert worksheet.column_dimensions["A"].width == 5
    assert worksheet.column_dimensions["B"].width == 8
    assert worksheet.column_dimensions["C"].width >= 56
    assert "A1:H3" in merged_ranges
    assert "A9:H9" in merged_ranges
    assert "A11:H11" in merged_ranges
    assert "A12:H12" in merged_ranges
    assert "A13:H13" in merged_ranges
    assert worksheet["A9"].value == "Detalle documental"
    assert worksheet["A9"].fill.fgColor.rgb == "00E9EEF3"
    assert worksheet["A10"].value is None
    assert worksheet["A10"].fill.fill_type is None
    assert worksheet["A11"].value == "Código EDT: 1.1 · Paquete documental"
    assert worksheet["A11"].fill.fgColor.rgb == "00EEF6FB"
    assert worksheet["A12"].value == "Responsable: Santiago Bermeo"
    assert worksheet["A13"].value == "Definición: Alcance documental de prueba"


def test_edt_templates_use_compact_two_row_document_identity():
    expected = [
        REPORT_TEMPLATES_DIR / "001 - EDT - listado.xlsx",
        REPORT_TEMPLATES_DIR / "001 - EDT - Diccionario.xlsx",
        REPORT_TEMPLATES_DIR / "001 - EDT - Valorada.xlsx",
    ]

    for path in expected:
        workbook = openpyxl.load_workbook(path)
        worksheet = workbook.active
        merged_ranges = {str(merged_range) for merged_range in worksheet.merged_cells.ranges}

        assert worksheet["A4"].value == "Cod del Proyecto:"
        assert worksheet["C4"].value == "#CODIGOPROYECTO"
        assert worksheet["D4"].value == "Rev:"
        assert worksheet["E4"].value == "#REVISION"
        assert worksheet["F4"].value == "Cod Referencial:"
        assert worksheet["G4"].value == "#CODIGOREFERENCIAL"
        assert worksheet["A5"].value == "Nombre del Proyecto:"
        assert worksheet["D5"].value == "#PRO_TITULO"
        assert {"A4:B4", "G4:H4", "A5:C5", "D5:H5"}.issubset(merged_ranges)
        assert "A1:H3" in merged_ranges
        assert not {"A5:D8", "E6:F8", "G6:H8"}.intersection(merged_ranges)
        assert worksheet["A4"].alignment.wrap_text is not True
        assert worksheet["A5"].alignment.wrap_text is not True
        assert worksheet.row_dimensions[6].hidden is True
        assert worksheet.row_dimensions[7].hidden is True
        assert worksheet.row_dimensions[8].hidden is True
        assert worksheet.column_dimensions["D"].width == 12
        assert worksheet.column_dimensions["E"].width == 8
        assert worksheet.column_dimensions["G"].width == 34
        assert worksheet.column_dimensions["H"].width == 36


def test_edt_generated_workbook_uses_compact_identity_header(db, sample_empresa):
    project = _create_project(db, sample_empresa)
    db.add(
        EdtNode(
            proyecto_id=project.id,
            empresa_id=sample_empresa.id,
            parent_id=None,
            tipo_nodo=TipoNodoEdt.CUENTA_PAQUETE,
            orden=1,
            codigo="1.1",
            nombre="Paquete documental",
        )
    )
    db.commit()

    for variant in ("listado", "diccionario", "valorada"):
        workbook = openpyxl.load_workbook(ReportingService().generate_edt_report(db, project.id, sample_empresa.id, variant))
        worksheet = workbook.active
        merged_ranges = {str(merged_range) for merged_range in worksheet.merged_cells.ranges}
        _assert_no_overlapping_merged_ranges(worksheet)

        assert worksheet["A4"].value == "Cod del Proyecto:"
        assert worksheet["C4"].value == "PRO-EDO-EDT"
        assert worksheet["D4"].value == "Rev:"
        assert worksheet["E4"].value == "R002"
        assert worksheet["F4"].value == "Cod Referencial:"
        assert worksheet["A5"].value == "Nombre del Proyecto:"
        assert worksheet["D5"].value == "Proyecto reportes edo edt"
        assert {"A4:B4", "G4:H4", "A5:C5", "D5:H5"}.issubset(merged_ranges)
        assert "A1:H3" in merged_ranges
        if variant == "listado":
            assert "A9:H9" in merged_ranges
            assert "A11:H11" in merged_ranges
            assert "A12:H12" in merged_ranges
            assert "A13:H13" in merged_ranges
            assert worksheet["A9"].value == "Detalle documental"
            assert worksheet["A10"].value is None
            assert worksheet["A11"].value.startswith("Código EDT:")
        assert worksheet.row_dimensions[6].hidden is True
        assert worksheet.column_dimensions["D"].width == 12
        assert worksheet.column_dimensions["E"].width == 8
        assert worksheet.column_dimensions["G"].width == 34
        assert worksheet.column_dimensions["H"].width == 36


def test_edt_template_signature_blocks_are_two_compact_lines():
    expected = [
        (REPORT_TEMPLATES_DIR / "001 - EDT - listado.xlsx", "A19:E20"),
        (REPORT_TEMPLATES_DIR / "001 - EDT - Diccionario.xlsx", "A23:E24"),
        (REPORT_TEMPLATES_DIR / "001 - EDT - Valorada.xlsx", "A26:E27"),
    ]
    for path, signature_range in expected:
        workbook = openpyxl.load_workbook(path)
        worksheet = workbook.active
        start_cell = signature_range.split(":")[0]

        assert worksheet[start_cell].value == "FIRMA DEL SPONSOR O\nPERSONA QUE AUTORIZA"
        assert worksheet[start_cell].font.sz <= 8
        assert worksheet[start_cell].alignment.wrap_text is True
        assert signature_range in [str(merged_range) for merged_range in worksheet.merged_cells.ranges]
