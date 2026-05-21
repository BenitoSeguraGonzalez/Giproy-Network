import os
import io

import openpyxl
import pytest

from app.services.reporting import ReportingService

try:
    from PIL import Image as PILImage
except Exception:  # pragma: no cover
    PILImage = None


def test_stakeholders_sheet_is_prepared_for_pdf_from_excel():
    service = ReportingService()
    workbook = openpyxl.load_workbook(os.path.join(service.templates_dir, "001 - Stakeholders.xlsx"))
    worksheet = workbook.active

    table_rows = [
        {
            "#ITEM": 1,
            "#CODIGO_STKR": "STK-0001",
            "#NOMBRES": "Nombre con datos completos",
            "#APELLIDOS": "Apellido de prueba",
            "#PROFESION": "Ingeniero civil · Institucion publica con nombre largo",
            "#DIRECCIONSTKR": "Avenida de prueba muy larga con edificio, piso y referencia adicional",
            "#CIUDADSTKR": "Quito / Distrito Metropolitano",
            "#PROVINCIASTKR": "Pichincha",
            "#PAISSTKR": "Ecuador",
            "#TELEFONOSTKR": "+593 999 999 999",
            "#CORREOSTKR": "correo.largo.de.prueba@example.com",
            "#ROL": "Fiscalizador principal del proyecto",
        },
        {
            "#ITEM": 2,
            "#CODIGO_STKR": "STK-0002",
            "#NOMBRES": "Segundo nombre",
            "#APELLIDOS": "Segundo apellido",
            "#PROFESION": "Arquitecto",
            "#DIRECCIONSTKR": "Calle corta",
            "#CIUDADSTKR": "Cuenca",
            "#PROVINCIASTKR": "Azuay",
            "#PAISSTKR": "Ecuador",
            "#TELEFONOSTKR": "+593 988 888 888",
            "#CORREOSTKR": "contacto@example.com",
            "#ROL": "Sin rol asignado",
        },
    ]

    service._apply_template_to_sheet(
        worksheet,
        {
            "#REVISION": "R001",
            "#PRO_TITULO": "Proyecto de prueba",
            "#CODIGOPROYECTO": "PR-001",
            "#CODIGOREFERENCIAL": "REF-001",
            "#CIUDAD": "Quito",
            "#PRO_FECHA": "07/04/2026",
        },
        table_rows,
        "#ITEM",
    )
    service._prepare_stakeholders_report_sheet(worksheet, 10, len(table_rows))
    output = service._save_workbook_buffer(workbook)
    rendered_workbook = openpyxl.load_workbook(output)
    rendered_sheet = rendered_workbook.active

    assert rendered_sheet.print_area.startswith("'Hoja1'!$A$1:$L$")
    assert rendered_sheet.max_row <= 22
    assert rendered_sheet.page_setup.orientation == "landscape"
    assert rendered_sheet.page_setup.fitToWidth == 1
    assert rendered_sheet.page_setup.fitToHeight == 0
    assert rendered_sheet.row_dimensions[10].height > 25
    assert rendered_sheet.column_dimensions["E"].width >= 22
    assert rendered_sheet.column_dimensions["F"].width >= 24
    assert rendered_sheet.column_dimensions["K"].width >= 28
    assert rendered_sheet["E9"].alignment.wrap_text is True
    assert rendered_sheet["K9"].alignment.wrap_text is True
    assert rendered_sheet["E10"].value == "Ingeniero civil · Institucion publica con nombre largo"
    assert rendered_sheet["K10"].value == "correo.largo.de.prueba@example.com"
    assert rendered_sheet["L11"].value == "Sin rol asignado"


def test_stakeholders_sheet_keeps_signature_in_two_line_wide_block():
    service = ReportingService()
    workbook = openpyxl.load_workbook(os.path.join(service.templates_dir, "001 - Stakeholders.xlsx"))
    worksheet = workbook.active

    service._prepare_stakeholders_report_sheet(worksheet, 10, 1)

    signature_cell = next(
        cell
        for row in worksheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and "FIRMA DEL" in cell.value
    )
    signature_range = f"A{signature_cell.row}:E{signature_cell.row + 1}"

    assert signature_cell.value == "FIRMA DEL SPONSOR O\nPERSONA QUE AUTORIZA"
    assert signature_cell.font.sz <= 8
    assert signature_cell.alignment.wrap_text is True
    assert signature_range in [str(merged_range) for merged_range in worksheet.merged_cells.ranges]


def test_direct_pdf_report_uses_portrait_a4_layout():
    service = ReportingService()

    output = service._build_pdf_report(
        "Reporte directo",
        [
            {
                "codigo": "PRO-001",
                "descripcion": "Documento formal de prueba",
                "categoria_base": "Proyecto",
                "unidad": "u",
                "costo_directo": 100,
                "costo_indirecto": 10,
                "precio_total": 110,
                "lineas": [
                    {
                        "codigo": "1",
                        "descripcion": "Rubro con descripcion suficiente para validar anchos verticales",
                        "unidad": "u",
                        "cantidad": 1,
                        "precio": 100,
                        "rendimiento": 1,
                        "subtotal": 100,
                    }
                ],
            }
        ],
    )
    pdf_bytes = output.getvalue()

    assert b"/MediaBox [ 0 0 595.2756 841.8898 ]" in pdf_bytes
    assert b"/MediaBox [ 0 0 841.8898 595.2756 ]" not in pdf_bytes


def test_image_placeholder_is_fitted_to_target_canvas_without_distortion():
    if PILImage is None:
        pytest.skip("Pillow no disponible")

    service = ReportingService()
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.merge_cells("A1:F12")
    worksheet["A1"] = "#IMAGENPROYECTO"
    for column_index in range(1, 7):
        worksheet.column_dimensions[openpyxl.utils.get_column_letter(column_index)].width = 13
    for row_index in range(1, 13):
        worksheet.row_dimensions[row_index].height = 15

    source = PILImage.new("RGB", (1600, 360), (243, 146, 0))
    source_buffer = io.BytesIO()
    source.save(source_buffer, format="PNG")

    expected_width, expected_height = service._estimate_range_pixels(
        worksheet,
        service._resolve_placeholder_range(worksheet, worksheet["A1"]),
    )
    service._insert_image_placeholder(worksheet, "#IMAGENPROYECTO", source_buffer.getvalue())

    assert worksheet["A1"].value == ""
    assert len(worksheet._images) == 1
    image = worksheet._images[0]
    assert image.width == expected_width
    assert image.height == expected_height

    fitted = PILImage.open(io.BytesIO(image._data()))
    assert fitted.size == (expected_width, expected_height)
    assert fitted.getpixel((0, 0))[:3] == (243, 146, 0)
    assert fitted.getpixel((expected_width - 1, expected_height - 1))[:3] == (243, 146, 0)


def test_excel_pdf_keeps_contractual_section_together():
    service = ReportingService()
    workbook = openpyxl.load_workbook(os.path.join(service.templates_dir, "001 - Acta de Constitucion - General.xlsx"))
    worksheet = workbook.active

    min_col, min_row, max_col, max_row = service._worksheet_print_bounds(worksheet)
    column_indexes = [
        index
        for index in range(min_col, max_col + 1)
        if service._worksheet_column_width_points(worksheet, index) > 0
    ]
    row_indexes = [
        index
        for index in range(min_row, max_row + 1)
        if service._worksheet_row_height_points(worksheet, index) > 0
    ]
    row_heights = {index: service._worksheet_row_height_points(worksheet, index) for index in row_indexes}
    section_groups = service._worksheet_pdf_keep_groups(
        worksheet,
        row_indexes,
        column_indexes,
        row_heights,
        10_000,
    )

    assert 31 in section_groups
    assert section_groups[31][0] == 37


def test_excel_pdf_acta_pagination_starts_contractual_section_on_new_page():
    service = ReportingService()
    workbook = openpyxl.load_workbook(os.path.join(service.templates_dir, "001 - Acta de Constitucion - General.xlsx"))
    worksheet = workbook.active

    min_col, min_row, max_col, max_row = service._worksheet_print_bounds(worksheet)
    column_indexes = [
        index
        for index in range(min_col, max_col + 1)
        if service._worksheet_column_width_points(worksheet, index) > 0
    ]
    row_indexes = [
        index
        for index in range(min_row, max_row + 1)
        if service._worksheet_row_height_points(worksheet, index) > 0
    ]
    column_widths = {index: service._worksheet_column_width_points(worksheet, index) for index in column_indexes}
    row_heights = {index: service._worksheet_row_height_points(worksheet, index) for index in row_indexes}
    total_width = sum(column_widths.values())
    page_width, page_height = (841.8897637795277, 595.2755905511812)
    margin = 10 * 72 / 25.4
    scale = min((page_width - (margin * 2)) / total_width, 1.0)
    usable_height = (page_height - (margin * 2)) / max(scale, 0.01)
    section_groups = service._worksheet_pdf_keep_groups(
        worksheet,
        row_indexes,
        column_indexes,
        row_heights,
        usable_height,
    )

    pages = []
    row_pointer = 0
    while row_pointer < len(row_indexes):
        page_start_pointer = row_pointer
        consumed_height = 0.0
        while row_pointer < len(row_indexes):
            candidate_row = row_indexes[row_pointer]
            candidate_height = row_heights[candidate_row]
            keep_group = section_groups.get(candidate_row)
            if keep_group is not None and row_pointer > page_start_pointer and consumed_height + keep_group[1] > usable_height:
                break
            if row_pointer > page_start_pointer and consumed_height + candidate_height > usable_height:
                break
            consumed_height += candidate_height
            row_pointer += 1
        if page_start_pointer == row_pointer:
            row_pointer = page_start_pointer + 1
        pages.append((row_indexes[page_start_pointer], row_indexes[row_pointer - 1]))

    assert pages[0] == (1, 30)
    assert pages[1][0] == 31


def test_excel_pdf_compacts_residual_tail_universally():
    service = ReportingService()
    workbook = openpyxl.load_workbook(os.path.join(service.templates_dir, "001 - Acta de Constitucion - General.xlsx"))
    worksheet = workbook.active

    min_col, min_row, max_col, max_row = service._worksheet_print_bounds(worksheet)
    column_indexes = [
        index
        for index in range(min_col, max_col + 1)
        if service._worksheet_column_width_points(worksheet, index) > 0
    ]
    row_indexes = [
        index
        for index in range(min_row, max_row + 1)
        if service._worksheet_row_height_points(worksheet, index) > 0
    ]
    column_widths = {index: service._worksheet_column_width_points(worksheet, index) for index in column_indexes}
    row_heights = {index: service._worksheet_row_height_points(worksheet, index) for index in row_indexes}
    total_width = sum(column_widths.values())
    page_width, page_height = (841.8897637795277, 595.2755905511812)
    margin = 10 * 72 / 25.4
    available_width = page_width - (margin * 2)
    available_height = page_height - (margin * 2)
    scale = min(available_width / total_width, 1.0)
    usable_height = available_height / max(scale, 0.01)
    section_groups = service._worksheet_pdf_keep_groups(
        worksheet,
        row_indexes,
        column_indexes,
        row_heights,
        usable_height,
    )
    initial_pages = service._paginate_worksheet_pdf_rows(
        row_indexes,
        row_heights,
        section_groups,
        usable_height,
    )

    compact_scale, compact_usable_height, compact_groups, compact_pages = service._compact_worksheet_pdf_residual_tail(
        worksheet,
        row_indexes,
        column_indexes,
        row_heights,
        available_height,
        scale,
        usable_height,
        section_groups,
        initial_pages,
    )

    assert len(initial_pages) == 3
    assert initial_pages[-1][2] <= usable_height * 0.18
    assert len(compact_pages) == 2
    assert compact_pages[0] == (1, 30, initial_pages[0][2])
    assert compact_pages[1][0] == 31
    assert compact_scale < scale
    assert compact_usable_height > usable_height
    assert 31 in compact_groups
