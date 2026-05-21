from pathlib import Path
import openpyxl


ROOT = Path(__file__).resolve().parents[2]
ADICIONALES_DIR = ROOT / "docs" / "adicionales"
REPORTES_DIR = ROOT / "docs" / "reportes"


def _keep_only_sheet(workbook, sheet_name: str):
    for current_name in list(workbook.sheetnames):
        if current_name != sheet_name:
            workbook.remove(workbook[current_name])
    workbook.active = 0
    return workbook[sheet_name]


def _apply_replacements(ws, replacements):
    for cell_ref, value in replacements.items():
        ws[cell_ref] = value


def rebuild_general_template():
    source_path = ADICIONALES_DIR / "Presupuesto ejemplo - General.xlsx"
    target_path = REPORTES_DIR / "001 - Analisis - APUS - General.xlsx"

    wb = openpyxl.load_workbook(source_path)
    ws = _keep_only_sheet(wb, "501139")
    ws.title = "Hoja1"

    replacements = {
        "B3": "#CODIGO_APU",
        "B4": "#DESCRIPCION",
        "B5": "#UNIDAD",
        "A11": "#CODCAT1",
        "B11": "#DESCRIPCION1",
        "C11": "#UNIDAD1",
        "D11": "#CANTIDAD1",
        "E11": "#PRECIO1",
        "F11": "#RENDIMIENTO1",
        "G11": "#SUBTOTAL1",
        "H11": "#PORCENT1",
        "G13": "#TOTAL1",
        "H13": "#TOTPORCENT1",
        "A17": "#CODCAT2",
        "B17": "#DESCRIPCION2",
        "C17": "#UNIDAD2",
        "D17": "#CANTIDAD2",
        "E17": "#PRECIO2",
        "G17": "#SUBTOTAL2",
        "H17": "#PORCENT2",
        "G20": "#TOTAL2",
        "H20": "#TOTPORCENT2",
        "A24": "#CODCAT3",
        "B24": "#DESCRIPCION3",
        "C24": "#UNIDAD3",
        "D24": "#CANTIDAD3",
        "E24": "#PRECIO3",
        "F24": "#RENDIMIENTO3",
        "G24": "#SUBTOTAL3",
        "H24": "#PORCENT3",
        "G25": "#TOTAL3",
        "H25": "#TOTPORCENT3",
        "A29": "#CODCAT4",
        "B29": "#DESCRIPCION4",
        "D29": "#CANTIDAD4",
        "E29": "#PRECIO4",
        "F29": "#RENDIMIENTO4",
        "G29": "#SUBTOTAL4",
        "H29": "#PORCENT4",
        "G32": "#TOTAL4",
        "H32": "#TOTPORCENT4",
        "H34": "#COSTODIRECTO",
        "A37": "#%INDIRECTO %",
        "H37": "#COSTOINDIRECTO",
        "H39": "#TTOTAL",
        "B41": "#TEXTOTOTAL",
    }
    _apply_replacements(ws, replacements)
    wb.save(target_path)
    print(f"OK: rebuilt {target_path.name}")


def rebuild_sercop_template():
    source_path = ADICIONALES_DIR / "Presupuesto ejemplo - SERCOP.xlsx"
    target_path = REPORTES_DIR / "002 - Analisis - APUS - SERCOP.xlsx"

    wb = openpyxl.load_workbook(source_path)
    ws = _keep_only_sheet(wb, "501139")
    ws.title = "Hoja1"

    replacements = {
        "A4": "#DESCRIPCION",
        "F4": "#UNIDAD",
        "A9": "#DESCRIPCION1",
        "B9": "#CANTIDAD1",
        "C9": "#PRECIO1",
        "E9": "#RENDIMIENTO1",
        "F9": "#SUBTOTAL1",
        "F11": "#TOTAL1",
        "A14": "#DESCRIPCION4",
        "B14": "#CANTIDAD4",
        "C14": "#PRECIO4",
        "E14": "#RENDIMIENTO4",
        "F14": "#SUBTOTAL4",
        "F17": "#TOTAL4",
        "A20": "#DESCRIPCION2",
        "C20": "#UNIDAD2",
        "D20": "#CANTIDAD2",
        "E20": "#PRECIO2",
        "F20": "#SUBTOTAL2",
        "F23": "#TOTAL2",
        "A26": "#DESCRIPCION3",
        "C26": "#UNIDAD3",
        "D26": "#CANTIDAD3",
        "E26": "#PRECIO3",
        "F26": "#SUBTOTAL3",
        "F27": "#TOTAL3",
        "F28": "#COSTODIRECTO",
        "E29": "#%INDIRECTO %",
        "F29": "#COSTOINDIRECTO",
        "F31": "#TTOTAL",
        "F32": "#TTOTAL",
    }
    _apply_replacements(ws, replacements)
    wb.save(target_path)
    print(f"OK: rebuilt {target_path.name}")


def main():
    rebuild_general_template()
    rebuild_sercop_template()


if __name__ == "__main__":
    main()
