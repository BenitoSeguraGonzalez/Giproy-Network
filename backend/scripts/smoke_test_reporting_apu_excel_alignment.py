import sys
from pathlib import Path

import openpyxl
from openpyxl.styles.borders import Border


ROOT = Path(__file__).resolve().parents[2]
BACKEND_DIR = ROOT / "backend"
EXAMPLES_DIR = ROOT / "docs" / "adicionales"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.core.database import SessionLocal
from app.models.apu import APU
from app.models.empresa import Empresa
from app.models.presupuesto import Presupuesto
from app.services.reporting import reporting_service


TARGET_EMPRESA = "Santiago Bermeo"


def _assert_equal(label, actual, expected):
    if actual != expected:
        raise AssertionError(f"{label}: esperado={expected!r} actual={actual!r}")


def _assert_contains(label, actual, expected_fragment):
    if expected_fragment not in (actual or ""):
        raise AssertionError(f"{label}: no contiene {expected_fragment!r} actual={actual!r}")


def _border_signature(border: Border):
    if border is None:
        return (None, None, None, None)
    return (
        getattr(border.left, "style", None),
        getattr(border.right, "style", None),
        getattr(border.top, "style", None),
        getattr(border.bottom, "style", None),
    )


def _normalize_cell_value(value):
    if isinstance(value, str):
        stripped = value.strip()
        try:
            return round(float(stripped), 6)
        except ValueError:
            return stripped
    if isinstance(value, (int, float)):
        return round(float(value), 6)
    return value


def _assert_row_matches_example(label, actual_ws, example_ws, row_idx, columns=8):
    actual_values = [_normalize_cell_value(actual_ws.cell(row_idx, col).value) for col in range(1, columns + 1)]
    example_values = [_normalize_cell_value(example_ws.cell(row_idx, col).value) for col in range(1, columns + 1)]
    if actual_values != example_values:
        raise AssertionError(
            f"{label} fila {row_idx}: valores distintos esperado={example_values!r} actual={actual_values!r}"
        )
    for col in range(1, columns + 1):
        actual_sig = _border_signature(actual_ws.cell(row_idx, col).border)
        example_sig = _border_signature(example_ws.cell(row_idx, col).border)
        if actual_sig != example_sig:
            raise AssertionError(
                f"{label} fila {row_idx} col {col}: borde distinto esperado={example_sig!r} actual={actual_sig!r}"
            )


def _assert_merge_matches_example(label, actual_ws, example_ws, row_min, row_max):
    actual = {str(rng) for rng in actual_ws.merged_cells.ranges if row_min <= rng.min_row and rng.max_row <= row_max}
    expected = {str(rng) for rng in example_ws.merged_cells.ranges if row_min <= rng.min_row and rng.max_row <= row_max}
    if actual != expected:
        raise AssertionError(f"{label}: merges distintos esperado={sorted(expected)!r} actual={sorted(actual)!r}")


def _find_rows_by_value(ws, value):
    rows = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == value:
                rows.append(cell.row)
                break
    return rows


def _assert_single_label_order(ws, labels):
    positions = []
    for label in labels:
        rows = _find_rows_by_value(ws, label)
        if len(rows) != 1:
            raise AssertionError(f"{label}: esperado una ocurrencia, actual={rows!r}")
        positions.append(rows[0])
    if positions != sorted(positions):
        raise AssertionError(f"Orden de bloques invalido: {list(zip(labels, positions))!r}")


def _find_single_row_by_value(ws, value):
    rows = _find_rows_by_value(ws, value)
    if len(rows) != 1:
        raise AssertionError(f"{value}: esperado una ocurrencia, actual={rows!r}")
    return rows[0]


def _assert_numeric_cell(label, actual, expected):
    if round(float(actual or 0), 2) != round(float(expected or 0), 2):
        raise AssertionError(f"{label}: esperado={expected!r} actual={actual!r}")


def _assert_light_section_fill(label, cell):
    fill = getattr(cell, "fill", None)
    color = getattr(getattr(fill, "fgColor", None), "rgb", None) or ""
    if not color.endswith("EAF3F8"):
        raise AssertionError(f"{label}: color de fondo esperado EAF3F8 actual={color!r}")


def main() -> int:
    db = SessionLocal()
    try:
        empresa = db.query(Empresa).filter(Empresa.nombre == TARGET_EMPRESA).first()
        if not empresa:
            raise RuntimeError(f"Empresa no encontrada: {TARGET_EMPRESA}")

        presupuesto = (
            db.query(Presupuesto)
            .filter(Presupuesto.empresa_id == empresa.id)
            .order_by(Presupuesto.id.asc())
            .first()
        )
        if not presupuesto:
            raise RuntimeError("No hay presupuesto para la empresa de prueba")

        general_bundle = reporting_service.generate_presupuesto_report_bundle(db, presupuesto.id, empresa.id, "001")
        general_wb = openpyxl.load_workbook(general_bundle)
        general_ws = general_wb[general_wb.sheetnames[1]]
        general_example_wb = openpyxl.load_workbook(EXAMPLES_DIR / "Presupuesto ejemplo - General.xlsx")
        general_example_ws = general_example_wb[general_example_wb.sheetnames[1]]

        apu = (
            db.query(APU)
            .filter(APU.empresa_id == empresa.id, APU.codigo == "5-001-0001")
            .first()
        )
        if not apu:
            raise RuntimeError("APU base de smoke test no encontrado: 5-001-0001")

        _assert_equal("General B3", general_ws["B3"].value, "5-001-0001")
        _assert_equal("General B4", general_ws["B4"].value, "Replanteo y nivelación")
        _assert_equal("General B5", general_ws["B5"].value, "m2")
        _assert_equal("General A11 uses internal code", general_ws["A11"].value, "1-0001-00003")
        _assert_equal("General A12 uses internal code", general_ws["A12"].value, "1-0001-00001")
        _assert_equal("General A23", general_ws["A23"].value, "Código")
        _assert_equal("General E23", general_ws["E23"].value, "Tarifa/U")
        _assert_equal("General F23", general_ws["F23"].value, "Distancia")
        _assert_equal("General A28", general_ws["A28"].value, "Código")
        _assert_equal("General D28", general_ws["D28"].value, "Número")
        _assert_contains("General B41", str(general_ws["B41"].value), "DÓLARES")
        _assert_equal("General sheet title", general_ws.title, "5-001-0001")
        _assert_equal("General transport header B23", general_ws["B23"].value, "Descripción")
        _assert_equal("General transport header E23", general_ws["E23"].value, "Tarifa/U")
        _assert_equal("General indirect percent label", general_ws["A37"].value, "21 %")
        _assert_equal("General mano header B28", general_ws["B28"].value, "Descripción")
        _assert_equal("General mano item B29", general_ws["B29"].value, "Peón")
        _assert_equal("General mano item B30", general_ws["B30"].value, "Topógrafo 2: título exper. mayor a 5 años (estr. oc. c1)")
        _assert_equal("General merge B29:C29", "B29:C29" in {str(rng) for rng in general_ws.merged_cells.ranges}, True)
        _assert_equal("General merge B30:C30", "B30:C30" in {str(rng) for rng in general_ws.merged_cells.ranges}, True)
        _assert_equal("General merge B31:C31", "B31:C31" in {str(rng) for rng in general_ws.merged_cells.ranges}, True)
        for row_idx in (21, 22, 23, 24, 25, 26, 27, 28, 32, 33, 34, 35, 36, 37, 38, 39, 40, 41):
            _assert_row_matches_example("General estructura", general_ws, general_example_ws, row_idx, columns=8)
        _assert_merge_matches_example("General merges estructura", general_ws, general_example_ws, 20, 41)

        general_apu = reporting_service.generate_apu_report(db, apu.id, empresa.id, "001")
        general_apu_wb = openpyxl.load_workbook(general_apu)
        _assert_equal("Standalone general sheet title", general_apu_wb.active.title, "5-001-0001")
        _assert_equal("Standalone general B3 uses internal code", general_apu_wb.active["B3"].value, "5-001-0001")
        _assert_equal("Standalone general A23", general_apu_wb.active["A23"].value, "Código")
        _assert_equal("Standalone general B23", general_apu_wb.active["B23"].value, "Descripción")
        _assert_equal("Standalone general B29", general_apu_wb.active["B29"].value, "Peón")
        _assert_equal("Standalone merge B30:C30", "B30:C30" in {str(rng) for rng in general_apu_wb.active.merged_cells.ranges}, True)
        _assert_equal("Standalone merge B31:C31", "B31:C31" in {str(rng) for rng in general_apu_wb.active.merged_cells.ranges}, True)

        problematic_apu = (
            db.query(APU)
            .filter(APU.empresa_id == empresa.id, APU.codigo == "5-017-0002")
            .first()
        )
        if not problematic_apu:
            raise RuntimeError("APU de validación no encontrado: 5-017-0002")

        problematic_apu_report = reporting_service.generate_apu_report(db, problematic_apu.id, empresa.id, "001")
        problematic_apu_wb = openpyxl.load_workbook(problematic_apu_report)
        problematic_ws = problematic_apu_wb.active
        for row_idx in (21, 22, 23, 24, 25, 26, 27, 28):
            _assert_row_matches_example("Standalone general estructura problemática", problematic_ws, general_example_ws, row_idx, columns=8)
        _assert_merge_matches_example("Standalone general merges problemáticos", problematic_ws, general_example_ws, 20, 28)

        sercop_bundle = reporting_service.generate_presupuesto_report_bundle(db, presupuesto.id, empresa.id, "002")
        sercop_wb = openpyxl.load_workbook(sercop_bundle)
        sercop_ws = sercop_wb[sercop_wb.sheetnames[1]]
        sercop_example_wb = openpyxl.load_workbook(EXAMPLES_DIR / "Presupuesto ejemplo - SERCOP.xlsx")
        sercop_example_ws = sercop_example_wb[sercop_example_wb.sheetnames[1]]

        _assert_equal("SERCOP A4", sercop_ws["A4"].value, "Replanteo y nivelación")
        _assert_equal("SERCOP F4", sercop_ws["F4"].value, "m2")
        _assert_equal("SERCOP A8", sercop_ws["A8"].value, "Descripción")
        _assert_equal("SERCOP A13", sercop_ws["A13"].value, "Descripción")
        _assert_equal("SERCOP A19", sercop_ws["A19"].value, "Descripción")
        _assert_equal("SERCOP A25", sercop_ws["A25"].value, "Descripción")
        _assert_equal("SERCOP A26", sercop_ws["A26"].value, None)
        _assert_equal("SERCOP F31", float(sercop_ws["F31"].value), 1.04)
        _assert_equal("SERCOP sheet title", sercop_ws.title, "5-001-0001")
        _assert_single_label_order(sercop_ws, [
            "EQUIPOS",
            "SUBTOTAL M",
            "MANO DE OBRA",
            "SUBTOTAL N",
            "MATERIALES",
            "SUBTOTAL O",
            "TRANSPORTE",
            "SUBTOTAL P",
            "TOTAL COSTO DIRECTO (M+N+O+P)",
        ])
        expected_sercop_subtotals = {
            "SUBTOTAL M": 0.18,
            "SUBTOTAL N": 0.30,
            "SUBTOTAL O": 0.38,
            "SUBTOTAL P": 0.00,
        }
        for label, expected_value in expected_sercop_subtotals.items():
            row_idx = _find_single_row_by_value(sercop_ws, label)
            _assert_numeric_cell(f"SERCOP {label}", sercop_ws.cell(row_idx, 6).value, expected_value)
        for label in ("EQUIPOS", "MANO DE OBRA", "MATERIALES", "TRANSPORTE"):
            row_idx = _find_single_row_by_value(sercop_ws, label)
            _assert_light_section_fill(f"SERCOP encabezado {label}", sercop_ws.cell(row_idx, 1))
        for row_idx in range(30, 38):
            _assert_row_matches_example("SERCOP estructura", sercop_ws, sercop_example_ws, row_idx, columns=6)
        _assert_merge_matches_example("SERCOP merges estructura", sercop_ws, sercop_example_ws, 30, 38)

        sercop_apu = reporting_service.generate_apu_report(db, apu.id, empresa.id, "002")
        sercop_apu_wb = openpyxl.load_workbook(sercop_apu)
        sercop_apu_ws = sercop_apu_wb.active
        _assert_equal("Standalone SERCOP sheet title", sercop_apu_ws.title, "5-001-0001")
        _assert_equal("Standalone SERCOP A8", sercop_apu_ws["A8"].value, "Descripción")
        _assert_equal("Standalone SERCOP A13", sercop_apu_ws["A13"].value, "Descripción")
        _assert_equal("Standalone SERCOP A19", sercop_apu_ws["A19"].value, "Descripción")
        _assert_equal("Standalone SERCOP A25", sercop_apu_ws["A25"].value, "Descripción")
        _assert_single_label_order(sercop_apu_ws, [
            "EQUIPOS",
            "SUBTOTAL M",
            "MANO DE OBRA",
            "SUBTOTAL N",
            "MATERIALES",
            "SUBTOTAL O",
            "TRANSPORTE",
            "SUBTOTAL P",
            "TOTAL COSTO DIRECTO (M+N+O+P)",
        ])
        for label, expected_value in expected_sercop_subtotals.items():
            row_idx = _find_single_row_by_value(sercop_apu_ws, label)
            _assert_numeric_cell(f"Standalone SERCOP {label}", sercop_apu_ws.cell(row_idx, 6).value, expected_value)
        for label in ("EQUIPOS", "MANO DE OBRA", "MATERIALES", "TRANSPORTE"):
            row_idx = _find_single_row_by_value(sercop_apu_ws, label)
            _assert_light_section_fill(f"Standalone SERCOP encabezado {label}", sercop_apu_ws.cell(row_idx, 1))

        print("SMOKE TEST APU ALIGNMENT: PASS")
        return 0
    except Exception as exc:
        print(f"SMOKE TEST APU ALIGNMENT: FAIL -> {exc}")
        return 1
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
