import sys
from pathlib import Path

import openpyxl
from openpyxl.worksheet.cell_range import CellRange


ROOT = Path(__file__).resolve().parents[2]
BACKEND_DIR = ROOT / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.core.database import SessionLocal
from app.models.apu import APU
from app.models.empresa import Empresa
from app.models.presupuesto import Presupuesto
from app.services.reporting import reporting_service


TARGET_EMPRESA = "Santiago Bermeo"


def _assert(condition, message):
    if not condition:
        raise AssertionError(message)


def _assert_sheet_print_setup(label, ws):
    header_blob = " ".join(
        [
            str(ws.oddHeader),
            str(ws.oddFooter),
            str(ws.evenHeader),
            str(ws.evenFooter),
            str(ws.firstHeader),
            str(ws.firstFooter),
        ]
    ).upper()
    _assert("INTERPRO" not in header_blob, f"{label}: contiene INTERPRO en headers/footers")

    _assert(ws.page_setup.fitToWidth == 1, f"{label}: fitToWidth esperado=1 actual={ws.page_setup.fitToWidth!r}")
    _assert(ws.page_setup.fitToHeight == 0, f"{label}: fitToHeight esperado=0 actual={ws.page_setup.fitToHeight!r}")
    _assert(getattr(getattr(ws.sheet_properties, "pageSetUpPr", None), "fitToPage", None) is True, f"{label}: fitToPage no activo")

    print_area = ws.print_area
    _assert(bool(print_area), f"{label}: print_area vacio")
    area_ref = str(print_area)
    if "!" in area_ref:
        area_ref = area_ref.split("!", 1)[1]
    area_ref = area_ref.replace("$", "").replace("'", "")
    used = CellRange(ws.calculate_dimension().replace("$", ""))
    area = CellRange(area_ref)
    _assert(area.max_col >= used.max_col, f"{label}: print_area no cubre columnas derechas ({area.max_col} < {used.max_col})")
    _assert(area.max_row >= used.max_row, f"{label}: print_area no cubre filas usadas ({area.max_row} < {used.max_row})")


def main() -> int:
    db = SessionLocal()
    try:
        empresa = db.query(Empresa).filter(Empresa.nombre == TARGET_EMPRESA).first()
        if not empresa:
            raise RuntimeError(f"Empresa no encontrada: {TARGET_EMPRESA}")

        presupuesto = (
            db.query(Presupuesto)
            .filter(Presupuesto.empresa_id == empresa.id)
            .order_by(Presupuesto.id.asc())
            .first()
        )
        if not presupuesto:
            raise RuntimeError("No hay presupuesto para la empresa de prueba")

        apu = (
            db.query(APU)
            .filter(APU.empresa_id == empresa.id)
            .order_by(APU.id.asc())
            .first()
        )
        if not apu:
            raise RuntimeError("No hay APU para la empresa de prueba")

        workbooks = {
            "presupuesto_001": openpyxl.load_workbook(reporting_service.generate_presupuesto_report(db, presupuesto.id, empresa.id, "001")),
            "presupuesto_apus_001": openpyxl.load_workbook(reporting_service.generate_presupuesto_report_bundle(db, presupuesto.id, empresa.id, "001")),
            "presupuesto_002": openpyxl.load_workbook(reporting_service.generate_presupuesto_report(db, presupuesto.id, empresa.id, "002")),
            "apu_001": openpyxl.load_workbook(reporting_service.generate_apu_report(db, apu.id, empresa.id, "001")),
            "apu_002": openpyxl.load_workbook(reporting_service.generate_apu_report(db, apu.id, empresa.id, "002")),
        }

        for workbook_label, wb in workbooks.items():
            for ws in wb.worksheets:
                _assert_sheet_print_setup(f"{workbook_label}:{ws.title}", ws)

        print("SMOKE TEST REPORT PRINT SETUP: PASS")
        return 0
    except Exception as exc:
        print(f"SMOKE TEST REPORT PRINT SETUP: FAIL -> {exc}")
        return 1
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
