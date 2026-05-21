import openpyxl
import os

def extract_excel_mapping(file_path):
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for sheetname in wb.sheetnames:
            print(f"--- Sheet: {sheetname} ---")
            ws = wb[sheetname]
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > 50: break
                print("\t".join([str(cell) if cell is not None else "" for cell in row]))
    except Exception as e:
        print(f"Error reading Excel: {e}")

if __name__ == "__main__":
    path = r"e:\Repositorios\GiProy Network\docs\reportes\001 - Presupuesto.xlsx"
    extract_excel_mapping(path)
