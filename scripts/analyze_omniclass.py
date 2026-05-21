import openpyxl
import json
import os

def analyze_sheet(sheet_name, wb):
    sheet = wb[sheet_name]
    data = []
    # Peek at first 10 rows
    for row in sheet.iter_rows(max_row=10, values_only=True):
        data.append(row)
    return data

def main():
    excel_path = 'e:/Repositorios/GiProy Network/tmp/omniclass.xlsx'
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    results = {}
    tables = ['OmniClass Table 21', 'OmniClass Table 22', 'OmniClass Table 23', 'OmniClass Table 34']
    
    for table in tables:
        if table in wb.sheetnames:
            results[table] = analyze_sheet(table, wb)
    
    print(json.dumps(results, indent=2))

if __name__ == "__main__":
    main()
