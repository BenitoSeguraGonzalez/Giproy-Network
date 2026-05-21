
import openpyxl
import os

def get_placeholders(path):
    if not os.path.exists(path): return []
    wb = openpyxl.load_workbook(path)
    phs = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "#" in cell.value:
                    import re
                    # Extract anything starting with # and containing uppercase/underscores
                    found = re.findall(r'#[A-Z0-9_%]+', cell.value)
                    phs.extend(found)
    return sorted(list(set(phs)))

print("VAE:", get_placeholders("e:/Repositorios/GiProy Network/docs/reportes/001 - VAE Proyecto - General.xlsx"))
print("Poly:", get_placeholders("e:/Repositorios/GiProy Network/docs/reportes/001 - Formula Polinomica.xlsx"))
