import openpyxl
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
import os

def patch_apu_template(path):
    print(f"Parcheando {path}...")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    
    # Headers en fila 10, Cols I, J
    ws.cell(row=10, column=9).value = "OmniClass Cod"
    ws.cell(row=10, column=10).value = "Clasificación"
    
    # Placeholders en fila 11, Cols I, J
    ws.cell(row=11, column=9).value = "#OMNICLASS_COD1"
    ws.cell(row=11, column=10).value = "#OMNICLASS_TIT1"
    
    # Copiar estilos de la columna H (Proporción %)
    ref_cell_hdr = ws.cell(row=10, column=8)
    ref_cell_data = ws.cell(row=11, column=8)
    
    from copy import copy
    for col in [9, 10]:
        # Header Style
        ws.cell(row=10, column=col).font = copy(ref_cell_hdr.font)
        ws.cell(row=10, column=col).border = copy(ref_cell_hdr.border)
        ws.cell(row=10, column=col).fill = copy(ref_cell_hdr.fill)
        ws.cell(row=10, column=col).alignment = copy(ref_cell_hdr.alignment)
        
        # Data Style
        ws.cell(row=11, column=col).font = copy(ref_cell_data.font)
        ws.cell(row=11, column=col).border = copy(ref_cell_data.border)
        ws.cell(row=11, column=col).alignment = copy(ref_cell_data.alignment)

    ws.column_dimensions['I'].width = 15
    ws.column_dimensions['J'].width = 40
    
    wb.save(path)
    print("APU template actualizado.")

def patch_budget_template(path):
    print(f"Parcheando {path}...")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    
    # Headers en fila 12, Cols I, J
    ws.cell(row=12, column=9).value = "OmniClass Cod"
    ws.cell(row=12, column=10).value = "Clasificación"
    
    # Placeholders en fila 13, Cols I, J
    ws.cell(row=13, column=9).value = "#OMNICLASS_COD"
    ws.cell(row=13, column=10).value = "#OMNICLASS_TIT"
    
    # Estilos similares
    ref_cell_hdr = ws.cell(row=12, column=8)
    ref_cell_data = ws.cell(row=13, column=8)
    
    from copy import copy
    for col in [9, 10]:
        ws.cell(row=12, column=col).font = copy(ref_cell_hdr.font)
        ws.cell(row=12, column=col).border = copy(ref_cell_hdr.border)
        ws.cell(row=12, column=col).fill = copy(ref_cell_hdr.fill)
        ws.cell(row=12, column=col).alignment = copy(ref_cell_hdr.alignment)
        
        ws.cell(row=13, column=col).font = copy(ref_cell_data.font)
        ws.cell(row=13, column=col).border = copy(ref_cell_data.border)
    
    ws.column_dimensions['I'].width = 15
    ws.column_dimensions['J'].width = 40
    
    wb.save(path)
    print("Budget template actualizado.")

def main():
    base_dir = "e:/Repositorios/GiProy Network/docs/reportes"
    patch_apu_template(os.path.join(base_dir, "001 - Analisis - APUS - General.xlsx"))
    patch_budget_template(os.path.join(base_dir, "001 - Presupuesto.xlsx"))

if __name__ == "__main__":
    main()
